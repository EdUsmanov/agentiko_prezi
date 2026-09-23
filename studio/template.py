from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import colorsys
import re
from PIL import Image
from pptx import Presentation
from defusedxml import ElementTree as SafeET
from .models import TemplateProfile, Pattern, Asset, Box
from .fonts import resolve_font
from .embedded_fonts import extract_embedded_font
from .security import validate_pptx, scan_text, digest, InputRejected

EMU = 12700
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

def luminance(color):
    values = [int(color[i:i+2], 16) / 255 for i in (1, 3, 5)]
    values = [v / 12.92 if v <= 0.04045 else ((v + .055) / 1.055) ** 2.4 for v in values]
    return sum(x * w for x, w in zip(values, [.2126, .7152, .0722]))

def contrast(a, b):
    light, dark = sorted([luminance(a), luminance(b)], reverse=True)
    return (light + .05) / (dark + .05)

def color_value(color):
    try:
        return "#" + str(color.rgb) if color.type and color.rgb else None
    except (AttributeError, ValueError, TypeError):
        return None

def walk_shapes(shapes, sx=1., sy=1., ox=0., oy=0.):
    for shape in shapes:
        box = Box(x=(ox + shape.left * sx) / EMU, y=(oy + shape.top * sy) / EMU,
                  w=shape.width * sx / EMU, h=shape.height * sy / EMU)
        yield shape, box
        if hasattr(shape, "shapes"):
            xf = shape._element.grpSpPr.xfrm
            if xf is not None and xf.chExt is not None:
                nsx, nsy = shape.width / max(xf.chExt.cx, 1), shape.height / max(xf.chExt.cy, 1)
                yield from walk_shapes(shape.shapes, sx * nsx, sy * nsy,
                    ox + sx * (shape.left - xf.chOff.x * nsx), oy + sy * (shape.top - xf.chOff.y * nsy))

def analyze_template(path: Path, artifact_dir: Path) -> TemplateProfile:
    warnings = validate_pptx(path)
    prs = Presentation(path)
    width, height = prs.slide_width / EMU, prs.slide_height / EMU
    if not (100 <= width <= 2500 and 100 <= height <= 2500):
        raise InputRejected("Неподдерживаемые размеры слайда")
    fonts, sizes, colors, text_colors = Counter(), Counter(), Counter(), Counter()
    theme_fonts, theme_colors = [], []
    with ZipFile(path) as z:
        for name in z.namelist():
            if re.fullmatch(r"ppt/theme/theme\d+\.xml", name):
                root = SafeET.fromstring(z.read(name))
                for node in root.iter(A + "latin"):
                    font = node.get("typeface", "")
                    if font and not font.startswith("+"):
                        theme_fonts.append(font)
                for scheme in root.iter(A + "clrScheme"):
                    for slot in scheme:
                        for node in slot:
                            c = node.get("lastClr") or node.get("val", "")
                            if re.fullmatch("[0-9a-fA-F]{6}", c):
                                theme_colors.append("#" + c.upper())
    patterns, counts, photos = [], Counter(), defaultdict(list)
    margins, title_sizes = [], []
    masters_layouts = [m for m in prs.slide_masters] + [l for m in prs.slide_masters for l in m.slide_layouts]
    for idx, surface in enumerate(list(prs.slides) + masters_layouts):
        zones = []
        for shape, box in walk_shapes(surface.shapes):
            if idx < len(prs.slides):
                counts["objects"] += 1
                counts["placeholders"] += bool(shape.is_placeholder)
            if counts["objects"] > 30000:
                raise InputRejected("Превышен лимит 30 000 объектов")
            try:
                c = color_value(shape.fill.fore_color)
                if c:
                    colors[c] += max(1, box.w * box.h / 500)
            except (AttributeError, TypeError):
                pass
            if shape.has_text_frame:
                _, suspicious = scan_text(shape.text, "template")
                if suspicious:
                    warnings.append("Инструкции внутри шаблона проигнорированы; текст шаблона не передаётся модели")
                if shape.text.strip() and box.w > 20 and box.h > 8 and 0 <= box.x < width and 0 <= box.y < height:
                    zones.append(box)
                    if .02 * width < box.x < .15 * width:
                        margins.append(box.x)
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.name and not r.font.name.startswith("+"):
                            fonts[r.font.name] += max(1, len(r.text))
                        if r.font.size:
                            size = round(r.font.size.pt, 1)
                            sizes[size] += max(1, len(r.text))
                            if box.y < height * .22 and size >= 20:
                                title_sizes.append(size)
                        c = color_value(r.font.color)
                        if c:
                            colors[c] += max(1, len(r.text) / 5)
                            if idx < len(prs.slides):
                                text_colors[c] += max(1,len(r.text))
            if idx < len(prs.slides) and hasattr(shape, "image"):
                raw = shape.image.blob
                key = digest(raw)
                # Small repeated assets at the canvas edge are candidate brand marks.
                if box.w * box.h < width * height * .04 and (box.y < height * .15 or box.y + box.h > height * .85):
                    photos[key].append((idx, box, raw))
        if idx < len(prs.slides) and zones:
            zones = sorted(zones, key=lambda b: (b.y, b.x))
            patterns.append(Pattern(id=f"slide-{idx+1}", source_slide=idx+1,
                source_layout=surface.slide_layout.name, text_zones=zones[:16],
                role="columns" if len(zones) >= 3 else "split" if len(zones) == 2 else "statement"))
    allowed_fonts = list(dict.fromkeys([f for f, _ in fonts.most_common()] + theme_fonts))
    if not allowed_fonts:
        raise InputRejected("В шаблоне не найден шрифт")
    font = allowed_fonts[0]
    font_file, font_origin, font_issues = extract_embedded_font(path, artifact_dir, font)
    warnings.extend(font_issues)
    if not font_file:
        font_file = resolve_font(font)
        if font_file:
            font_origin = {"kind":"local", "sha256":digest(Path(font_file).read_bytes()),
                           "template_embedding":font_origin["kind"]}
            warnings.append(f"Шрифт {font}: использовано точное локальное начертание; подходящий встроенный шрифт в PPTX отсутствует или недоступен.")
    if not font_file:
        raise InputRejected(f"Шрифт {font} указан в шаблоне, но его доступных данных нет ни в PPTX, ни в локальном каталоге. Сохраните PPTX с встраиванием всех символов шрифта либо добавьте его TTF в fonts/ и повторите анализ. " + " ".join(font_issues))
    palette = list(dict.fromkeys([c for c, _ in colors.most_common(64)] + theme_colors))
    if not palette:
        raise InputRejected("В шаблоне не найдена палитра")
    primary_text=text_colors.most_common(1)[0][0] if text_colors else min(palette,key=luminance)
    background = max(palette, key=lambda c: contrast(c,primary_text))
    foreground = max(palette, key=lambda c: contrast(c, background))
    saturated = [c for c in palette if colorsys.rgb_to_hsv(*[int(c[i:i+2], 16)/255 for i in (1,3,5)])[1] > .15]
    accent = saturated[0] if saturated else foreground
    scale = sorted(s for s in sizes if 8 <= s <= 80)
    if not scale:
        raise InputRejected("Не удалось извлечь типографическую шкалу")
    title_size = min(scale, key=lambda s: abs(s - min(36, width * .038)))
    body_size = min(scale, key=lambda s: abs(s - min(20, width * .022)))
    margin = sorted(margins)[len(margins)//2] if margins else width * .05
    margin = max(width * .04, min(margin, width * .08))
    assets = []
    asset_dir = artifact_dir / "assets"
    asset_dir.mkdir(parents=True, exist_ok=True)
    for key, entries in photos.items():
        if len({entry[0] for entry in entries}) < 3:
            continue
        _, box, raw = entries[0]
        # A logo is only reused if repeated at essentially the same location.
        if any(abs(b.x-box.x) > 4 or abs(b.y-box.y) > 4 for _, b, _ in entries):
            continue
        try:
            with Image.open(BytesIO(raw)) as im:
                if im.width * im.height > 20_000_000:
                    continue
                if box.h<=0 or abs((im.width/im.height)/(box.w/box.h)-1)>.03:
                    # Cropped or masked source images need a dedicated crop adapter.
                    # Skip rather than stretch the full image into its source frame.
                    continue
                target = asset_dir / f"{key[:20]}.png"
                im.convert("RGBA").save(target)
            assets.append(Asset(id=key[:20], path=str(target), box=box, occurrences=len(entries), role="brand_candidate"))
        except (OSError, ValueError):
            continue
        if len(assets) == 2:
            break
    layouts = list(prs.slide_layouts)
    layout_index = min(range(len(layouts)), key=lambda i: len(layouts[i].shapes)) if layouts else 0
    ratio = counts["placeholders"] / max(counts["objects"], 1)
    return TemplateProfile(sha256=digest(path.read_bytes()), name=path.name, width=width, height=height,
        slide_count=len(prs.slides), master_count=len(prs.slide_masters),
        layout_count=sum(len(m.slide_layouts) for m in prs.slide_masters),
        object_count=counts["objects"], placeholder_count=counts["placeholders"],
        fonts=allowed_fonts, font=font, font_file=font_file, font_origin=font_origin, font_sizes=scale, title_size=title_size, body_size=body_size,
        colors=palette, background=background, foreground=foreground, accent=accent, margin=margin,
        patterns=patterns, assets=assets, warnings=sorted(set(warnings)),
        source_kind="layout_rich" if ratio > .25 else "example_deck", layout_index=layout_index)
