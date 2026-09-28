"""Connect the supplied standalone font extractor to Studio's template profile."""

from __future__ import annotations

import asyncio
import importlib
import hashlib
import shutil
import sys
import tempfile
from collections import Counter
from pathlib import Path

from .config import ROOT
from .fonts import resolve_font
from .powerpoint import open_presentation

KIT = Path(__file__).resolve().parent / "_vendor" / "font_extraction"


def _kit(module: str):
    # The supplied bundle uses absolute imports between its top-level scripts and
    # pyapi package. Keep those sources intact and load them from their own root.
    if str(KIT) not in sys.path:
        sys.path.insert(0, str(KIT))
    return importlib.import_module(module)


class _StudioResolver:
    """Offer Studio's user-provided fonts before the bundle's public providers."""

    def __init__(self, downloads, directory: Path):
        self.downloads = downloads
        self.directory = directory

    async def resolve_variants(self, requests):
        match = _kit("pyapi.infrastructure.font_downloads")._matches
        found = []
        pending = {family: set(variants) for family, variants in requests.items()}
        roots = (
            ROOT / "fonts",
            ROOT / "data" / "local-fonts",
            self.directory.parent.parent / "local-fonts",
        )
        files = sorted(
            path
            for root in roots
            if root.is_dir()
            for path in root.rglob("*")
            if path.is_file() and path.suffix.lower() in {".ttf", ".otf"}
        )
        for path in files:
            try:
                if path.stat().st_size > 8 * 1024 * 1024:
                    continue
                data = path.read_bytes()
            except OSError:
                continue
            for family, variants in pending.items():
                for weight, style in tuple(variants):
                    if match(data, family, weight, style):
                        found.append(
                            {
                                "family": family,
                                "resolvedFamily": family,
                                "weight": weight,
                                "style": style,
                                "source": "installed",
                                "data": data,
                            }
                        )
                        variants.remove((weight, style))
        # The previous pipeline also searched user and system font directories.
        # Resolve exact faces there when Fontconfig is absent (notably on macOS).
        for family, variants in pending.items():
            for weight, style in tuple(variants):
                path = resolve_font(_display_face(family, weight, style))
                if not path:
                    continue
                try:
                    data = Path(path).read_bytes()
                except OSError:
                    continue
                if len(data) <= 8 * 1024 * 1024 and match(data, family, weight, style):
                    found.append(
                        {
                            "family": family,
                            "resolvedFamily": family,
                            "weight": weight,
                            "style": style,
                            "source": "installed",
                            "data": data,
                        }
                    )
                    variants.remove((weight, style))
        missing = {family: variants for family, variants in pending.items() if variants}
        if missing and self.downloads is not None:
            found.extend(await self.downloads.resolve_variants(missing))
        return found


def _display_face(family: str, weight: int, style: str) -> str:
    label = family
    lower = label.casefold()
    names = {
        100: "Thin",
        200: "ExtraLight",
        300: "Light",
        500: "Medium",
        600: "SemiBold",
        700: "Bold",
        800: "ExtraBold",
        900: "Black",
    }
    suffix = names.get(weight)
    if suffix and suffix.casefold() not in lower:
        label += " " + suffix
    if style == "italic" and "italic" not in lower and "oblique" not in lower:
        label += " Italic"
    return label


def _role_bindings(model, assets):
    by_id = {asset["id"]: asset for asset in assets}
    counts = {}
    for slide in model.slides:
        for role, uses in (slide.elements or {}).items():
            target = "body" if role in {"text", "other"} else role
            if target not in {"title", "body", "table", "chart", "footer", "subtitle"}:
                continue
            bucket = counts.setdefault(target, Counter())
            for use in uses:
                key = (use.family, use.weight, use.style)
                bucket[key] += max(1, use.characters or 0)
    selected = {}
    for role, bucket in counts.items():
        if bucket:
            selected[role] = bucket.most_common(1)[0][0]
    if "body" not in selected and selected:
        selected["body"] = next(iter(selected.values()))
    if "title" not in selected and "body" in selected:
        selected["title"] = selected["body"]
    asset_by_key = {(a["family"], a["weight"], a["style"]): a for a in assets}
    roles = {role: asset_by_key[key]["id"] for role, key in selected.items() if key in asset_by_key}
    primary = by_id.get(roles.get("body")) or next(iter(assets), None)
    return roles, selected, primary


def _bundled_fallback(weight: int, assets: list[dict]) -> dict:
    label = "Bold" if weight >= 600 else "Medium" if weight >= 450 else "Regular"
    path = ROOT / "fonts" / f"Montserrat-{label}.ttf"
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    existing = next((asset for asset in assets if asset["sha256"] == digest), None)
    if existing:
        return existing
    asset = {
        "id": digest[:20],
        "requested": "Montserrat" if label == "Regular" else f"Montserrat {label}",
        "family": "Montserrat",
        "path": str(path),
        "sha256": digest,
        "bytes": len(raw),
        "status": "substituted",
        "weight": {"Regular": 400, "Medium": 500, "Bold": 700}[label],
        "style": "normal",
        "italic": False,
        "redistributable": True,
        "origin": {"kind": "missing_font_fallback", "source": "bundled", "sha256": digest},
    }
    assets.append(asset)
    return asset


def _apply_missing_font_fallbacks(model, assets, roles, selected, allow_download):
    """Keep the kit's exact report intact and bind unresolved output faces explicitly."""
    by_key = {(a["family"], a["weight"], a["style"]): a for a in assets}
    replacements = []
    fallback_by_key = {}
    for item in model.unresolved:
        key = (item.family, item.weight, item.style)
        same_family = [a for a in assets if a["family"].casefold() == item.family.casefold()]
        if same_family:
            fallback = min(
                same_family,
                key=lambda a: (a["style"] != item.style, abs(a["weight"] - item.weight)),
            )
        else:
            preferred = next(
                (
                    by_key[choice]
                    for role in ("body", "title")
                    if (choice := selected.get(role)) in by_key
                ),
                None,
            )
            if preferred is None:
                fallback = _bundled_fallback(item.weight, assets)
            else:
                fallback = min(
                    (a for a in assets if a["family"] == preferred["family"]),
                    key=lambda a: (a["style"] != item.style, abs(a["weight"] - item.weight)),
                )
        requested = _display_face(*key)
        fallback["template_aliases"] = list(
            dict.fromkeys([*fallback.get("template_aliases", []), requested])
        )
        fallback_by_key[key] = fallback
        replacements.append(
            {
                "scope": "font",
                "template_font": requested,
                "fallback_font": fallback["requested"],
                "reason": "missing_font",
                "download_attempted": allow_download,
                "style_changed": (item.weight, item.style)
                != (fallback["weight"], fallback["style"]),
            }
        )
    for role, key in selected.items():
        if role not in roles and key in fallback_by_key:
            roles[role] = fallback_by_key[key]["id"]
    by_id = {a["id"]: a for a in assets}
    primary = by_id.get(roles.get("body")) or by_id.get(roles.get("title"))
    if primary is None:
        primary = _bundled_fallback(400, assets)
        roles.update(body=primary["id"], title=primary["id"])
    return primary, replacements, fallback_by_key


async def _extract(path: Path, directory: Path, allow_download: bool, layout_face: str | None):
    delivery = _kit("delivery")
    decoder = _kit("decoder")
    providers = _kit("pyapi.infrastructure.font_downloads")
    downloads = providers.DownloadableFontResolver() if allow_download else None
    try:
        service = delivery.FontExtractionService(
            decoder.EmbeddedFontDecoder(), _StudioResolver(downloads, directory)
        )
        model = await service.extract_file(path, directory)
        if model.slides:
            return model
        # POTX files can consist only of layouts. Give the same extractor
        # representative slides so it can resolve inherited layout typography.
        with tempfile.TemporaryDirectory(prefix=".font-layout-", dir=directory) as scratch:
            workspace = Path(scratch)
            presentation = open_presentation(path)
            for layout in presentation.slide_layouts:
                slide = presentation.slides.add_slide(layout)
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        shape.text = "Aa"
            if layout_face:
                from pptx.util import Pt

                probe = presentation.slides.add_slide(presentation.slide_layouts[-1])
                for y, size in ((0, 32), (presentation.slide_height / 2, 20)):
                    shape = probe.shapes.add_textbox(0, int(y), presentation.slide_width, Pt(80))
                    run = shape.text_frame.paragraphs[0].add_run()
                    run.text = "Aa" * 32
                    run.font.name = layout_face
                    run.font.size = Pt(size)
            if not len(presentation.slides):
                return model
            reference = workspace / "layout-reference.pptx"
            presentation.save(reference)
            layout_output = workspace / "analysis"
            layout_model = await service.extract_file(reference, layout_output)
            for asset in layout_model.fontAssets:
                destination = directory / asset.path
                destination.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(layout_output / asset.path, destination)
                if hashlib.sha256(destination.read_bytes()).hexdigest() != asset.sha256:
                    raise ValueError("Layout font asset checksum mismatch")
            (directory / "layout-font-model.json").write_text(
                layout_model.model_dump_json(indent=2, exclude_none=True)
            )
            return layout_model
    finally:
        if downloads is not None:
            await downloads.close()


def extract_template_fonts(
    path: Path, directory: Path, allow_download=False, progress=None, layout_face=None
):
    """Run bundle v2 and adapt its validated output to Studio's rendering model."""
    if progress:
        progress("Ищем шрифты в шаблоне, системе и открытых каталогах")
    model = asyncio.run(_extract(path, directory, allow_download, layout_face))
    assets = []
    for item in model.fontAssets:
        file = directory / item.path
        local_only = item.origin == "microsoft-aptos" or item.family.casefold().startswith("aptos")
        assets.append(
            {
                "id": item.id,
                "requested": _display_face(item.family, item.weight, item.style),
                "family": item.family,
                "path": str(file),
                "sha256": item.sha256,
                "bytes": item.bytes,
                "status": item.source,
                "weight": item.weight,
                "style": item.style,
                "italic": item.style == "italic",
                "redistributable": not local_only,
                "origin": {
                    "kind": item.source,
                    "provider": item.origin,
                    "sha256": item.sha256,
                },
            }
        )
    roles, selected, _ = _role_bindings(model, assets)
    primary, replacements, fallback_by_key = _apply_missing_font_fallbacks(
        model, assets, roles, selected, allow_download
    )
    missing = [
        {
            "requested": _display_face(item.family, item.weight, item.style),
            "family": item.family,
            "weight": item.weight,
            "style": item.style,
            "reason": "Точное начертание недоступно; использована подстановка",
            "required_for_generation": False,
            "substituted_by": fallback_by_key[(item.family, item.weight, item.style)]["requested"],
        }
        for item in model.unresolved
    ]
    return {
        "assets": assets,
        "roles": roles,
        "primary": primary,
        "unresolved": missing,
        "replacements": replacements,
        "warnings": [
            warning
            for warning in model.warnings
            if not any(item["family"] in warning for item in missing)
        ],
    }
