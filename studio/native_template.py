"""Source geometry and native artwork, without copying source narrative into new slides."""

from .shape_geometry import box as box, intersects as intersects
from .native_surface import (
    scrub_surface as scrub_surface,
    copy_node as copy_node,
    source_slide as source_slide,
)

from collections import Counter
from copy import deepcopy
from pathlib import Path
import re
from pptx.enum.shapes import PP_PLACEHOLDER, MSO_SHAPE_TYPE
from .models import Pattern, Box, SlideScene
from .pictures import is_picture

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def text_bounds(shape):
    """Respect source text insets: card icons often occupy the top padding."""
    bounds = box(shape)
    if not shape.has_text_frame:
        return bounds
    frame = shape.text_frame
    left, right, top, bottom = [
        (getattr(frame, "margin_" + side) or 0) / 12700
        for side in ("left", "right", "top", "bottom")
    ]
    return Box(
        x=bounds.x + left, y=bounds.y + top, w=bounds.w - left - right, h=bounds.h - top - bottom
    )


def evidence_placeholder(shape):
    # An authored empty illustration container can hold a chart/table. A photo
    # label or existing picture is not permission to cover source artwork.
    return shape.has_text_frame and shape.text.strip().casefold() in (
        "иллюстрация",
        "изображение",
        "illustration",
        "image",
    )


def image_regions(surface):
    return [
        box(sh)
        for sh in surface.shapes
        if evidence_placeholder(sh)
        or sh.is_placeholder
        and sh.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    ]


def title_bounds(shape, surface, layout):
    """A text placeholder can be wider than its visible title badge."""
    bounds = text_bounds(shape)
    # A single-line title may use its lower inset for line leading, while its
    # actual glyphs still remain inside the original shape/container.
    bounds.h = box(shape).y + box(shape).h - bounds.y
    for owner in (surface, layout):
        for sh in owner.shapes:
            if sh.is_placeholder or sh.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or sh == shape:
                continue
            b = box(sh)
            try:
                filled = sh.fill.type is not None and sh.fill.type != 5
            except (AttributeError, ValueError):
                filled = False
            if (
                filled
                and b.h < bounds.h * 3
                and b.w >= bounds.w * 0.2
                and b.x <= bounds.x <= b.x + b.w * 0.5
                and b.y <= bounds.y <= b.y + b.h
                and b.y + b.h >= bounds.y + bounds.h * 0.8
            ):
                right = min(bounds.x + bounds.w, b.x + b.w - 10)
                bottom = min(bounds.y + bounds.h, b.y + b.h - 4)
                if right > bounds.x + 40 and bottom > bounds.y + 8:
                    bounds = Box(x=bounds.x, y=bounds.y, w=right - bounds.x, h=bottom - bounds.y)
    # Empty auto-shapes still have a text frame. Large decorative shapes are
    # protected artwork too, not free space inside an oversized placeholder.
    from .template_geometry import walk_shapes

    for owner in (surface, layout, layout.slide_master):
        for sh, b in walk_shapes(owner.shapes):
            if (
                sh.is_placeholder
                or (sh.has_text_frame and sh.text.strip())
                or not intersects(bounds, b)
            ):
                continue
            try:
                visible = is_picture(sh) or sh.fill.type is not None and sh.fill.type != 5
            except (AttributeError, ValueError):
                visible = is_picture(sh)
            if visible and b.x > bounds.x + 60 and b.w < bounds.w * 0.9:
                bounds.w = min(bounds.w, b.x - bounds.x - 10)
    return bounds


def inherited_title_size(shape, layout):
    candidates = [shape]
    if shape.is_placeholder:
        candidates.extend(
            s
            for s in layout.placeholders
            if s.placeholder_format.idx == shape.placeholder_format.idx
        )
    for candidate in candidates:
        for paragraph in candidate.text_frame.paragraphs:
            sizes = [r.font.size.pt for r in paragraph.runs if r.font.size]
            if paragraph.font.size:
                sizes.append(paragraph.font.size.pt)
            if sizes:
                return max(sizes)
    return 0


def geometry_surface(surface):
    """Read grouped fields in slide coordinates; never mutate the uploaded OOXML."""
    from types import SimpleNamespace
    from pptx.shapes.shapetree import BaseShapeFactory
    from pptx.util import Pt
    from .template_geometry import walk_shapes

    shapes = []
    for shape, bounds in walk_shapes(surface.shapes):
        if hasattr(shape, "shapes"):
            continue
        clone = BaseShapeFactory(deepcopy(shape._element), surface.shapes)
        clone.left = Pt(bounds.x)
        clone.top = Pt(bounds.y)
        clone.width = Pt(bounds.w)
        clone.height = Pt(bounds.h)
        # Scale local text insets along with a scaled group.
        if shape.has_text_frame and shape.width and shape.height:
            for side in ("left", "right", "top", "bottom"):
                ratio = (
                    bounds.w * 12700 / shape.width
                    if side in ("left", "right")
                    else bounds.h * 12700 / shape.height
                )
                setattr(
                    clone.text_frame,
                    "margin_" + side,
                    int((getattr(shape.text_frame, "margin_" + side) or 0) * ratio),
                )
        shapes.append(clone)
    return SimpleNamespace(shapes=shapes, part=surface.part)


def native_patterns(prs, styles=None):
    w, h = prs.slide_width / 12700, prs.slide_height / 12700
    layouts = {
        str(layout.part.partname): (mi, li)
        for mi, m in enumerate(prs.slide_masters)
        for li, layout in enumerate(m.slide_layouts)
    }
    surfaces = [(i + 1, s, s.slide_layout) for i, s in enumerate(prs.slides)]
    surfaces += [(0, layout, layout) for m in prs.slide_masters for layout in m.slide_layouts]
    result = []
    for number, surface, layout in surfaces:
        surface = geometry_surface(surface)

        def style(shape):
            return (styles or {}).get((str(surface.part.partname).lstrip("/"), shape.shape_id), {})

        def title_size(shape):
            return style(shape).get("size") or inherited_title_size(shape, layout)

        titles, bodies = [], []
        for sh in surface.shapes:
            if not sh.is_placeholder:
                continue
            if sh.has_text_frame and sh.text_frame._txBody.bodyPr.get("vert", "horz") != "horz":
                continue  # Vertical writing needs a different layout engine.
            kind = sh.placeholder_format.type
            if kind in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                titles.append(sh)
            elif kind in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE):
                bodies.append(sh)
        if number and not titles and not bodies:
            # Example decks may have no placeholders: use small, unambiguous text structures.
            text_shapes = [
                sh
                for sh in surface.shapes
                if sh.has_text_frame
                and sh.text.strip()
                and not evidence_placeholder(sh)
                and sh.width / 12700 > w * 0.18
                and sh.height / 12700 > 8
                and sh.top / 12700 < h * 0.85
            ]
            if 2 <= len(text_shapes) <= 20:
                # A small brand label at the top is not the slide headline.
                main = max(text_shapes, key=lambda sh: (title_size(sh), sh.height))
                rest = [
                    sh
                    for sh in text_shapes
                    if sh != main
                    and sh.top >= main.top - 2 * 12700
                    and not intersects(text_bounds(sh), text_bounds(main))
                    and sh.height / 12700 >= max(8, h * 0.03)
                ]
                if rest:
                    titles, bodies = [main], rest
                elif main.height / 12700 >= h * 0.07:
                    titles = [main]
            elif len(text_shapes) == 1 and text_shapes[0].height / 12700 >= h * 0.07:
                titles = text_shapes
        if number and len(titles) == 1 and not bodies:
            # Mixed examples: title is a placeholder, content lives in ordinary
            # cards/text boxes. Ignore tiny labels and image instructions.
            current_title = box(titles[0])
            candidates = [
                sh
                for sh in surface.shapes
                if not sh.is_placeholder
                and sh.has_text_frame
                and sh.text_frame._txBody.bodyPr.get("vert", "horz") == "horz"
                and sh.text.strip()
                and sh.width / 12700 > w * 0.12
                and sh.height / 12700 >= 8
                and sh.top / 12700 >= current_title.y + current_title.h - 2
                and sh.text.strip().casefold()
                not in ("иллюстрация", "изображение", "фото", "image", "illustration", "photo")
            ]
            selected = []
            for sh in sorted(candidates, key=lambda s: s.width * s.height, reverse=True):
                if not any(intersects(box(sh), box(other)) for other in selected):
                    selected.append(sh)
            if 1 <= len(selected) <= 24:
                bodies = selected
        if len(titles) != 1:
            continue
        title = title_bounds(titles[0], surface, layout)
        divider_name = bool(re.search(r"section|divider|раздел|разделител", layout.name, re.I))
        if not bodies or divider_name:
            # Title-only / named section layouts were previously dropped because
            # the content catalog required a body zone. Keep them separately.
            if (
                title.w > 40
                and title.h > 8
                and title.x >= 0
                and title.y >= 0
                and title.x + title.w <= w + 1
                and title.y + title.h <= h + 1
            ):
                mi, li = layouts[str(layout.part.partname)]
                result.append(
                    Pattern(
                        id=f"native-slide-{number}" if number else f"native-layout-{mi}-{li}",
                        source_slide=number,
                        source_layout=layout.name,
                        master_index=mi,
                        layout_index=li,
                        title_zone=title,
                        text_zones=[title],
                        role="divider",
                        image_zones=image_regions(surface),
                        title_size=title_size(titles[0]),
                        title_foreground=style(titles[0]).get("color", ""),
                    )
                )
            continue
        # Content-rich layouts contain labels/numbers PLUS large body placeholders.
        # Pair the short heading above each body instead of rejecting the whole layout.
        # Large short labels above smaller copy are native subheadings (years,
        # card titles), not extra paragraphs. Preserve them as paired fields.
        heading_shapes = [
            s
            for s in bodies
            if s.has_text_frame
            and len(s.text.strip()) <= 60
            and any(
                other != s
                and abs(text_bounds(s).x - text_bounds(other).x) < 15
                and 0 <= text_bounds(other).y - (text_bounds(s).y + text_bounds(s).h) <= 40
                and text_bounds(s).w >= text_bounds(other).w * 0.65
                and (
                    (title_size(s) > title_size(other) * 1.3 and title_size(other) > 0)
                    or text_bounds(s).h < text_bounds(other).h * 0.45
                )
                for other in bodies
            )
        ]
        body_shapes = [
            s
            for s in bodies
            if s not in heading_shapes
            and text_bounds(s).w >= w * 0.12
            and text_bounds(s).h >= max(8, h * 0.03)
        ]
        body_shapes = sorted(body_shapes, key=lambda s: (s.top, s.left))
        selected = []
        for sh in sorted(body_shapes, key=lambda s: s.height * s.width):
            if not any(intersects(text_bounds(sh), text_bounds(other)) for other in selected):
                selected.append(sh)
        selected.sort(key=lambda s: (round(s.top / 12700 / 20), s.left))
        if not 1 <= len(selected) <= 12:
            continue
        title = title_bounds(titles[0], surface, layout)
        zones = [text_bounds(s) for s in selected]
        from .artwork import picture_safe_body_zone

        zones = [picture_safe_body_zone(surface, zone, w, h) for zone in zones]
        # Some source layouts contain placeholders extending beyond the canvas
        # or under the title. Retain their authored columns, using only the safe
        # intersection of their rectangles (never expand into unrelated space).
        for zone in zones:
            zone.w = min(zone.w, w - zone.x)
            zone.h = min(zone.h, h - zone.y - 8)
        underneath = [z.y for z in zones if z.y > title.y + 20 and intersects(title, z)]
        if underneath:
            title.h = min(title.h, min(underneath) - title.y - 8)
        headings = []
        numbers = []
        for body in selected:
            b = text_bounds(body)
            labels = [
                text_bounds(s)
                for s in bodies
                if s not in selected
                and abs(text_bounds(s).x - b.x) < 15
                and text_bounds(s).w >= b.w * 0.65
                and 0 <= b.y - (text_bounds(s).y + text_bounds(s).h) <= 50
            ]
            headings.append(max(labels, key=lambda z: z.y) if labels else None)
            number_slots = [
                text_bounds(s)
                for s in bodies
                if s not in selected
                and abs(text_bounds(s).x - b.x) < 15
                and 30 <= text_bounds(s).w < b.w * 0.6
                and 0 < b.y - (text_bounds(s).y + text_bounds(s).h) < 140
            ]
            numbers.append(min(number_slots, key=lambda z: z.y) if number_slots else None)
        if any(
            b.x < 0
            or b.y < 0
            or b.w < w * 0.12
            or b.h < 8
            or b.x + b.w > w + 1
            or b.y + b.h > h + 1
            for b in [title] + zones
        ):
            continue
        if any(
            intersects(a, b)
            for i, a in enumerate([title] + zones)
            for b in ([title] + zones)[i + 1 :]
        ):
            continue
        # Picture-driven/team/grid layouts need a semantic image adapter, not empty photo slots.
        is_cover = (
            titles[0].is_placeholder
            and titles[0].placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE
        )
        if is_cover:
            for region in [title] + zones:
                if region.x < w * 0.04:
                    right = region.x + region.w
                    region.x = w * 0.04
                    region.w = right - region.x
        if not is_cover and any(
            sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.PICTURE
            for sh in surface.shapes
        ):
            continue
        mi, li = layouts[str(layout.part.partname)]
        _sizes = [
            r.font.size.pt for p in titles[0].text_frame.paragraphs for r in p.runs if r.font.size
        ]
        result.append(
            Pattern(
                id=f"native-slide-{number}" if number else f"native-layout-{mi}-{li}",
                source_slide=number,
                source_layout=layout.name,
                master_index=mi,
                layout_index=li,
                title_zone=title,
                body_zones=zones,
                text_zones=[title] + zones,
                title_size=title_size(titles[0]),
                heading_zones=headings,
                number_zones=numbers,
                title_foreground=style(titles[0]).get("color", ""),
                image_zones=image_regions(surface),
                zone_foregrounds=[style(s).get("color", "") for s in selected],
                foreground=style(selected[0]).get("color", ""),
                graphic_count=sum(
                    not s.is_placeholder
                    and not s.has_text_frame
                    or not s.is_placeholder
                    and not s.text.strip()
                    for s in surface.shapes
                ),
                role="cover"
                if is_cover
                else "columns"
                if len(zones) > 2
                else "split"
                if len(zones) == 2
                else "statement",
            )
        )
    # Persist real object IDs as well as rectangles. A layout-only pattern uses
    # IDs from its layout; a slide exemplar uses IDs from that exact slide.
    for pattern in result:
        surface = (
            prs.slides[pattern.source_slide - 1]
            if pattern.source_slide
            else prs.slide_masters[pattern.master_index].slide_layouts[pattern.layout_index]
        )
        layout = surface.slide_layout if pattern.source_slide else surface
        surface = geometry_surface(surface)

        def style(shape):
            return (styles or {}).get((str(surface.part.partname).lstrip("/"), shape.shape_id), {})

        roles = [
            ("title", [pattern.title_zone]),
            ("body", pattern.body_zones),
            ("heading", pattern.heading_zones),
            ("number", pattern.number_zones),
            ("image", pattern.image_zones),
        ]
        for role, zones in roles:
            for index, zone in enumerate(zones):
                if zone is None:
                    continue
                matches = []
                for sh in surface.shapes:
                    if role != "image" and not sh.has_text_frame:
                        continue
                    b = box(sh)
                    if (
                        b.x <= zone.x + 1
                        and b.y <= zone.y + 1
                        and b.x + b.w >= zone.x + zone.w - 1
                        and b.y + b.h >= zone.y + zone.h - 1
                    ):
                        matches.append(sh)
                if matches:
                    sh = min(matches, key=lambda s: s.width * s.height)
                    size = style(sh).get("size") or inherited_title_size(sh, layout)
                    pattern.fields.append(
                        {
                            "role": role,
                            "index": index,
                            "shape_id": sh.shape_id,
                            "box": zone.model_dump(),
                            "font_size": size or 0,
                            "style": style(sh),
                        }
                    )
                    if role == "image" and evidence_placeholder(sh):
                        pattern.fields[-1]["evidence_placeholder"] = True
        # Non-field sample wording is not immutable brand artwork. Explicit
        # empty assignments prevent old labels from leaking through PPTAgent.
        assigned = {f["shape_id"] for f in pattern.fields}
        for sh in surface.shapes:
            if sh.shape_id not in assigned and sh.has_text_frame and sh.text.strip():
                pattern.fields.append(
                    {
                        "role": "unused",
                        "index": len(pattern.fields),
                        "shape_id": sh.shape_id,
                        "box": text_bounds(sh).model_dump(),
                    }
                )
    return result


def compile_backgrounds(profile, source, directory):
    """Preparation has no five-minute limit: render sanitized artwork once and cache it."""
    from .office import executable, to_pdf
    from .render import render_pptx, _pdfium_lock
    from .template_geometry import contrast
    from .portable_templates import extract_backgrounds, inspect_text_zone

    model = extract_backgrounds(profile, source, directory)
    zone_report = {}
    if not executable() or not profile.patterns:
        profile.warnings.append(
            "LibreOffice недоступен: фон HTML/PDF может отличаться от PPTX; требуется визуальная проверка."
        )
        return
    folder = Path(directory) / "template-layers"
    folder.mkdir(exist_ok=True)
    scenes = [
        SlideScene(
            title="",
            background=profile.background,
            elements=[],
            source_ids=[],
            layout=p.role,
            pattern_id=p.id,
            strategy="native_template",
        )
        for p in profile.patterns
    ]
    render_pptx(scenes, profile, source, folder / "artwork.pptx", verify_text=False)
    from .fonts import profile_font_files

    to_pdf(folder / "artwork.pptx", folder, timeout=120, font_files=profile_font_files(profile))
    import pypdfium2 as pdfium

    with _pdfium_lock, pdfium.PdfDocument(str(folder / "artwork.pdf")) as doc:
        if len(doc) != len(profile.patterns):
            raise ValueError("Неверное число отрисованных шаблонных паттернов")
        for i, pattern in enumerate(profile.patterns):
            page = doc[i]
            bitmap = page.render(scale=1.5)
            im = bitmap.to_pil()
            target = folder / (pattern.id + ".png")
            im.save(target)
            pattern.background_image = str(target)
            # Use representative pixels INSIDE the actual body zone, not palette contrast guesses.
            rgb = im.convert("RGB")
            zone_report[pattern.id] = inspect_text_zone(rgb, pattern, profile, model)
            # Keep every authored field separate, but constrain its text/media
            # viewport against the COMPLETE rendered art (including masters).
            # A global free-zone result may be unknown on complex templates.
            from .artwork import constrain_body_zones

            constrain_body_zones(pattern, rgb)
            backgrounds = []
            foregrounds = []
            schemes = {scheme["id"]: scheme for scheme in profile.color_schemes}
            paired = schemes.get(pattern.color_scheme_id, {}).get("foreground", "")
            source_colors = [pattern.title_foreground] + pattern.zone_foregrounds
            for zi, zone in enumerate([pattern.title_zone] + pattern.body_zones):
                samples = [
                    rgb.getpixel(
                        (
                            min(im.width - 1, int((zone.x + zone.w * x) * 1.5)),
                            min(im.height - 1, int((zone.y + zone.h * y) * 1.5)),
                        )
                    )
                    for x in (0.2, 0.5, 0.8)
                    for y in (0.2, 0.5, 0.8)
                ]
                color = Counter(samples).most_common(1)[0][0]
                bg = "#%02X%02X%02X" % color
                # Contrast is a diagnostic, not permission to repaint the brand.
                original = source_colors[zi] if zi < len(source_colors) else ""
                candidate = original or paired
                foregrounds.append(
                    candidate
                    if candidate and contrast(candidate, bg) >= 3
                    else max(profile.colors, key=lambda c: contrast(c, bg))
                )
                backgrounds.append(bg)
            pattern.title_background = backgrounds[0]
            pattern.title_foreground = foregrounds[0]
            pattern.zone_backgrounds = backgrounds[1:]
            pattern.zone_foregrounds = foregrounds[1:]
            pattern.background = backgrounds[-1]
            pattern.foreground = foregrounds[-1]
            profile.colors = list(dict.fromkeys(profile.colors + backgrounds + foregrounds))
            bitmap.close()
            page.close()
    import json

    (Path(directory) / "text-zones.json").write_text(
        json.dumps(zone_report, ensure_ascii=False, indent=2)
    )
    from .colors import resolve_rendered_schemes

    resolve_rendered_schemes(profile)
