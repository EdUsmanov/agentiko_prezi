"""Adapters for the colleague's unmodified portable background and zone packages.

No optional VL client is started here: provider policy stays with ModelGateway.
Authored fields remain the editable contract; a free rectangle is evidence, not
permission to replace a multi-column layout with invented coordinates.
"""

from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile
from defusedxml import ElementTree as SafeET
from defusedxml.common import DefusedXmlException
from pptx.oxml.ns import qn

from studio._vendor.portable_background_extractor.bgextract import (
    inspect_template_backgrounds,
    extract_background_pptx,
)
from studio._vendor.portable_background_extractor.bgextract.roles import (
    protected_background_regions,
    apply_background_roles,
)
from studio._vendor.portable_text_zone_finder.textzone import analyze_image
from studio.composition.powerpoint import open_presentation
from studio.security import digest, validate_pptx


def _use_ole_previews(prs):
    """Classify embedded OLE preview images as pictures, without opening OLE data."""
    changed = False
    for part in list(prs.part.package.iter_parts()):
        root = getattr(part, "_element", None)
        if root is None:
            continue
        for frame in list(root.iter(qn("p:graphicFrame"))):
            preview = frame.find(f".//{qn('p:oleObj')}/{qn('p:pic')}")
            if preview is None:
                continue
            blip = preview.find(f"{qn('p:blipFill')}/{qn('a:blip')}")
            rel = part.rels.get(blip.get(qn("r:embed"), "")) if blip is not None else None
            if rel is None or rel.is_external or not rel.reltype.endswith("/image"):
                continue
            picture = deepcopy(preview)
            identity = picture.find(f"{qn('p:nvPicPr')}/{qn('p:cNvPr')}")
            original = frame.find(f"{qn('p:nvGraphicFramePr')}/{qn('p:cNvPr')}")
            if identity is None or original is None:
                continue
            identity.getparent().replace(identity, deepcopy(original))
            transform = frame.find(qn("p:xfrm"))
            props = picture.find(qn("p:spPr"))
            if transform is None or props is None:
                continue
            current = props.find(qn("a:xfrm"))
            if current is not None:
                props.remove(current)
            transform = deepcopy(transform)
            transform.tag = qn("a:xfrm")
            props.insert(0, transform)
            frame.getparent().replace(frame, picture)
            changed = True
    return changed


def template_reference(source):
    """Normalize OLE previews and add an exemplar for layout-only POTX files."""
    raw = Path(source).read_bytes()
    prs = open_presentation(BytesIO(raw))
    changed = _use_ole_previews(prs)
    if not len(prs.slides):
        # POTX may contain only authored layouts. A temporary exemplar lets the
        # extractor inspect its inheritance; it is never added to the user file.
        prs.slides.add_slide(prs.slide_masters[0].slide_layouts[0])
        changed = True
    if changed:
        stream = BytesIO()
        prs.save(stream)
        raw = stream.getvalue()
    return raw


def extract_backgrounds(profile, source, directory):
    validate_pptx(Path(source))
    raw = template_reference(source)
    prs = open_presentation(BytesIO(raw))
    model = inspect_template_backgrounds(raw)
    protect_field_surfaces(model, prs, profile.patterns, raw)
    review_path = Path(directory) / "raster-review.json"
    if review_path.is_file():
        review = json.loads(review_path.read_text())
        if review.get("source_sha256") != digest(Path(source).read_bytes()):
            raise ValueError("VL-анализ фона не соответствует исходному шаблону")
        model["rasterReview"] = {
            "status": review["status"],
            "model": review["model"],
            "assets": review["assets"],
        }
        model["rasterRegions"] = review["regions"]
    cleaned = extract_background_pptx(raw, model)
    folder = Path(directory) / "template-layers"
    folder.mkdir(exist_ok=True)
    target = folder / "background-source.pptx"
    target.write_bytes(cleaned)
    validate_pptx(target)
    profile.background_source = str(target)
    report = Path(directory) / "background-model.json"
    report.write_text(json.dumps(model, ensure_ascii=False, indent=2))
    return model


def zone_metadata(pattern, profile, model):
    if pattern.source_slide:
        slide = model["slides"][pattern.source_slide - 1]
    else:
        prs = open_presentation(profile.background_source)
        layout = prs.slide_masters[pattern.master_index].slide_layouts[pattern.layout_index]
        parts = [
            str(layout.part.partname).lstrip("/"),
            str(layout.slide_master.part.partname).lstrip("/"),
        ]
        objects = [row for part in parts for row in model["parts"].get(part, [])]
        slide = {
            "objects": objects,
            "protectedRegions": protected_background_regions(
                objects, profile.width * 12700, profile.height * 12700
            ),
        }
    return slide


def inspect_text_zone(image, pattern, profile, model, vl_cells=None):
    size = (1280, round(1280 * profile.height / profile.width))
    image = image.convert("RGB").resize(size)
    slide = zone_metadata(pattern, profile, model)
    result = analyze_image(image, slide, vl_cells=vl_cells)
    result["coordinate_space"] = {
        "width": size[0],
        "height": size[1],
        "unit": "pixels",
        "box": "left,top,right,bottom",
    }
    result["application"] = "advisory; authored editable fields retained"
    pattern.safe_text_zone = result
    return result


def clean_editable_source(prs, patterns):
    """Use the SAME peer classification for PPTAgent without flattening groups.

    The full extractor is used for rendered backgrounds. On the editing source,
    apply its decisions directly to native XML so group transforms and object
    identities remain stable. Authored contract fields are protected explicitly.
    """
    _use_ole_previews(prs)
    stream = BytesIO()
    prs.save(stream)
    raw = stream.getvalue()
    model = inspect_template_backgrounds(raw)
    protect_field_surfaces(model, prs, patterns, raw)
    protected = {}
    for pattern in patterns:
        surface = (
            prs.slides[pattern.source_slide - 1]
            if pattern.source_slide
            else prs.slide_masters[pattern.master_index].slide_layouts[pattern.layout_index]
        )
        key = str(surface.part.partname).lstrip("/")
        protected.setdefault(key, set()).update(
            str(f["shape_id"]) for f in pattern.fields if f["role"] != "image"
        )
    for part in prs.part.package.iter_parts():
        key = str(part.partname).lstrip("/")
        if key not in model["parts"] or not hasattr(part, "_element"):
            continue
        roles = deepcopy(model["parts"][key])
        for row in roles:
            if row["id"] in protected.get(key, set()):
                row["action"] = "keep"
                row.pop("replacementText", None)
        apply_background_roles(part._element, roles)
    return model


def _solid_svg_panel(archive, name):
    """Accept only a blank single-color SVG rectangle, never source imagery."""
    if not name.lower().endswith(".svg"):
        return False
    try:
        root = SafeET.fromstring(archive.read(name))
    except (KeyError, ValueError, SafeET.ParseError, DefusedXmlException):
        return False
    if root.tag != "{http://www.w3.org/2000/svg}svg" or len(root) != 1:
        return False
    if set(root.attrib) - {"width", "height", "viewBox", "fill"}:
        return False
    rect = root[0]
    if rect.tag != "{http://www.w3.org/2000/svg}rect" or len(rect):
        return False
    if set(rect.attrib) - {"width", "height", "x", "y", "rx", "ry", "fill"}:
        return False
    fill = rect.get("fill", "")
    if not re.fullmatch(r"#[0-9a-fA-F]{6}|white|black", fill, re.I):
        return False
    return (
        rect.get("x", "0") == "0"
        and rect.get("y", "0") == "0"
        and (rect.get("width") == root.get("width") and rect.get("height") == root.get("height"))
    )


def protect_field_surfaces(model, prs, patterns, raw=None):
    """Clear sample wording, but retain the authored field's fill and geometry.

    A white text card is both a content placeholder and a reusable surface.
    Dropping the whole object would put its black text directly on a dark brand
    background. Keep that native surface; never keep sample pictures/charts.
    """
    protected = {}
    graphic_ids = {}
    panel_assets = {}
    archive = ZipFile(BytesIO(raw)) if raw is not None else None
    for pattern in patterns:
        surface = (
            prs.slides[pattern.source_slide - 1]
            if pattern.source_slide
            else prs.slide_masters[pattern.master_index].slide_layouts[pattern.layout_index]
        )
        if pattern.source_slide and pattern.graphic_order_verified and len(pattern.body_zones) >= 2:
            key = str(surface.part.partname).lstrip("/")
            if pattern.graphic_kind != "none":
                graphic_ids.setdefault(key, set()).update(map(str, pattern.graphic_shape_ids))
            # No generic 'keep all vectors' fallback: unverified shapes may encode old data.
        protected.setdefault(str(surface.part.partname).lstrip("/"), set()).update(
            str(f["shape_id"]) for f in pattern.fields if f["role"] != "image"
        )
    for part in prs.part.package.iter_parts():
        key = str(part.partname).lstrip("/")
        if key not in model["parts"] or not hasattr(part, "_element"):
            continue
        ids = protected.setdefault(key, set())
        if not key.startswith("ppt/slides/"):
            for shape in part._element.iter(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}sp"
            ):
                if (
                    shape.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}ph")
                    is not None
                ):
                    props = shape.find(
                        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
                    )
                    if props is not None:
                        ids.add(props.get("id"))
        for row in model["parts"][key]:
            # Only semantically verified compositions reuse vector paths,
            # nodes and connectors. Photos, charts, embedded data stay excluded.
            if row["id"] in graphic_ids.get(key, set()) and row["type"] in ("sp", "cxnSp", "grpSp"):
                row.update(
                    action="clear-text" if row["type"] in ("sp", "grpSp") else "keep",
                    role="decoration",
                    reason="verified_composition_vector_graphics",
                )
                row.pop("replacementText", None)
            if row["id"] in ids and row["type"] == "sp":
                row.update(
                    action="clear-text", role="background", reason="authored_editable_field_surface"
                )
                row.pop("replacementText", None)
            if (
                archive is not None
                and row["type"] == "pic"
                and row["reason"] == "container_of_sample_text"
            ):
                asset = row.get("asset") or ""
                if asset not in panel_assets:
                    panel_assets[asset] = _solid_svg_panel(archive, asset)
                if panel_assets[asset]:
                    row.update(action="keep", role="background", reason="reusable_solid_svg_panel")
    if archive is not None:
        archive.close()
    by_id = {(part, row["id"]): row for part, rows in model["parts"].items() for row in rows}
    for slide in model["slides"]:
        slide["objects"] = [
            {**row, **by_id.get((row["part"], row["id"]), {})} for row in slide["objects"]
        ]
        slide["protectedRegions"] = protected_background_regions(
            slide["objects"], prs.slide_width, prs.slide_height
        )
    return model
