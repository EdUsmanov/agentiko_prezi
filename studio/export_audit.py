"""Read/repair actual native PPTX objects, not planned source-ID metadata."""
from copy import deepcopy
import re
from pptx.util import Pt
from .fonts import role_font, resolve_font, wrap_text, font_runs, table_cell_fits, text_width
from .template import walk_shapes


def text_key(text):
    from .text_layout import WORD_JOINERS
    text=''.join(c for c in text if c not in WORD_JOINERS)
    return re.sub(r'\s+', ' ', text.replace('**','').replace('`','')).strip().casefold()


def contains_text(actual, required):
    # Prevent "рост 1" from being accepted as "рост 15".
    # A fact promoted to a headline may omit its final sentence punctuation.
    value=text_key(required).rstrip('.!?;:')
    tail = r'(?!\w|[.,]\d)' if value and value[-1].isdigit() else r'(?!\w)'
    return not value or re.search(r'(?<!\w)'+re.escape(value)+tail,text_key(actual)) is not None


def native_chart_text(chart):
    """Read evidence from the exported chart, never from its planned table."""
    values=[]
    for plot in chart.plots:
        try:
            categories=[str(category.label) for category in plot.categories]
        except (AttributeError, ValueError):
            continue  # XY charts have no categorical rows; do not invent labels.
        for series in plot.series:
            values.append(str(series.name))
            values.extend(f'{label} — {value}' for label,value in zip(categories,series.values))
    return '\n'.join(values)


def slide_text(slide):
    values=[]
    for shape,_ in walk_shapes(slide.shapes):
        if shape.has_text_frame:
            values.append(shape.text)
        if shape.has_table:
            values.extend(cell.text for row in shape.table.rows for cell in row.cells)
        if shape.has_chart:
            values.append(native_chart_text(shape.chart))
    return '\n'.join(values)


def content_geometry_signature(prs, titles=()):
    """Measure actual content viewports; names, colors and template IDs do not count."""
    import hashlib
    import json
    title_keys = {text_key(t) for t in titles}
    pages = []
    for slide in prs.slides:
        objects = []
        for shape, box in walk_shapes(slide.shapes):
            kind = 'table' if shape.has_table else 'chart' if shape.has_chart else (
                'image' if shape.name.startswith('forma_image_') else 'text')
            if kind == 'text':
                if not shape.has_text_frame or not shape.text.strip() or text_key(shape.text) in title_keys:
                    continue
                frame = shape.text_frame
                sx = box.w / (shape.width / 12700) if shape.width else 1
                sy = box.h / (shape.height / 12700) if shape.height else 1
                box = box.model_copy(update={
                    'x':box.x+frame.margin_left/12700*sx,
                    'y':box.y+frame.margin_top/12700*sy,
                    'w':box.w-(frame.margin_left+frame.margin_right)/12700*sx,
                    'h':box.h-(frame.margin_top+frame.margin_bottom)/12700*sy})
            objects.append((kind, *(round(getattr(box,k),1) for k in ('x','y','w','h'))))
        pages.append(sorted(objects))
    return hashlib.sha256(json.dumps(pages).encode()).hexdigest()


def _face(profile, name, role='body', bold=False, italic=False):
    # Exported OOXML uses a family plus B/I flags, unlike requested face names.
    from .font_identity import ooxml_face
    for asset in profile.font_assets:
        if ooxml_face(asset['path']) == (name, bool(bold), bool(italic)):
            return asset['path']
    exact=next((a['path'] for a in profile.font_assets if a['requested']==name),None)
    return exact or resolve_font(name) or role_font(profile,role)[1]


def repair_symbols(prs,profile):
    changed=0
    for slide in prs.slides:
        for shape,_ in walk_shapes(slide.shapes):
            frames=[shape.text_frame] if shape.has_text_frame else [c.text_frame for row in shape.table.rows for c in row.cells] if shape.has_table else []
            for frame in frames:
                for paragraph in frame.paragraphs:
                    for run in list(paragraph.runs):
                        path=_face(profile,run.font.name or paragraph.font.name or profile.font)
                        parts=list(font_runs(run.text,path))
                        if not any(face!=path for _,face in parts):
                            continue
                        anchor=run._r
                        for text,face in parts:
                            added=paragraph.add_run();added._r.getparent().remove(added._r)
                            anchor.addnext(added._r);anchor=added._r
                            if run._r.rPr is not None:
                                added._r.insert(0,deepcopy(run._r.rPr))
                            added.text=text
                            if face!=path:
                                added.font.name='Montserrat'
                        run._r.getparent().remove(run._r);changed+=1
    return changed


def explicit_line_widths(paragraph, profile, default_size):
    """Measure saved OOXML lines and each run's actual face; never reflow no-wrap text."""
    ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    widths = [0.0]
    runs = {run._r: run for run in paragraph.runs}
    for node in paragraph._p:
        if node.tag == ns+'br':
            widths.append(0.0)
        elif node.tag in (ns+'r', ns+'fld'):
            run = runs.get(node)
            if run is not None:
                font = run.font
                value = run.text
                name = font.name or paragraph.font.name or profile.font
                bold = font.bold if font.bold is not None else paragraph.font.bold
                italic = font.italic if font.italic is not None else paragraph.font.italic
                size = font.size.pt if font.size is not None else (paragraph.font.size.pt if paragraph.font.size is not None else default_size)
            else:
                # Field runs are not exposed by Paragraph.runs; read their direct style.
                text_node, props = node.find(ns+'t'), node.find(ns+'rPr')
                value = text_node.text or '' if text_node is not None else ''
                latin = props.find(ns+'latin') if props is not None else None
                name = latin.get('typeface') if latin is not None else paragraph.font.name or profile.font
                bold = props.get('b') in ('1','true') if props is not None and props.get('b') is not None else paragraph.font.bold
                italic = props.get('i') in ('1','true') if props is not None and props.get('i') is not None else paragraph.font.italic
                size = int(props.get('sz'))/100 if props is not None and props.get('sz') else default_size
            path = _face(profile, name, bold=bold, italic=italic)
            widths[-1] += text_width(value, path, size)
    return widths


def geometry(prs, profile, repair=False):
    """Conservative text-fit bounds. Does not claim to replace rendered VLM QA.

    Only reduces font sizes inside the same object; never moves text over art.
    Rotated text requires rendered inspection, not guessed axis-aligned repairs.
    """
    findings=[];repairs=[]
    for si,slide in enumerate(prs.slides,1):
        occupied=[]
        for shape,box in walk_shapes(slide.shapes):
            if shape.has_table:
                occupied.append((shape.shape_id,box))
                for ri,row in enumerate(shape.table.rows):
                    for ci,cell in enumerate(row.cells):
                        for p in cell.text_frame.paragraphs:
                            if not p.text.strip():
                                continue
                            size=max([r.font.size.pt for r in p.runs if r.font.size] or [p.font.size.pt if p.font.size else profile.body_size])
                            name=next((r.font.name for r in p.runs if r.font.name),p.font.name or profile.font)
                            width=(shape.table.columns[ci].width-cell.margin_left-cell.margin_right)/12700
                            height=(row.height-cell.margin_top-cell.margin_bottom)/12700
                            if not table_cell_fits(p.text,_face(profile,name,'table'),size,max(1,width),max(1,height),ri==0):
                                findings.append({'code':'pptx_table_overflow','severity':'error','slide':si,'element':shape.shape_id,
                                    'message':'Ячейка таблицы не помещается по метрикам шрифта.'})
            elif shape.has_chart or shape.name.startswith('forma_image_'):
                occupied.append((shape.shape_id,box))
            if not shape.has_text_frame or not shape.text.strip():
                continue
            def add(code,message):
                findings.append({'code':code,'message':message,'slide':si,'element':shape.shape_id,'severity':'error'})
            if box.x<-.5 or box.y<-.5 or box.x+box.w>profile.width+.5 or box.y+box.h>profile.height+.5:
                add('pptx_out_of_bounds','Объект в готовом PPTX выходит за границы слайда.')
            if shape.rotation:
                continue
            frame=shape.text_frame
            sx=box.w/(shape.width/12700) if shape.width else 1
            sy=box.h/(shape.height/12700) if shape.height else 1
            width=box.w-(frame.margin_left+frame.margin_right)/12700*sx
            height=box.h-(frame.margin_top+frame.margin_bottom)/12700*sy
            from .models import Box
            occupied.append((shape.shape_id,Box(x=box.x+frame.margin_left/12700*sx,
                y=box.y+frame.margin_top/12700*sy,w=max(0,width),h=max(0,height))))
            paragraphs=[]
            no_wrap = frame.word_wrap is False
            for p in frame.paragraphs:
                sizes=[r.font.size.pt for r in p.runs if r.font.size]
                size=max(sizes or [p.font.size.pt if p.font.size else profile.body_size])
                name=next((r.font.name for r in p.runs if r.font.name),p.font.name or profile.font)
                run=next((r for r in p.runs if r.font.name),None)
                bold=run.font.bold if run is not None and run.font.bold is not None else p.font.bold
                italic=run.font.italic if run is not None and run.font.italic is not None else p.font.italic
                paragraphs.append((p,size,_face(profile,name,bold=bold,italic=italic)))
            explicit_widths = {id(p): explicit_line_widths(p, profile, size)
                               for p,size,_ in paragraphs} if no_wrap else {}
            def needed(scale):
                total=0
                widest=0
                for p,size,path in paragraphs:
                    points=size*scale
                    spacing=p.line_spacing
                    line_height=spacing.pt*scale*sy if hasattr(spacing,'pt') else points*(spacing if isinstance(spacing,float) else 1.2)*sy
                    # Respect explicit paragraph spacing; do not silently erase it.
                    if no_wrap:
                        lines=explicit_widths[id(p)]
                        props=p._p.pPr
                        margin=int(props.get('marL','0'))/12700 if props is not None else 0
                        indent=int(props.get('indent','0'))/12700 if props is not None else 0
                        # Text indents stay fixed when repair reduces the font size.
                        widest=max(widest,max((value*scale+margin+(max(0,indent) if i==0 else 0))*sx
                                              for i,value in enumerate(lines)))
                        total+=len(lines)*line_height
                    else:
                        # Soft breaks are fixed line boundaries even when wrapping is enabled.
                        total+=len(wrap_text(p.text.replace('\v','\n'),path,points*sx,max(1,width*.96)))*line_height
                    total+=sum(v.pt*sy for v in (p.space_before,p.space_after) if v)
                return total,widest
            def fits(scale):
                required_height,required_width=needed(scale)
                return required_height<=height+1 and required_width<=width+1
            if fits(1):
                continue
            minimum=min(size for _,size,_ in paragraphs)
            floor=min(1,12/max(minimum,.1))
            scale=1
            if repair and floor<1:
                for step in range(1,21):
                    candidate=max(floor,1-step*.035)
                    if fits(candidate):
                        scale=candidate;break
                if scale<1:
                    for p,size,_ in paragraphs:
                        if hasattr(p.line_spacing,'pt'):
                            p.line_spacing=Pt(p.line_spacing.pt*scale)
                        p.font.size=Pt(size*scale)
                        for run in p.runs:
                            run.font.size=Pt((run.font.size.pt if run.font.size else size)*scale)
                    repairs.append({'code':'pptx_text_fit','slide':si,'element':shape.shape_id,
                                    'message':'Уменьшено начертание внутри исходного объекта без изменения текста и геометрии.','scale':round(scale,3)})
            if scale==1:
                add('pptx_text_overflow','Текст в готовом PPTX не помещается по метрической проверке; требуется другой макет или больше слайдов.')
        from .audit import overlaps
        for i,(object_id,a) in enumerate(occupied):
            for other_id,b in occupied[i+1:]:
                if object_id!=other_id and overlaps(a,b):
                    findings.append({'code':'pptx_content_overlap','severity':'error','slide':si,'element':object_id,
                        'message':f'Области содержимого объектов {object_id} и {other_id} пересекаются в готовом PPTX.'})
    return findings,repairs


def evidence(prs, variant, package):
    facts={f.id:f for f in package.content.facts}; findings=[]
    for index,(slide,plan) in enumerate(zip(prs.slides,variant.slides),1):
        actual=slide_text(slide)
        table_ids={t.id for t in package.content.tables}
        for fid in plan.fact_ids:
            fact=facts[fid]
            if fact.source in table_ids:
                continue  # Table/chart cells are checked by their own contract.
            if not contains_text(actual,fact.text):
                findings.append({'code':'pptx_fact_coverage','severity':'error','slide':index,
                    'message':f'В видимом тексте PPTX не найден полный исходный факт {fid}.'})
    return findings


def inspect_content(prs, variant, package):
    from .security import digest
    normalized=text_key
    from .template import walk_shapes
    from .pictures import embedded_picture_blob
    actual, findings, images = [], [], set()

    def add(code, message, slide=0, severity='error'):
        findings.append(dict(code=code, message=message, slide=slide, severity=severity))

    for index, (plan, slide) in enumerate(zip(variant.slides, prs.slides), 1):
        texts, tables = [], []
        for shape, _ in walk_shapes(slide.shapes):
            if shape.has_text_frame:
                texts.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if shape.has_table:
                rows = [[c.text for c in row.cells] for row in shape.table.rows]
                tables.append(rows)
                texts.extend(cell for row in rows for cell in row)
            if shape.has_chart:
                texts.append(native_chart_text(shape.chart))
            blob = embedded_picture_blob(shape)
            if blob:
                images.add(digest(blob))
        text = '\n'.join(texts)
        actual.append({'actual_text': text, 'native_tables': tables})
        title = normalized(plan.title)
        matches = sum(normalized(t) == title for t in texts)
        if not matches:
            add('title', 'Точный заголовок сценария не найден в тексте PPTX.', index, 'warning')
        if matches > 1:
            add('duplicate_title', 'Заголовок повторяется в тексте слайда.', index)
        if not text.strip():
            add('empty_slide', 'На слайде нет редактируемого текста.', index)
        if plan.table_id and plan.layout != 'chart':
            table = next(t for t in package.content.tables if t.id == plan.table_id)
            expected = [[normalized(c) for c in row] for row in [table.headers] + table.rows]
            found = [[[normalized(c) for c in row] for row in t] for t in tables]
            if table.visualization=='metrics':
                if any(not contains_text(text,cell) for row in [table.headers]+table.rows for cell in row):
                    add('metric_coverage','Не все исходные подписи и значения карточек найдены в редактируемом тексте PPTX.',index)
            elif expected not in found:
                add('table_coverage', 'Исходная таблица не воспроизведена полностью как редактируемая таблица; проверьте данные.', index)
        if plan.chart_type not in ('auto',) or plan.layout == 'chart':
            charts=[shape.chart for shape,_ in walk_shapes(slide.shapes) if shape.has_chart]
            if not charts:
                add('chart_coverage', 'В сценарии предусмотрена диаграмма, но нативная диаграмма в PPTX отсутствует.', index)
            elif plan.table_id:
                from .charts import table_series, chart_projection
                table=next(t for t in package.content.tables if t.id==plan.table_id)
                projected,supplement=chart_projection(table)
                expected=table_series(projected)
                if not any([list(s.values) for s in chart.series]==expected and
                           [' '.join(str(c.label).split()) for c in chart.plots[0].categories]==[' '.join(row[0].split()) for row in projected.rows] for chart in charts):
                    add('chart_values','Числа или категории нативной диаграммы не совпадают с исходной таблицей.',index)
                _,compact=chart_projection(table,compact_captions=True)
                # Both layouts must retain the entire category/value mapping.
                if not any(all(contains_text(text,caption) for caption in captions)
                           for captions in (supplement,compact)):
                    add('chart_supplement','Итоги или значения с другой единицей измерения потеряны при построении диаграммы.',index)
    for asset in package.images:
        if asset.sha256 not in images:
            add('image_coverage', f'Загруженное изображение «{asset.name}» не найдено в готовом PPTX.')
    return actual, findings



def audit_export(path, variant, package):
    from .powerpoint import open_presentation
    prs=open_presentation(path)
    _,content_findings=inspect_content(prs,variant,package)
    object_findings,_=geometry(prs,package.template)
    return content_findings+object_findings+evidence(prs,variant,package)+readability(prs,variant,package)


def readability(prs, variant, package):
    """Read actual exported type sizes after native object repair."""
    findings=[]
    facts={f.id:f.text for f in package.content.facts}
    for index,(plan,slide) in enumerate(zip(variant.slides,prs.slides),1):
        for shape,_ in walk_shapes(slide.shapes):
            frames=[];floor=16
            if shape.has_table:
                frames=[c.text_frame for row in shape.table.rows for c in row.cells]
            elif shape.has_text_frame:
                if text_key(shape.text)==text_key(plan.title):
                    frames=[shape.text_frame];floor=18
                elif any(contains_text(shape.text,facts[fid]) for fid in plan.fact_ids):
                    frames=[shape.text_frame]
            sizes=[r.font.size.pt if r.font.size else p.font.size.pt
                   for frame in frames for p in frame.paragraphs for r in p.runs
                   if r.text.strip() and (r.font.size or p.font.size)]
            if shape.has_chart:
                chart=shape.chart
                if chart.font.size: sizes.append(chart.font.size.pt)
                for name in ('category_axis','value_axis'):
                    try:
                        size=getattr(chart,name).tick_labels.font.size
                        if size: sizes.append(size.pt)
                    except (ValueError,AttributeError):pass
            if sizes and min(sizes)<floor-.1:
                findings.append({'code':'pptx_readability','severity':'warning','slide':index,'element':shape.shape_id,
                    'message':f'В готовом PPTX текст {min(sizes):g} pt ниже порога {floor} pt.'})
    return findings


def content_scenes(prs, variant, package):
    """Source-linked geometry read from PPTX for final diversity verification."""
    from .models import Element, SlideScene
    facts={f.id:f for f in package.content.facts}
    scenes=[]
    for plan, slide in zip(variant.slides, prs.slides):
        elements=[]
        for shape, box in walk_shapes(slide.shapes):
            ids=[];kind=None
            if shape.has_table or shape.has_chart:
                ids=[fid for fid in plan.fact_ids if facts[fid].source==plan.table_id]
                kind='table' if shape.has_table else 'chart'
            elif shape.has_text_frame and text_key(shape.text)!=text_key(plan.title):
                ids=[fid for fid in plan.fact_ids if contains_text(shape.text,facts[fid].text)]
                kind='text'
            if kind and ids:
                elements.append(Element(kind=kind,box=box,source_ids=ids))
        scenes.append(SlideScene(title=plan.title,background=package.template.background,
            elements=elements,source_ids=plan.fact_ids,layout=plan.layout,purpose=plan.purpose))
    return scenes
