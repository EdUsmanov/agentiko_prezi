"""Recognize visible empty object panels without erasing authored artwork."""

from functools import lru_cache
from pathlib import Path
from zipfile import BadZipFile

from studio.models import Box
from studio.templates.template_geometry import contrast, walk_shapes


@lru_cache(maxsize=8)
def _source_panels(path, modified, size):
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError

    try:
        prs = Presentation(path)
    except (OSError, BadZipFile, PackageNotFoundError):
        # This optional visual heuristic must not break quality checks when a
        # cached background source is unavailable. Export validates its source.
        return {}

    def rectangles(surface):
        result = []
        for shape, box in walk_shapes(surface.shapes):
            geometry = shape._element.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}spPr/"
                "{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom"
            )
            if geometry is None or geometry.get("prst") not in ("rect", "roundRect"):
                continue
            if getattr(shape, "fill", None) is None or shape.fill.type != 1:
                continue
            result.append((box.x, box.y, box.w, box.h))
        return tuple(result)

    layouts = {}
    regions = {}
    for mi, master in enumerate(prs.slide_masters):
        for li, layout in enumerate(master.slide_layouts):
            layouts[str(layout.part.partname)] = (mi, li)
            regions[(mi, li, 0)] = rectangles(master) + rectangles(layout)
    for index, slide in enumerate(prs.slides, 1):
        mi, li = layouts[str(slide.slide_layout.part.partname)]
        regions[(mi, li, index)] = regions[(mi, li, 0)] + rectangles(slide)
    return regions


def empty_content_panels(scene, profile, pattern):
    """Large, visibly blank panels are omissions even without a text placeholder.

    Geometry comes from the sanitized PPTX and uniform paint is confirmed in
    its rendered background. Photos, gradients, full-slide backgrounds, small
    ornaments and panels containing actual output are not empty object slots.
    """
    source = getattr(profile, "background_source", "")
    image = getattr(pattern, "background_image", "")
    if not source or not image or not Path(source).is_file() or not Path(image).is_file():
        return []
    from studio.composition.background_selection import flat_region

    stat = Path(source).stat()
    regions = _source_panels(source, stat.st_mtime_ns, stat.st_size)
    width, height = profile.width, profile.height
    content = [
        e.box
        for e in scene.elements
        if (e.source_ids or e.image_id or e.role == "title")
        and e.role not in ("footer", "brand", "template_background")
    ]
    empty = []
    for x, y, w, h in set(
        regions.get((pattern.master_index, pattern.layout_index, pattern.source_slide), ())
    ):
        if not (
            0.08 <= w * h / (width * height) <= 0.8 and w >= width * 0.2 and h >= height * 0.18
        ):
            continue
        zone = Box(x=x, y=y, w=w, h=h)
        if any(
            max(0, min(x + w, b.x + b.w) - max(x, b.x))
            * max(0, min(y + h, b.y + b.h) - max(y, b.y))
            >= 0.5 * min(w * h, b.w * b.h)
            for b in content
            if b.w > 0 and b.h > 0
        ):
            continue
        # Ignore outlines/rounded corners; inspect the visible interior. A
        # source rectangle hidden by a photo must not penalize that photo.
        inset = min(w, h) * 0.06
        interior = Box(x=x + inset, y=y + inset, w=w - inset * 2, h=h - inset * 2)
        color = flat_region(image, interior, width, height)
        if color and contrast(color, scene.background) >= 1.15:
            empty.append(zone)
    return empty
