"""Measured stacked columns with editable legends and collision-free value labels."""

import math
import re

from studio.templates.fonts import element_font, text_width, wrap_text


def value_suffix(e):
    suffixes = {re.sub(r"^[\s−\-+\d.,]+", "", v).strip() for v in e.value_labels}
    suffix = next(iter(suffixes)) if len(suffixes) == 1 else ""
    return suffix if suffix == "%" else " " + suffix if suffix else ""


def legend_layout(names, font, size, width):
    """Pack measured legend entries in rows without dropping totals or labels."""
    best = None
    for columns in range(1, min(4, len(names)) + 1):
        cell_width = width / columns
        lines = [wrap_text(name, font, size, max(1, cell_width - 20)) for name in names]
        if columns > 1 and any(
            text_width(word, font, size) > cell_width - 20
            for name in names
            for word in name.split()
        ):
            continue
        cells, y = [], 0
        for start in range(0, len(names), columns):
            row = lines[start : start + columns]
            height = max(max(1, len(entry)) * size * 1.25 for entry in row)
            cells.extend(
                dict(x=col * cell_width, y=y, w=cell_width, h=height) for col in range(len(row))
            )
            y += height + 3
        if best is None or y < best[2]:
            best = lines, cells, y
    return best


def stacked_layout(e, profile):
    font = element_font(profile, e)[1]
    series = e.series_values or [e.values]
    names = e.series_names or [e.unit]
    if not e.labels or len(series) != len(names) or any(len(s) != len(e.labels) for s in series):
        return {"fits": False}
    totals = [row for row in e.rows[1:] if re.match(r"^(всего|итого|total)\b", row[0], re.I)]
    legend = [
        name + "".join(f" — {row[0]}: {row[i + 1]}" for row in totals if i + 1 < len(row))
        for i, name in enumerate(names)
    ]
    legend_size = min(14, e.size)
    if totals or len(names) > 2:
        legend_size = min(12, legend_size)
    legend_lines, legend_cells, legend_h = legend_layout(legend, font, legend_size, e.box.w)
    positive = [sum(max(0, s[i]) for s in series) for i in range(len(e.labels))]
    negative = [sum(min(0, s[i]) for s in series) for i in range(len(e.labels))]

    def limit(value):
        if not value:
            return 0
        unit = 10 ** math.floor(math.log10(abs(value)))
        return math.ceil(abs(value) / unit) * unit

    maximum = limit(max(positive, default=0)) or 1
    minimum = -limit(min(negative, default=0))
    left = max(text_width(f"{v:g}", font, 16) for v in (minimum, maximum)) + 16
    width = e.box.w - left - 8
    band = width / max(1, len(e.labels))
    heading = wrap_text(e.category_title, font, 14, max(1, e.box.w)) if e.category_title else []
    heading_h = len(heading) * 17.5 + (6 if heading else 0)
    top = 10
    label_size = min(16, e.size)
    for size in (label_size, min(14, label_size), min(12, label_size)):
        label_size = size
        from studio.composition.chart_layout import wrap_category

        labels = [
            wrap_category(label, font, size, max(1, band - 2)).splitlines() for label in e.labels
        ]
        label_h = max((len(lines) for lines in labels), default=1) * label_size * 1.25 + 6
        height = e.box.h - top - label_h - heading_h - legend_h - 8
        if height >= max(110, len(series) * 20) and all(
            text_width(word, font, size) <= band - 2 for label in e.labels for word in label.split()
        ):
            label_size = size
            break
    values_fit = True
    if height >= max(110, len(series) * 20):
        for col in range(len(e.labels)):
            values = [s[col] for s in series]
            ideal, placed = value_label_positions(values, minimum, maximum, height)
            for value, a, b in zip(values, ideal, placed):
                available = band * (0.24 if abs(a - b) > 1 else 0.48)
                values_fit = (
                    values_fit
                    and text_width(f"{value:g}" + value_suffix(e), font, 16) <= available * 0.96
                )
    fits = (
        values_fit
        and bool(e.labels)
        and len(e.labels) <= 30
        and len(series) == len(names)
        and all(len(s) == len(e.labels) for s in series)
        and e.box.w >= 260
        and height >= max(110, len(series) * 20)
        and band >= 32
        and all(
            text_width(word, font, label_size) <= band - 2
            for label in e.labels
            for word in label.split()
        )
        and all(
            text_width(word, font, legend_size) <= e.box.w - 20
            for name in legend
            for word in name.split()
        )
    )
    return dict(
        left=left,
        top=top,
        width=width,
        height=height,
        labels=labels,
        label_size=label_size,
        label_h=label_h,
        heading=heading,
        heading_h=heading_h,
        legend=legend_lines,
        legend_cells=legend_cells,
        legend_size=legend_size,
        legend_h=legend_h,
        minimum=minimum,
        maximum=maximum,
        fits=fits,
    )


def value_label_positions(values, minimum, maximum, height, size=16):
    """Keep every label in its category band, moving only overlapping labels."""
    positive = negative = 0
    ideal = []
    for value in values:
        start = positive if value >= 0 else negative
        center = start + value / 2
        if value >= 0:
            positive += value
        else:
            negative += value
        ideal.append((maximum - center) / (maximum - minimum) * height)
    order = sorted(range(len(values)), key=ideal.__getitem__)
    positions = ideal[:]
    half = size * 0.625
    gap = size * 1.25
    last = -half
    for i in order:
        positions[i] = max(half, ideal[i], last + gap)
        last = positions[i]
    if order and last > height - half:
        positions[order[-1]] = height - half
        for a, b in zip(reversed(order[:-1]), reversed(order[1:])):
            positions[a] = min(positions[a], positions[b] - gap)
    return ideal, positions


def render_stacked_chart(slide, e, profile):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Pt
    from studio.templates.font_identity import apply_ooxml_font
    from studio.models import Box, Element
    from studio.composition.pptx_text import rgb, set_text
    from studio.templates.template_geometry import contrast

    layout = stacked_layout(e, profile)
    if not layout["fits"]:
        raise ValueError(
            "Накопительная диаграмма: недостаточно места для столбцов, подписей и легенды"
        )
    b = e.box
    data = CategoryChartData()
    data.categories = e.labels
    series = e.series_values or [e.values]
    for name, values in zip(e.series_names or [e.unit], series):
        data.add_series(name, values)
    group = slide.shapes.add_group_shape()
    group.name = "forma_chart_annotations"
    shapes = group.shapes
    chart = shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED, Pt(b.x), Pt(b.y), Pt(b.w), Pt(b.h), data
    ).chart
    chart._element.set("{urn:vk-forma:chart}category-heading", e.category_title)
    chart._element.set("{urn:vk-forma:chart}value-suffix", value_suffix(e))
    chart.has_title = False
    chart.has_legend = False
    font = element_font(profile, e)[1]
    apply_ooxml_font(chart.font, font)
    chart.font.size = Pt(16)
    chart.font.color.rgb = rgb(e.color)
    plot = chart.plots[0]
    plot.has_data_labels = False
    plot.gap_width = 100
    plot.overlap = 100
    axis = chart.value_axis
    axis.minimum_scale = layout["minimum"]
    axis.maximum_scale = layout["maximum"]
    axis.has_title = False
    axis.has_major_gridlines = True
    apply_ooxml_font(axis.tick_labels.font, font)
    axis.tick_labels.font.size = Pt(16)
    axis.tick_labels.font.color.rgb = rgb(e.color)
    chart.category_axis.has_title = False
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    chart.category_axis.reverse_order = False
    area = chart._chartSpace.chart.plotArea
    node = area.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}layout")
    if node is None:
        node = OxmlElement("c:layout")
        area.insert(0, node)
    manual = OxmlElement("c:manualLayout")
    for key, val in [
        ("layoutTarget", "inner"),
        ("xMode", "edge"),
        ("yMode", "edge"),
        ("wMode", "factor"),
        ("hMode", "factor"),
        ("x", layout["left"] / b.w),
        ("y", layout["top"] / b.h),
        ("w", layout["width"] / b.w),
        ("h", layout["height"] / b.h),
    ]:
        item = OxmlElement("c:" + key)
        item.set("val", str(val))
        manual.append(item)
    node.append(manual)
    palette = [
        c
        for c in dict.fromkeys([profile.accent] + profile.colors)
        if contrast(c, e.background_hint or profile.background) > 1.6
    ] or [profile.accent]
    for i, item in enumerate(chart.series):
        item.format.fill.solid()
        item.format.fill.fore_color.rgb = rgb(palette[i % len(palette)])
        item.format.line.color.rgb = rgb(palette[i % len(palette)])

    def text(value, x, y, w, h, size, color=e.color, align=PP_ALIGN.LEFT):
        shape = shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        element = Element(
            kind="text",
            text=value,
            box=Box(x=x, y=y, w=w, h=h),
            font=e.font,
            size=size,
            color=color,
            field_style=e.field_style,
        )
        set_text(shape.text_frame, value, element, profile)
        shape.text_frame.margin_left = shape.text_frame.margin_right = 0
        shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = align
        return shape

    band = layout["width"] / len(e.labels)
    y = b.y + layout["top"] + layout["height"] + 3
    for i, lines in enumerate(layout["labels"]):
        text(
            "\n".join(lines),
            b.x + layout["left"] + i * band,
            y,
            band,
            layout["label_h"],
            layout["label_size"],
            align=PP_ALIGN.CENTER,
        )
    y += layout["label_h"]
    if layout["heading"]:
        text(
            "\n".join(layout["heading"]),
            b.x,
            y,
            b.w,
            layout["heading_h"],
            14,
            align=PP_ALIGN.CENTER,
        )
        y += layout["heading_h"]
    for i, lines in enumerate(layout["legend"]):
        cell = layout["legend_cells"][i]
        x, ly = b.x + cell["x"], y + cell["y"]
        swatch = shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x + 2), Pt(ly + 3), Pt(9), Pt(9))
        swatch.fill.solid()
        swatch.fill.fore_color.rgb = rgb(palette[i % len(palette)])
        swatch.line.fill.background()
        text("\n".join(lines), x + 18, ly, cell["w"] - 18, cell["h"], layout["legend_size"])
    for col in range(len(e.labels)):
        values = [s[col] for s in series]
        ideal, positions = value_label_positions(
            values, layout["minimum"], layout["maximum"], layout["height"]
        )
        center = b.x + layout["left"] + (col + 0.5) * band
        for i, (value, original, placed) in enumerate(zip(values, ideal, positions)):
            displaced = abs(original - placed) > 1
            # Move cramped segment labels beside their stack with a leader.
            x = center + band * 0.26 if displaced else center - band * 0.25
            w = band * (0.24 if displaced else 0.48)
            color = (
                e.color
                if displaced
                else max(
                    [e.color, "#000000", "#FFFFFF"],
                    key=lambda c: contrast(c, palette[i % len(palette)]),
                )
            )
            if displaced:
                line = shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT,
                    Pt(center + band * 0.23),
                    Pt(b.y + layout["top"] + original),
                    Pt(x),
                    Pt(b.y + layout["top"] + placed),
                )
                line.line.color.rgb = rgb(e.color)
                line.line.width = Pt(0.5)
            text(
                f"{value:g}" + value_suffix(e),
                x,
                b.y + layout["top"] + placed - 10,
                w,
                20,
                16,
                color,
                PP_ALIGN.CENTER,
            )
    return chart
