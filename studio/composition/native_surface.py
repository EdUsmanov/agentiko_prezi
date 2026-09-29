"""Copy only permitted source artwork into editable slides."""

from copy import deepcopy
from hashlib import sha256
from pptx.enum.shapes import MSO_SHAPE_TYPE
from studio.composition.pictures import is_picture, embedded_picture_blob, embedded_blip_blob
from studio.composition.shape_geometry import box, intersects
from pptx.util import Pt

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def scrub_surface(surface):
    # Masters/layouts keep their artwork and geometry; sample placeholder copy is not content.
    for node in surface._element.iter(A + "t"):
        node.text = ""
    for node in list(surface._element.iter()):
        if node.tag in (A + "hlinkClick", A + "hlinkMouseOver"):
            node.getparent().remove(node)


def copy_node(node, source_part, target_part):
    copied = deepcopy(node)
    for child in list(copied.iter()):
        if child.tag in (A + "hlinkClick", A + "hlinkMouseOver"):
            child.getparent().remove(child)
            continue
        for attr, value in list(child.attrib.items()):
            if not attr.startswith(R):
                continue
            rel = source_part.rels.get(value)
            if rel and not rel.is_external and rel.reltype.endswith("/image"):
                child.set(attr, target_part.relate_to(rel.target_part, rel.reltype))
            else:
                del child.attrib[attr]
    return copied


def source_slide(prs, pattern):
    layout = prs.slide_masters[pattern.master_index].slide_layouts[pattern.layout_index]
    slide = prs.slides.add_slide(layout)
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)
    if pattern.source_slide:
        original = prs._studio_sources[pattern.source_slide - 1]
        bg = original._element.find(P + "cSld/" + P + "bg")
        if bg is not None:
            slide._element.cSld.insert(0, copy_node(bg, original.part, slide.part))
        if "showMasterSp" in original._element.attrib:
            slide._element.set("showMasterSp", original._element.get("showMasterSp"))
        for sh in original.shapes:
            if getattr(prs, "_studio_background_clean", False):
                copied = copy_node(sh._element, original.part, slide.part)
                slide.shapes._spTree.insert_element_before(copied, "p:extLst")
                continue
            # Ordinary slide pictures/charts may be prior content. Do not silently reuse them.
            if sh.is_placeholder or sh.has_chart or sh.has_table:
                continue
            if is_picture(sh):
                raw = embedded_picture_blob(sh)
                if raw is None or sha256(raw).hexdigest()[:20] not in prs._studio_brand_hashes:
                    continue
            # Pictures nested in groups must obey the same rule as top-level
            # pictures; otherwise an old private photo can bypass the guard.
            untrusted_image = False
            for node in sh._element.iter(A + "blip"):
                raw = embedded_blip_blob(node, original.part)
                if raw is None or sha256(raw).hexdigest()[:20] not in prs._studio_brand_hashes:
                    untrusted_image = True
                    break
            if untrusted_image:
                continue
            has_text = any((n.text or "").strip() for n in sh._element.iter(A + "t"))
            if has_text and not (
                sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and any(intersects(box(sh), b) for b in pattern.body_zones)
            ):
                continue
            copied = copy_node(sh._element, original.part, slide.part)
            if has_text:
                for node in copied.iter(A + "t"):
                    node.text = ""
            slide.shapes._spTree.insert_element_before(copied, "p:extLst")
    return slide


def place_template_resource(slide, element, profile, resource_source):
    """Copy allowlisted native shapes from the original source at scene coordinates."""
    resource = next((r for r in profile.resources if r.id == element.resource_id), None)
    if resource is None or not 1 <= resource.source_slide <= len(resource_source.slides):
        raise ValueError("Неизвестный ресурс исходного шаблона")
    original = resource_source.slides[resource.source_slide - 1]
    selected = [s for s in original.shapes if s.shape_id in set(resource.shape_ids)]
    if len(selected) != len(resource.shape_ids) or resource.box.w <= 0 or resource.box.h <= 0:
        raise ValueError("Ресурс исходного шаблона неполон")
    used_ids = [
        int(node.get("id")) for node in slide.shapes._spTree.iter(P + "cNvPr") if node.get("id")
    ]
    next_id = max(used_ids, default=0) + 1
    for shape in selected:
        if any((node.text or "").strip() for node in shape._element.iter(A + "t")):
            raise ValueError("Ресурс содержит старый текст")
        copied = copy_node(shape._element, original.part, slide.part)
        xfrm = copied.find(".//" + A + "xfrm")
        if xfrm is None or xfrm.find(A + "off") is None or xfrm.find(A + "ext") is None:
            raise ValueError("У ресурса нет редактируемой геометрии")
        scale_x = element.box.w / resource.box.w
        scale_y = element.box.h / resource.box.h
        x = element.box.x + (shape.left / Pt(1) - resource.box.x) * scale_x
        y = element.box.y + (shape.top / Pt(1) - resource.box.y) * scale_y
        w, h = shape.width / Pt(1) * scale_x, shape.height / Pt(1) * scale_y
        xfrm.find(A + "off").set("x", str(Pt(x)))
        xfrm.find(A + "off").set("y", str(Pt(y)))
        xfrm.find(A + "ext").set("cx", str(Pt(w)))
        xfrm.find(A + "ext").set("cy", str(Pt(h)))
        for node in copied.iter(P + "cNvPr"):
            node.set("id", str(next_id))
            node.set("name", f"forma_resource_{resource.id}_{next_id}")
            next_id += 1
        slide.shapes._spTree.insert_element_before(copied, "p:extLst")
