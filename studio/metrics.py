"""Editable metric cards grounded in every cell of a source table."""

import math
from .models import Element, Box
from .fonts import role_font, wrap_text


def metric_elements(element, profile):
    headers, *rows = element.rows
    if not 1 <= len(rows) <= 6 or not 2 <= len(headers) <= 3:
        raise ValueError("Карточки показателей требуют 1–6 строк и 2–3 столбца")
    b = element.box
    gap = 18
    columns = 2 if b.w >= 360 and len(rows) > 1 else 1
    caption = " · ".join(headers)
    font_file = role_font(profile, "body")[1]
    cap_h = len(wrap_text(caption, font_file, 12, b.w * 0.94)) * 15 + 8
    width = (b.w - gap * (columns - 1)) / columns
    height = (b.h - cap_h - gap * (math.ceil(len(rows) / columns) - 1)) / math.ceil(
        len(rows) / columns
    )
    result = [
        Element(
            kind="text",
            role="metric_header",
            text=caption,
            box=Box(x=b.x, y=b.y, w=b.w, h=cap_h),
            font=element.font,
            size=12,
            color=element.color,
            source_ids=element.source_ids,
            background_hint=element.background_hint,
        )
    ]
    for index, row in enumerate(rows):
        x = b.x + (index % columns) * (width + gap)
        y = b.y + cap_h + (index // columns) * (height + gap)
        result.append(
            Element(
                kind="line",
                role="metric_rule",
                box=Box(x=x, y=y, w=width, h=1.5),
                fill=element.fill,
            )
        )
        parts = [
            (row[-1], min(30, profile.title_size), True),
            (row[0], min(16, profile.body_size), True),
        ]
        if len(row) > 2:
            parts.append((row[1], 12, False))
        cy = y + 9
        for text, size, bold in parts:
            h = len(wrap_text(text, font_file, size, width * (0.94 if bold else 1))) * size * 1.25
            result.append(
                Element(
                    kind="text",
                    role="metric_value" if text == row[-1] else "metric_label",
                    box=Box(x=x, y=cy, w=width, h=h),
                    text=text,
                    font=element.font,
                    size=size,
                    bold=bold,
                    color=element.color,
                    source_ids=element.source_ids,
                    background_hint=element.background_hint,
                )
            )
            cy += h + 6
        if cy - 6 > y + height + 0.5:
            raise ValueError("Карточки показателей не помещаются в область макета")
    return result


def render_metric_cards(slide, element, profile):
    from pptx.util import Pt
    from pptx.enum.shapes import MSO_SHAPE
    from .pptx_text import set_text, rgb

    for item in metric_elements(element, profile):
        b = item.box
        if item.kind == "text":
            shape = slide.shapes.add_textbox(Pt(b.x), Pt(b.y), Pt(b.w), Pt(b.h))
            set_text(shape.text_frame, item.text, item, profile)
        else:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(b.x), Pt(b.y), Pt(b.w), Pt(b.h))
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(item.fill)
            shape.line.fill.background()
