"""All exports consume the same scene; PPTX remains native and editable."""
from .table_style import column_widths
import base64
from copy import deepcopy
from html import escape
from pathlib import Path
import json
import re
import threading
from zipfile import ZipFile
from pptx import Presentation
from .powerpoint import open_presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.colors import HexColor
from .models import Element, Box
from .fonts import pdf_font, wrap_text, element_font, profile_font_files, font_runs, symbol_font
from .template import contrast
from .embedded_fonts import check_glyphs, P, R
from .native_template import scrub_surface, source_slide
from .font_identity import apply_ooxml_font, ooxml_face

_pdfium_lock=threading.Lock()

def primitives(element,profile):
    if element.kind!="chart":
        return [element]
    e=element; b=e.box
    label_w=b.w*.3
    row_h=(b.h-e.size*2)/len(e.values)
    values_w=b.w-label_w-65
    scale=max(e.values)
    out=[]
    for i,(label,value) in enumerate(zip(e.labels,e.values)):
        y=b.y+i*row_h
        out.append(Element(kind="text",box=Box(x=b.x,y=y,w=label_w-10,h=row_h-6),text=label,font=e.font,size=e.size,color=e.color,role="chart"))
        bar_w=values_w*value/scale
        if value>0:
            out.append(Element(kind="rect",box=Box(x=b.x+label_w,y=y+3,w=bar_w,h=max(8,min(row_h-12,25))),fill=e.fill))
        label=e.value_labels[i] if len(e.value_labels)==len(e.values) else format(value,'.17g')
        out.append(Element(kind="text",box=Box(x=b.x+label_w+bar_w+7,y=y,w=58,h=row_h-6),text=label,font=e.font,size=e.size,color=e.color,role="chart"))
    out.append(Element(kind="text",box=Box(x=b.x,y=b.y+b.h-e.size*1.6,w=b.w,h=e.size*1.6),text=e.unit,font=e.font,size=e.size,color=e.color,role="chart"))
    return out

def chart_fits(element,profile):
    """Validate actual chart label/bar geometry, not just its enclosing box."""
    if element.series_values:
        # Native charts own axis layout. This is only a minimum-size guard;
        # actual label clipping is checked on the rendered PPTX by vision.
        from .charts import chart_caption_layout
        text,height=chart_caption_layout(element.rows,element.box.w,profile)
        available=element.box.h-height-(12 if text else 0)
        return (element.box.w>=260 and available>=160 and
                len(element.labels)<=30 and all(len(v)==len(element.labels) for v in element.series_values))
    for item in primitives(element,profile):
        b=item.box;outer=element.box
        if b.w<=0 or b.h<=0 or b.x<outer.x-.5 or b.y<outer.y-.5 or b.x+b.w>outer.x+outer.w+.5 or b.y+b.h>outer.y+outer.h+.5:
            return False
        if item.kind=='text' and len(wrap_text(item.text,element_font(profile,element)[1],item.size,b.w))*item.size*1.25>b.h+.5:
            return False
    return True

def rgb(value):
    return RGBColor.from_string(value.lstrip("#"))

def set_text(frame,text,e,profile,width=None):
    frame.clear()
    frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
    frame.word_wrap=False
    frame.vertical_anchor=MSO_ANCHOR.TOP
    inset=e.size*1.4 if e.bullet else 0
    family, font_file = element_font(profile, e)
    lines=wrap_text(text,font_file,e.size,((width or e.box.w)-inset)*(.94 if e.bold or e.bold_prefix else 1))
    p=frame.paragraphs[0]
    p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(e.size*1.25)
    if e.bullet:
        props=p._p.get_or_add_pPr();props.set("marL",str(int(Pt(inset))));props.set("indent",str(-int(Pt(e.size))))
        bullet_color=OxmlElement('a:buClr');color=OxmlElement('a:srgbClr');color.set('val',e.color.lstrip('#'));bullet_color.append(color);props.append(bullet_color)
        bullet_size=OxmlElement('a:buSzPct');bullet_size.set('val','100000');props.append(bullet_size)
        bullet_face=next(font_runs('•',font_file))[1]
        bullet_font=OxmlElement("a:buFont");bullet_font.set("typeface",ooxml_face(bullet_face)[0]);props.append(bullet_font)
        bullet=OxmlElement("a:buChar");bullet.set("char","•");props.append(bullet)
    # Bullet glyphs use paragraph defaults, not necessarily the first run.
    apply_ooxml_font(p.font,font_file,e.bold)
    p.font.size=Pt(e.size);p.font.color.rgb=rgb(e.color)
    for i,line in enumerate(lines):
        if i:
            p._p.append(OxmlElement("a:br"))
        segments=[(line,e.bold)]
        if i==0 and e.bold_prefix and line.startswith(e.bold_prefix):
            segments=[(e.bold_prefix,True),(line[len(e.bold_prefix):],e.bold)]
        for value,bold in segments:
            for fragment,face in font_runs(value,font_file):
                run=p.add_run();run.text=fragment
                apply_ooxml_font(run.font,face,bold)
                run.font.size=Pt(e.size)
                run.font.color.rgb=rgb(e.color)

def clean_base(source,profile):
    prs=open_presentation(profile.background_source or source)
    prs._studio_background_clean = bool(profile.background_source)
    prs._studio_sources=list(prs.slides)
    prs._studio_brand_hashes={a.id for a in profile.assets}
    # Keep only the validated font part actually selected for this package.
    selected={a["origin"].get("relationship_id") for a in profile.font_assets
        if a["origin"].get("kind")=="embedded" and a.get("redistributable",True)}
    if not profile.font_assets and profile.font_origin.get("kind")=="embedded":
        selected.add(profile.font_origin.get("relationship_id"))
    font_list=prs.part._element.find(P+"embeddedFontLst")
    if font_list is not None:
        for item in list(font_list):
            for node in list(item):
                if node.tag!=P+"font" and node.get(R+"id") not in selected:
                    item.remove(node)
            if not any(node.get(R+"id") in selected for node in item if node.tag!=P+"font"):
                font_list.remove(item)
        if not len(font_list):
            prs.part._element.remove(font_list)
    for rel in list(prs.part.rels.values()):
        if rel.reltype.endswith("/font") and rel.rId not in selected:
            prs.part.drop_rel(rel.rId)
    for entry in list(prs.slides._sldIdLst):
        prs.part.drop_rel(entry.rId)
        prs.slides._sldIdLst.remove(entry)
    # Keep actual source master/layout/theme relationships, but never old slide copy.
    for surface in list(prs.slide_masters)+[l for m in prs.slide_masters for l in m.slide_layouts]:
        if not prs._studio_background_clean:
            scrub_surface(surface)
        for rel in list(surface.part.rels.values()):
            if rel.is_external or rel.reltype.rsplit("/",1)[-1] not in ("theme","slideLayout","slideMaster","image"):
                surface.part.drop_rel(rel.rId)
    for rel in list(prs.part.rels.values()):
        if rel.is_external:
            prs.part.drop_rel(rel.rId)
    props=prs.core_properties
    props.author="VK Forma Presentation Studio";props.last_modified_by="VK Forma Presentation Studio"
    props.title="Generated presentation";props.subject="";props.comments="";props.keywords=""
    return prs

def render_pptx(scenes,profile,source,path,verify_text=True):
    prs=clean_base(source,profile)
    for scene in scenes:
        pattern=next((p for p in profile.patterns if p.id==scene.pattern_id and p.title_zone),None)
        if pattern:
            slide=source_slide(prs,pattern)
        else:
            slide=prs.slides.add_slide(prs.slide_layouts[profile.layout_index])
            for shape in list(slide.shapes):
                shape._element.getparent().remove(shape._element)
            slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(scene.background)
            # Token composition has no authored source geometry. Do not inherit
            # cover artwork from the arbitrarily selected base slide layout.
            slide._element.set('showMasterSp', '0')
        for original in scene.elements:
            if original.role=="template_background":
                continue  # Native source artwork is already present, never rasterize it in PPTX.
            if original.kind=='chart' and original.series_values:
                from .charts import render_chart
                render_chart(slide,original,profile)
                continue
            for e in primitives(original,profile):
                b=e.box
                if e.kind=="text":
                    shape=slide.shapes.add_textbox(Pt(b.x),Pt(b.y),Pt(b.w),Pt(b.h))
                    set_text(shape.text_frame,e.text,e,profile)
                elif e.kind in ("rect","line"):
                    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(b.x),Pt(b.y),Pt(b.w),Pt(max(b.h,.5)))
                    shape.fill.solid();shape.fill.fore_color.rgb=rgb(e.fill or e.color)
                    shape.line.fill.background()
                elif e.kind=="image":
                    picture=slide.shapes.add_picture(e.image_path,Pt(b.x),Pt(b.y),Pt(b.w),Pt(b.h))
                    if e.image_id:
                        picture.name='forma_image_'+e.image_id
                        picture._element.nvPicPr.cNvPr.set('descr',e.text or 'Загруженная иллюстрация')
                elif e.kind=="table":
                    rows,cols=len(e.rows),len(e.rows[0])
                    table=slide.shapes.add_table(rows,cols,Pt(b.x),Pt(b.y),Pt(b.w),Pt(b.h)).table
                    table.first_row=False;table.horz_banding=False
                    widths=column_widths(e.rows,b.w,element_font(profile,e)[1],e.size)
                    for ci,width in enumerate(widths):table.columns[ci].width=Pt(width)
                    for ri,row in enumerate(e.rows):
                        table.rows[ri].height=Pt(b.h/rows)
                        for ci,value in enumerate(row):
                            cell=table.cell(ri,ci)
                            from .table_style import cell_paint
                            fill,opacity,color=cell_paint(e,ri,scene,profile)
                            if opacity:
                                cell.fill.solid();cell.fill.fore_color.rgb=rgb(fill)
                                if opacity<1:
                                    alpha=OxmlElement("a:alpha");alpha.set("val",str(round(opacity*100000)))
                                    cell._tc.get_or_add_tcPr().find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr').append(alpha)
                            else:
                                cell.fill.background()
                            cell.margin_left=cell.margin_right=Pt(8);cell.margin_top=cell.margin_bottom=Pt(6)
                            te=e.model_copy(update={"color":color,"bold":ri==0})
                            set_text(cell.text_frame,value,te,profile,widths[ci]-16)
                            # tcPr already owns cell padding. Repeating it in
                            # bodyPr makes LibreOffice apply the inset twice.
                            from .table_style import cell_borders
                            cell_borders(cell,color)
        slide.notes_slide.notes_text_frame.text=scene.notes
    prs.save(path)
    # Reopen native output and check its object content rather than trusting save().
    check=Presentation(path)
    if len(check.slides)!=len(scenes) or (verify_text and any(not any(sh.has_text_frame for sh in s.shapes) for s in check.slides)):
        raise ValueError("PPTX не прошёл проверку редактируемости")
    with ZipFile(path) as z:
        if any(b'TargetMode="External"' in z.read(n) for n in z.namelist() if n.endswith(".rels")):
            raise ValueError("В выходном PPTX осталась внешняя ссылка")

def render_pdf(scenes,profile,path):
    name=pdf_font(profile.font_file)
    c=Canvas(str(path),pagesize=(profile.width,profile.height),pageCompression=1)
    c.setTitle("VK Forma Presentation Studio");c.setAuthor("VK Forma Presentation Studio")
    def draw_text(text,x,y,w,size,color):
        font_file=element_font(profile,e)[1]
        c.setFillColor(HexColor(color));c.setFont(pdf_font(font_file),size)
        for li,line in enumerate(wrap_text(text,font_file,size,w)):
            cursor=x
            for fragment,face in font_runs(line,font_file):
                name=pdf_font(face);c.setFont(name,size)
                c.drawString(cursor,profile.height-y-size-li*size*1.25,fragment)
                cursor+=c.stringWidth(fragment,name,size)
    for scene in scenes:
        c.setFillColor(HexColor(scene.background));c.rect(0,0,profile.width,profile.height,stroke=0,fill=1)
        for original in scene.elements:
            for e in primitives(original,profile):
                b=e.box
                if e.kind=="text":
                    inset=e.size*1.4 if e.bullet else 0
                    if e.bullet:
                        draw_text("•",b.x,b.y,inset,e.size,e.color)
                    draw_text(e.text,b.x+inset,b.y,(b.w-inset)*(.94 if e.bold or e.bold_prefix else 1),e.size,e.color)
                elif e.kind in ("rect","line"):
                    c.setFillColor(HexColor(e.fill or e.color));c.rect(b.x,profile.height-b.y-b.h,b.w,b.h,stroke=0,fill=1)
                elif e.kind=="image":
                    c.drawImage(e.image_path,b.x,profile.height-b.y-b.h,b.w,b.h,mask="auto")
                elif e.kind=="table":
                    rw=b.h/len(e.rows);widths=column_widths(e.rows,b.w,element_font(profile,e)[1],e.size)
                    for ri,row in enumerate(e.rows):
                        for ci,value in enumerate(row):
                            cw=widths[ci];left=b.x+sum(widths[:ci])
                            from .table_style import cell_paint
                            fill,opacity,color=cell_paint(e,ri,scene,profile)
                            if opacity:
                                c.saveState();c.setFillAlpha(opacity)
                                c.setFillColor(HexColor(fill));c.rect(left,profile.height-b.y-(ri+1)*rw,cw,rw,stroke=0,fill=1)
                                c.restoreState()
                            draw_text(value,left+8,b.y+ri*rw+6,cw-16,e.size,color)
        c.showPage()
    c.save()

def render_html(scenes,profile,path):
    # No remote resources, scripts or user CSS. Every user string is escaped.
    if any(not a.get("redistributable",True) for a in profile.font_assets) or any(e.series_values for s in scenes for e in s.elements):
        pages=[]
        for i,scene in enumerate(scenes,1):
            data=base64.b64encode((path.parent/f"slide-{i}.png").read_bytes()).decode()
            pages.append(f'<img alt="{escape(scene.title,quote=True)}" style="display:block;max-width:100%;margin:24px auto" src="data:image/png;base64,{data}">')
        path.write_text('<!doctype html><html lang="ru"><meta charset="utf-8"><title>Презентация</title><body>'+''.join(pages)+'</body></html>')
        return
    font_data=base64.b64encode(Path(profile.font_file).read_bytes()).decode()
    css=f"""@font-face{{font-family:DeckFont;src:url(data:font/ttf;base64,{font_data})}}*{{box-sizing:border-box}}body{{margin:0;background:#e8e8ec;display:grid;gap:24px;padding:24px;justify-content:center}}.slide{{position:relative;width:{profile.width}px;height:{profile.height}px;overflow:hidden;box-shadow:0 4px 18px #0002}}.el{{position:absolute;font-family:DeckFont;white-space:pre;line-height:1.25}}table{{border-collapse:collapse;table-layout:fixed;width:100%;height:100%;font-weight:normal}}td{{padding:6px 8px;vertical-align:top;white-space:pre-wrap;line-height:1.25}}@media print{{body{{padding:0;gap:0}}.slide{{break-after:page;box-shadow:none}}}}"""
    font_names = {}
    for i,file in enumerate(profile_font_files(profile)):
        font_names[file]=f"DeckFace{i}"
        data=base64.b64encode(Path(file).read_bytes()).decode()
        css+=f'@font-face{{font-family:DeckFace{i};src:url(data:font/ttf;base64,{data})}}'
    out=['<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Презентация</title><style>'+css+'</style><body>']
    for scene in scenes:
        out.append(f'<section class="slide" aria-label="{escape(scene.title,quote=True)}" style="background:{scene.background}">')
        for original in scene.elements:
            for e in primitives(original,profile):
                b=e.box
                style=f"left:{b.x}px;top:{b.y}px;width:{b.w}px;height:{b.h}px;font-size:{e.size}px;color:{e.color or profile.foreground};font-weight:{700 if e.bold else 400};"
                file=element_font(profile,e)[1]
                faces=','.join(dict.fromkeys([font_names[file],font_names.get(symbol_font(),font_names[file])]))
                style+=f'font-family:{faces};'
                if e.kind=="text":
                    inset=e.size*1.4 if e.bullet else 0
                    content=escape("\n".join(wrap_text(e.text,file,e.size,(b.w-inset)*(.94 if e.bold or e.bold_prefix else 1))))
                    if e.bold_prefix and content.startswith(escape(e.bold_prefix)):
                        prefix=escape(e.bold_prefix);content='<strong>'+prefix+'</strong>'+content[len(prefix):]
                    if e.bullet:
                        out.append(f'<ul class="el" style="{style}margin:0;padding:0 0 0 {inset}px"><li>{content}</li></ul>')
                    else:
                        out.append(f'<div class="el" style="{style}">{content}</div>')
                elif e.kind in ("rect","line"):
                    out.append(f'<div class="el" aria-hidden="true" style="{style}background:{e.fill or e.color}"></div>')
                elif e.kind=="image":
                    data=base64.b64encode(Path(e.image_path).read_bytes()).decode()
                    alt=escape(e.text or ('Загруженная иллюстрация' if e.image_id else 'Элемент шаблона'),quote=True)
                    out.append(f'<img class="el" alt="{alt}" style="{style}" src="data:image/png;base64,{data}">')
                elif e.kind=="table":
                    widths=column_widths(e.rows,b.w,file,e.size)
                    out.append(f'<div class="el" style="{style}"><table><colgroup>'+''.join(f'<col style="width:{width}px">' for width in widths)+'</colgroup>')
                    for ri,row in enumerate(e.rows):
                        from .table_style import cell_paint
                        fill,opacity,color=cell_paint(e,ri,scene,profile)
                        paint='transparent' if not opacity else f'rgba({int(fill[1:3],16)},{int(fill[3:5],16)},{int(fill[5:7],16)},{opacity})'
                        out.append(f'<tr style="height:{b.h/len(e.rows)}px;background:{paint};color:{color}">')
                        for ci,cell in enumerate(row):
                            wrapped="\n".join(wrap_text(cell,file,e.size,widths[ci]-16))
                            out.append('<td>'+escape(wrapped)+'</td>')
                        out.append('</tr>')
                    out.append('</table></div>')
        out.append('</section>')
    out.append('</body></html>')
    path.write_text("".join(out))

def render_variant(scenes,profile,source,directory):
    rendered_text=[]
    for scene in scenes:
        for original in scene.elements:
            for element in primitives(original,profile):
                if element.kind not in ('text','table'):
                    continue
                text=element.text+"\n"+"\n".join(cell for row in element.rows for cell in row)
                if element.bullet:
                    text+="•"
                check_glyphs(element_font(profile,element)[1],text)
    directory.mkdir(parents=True,exist_ok=True)
    render_pptx(scenes,profile,source,directory/"deck.pptx")
    from .export_audit import repair_symbols, geometry
    actual=open_presentation(directory/'deck.pptx')
    repair_symbols(actual,profile)
    # Native scenes were already repaired before rendering. Keep scene-based
    # HTML/PDF fallbacks consistent; actual-object repairs belong to PPTAgent.
    actual_findings,actual_repairs=geometry(actual,profile,repair=False)
    actual.save(directory/'deck.pptx')
    from .font_disclosure import exported_substitutions
    font_substitutions = exported_substitutions(actual, profile)
    from .office import to_pdf
    native_preview=to_pdf(directory/"deck.pptx",directory,font_files=profile_font_files(profile))
    if not native_preview:
        if any(e.series_values for s in scenes for e in s.elements):
            raise ValueError('Для проверки нативных диаграмм требуется LibreOffice; приближённый рендер не выдаётся за результат PPTX')
        render_pdf(scenes,profile,directory/"deck.pdf")
    (directory/"slides.json").write_text(json.dumps([s.model_dump() for s in scenes],ensure_ascii=False,indent=2))
    import pypdfium2 as pdfium
    # PDFium calls are serialized; its C API is not thread-safe.
    with _pdfium_lock, pdfium.PdfDocument(str(directory/"deck.pdf")) as doc:
        if len(doc)!=len(scenes):
            raise ValueError("Неверное число страниц PDF")
        for i in range(len(doc)):
            page=doc[i]
            bitmap=page.render(scale=1)
            bitmap.to_pil().save(directory/f"slide-{i+1}.png")
            bitmap.close();page.close()
    if any(not a.get("redistributable",True) for a in profile.font_assets):
        # Local-only font licenses: distribute rendered pixels, never font programs.
        c=Canvas(str(directory/"deck.pdf"),pagesize=(profile.width,profile.height))
        for i in range(1,len(scenes)+1):
            c.drawImage(str(directory/f"slide-{i}.png"),0,0,profile.width,profile.height)
            c.showPage()
        c.save()
    render_html(scenes,profile,directory/"deck.html")
    licenses=[]
    for file in profile_font_files(profile):
        license_path=Path(file).parent/"OFL.txt"
        if license_path.is_file():
            licenses.append(license_path.read_text())
    if licenses:
        (directory/"FONT_LICENSES.txt").write_text("\n\n".join(dict.fromkeys(licenses)))
    return {"preview_source":"libreoffice_pptx" if native_preview else "scene_model", "native_render":native_preview,
            'object_findings':actual_findings,'object_repairs':actual_repairs,
            'font_substitutions':font_substitutions}
