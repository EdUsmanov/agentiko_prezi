"""Native editable charts, grounded exclusively in source table cells."""
import math
import re
from .models import Element


def table_series(table):
    series=[]
    if len(table.rows)>30:
        return None
    for ci in range(1,len(table.headers)):
        values=[]; units=set()
        for row in table.rows:
            match=re.fullmatch(r'\s*([−\-+]?\d[\d\s]*(?:[.,]\d+)?)\s*(%|₽|руб\.?|млн|тыс\.?)?\s*',row[ci])
            if not match:
                return None
            number=float(match[1].replace(' ','').replace(',','.').replace('−','-'))
            if not math.isfinite(number): return None
            values.append(number);units.add(match[2] or '')
        if len(units)>1: return None
        series.append(values)
    return series or None


def make_chart(table, slide, box, profile, color, source_ids):
    source_table=table
    table,supplement=chart_projection(table)
    values=table_series(table)
    chart_type=slide.chart_type if slide.chart_type!='auto' else 'bar'
    suffixes={re.sub(r'^[\s−\-+\d.,]+','',cell).strip() for row in table.rows for cell in row[1:]}
    suffix=next(iter(suffixes)) if len(suffixes)==1 else ''
    if len(suffixes)>1:
        values=None  # Different physical units must not share an unlabeled axis.
    if not values or chart_type=='pie' and (len(values)!=1 or any(v<0 for v in values[0]) or sum(values[0])<=0):
        return Element(kind='table',box=box,rows=[source_table.headers]+source_table.rows,font=profile.font,
            size=profile.body_size,color=color,fill=profile.accent,source_ids=source_ids)
    return Element(kind='chart',box=box,labels=[row[0] for row in table.rows],values=values[0],
        value_labels=[row[1] for row in table.rows],unit=table.headers[1]+(f' ({suffix})' if suffix else ''),category_title=table.headers[0],
        series_names=table.headers[1:],series_values=values,chart_type=chart_type,chart_style=slide.chart_style,
        rows=[source_table.headers]+source_table.rows,font=profile.font,size=max(16,min(profile.body_size,20)),
        color=color,fill=profile.accent,source_ids=source_ids)


def chart_projection(table, *, compact_captions=False):
    """Comparable series only; totals and other units stay visible as captions."""
    from .content import plain_inline
    headers=[plain_inline(c) for c in table.headers]
    rows=[[plain_inline(c) for c in row] for row in table.rows]
    units=[]
    for ci in range(1,len(headers)):
        suffixes={re.sub(r'^[\s−\-+\d.,]+','',row[ci]).strip() for row in rows}
        units.append(next(iter(suffixes)) if len(suffixes)==1 else None)
    selected=[i+1 for i,u in enumerate(units) if u==units[0]] if units else []
    omitted=[i for i in range(1,len(headers)) if i not in selected]
    plotted=[];supplement=[]
    for row in rows:
        total=bool(re.match(r'^(?:всего|итого|total)\b',row[0],re.I))
        columns=list(range(1,len(headers))) if total else omitted
        if columns and (not compact_captions or total):
            supplement.append(f'{headers[0]} {row[0]}: '+ '; '.join(f'{headers[i]}: {row[i]}' for i in columns))
        if not total:
            plotted.append([row[0]]+[row[i] for i in selected])
    if compact_captions:
        regular=[row for row in rows if not re.match(r'^(?:всего|итого|total)\b',row[0],re.I)]
        supplement=[f'{headers[i]} ({headers[0]}): '+ '; '.join(f'{row[0]} — {row[i]}' for row in regular)
                    for i in omitted if regular]+supplement
    return table.model_copy(update={'headers':[headers[0]]+[headers[i] for i in selected],'rows':plotted}),supplement


def chart_caption_layout(rows,width,profile):
    """One measured caption contract for selection, audit and native export."""
    from .models import TableData
    from .fonts import role_font,wrap_text
    if not rows:return '',0
    _,lines=chart_projection(TableData(id='chart',headers=rows[0],rows=rows[1:]),compact_captions=True)
    if not lines:return '',0
    height=sum(len(wrap_text(line,role_font(profile,'body')[1],16,width*.96))*20 for line in lines)+12
    return '\n'.join(lines),height


def render_chart(slide,e,profile):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_LABEL_POSITION
    from pptx.util import Pt
    from .render import rgb
    data=CategoryChartData()
    readable=e.chart_style=='readable'
    if readable and e.chart_type in ('column','line'):
        import textwrap
        per_category=max(24,(e.box.w-90)/max(1,len(e.labels)))
        characters=max(5,min(14,int(per_category/(min(e.size,16)*.55))))
        data.categories=['\n'.join(textwrap.wrap(label,width=characters,break_long_words=False,break_on_hyphens=False)) for label in e.labels]
    else:
        data.categories=e.labels
    for name,values in zip(e.series_names or [e.unit], e.series_values or [e.values]):
        data.add_series(name,values)
    types={'bar':XL_CHART_TYPE.BAR_CLUSTERED,'column':XL_CHART_TYPE.COLUMN_CLUSTERED,
           'column_stacked':XL_CHART_TYPE.COLUMN_STACKED,
           'line':XL_CHART_TYPE.LINE_MARKERS,'pie':XL_CHART_TYPE.PIE}
    b=e.box.model_copy()
    # Use the exact same caption height as selection and pre-export audit.
    text,height=chart_caption_layout(e.rows,b.w,profile)
    if text:
        from .render import set_text
        from .models import Box
        if b.h-height-12<160:
            raise ValueError('Диаграмма и пояснения единиц не помещаются в поле; требуется более вместительный макет')
        caption=Element(kind='text',text=text,box=Box(x=b.x,y=b.y+b.h-height,w=b.w,h=height),
            font=e.font,size=16,color=e.color,field_style=e.field_style)
        shape=slide.shapes.add_textbox(Pt(caption.box.x),Pt(caption.box.y),Pt(caption.box.w),Pt(caption.box.h))
        set_text(shape.text_frame,text,caption,profile)
        b.h-=height+12
    chart=slide.shapes.add_chart(types[e.chart_type],Pt(b.x),Pt(b.y),Pt(b.w),Pt(b.h),data).chart
    chart.has_title=False
    from .fonts import element_font
    from .font_identity import apply_ooxml_font
    font_file=element_font(profile,e)[1]
    apply_ooxml_font(chart.font,font_file)
    chart.font.size=Pt(e.size);chart.font.color.rgb=rgb(e.color)
    chart.has_legend=len(e.series_values)>1 or e.chart_type=='pie' and not readable
    if chart.has_legend:
        chart.legend.position=XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout=False
        apply_ooxml_font(chart.legend.font,font_file)
        chart.legend.font.size=Pt(min(16,e.size))
        chart.legend.font.color.rgb=rgb(e.color)
    palette=list(dict.fromkeys([profile.accent]+profile.colors))
    from .template import contrast
    palette=[c for c in palette if contrast(c,e.background_hint or profile.background)>1.6] or [profile.accent]
    for i,series in enumerate(chart.series):
        color=palette[i%len(palette)]
        series.format.line.color.rgb=rgb(color)
        if e.chart_type!='line':
            series.format.fill.solid();series.format.fill.fore_color.rgb=rgb(color)
        if e.chart_type=='column_stacked':
            data_labels=series.data_labels
            data_labels.show_value=True
            data_labels.position=XL_LABEL_POSITION.CENTER
            apply_ooxml_font(data_labels.font,font_file)
            data_labels.font.size=Pt(min(16,e.size))
            data_labels.font.color.rgb=rgb(max(profile.colors,key=lambda ink:contrast(ink,color)))
        if e.chart_type=='pie':
            for j,point in enumerate(series.points):
                point.format.fill.solid();point.format.fill.fore_color.rgb=rgb(palette[j%len(palette)])
    plot=chart.plots[0]
    plot.has_data_labels=True
    labels=plot.data_labels
    label_size=min(e.size,16)
    apply_ooxml_font(labels.font,font_file)
    labels.font.size=Pt(label_size);labels.font.color.rgb=rgb(e.color)
    labels.show_value=True
    suffixes={re.sub(r'^[\s−\-+\d.,]+','',v).strip() for v in e.value_labels}
    if len(suffixes)==1 and next(iter(suffixes)):
        suffix=next(iter(suffixes))
        labels.number_format='0.########" '+suffix+'"'
    if e.chart_type=='pie':
        labels.position=XL_LABEL_POSITION.OUTSIDE_END
        if readable:
            labels.show_category_name=True
            labels.position=XL_LABEL_POSITION.OUTSIDE_END
    else:
        labels.position=(XL_LABEL_POSITION.CENTER if e.chart_type=='column_stacked' else
            XL_LABEL_POSITION.ABOVE if e.chart_type=='line' else XL_LABEL_POSITION.OUTSIDE_END)
        chart.value_axis.has_minor_gridlines=False
        chart.value_axis.has_major_gridlines=True
        apply_ooxml_font(chart.category_axis.tick_labels.font,font_file)
        category_size=label_size
        if readable and e.chart_type in ('column','line'):
            longest_word=max((len(word) for label in e.labels for word in label.split()),default=1)
            category_size=max(10,min(label_size,per_category/(longest_word*.65)))
        chart.category_axis.tick_labels.font.size=Pt(category_size)
        chart.category_axis.tick_labels.font.color.rgb=rgb(e.color)
        if readable:
            axis=chart.category_axis._element
            tx=axis.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}txPr')
            if tx is not None:
                body=tx.find('{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr')
                if body is not None:body.set('rot','0')
        apply_ooxml_font(chart.value_axis.tick_labels.font,font_file)
        chart.value_axis.tick_labels.font.size=Pt(label_size)
        chart.value_axis.tick_labels.font.color.rgb=rgb(e.color)
        for axis,title in ((chart.category_axis,e.category_title),(chart.value_axis,e.unit if len(e.series_values)<=1 else '')):
            axis.has_title=bool(title)
            if title:
                frame=axis.axis_title.text_frame;frame.text=title
                for p in frame.paragraphs:
                    apply_ooxml_font(p.font,font_file)
                    p.font.size=Pt(e.size);p.font.color.rgb=rgb(e.color)
    return chart
