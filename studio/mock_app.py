"""In-memory API that replaces the production backend responses.

Run with ``python -m studio.mock_app --port 8765``. This module deliberately does
not import the production app, pipeline or store and never opens ``data/``.
"""

import argparse
from copy import deepcopy
from html import escape
from io import BytesIO
import json
from pathlib import Path
import re
import time
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field, field_validator


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = {"id": "mock-vk-forma", "name": "Демо-шаблон VK Forma.pptx"}
VARIANTS = (
    ("executive", "Главное и решения", "#2457EA"),
    ("analytical", "Данные и доказательства", "#101828"),
    ("story", "Контекст и развитие", "#11BFA5"),
)
PRESETS = {"mini": 4, "standard": 7, "large": 12}


class GenerateRequest(BaseModel):
    package_id: str
    accept_adjusted_slide_count: bool = False
    variant_count: int = 3

    @field_validator("variant_count", mode="before")
    @classmethod
    def strict_variant_count(cls, value):
        if type(value) is not int or value not in (1, 3):
            raise ValueError("Выберите одну или три презентации")
        return value


class ReviseRequest(BaseModel):
    instructions: str = Field(min_length=1, max_length=5000)
    slides: int | None = Field(default=None, ge=1, le=30)
    size_preset: str | None = None


def create_app(*, preparation_seconds: float = 1.0, generation_seconds: float = 1.5):
    app = FastAPI(title="VK Forma Mock API", version="0.1.0-mock")
    jobs: dict[str, dict] = {}
    events: dict[str, list[dict]] = {}
    files: dict[tuple[str, str], bytes] = {}
    texts: dict[str, str] = {}
    app.state.jobs = jobs

    def log(job: dict, event: str, **data):
        history = events[job["id"]]
        history.append(
            {
                "id": history[-1]["id"] + 1 if history else 1,
                "created": time.time(),
                "level": "info",
                "event": event,
                "data": data,
            }
        )

    def create(kind: str, **fields):
        job = {
            "id": uuid4().hex,
            "kind": kind,
            "state": "accepted",
            "created": time.time(),
            "phase": "Ожидание",
            "progress": 0,
            **fields,
        }
        jobs[job["id"]] = job
        events[job["id"]] = []
        log(job, "job.created", kind=kind)
        return job

    def get(jid: str):
        job = jobs.get(jid)
        if job is None:
            raise HTTPException(404, "Задание не найдено")
        refresh(job)
        return job

    def start_generation(package: dict, variant_count: int = 3):
        existing_id = package.get("generation_id")
        if existing_id:
            if package.get("variant_count", 3) != variant_count:
                raise HTTPException(
                    409, "Для этого пакета уже запущено другое количество презентаций"
                )
            return get(existing_id)
        job = create(
            "generation",
            package_id=package["id"],
            variant_count=variant_count,
            progress=0,
            phase="Планирование одного варианта"
            if variant_count == 1
            else "Планирование трёх вариантов",
            analysis_seconds=package.get("analysis_seconds"),
            slide_count_decision={
                "mode": "confirmed",
                "count": package["analysis"]["planned_slides"],
                "message": "План подтверждён",
            },
        )
        package["generation_id"] = job["id"]
        package["auto_generation"] = "started"
        package["variant_count"] = variant_count
        log(package, "generation.started", generation_id=job["id"], automatic=False)
        return job

    def refresh(job: dict):
        age = time.time() - job["created"]
        if job["state"] in ("accepted", "running"):
            duration = preparation_seconds if job["kind"] == "preparation" else generation_seconds
            if age >= duration:
                job["state"] = "ready" if job["kind"] == "preparation" else "completed"
                job["phase"] = "Готово к генерации" if job["kind"] == "preparation" else "Готово"
                job["progress"] = 100
                if job["kind"] == "preparation":
                    job["analysis_seconds"] = round(duration, 1)
                    job["auto_generation"] = "manual"
                else:
                    job["elapsed_seconds"] = round(duration, 1)
                    package = jobs[job["package_id"]]
                    count = package["analysis"]["planned_slides"]
                    job["variants"] = [
                        {"key": key, "title": title, "slides": count}
                        for key, title, _ in VARIANTS[: job.get("variant_count", 3)]
                    ]
                    job["model_mode"] = "extractive"
                    job["engine"] = {"engine": "mock"}
                    job["font_substitutions"] = []
                log(job, "job.state", state=job["state"], phase=job["phase"])
            else:
                job["state"] = "running"
                job["progress"] = min(95, round(age / max(duration, 0.001) * 100))
                job["phase"] = (
                    "Факты, таблицы и ограничения"
                    if job["kind"] == "preparation"
                    else "Вёрстка, аудит и экспорт"
                )

    def copy(job: dict):
        return deepcopy(job)

    @app.get("/api/health")
    def health():
        return {
            "status": "ok",
            "model_mode": "extractive",
            "model_id": "mock",
            "engine": "mock",
            "features": {"vlm": False, "download_fonts": False, "image_uploads": True},
            "mock": True,
        }

    @app.get("/api/runtime")
    def runtime():
        return {"restart_required": False, "organizer_preanalysis": False}

    @app.get("/api/references")
    def references():
        return [REFERENCE]

    @app.get("/api/jobs")
    def list_jobs():
        return [
            copy(get(jid))
            for jid in sorted(jobs, key=lambda key: jobs[key]["created"], reverse=True)[:30]
        ]

    @app.get("/api/jobs/{jid}")
    def job(jid: str):
        return copy(get(jid))

    @app.delete("/api/jobs/{jid}")
    def delete_job(jid: str):
        if jid not in jobs:
            raise HTTPException(404, "Задание не найдено")
        targets = {jid}
        while True:
            linked = {
                key
                for key, current in jobs.items()
                if key not in targets
                and (
                    (current["kind"] == "generation" and current.get("package_id") in targets)
                    or (
                        current["kind"] == "preparation"
                        and current.get("parent_package") in targets
                    )
                )
            }
            if not linked:
                break
            targets.update(linked)
        if any(jobs[key]["state"] in ("accepted", "running") for key in targets):
            raise HTTPException(409, "Дождитесь завершения запуска перед удалением")
        for current in jobs.values():
            if current.get("generation_id") in targets:
                current.pop("generation_id", None)
                current["auto_generation"] = "cancelled"
        for key in targets:
            jobs.pop(key)
            events.pop(key, None)
            texts.pop(key, None)
        for key in list(files):
            if key[0] in targets:
                del files[key]
        return {"deleted": sorted(targets)}

    @app.post("/api/prepare", status_code=202)
    async def prepare(
        text: str = Form(..., max_length=120000),
        audience: str = Form("", max_length=2000),
        instructions: str = Form("", max_length=5000),
        slides: int | None = Form(None, ge=1, le=30),
        size_preset: str = Form("standard"),
        reference_id: str = Form(""),
        source_package_id: str = Form(""),
        template: UploadFile | None = File(None),
        images: list[UploadFile] | None = File(None),
    ):
        if not text.strip():
            raise HTTPException(422, "Добавьте текст")
        if (
            sum((bool(reference_id), bool(source_package_id), bool(template and template.filename)))
            != 1
        ):
            raise HTTPException(422, "Выберите один шаблон: файл, пример или изученный пакет")
        if reference_id and reference_id != REFERENCE["id"]:
            raise HTTPException(404, "Неизвестный пример")
        if template and not template.filename.lower().endswith((".pptx", ".potx")):
            raise HTTPException(422, "Поддерживаются PPTX и POTX")
        if size_preset not in PRESETS:
            raise HTTPException(422, "Неизвестный размер презентации")
        image_names = [image.filename for image in images or [] if image.filename]
        if len(image_names) > 12:
            raise HTTPException(422, "Допускается до 12 изображений")
        source = get(source_package_id) if source_package_id else None
        if source and (source["kind"] != "preparation" or source["state"] != "ready"):
            raise HTTPException(409, "Дизайн-система ещё не подготовлена")
        name = (
            source["template"]["name"]
            if source
            else REFERENCE["name"]
            if reference_id
            else Path(template.filename).name
        )
        count = slides or instruction_slide_count(instructions) or PRESETS[size_preset]
        titles = outline_titles(text, count)
        package = create(
            "preparation",
            template_name=name,
            **({"source_package_id": source_package_id} if source else {}),
            template=deepcopy(source["template"])
            if source
            else {
                "name": name,
                "colors": ["#2457EA", "#101828", "#F3F6FF", "#11BFA5"],
                "slide_count": 12,
                "patterns": ["cover", "two-column", "data", "quote", "summary"],
                "layout_count": 5,
                "font": "Montserrat",
                "font_origin": {"kind": "local"},
                "font_roles": {},
            },
            content={
                "title": titles[0],
                "facts": max(1, len(text.splitlines())),
                "tables": text.count("|---"),
                "images": len(image_names),
            },
            constraints={
                "slides": count,
                "audience": audience,
                "instructions": instructions,
                "size_preset": size_preset,
            },
            analysis={
                "planned_slides": count,
                "slide_budget": {"status": "ok", "planned": count},
                "canonical_storyboard": [{"title": title} for title in titles],
                "planning_source": "mock",
                "template_semantics": {"status": "mock"},
                "document_structure": {"status": "mock"},
                "native_render": {"patterns": 5},
                "composition_preview": {"verified": True},
            },
        )
        texts[package["id"]] = text
        return copy(package)

    @app.post("/api/generate", status_code=202)
    def generate(body: GenerateRequest):
        package = get(body.package_id)
        if package["kind"] != "preparation" or package["state"] != "ready":
            raise HTTPException(409, "Пакет не готов к генерации")
        return copy(start_generation(package, body.variant_count))

    @app.post("/api/packages/{pid}/auto-generation/cancel")
    def cancel_auto_generation(pid: str):
        package = get(pid)
        if package["kind"] != "preparation":
            raise HTTPException(409, "Это не пакет анализа")
        if package.get("auto_generation") == "scheduled":
            package["auto_generation"] = "cancelled"
            log(package, "generation.autostart_cancelled")
        return copy(package)

    @app.post("/api/packages/{pid}/retry-fonts", status_code=202)
    def retry_fonts(pid: str):
        get(pid)
        raise HTTPException(409, "Мок не создаёт состояние waiting_fonts")

    @app.post("/api/packages/{pid}/revise", status_code=202)
    def revise(pid: str, body: ReviseRequest):
        original = get(pid)
        if original["kind"] != "preparation" or original["state"] != "ready":
            raise HTTPException(409, "Пакет не готов к изменению")
        if body.size_preset is not None and body.size_preset not in PRESETS:
            raise HTTPException(422, "Неизвестный размер презентации")
        count = (
            body.slides
            or instruction_slide_count(body.instructions)
            or PRESETS.get(body.size_preset, original["constraints"]["slides"])
        )
        revised = deepcopy(original)
        revised.pop("id")
        revised.pop("created")
        revised.pop("generation_id", None)
        revised.pop("auto_generation", None)
        revised.pop("auto_generate_at", None)
        revised.pop("analysis_seconds", None)
        revised["parent_package"] = pid
        revised["constraints"]["slides"] = count
        revised["constraints"]["instructions"] = body.instructions
        revised["analysis"]["planned_slides"] = count
        revised["analysis"]["slide_budget"] = {"status": "ok", "planned": count}
        revised["analysis"]["canonical_storyboard"] = [
            {"title": title} for title in outline_titles(texts[pid], count)
        ]
        revised["state"] = "accepted"
        revised["phase"] = "Анализ содержания и подготовка плана"
        revised["progress"] = 0
        original["auto_generation"] = "cancelled"
        package = create(
            "preparation", **{k: v for k, v in revised.items() if k not in ("kind", "state")}
        )
        texts[package["id"]] = texts[pid]
        return copy(package)

    @app.get("/api/jobs/{jid}/diagnostics")
    def diagnostics(jid: str, after: int = -1, download: bool = False):
        current = get(jid)
        history = events[jid]
        selected = (
            history if download else [event for event in history if event["id"] > after][-500:]
        )
        result = {
            "job_id": jid,
            "state": current["state"],
            "phase": current["phase"],
            "error": current.get("error", ""),
            "events": selected,
            "checks": {"mock": True},
            "attachments": {},
            "retention": "Моковые события хранятся только до остановки сервера.",
        }
        headers = (
            {"Content-Disposition": f'attachment; filename="diagnostics-{jid}.json"'}
            if download
            else {}
        )
        return Response(
            json.dumps(result, ensure_ascii=False), media_type="application/json", headers=headers
        )

    @app.get("/api/jobs/{jid}/files/{filename:path}")
    def artifact(jid: str, filename: str):
        current = get(jid)
        if current["state"] not in ("ready", "completed", "needs_review"):
            raise HTTPException(409, "Артефакты ещё не готовы")
        if current["kind"] == "preparation":
            allowed = {"analysis.json", "DESIGN.md", "font-model.json", "color-model.json"}
        else:
            allowed = {"manifest.json", "presentations.zip"}
            selected = VARIANTS[: current.get("variant_count", 3)]
            allowed.update(
                f"{key}/deck.{ext}" for key, _, _ in selected for ext in ("pptx", "pdf", "html")
            )
            allowed.update(
                f"{key}/slide-{index}.png"
                for key, _, _ in selected
                for index in range(1, current["variants"][0]["slides"] + 1)
            )
        if filename not in allowed:
            raise HTTPException(404, "Файл не найден")
        key = (jid, filename)
        if key not in files:
            files[key] = make_file(current, filename, jobs)
        suffix = filename.rsplit(".", 1)[-1]
        media = {
            "png": "image/png",
            "pdf": "application/pdf",
            "html": "text/html",
            "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "zip": "application/zip",
            "md": "text/markdown",
            "json": "application/json",
        }[suffix]
        headers = (
            {"Content-Disposition": f'attachment; filename="{Path(filename).name}"'}
            if suffix in ("pptx", "zip", "json", "md")
            else {}
        )
        return Response(files[key], media_type=media, headers=headers)

    return app


def outline_titles(text: str, count: int):
    headings = [
        match.group(1).strip()
        for line in text.splitlines()
        if (match := re.match(r"^#{1,3}\s+(.+)", line))
    ]
    lines = headings or [re.sub(r"^[*\-\d.\s]+", "", line).strip() for line in text.splitlines()]
    lines = [line[:80] for line in lines if len(line) > 4 and not line.startswith("|")]
    first = lines[0] if lines else "Презентация"
    middle = lines[1:] or ["Ключевые данные", "Основные выводы", "Следующие шаги"]
    return [first] + [middle[(index - 1) % len(middle)] for index in range(1, count)]


def instruction_slide_count(instructions: str):
    match = re.search(r"\b(\d{1,2})\s*слайд(?:а|ов)?\b", instructions.lower())
    if match:
        count = int(match.group(1))
        if 1 <= count <= 30:
            return count
    return None


def make_file(job: dict, filename: str, jobs: dict[str, dict]):
    if job["kind"] == "preparation":
        if filename == "analysis.json":
            return json_bytes(job["analysis"])
        if filename == "DESIGN.md":
            return f"# Моковая дизайн-система\n\nШаблон: {job['template']['name']}\n\nЦвета: синий, мятный, белый.\n".encode()
        if filename == "font-model.json":
            return json_bytes({"mock": True, "font": "Montserrat", "substitutions": []})
        return json_bytes({"mock": True, "colors": job["template"]["colors"]})

    package = jobs[job["package_id"]]
    if filename == "manifest.json":
        return json_bytes({"mock": True, "job_id": job["id"], "variants": job["variants"]})
    if filename == "presentations.zip":
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as archive:
            archive.writestr("MOCK.txt", "Демонстрационные файлы мокового сервера VK Forma.\n")
            for key, _, _ in VARIANTS[: job.get("variant_count", 3)]:
                for ext in ("pptx", "pdf", "html"):
                    name = f"{key}/deck.{ext}"
                    archive.writestr(name, make_file(job, name, jobs))
        return output.getvalue()

    variant_key, name = filename.split("/", 1)
    variant = next(item for item in VARIANTS if item[0] == variant_key)
    titles = [slide["title"] for slide in package["analysis"]["canonical_storyboard"]]
    if name.endswith(".png"):
        index = int(re.fullmatch(r"slide-(\d+)\.png", name).group(1))
        return png_slide(titles[index - 1], index, len(titles), variant)
    if name == "deck.pptx":
        return pptx_deck(titles, variant)
    if name == "deck.pdf":
        return pdf_deck(titles, variant)
    return html_deck(titles, variant).encode("utf-8")


def json_bytes(value):
    return json.dumps(value, ensure_ascii=False, indent=2).encode("utf-8")


def png_slide(title, index, total, variant):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1280, 720), "#F4F7FC")
    draw = ImageDraw.Draw(image)
    color = variant[2]
    draw.rectangle((0, 0, 1280, 22), fill=color)
    draw.rounded_rectangle((88, 92, 1192, 625), radius=24, fill="white")
    draw.rectangle((126, 148, 137, 500), fill=color)
    font_path = ROOT / "fonts/Montserrat-Bold.ttf"
    regular_path = ROOT / "fonts/Montserrat-Regular.ttf"
    heading = ImageFont.truetype(str(font_path), 50)
    small = ImageFont.truetype(str(regular_path), 25)
    words, rows, line = title.split(), [], ""
    for word in words:
        candidate = f"{line} {word}".strip()
        if draw.textlength(candidate, font=heading) > 900 and line:
            rows.append(line)
            line = word
        else:
            line = candidate
    rows.append(line)
    for offset, row in enumerate(rows[:4]):
        draw.text((177, 205 + offset * 72), row, font=heading, fill="#101828")
    draw.text((177, 560), f"{variant[1]}  ·  ДЕМО", font=small, fill=color)
    draw.text((1100, 560), f"{index} / {total}", font=small, fill="#667085")
    output = BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def pptx_deck(titles, variant):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    color = RGBColor.from_string(variant[2].lstrip("#"))
    for index, title in enumerate(titles, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        bar = slide.shapes.add_shape(1, 0, 0, deck.slide_width, Inches(0.18))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()
        heading = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.2), Inches(3.2))
        paragraph = heading.text_frame.paragraphs[0]
        paragraph.text = title
        paragraph.font.name = "Montserrat"
        paragraph.font.size = Pt(38)
        paragraph.font.bold = True
        footer = slide.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.2), Inches(0.4))
        footer.text_frame.text = f"{variant[1]} · ДЕМО · {index}/{len(titles)}"
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def pdf_deck(titles, variant):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    if "MockMontserrat" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(
            TTFont("MockMontserrat", str(ROOT / "fonts/Montserrat-Regular.ttf"))
        )
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=(960, 540))
    for index, title in enumerate(titles, 1):
        pdf.setFillColor(variant[2])
        pdf.rect(0, 525, 960, 15, fill=1, stroke=0)
        pdf.setFillColor("#101828")
        pdf.setFont("MockMontserrat", 30)
        for offset, line in enumerate(
            [title[pos : pos + 48] for pos in range(0, len(title), 48)][:4]
        ):
            pdf.drawString(72, 370 - offset * 48, line)
        pdf.setFont("MockMontserrat", 14)
        pdf.drawString(72, 48, f"{variant[1]} · ДЕМО · {index}/{len(titles)}")
        pdf.showPage()
    pdf.save()
    return output.getvalue()


def html_deck(titles, variant):
    slides = "".join(
        f"<section><h1>{escape(title)}</h1><footer>{escape(variant[1])} · ДЕМО · {index}/{len(titles)}</footer></section>"
        for index, title in enumerate(titles, 1)
    )
    return (
        '<!doctype html><html lang="ru"><meta charset="utf-8"><title>VK Forma · Демо</title>'
        f"<style>body{{margin:0;background:#eef2f8;font-family:Arial,sans-serif}}section{{box-sizing:border-box;width:960px;height:540px;margin:24px auto;padding:100px 70px;background:white;border-top:14px solid {variant[2]};page-break-after:always}}h1{{font-size:40px;color:#101828}}footer{{margin-top:210px;color:{variant[2]}}}</style>"
        f"{slides}</html>"
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Run the standalone VK Forma mock backend")
    parser.add_argument("--port", type=int, default=8765)
    args = parser.parse_args()
    import uvicorn
    from .frontend import create_frontend_app

    uvicorn.run(create_frontend_app(create_app()), host="127.0.0.1", port=args.port)
