"""One real template through isolated HTTP + workers; artifacts stay under test-results."""

import argparse
from contextlib import ExitStack, nullcontext
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import mimetypes
from pathlib import Path
import sys
import time
from zipfile import ZipFile

import httpx
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from test_support.app_server import application  # noqa: E402


CONTENT = """# Обзор проекта
## Контекст
Команда готовит единый сервис для обработки заявок.
## Процесс
Сотрудник регистрирует заявку и назначает ответственного.
## Проверка
Пилот длится 12 недель; команда проверяет 40 заявок.
## Результат
Руководитель сравнивает время обработки до и после пилота.
## Решение
После проверки команда решает, расширять ли сервис.
"""
TERMINAL = {"ready", "needs_review", "completed", "failed", "cancelled", "waiting_fonts"}
VARIANTS = ("executive", "analytical", "story")


class RunFailure(RuntimeError):
    pass


def save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str))


def request_json(client, method, url, **kwargs):
    response = client.request(method, url, **kwargs)
    try:
        payload = response.json()
    except ValueError:
        payload = response.text[:3000]
    if response.is_error:
        raise RunFailure(f"{method} {url}: HTTP {response.status_code}: {payload}")
    return payload


def wait_job(client, url, jid, stage, directory, timeout, poll):
    deadline = time.monotonic() + timeout
    last = None
    while time.monotonic() < deadline:
        last = request_json(client, "GET", f"{url}/api/jobs/{jid}")
        save_json(directory / f"{stage}.json", last)
        if last["state"] in TERMINAL:
            return last
        time.sleep(poll)
    try:
        request_json(client, "POST", f"{url}/api/jobs/{jid}/cancel")
    except (RunFailure, httpx.HTTPError):
        pass
    raise RunFailure(f"{stage} exceeded {timeout}s; last state: {last}")


def require_ready(job, stage):
    if job["state"] != "ready":
        raise RunFailure(f"{stage}: {job['state']}: {job.get('error') or job.get('missing_fonts')}")


def download(client, url, relative, target):
    target.parent.mkdir(parents=True, exist_ok=True)
    with client.stream("GET", f"{url}{relative}") as response:
        if response.is_error:
            raise RunFailure(
                f"GET {relative}: HTTP {response.status_code}: {response.read()[:3000]!r}"
            )
        with target.open("wb") as output:
            for chunk in response.iter_bytes():
                output.write(chunk)
    if target.stat().st_size == 0:
        raise RunFailure(f"Empty artifact: {relative}")
    return target


def verify_artifacts(client, url, job, directory):
    variants = job.get("variants") or []
    if {v["key"] for v in variants} != set(VARIANTS):
        raise RunFailure(f"Expected 3 variants, got {variants}")
    if job.get("quality_report", {}).get("errors") != 0:
        raise RunFailure(f"Quality gate reported errors: {job.get('quality_report')}")
    gid = job["id"]
    counts = {}
    for variant in variants:
        key = variant["key"]
        count = variant["slides"]
        counts[key] = count
        if not variant.get("rendering", {}).get("native_render"):
            raise RunFailure(f"{key}: native PPTX render was not verified")
        base = f"/api/jobs/{gid}/files/{key}"
        pptx = download(client, url, base + "/deck.pptx", directory / key / "deck.pptx")
        pdf = download(client, url, base + "/deck.pdf", directory / key / "deck.pdf")
        html = download(client, url, base + "/deck.html", directory / key / "deck.html")
        with ZipFile(pptx) as archive:
            if archive.testzip() is not None:
                raise RunFailure(f"{key}: corrupt PPTX ZIP")
        opened = Presentation(pptx)
        if len(opened.slides) != count:
            raise RunFailure(f"{key}: PPTX has {len(opened.slides)} slides, expected {count}")
        if not pdf.read_bytes().startswith(b"%PDF-"):
            raise RunFailure(f"{key}: invalid PDF header")
        if "<html" not in html.read_text(errors="replace").lower():
            raise RunFailure(f"{key}: HTML document missing")
        preview = download(
            client,
            url,
            f"/api/generations/{gid}/preview/{key}/1",
            directory / key / "preview-1.png",
        )
        with Image.open(preview) as image:
            image.verify()
    evidence = download(
        client, url, f"/api/jobs/{gid}/files/evidence.html", directory / "evidence.html"
    )
    if "<!doctype html" not in evidence.read_text(errors="replace").lower():
        raise RunFailure("Evidence report content missing")
    bundle = download(
        client, url, f"/api/jobs/{gid}/files/presentations.zip", directory / "presentations.zip"
    )
    with ZipFile(bundle) as archive:
        if archive.testzip() is not None:
            raise RunFailure("Corrupt presentations.zip")
        names = set(archive.namelist())
        if {f"{key}.pptx" for key in VARIANTS} - names or "evidence.html" not in names:
            raise RunFailure(f"ZIP missing PPTX/evidence: {sorted(names)}")
        for key in VARIANTS:
            if len(Presentation(BytesIO(archive.read(f"{key}.pptx"))).slides) != counts[key]:
                raise RunFailure(f"ZIP {key}.pptx slide count differs")
    return counts


def run(
    template,
    directory,
    *,
    url,
    timeout,
    poll,
    content=CONTENT,
    images=(),
    image_presentation="plain",
    input_mode="content",
    slides=5,
    allow_model_calls=False,
):
    summary = {
        "template": str(template),
        "source_sha256": sha256(template.read_bytes()).hexdigest(),
        "content_sha256": sha256(content.encode()).hexdigest(),
        "images": {str(path): sha256(path.read_bytes()).hexdigest() for path in images},
        "image_presentation": image_presentation,
        "input_mode": input_mode,
        "requested_slides": slides,
        "steps": {},
    }
    save_json(directory / "summary.json", summary)
    ids = []
    try:
        with httpx.Client(timeout=60, trust_env=False) as client:
            health = request_json(client, "GET", url + "/api/health")
            summary["health"] = health
            if health["model_mode"] != "extractive" and not allow_model_calls:
                raise RunFailure(
                    "Server uses model calls; pass --allow-model-calls after budget approval"
                )
            if health["engine"] != "native" and not allow_model_calls:
                raise RunFailure("Isolated runner requires native engine")
            for number in (1, 2):
                with template.open("rb") as source:
                    uploaded = request_json(
                        client,
                        "POST",
                        url + "/api/templates/analyze",
                        files={"template": (template.name, source, "application/octet-stream")},
                    )
                jid = uploaded["id"]
                ids.append(jid)
                stage = f"template-{number}"
                analyzed = wait_job(client, url, jid, stage, directory, timeout, poll)
                summary["steps"][stage] = {
                    "id": jid,
                    "state": analyzed["state"],
                    "cache": analyzed.get("template_cache"),
                }
                require_ready(analyzed, stage)
            cache_saved = summary["steps"]["template-1"]["cache"].get("saved") is True
            summary["cache_expected"] = cache_saved
            if cache_saved and summary["steps"]["template-2"]["cache"].get("hit") is not True:
                raise RunFailure("Repeated template analysis missed the cache")
            with ExitStack() as stack:
                files = [
                    (
                        "images",
                        (
                            path.name,
                            stack.enter_context(path.open("rb")),
                            mimetypes.guess_type(path.name)[0] or "application/octet-stream",
                        ),
                    )
                    for path in images
                ]
                prepared = request_json(
                    client,
                    "POST",
                    url + "/api/prepare",
                    data={
                        "text": content,
                        "slides": str(slides),
                        "size_preset": "mini",
                        "template_job_id": ids[-1],
                        "input_mode": input_mode,
                        "image_presentation": image_presentation,
                    },
                    files=files or None,
                )
            pid = prepared["id"]
            ids.append(pid)
            ready = wait_job(client, url, pid, "preparation", directory, timeout, poll)
            summary["steps"]["preparation"] = {
                "id": pid,
                "state": ready["state"],
                "cache": ready.get("analysis", {}).get("template_cache"),
            }
            require_ready(ready, "preparation")
            if ready.get("content", {}).get("images") != len(images):
                raise RunFailure("Preparation did not retain every uploaded image")
            if health["model_mode"] == "extractive" and ready.get("analysis", {}).get(
                "model_calls"
            ):
                raise RunFailure("Unexpected model calls during extractive preparation")
            if cache_saved and summary["steps"]["preparation"]["cache"].get("hit") is not True:
                raise RunFailure("Preparation missed the template cache")
            budget = ready.get("control", {}).get("slide_budget") or {}
            accept = budget.get("status") == "adjusted" or ready.get("constraints", {}).get(
                "confirm_plan", False
            )
            summary["explicit_slide_count_acceptance"] = bool(accept)
            if input_mode == "brief":
                blocked = client.post(
                    url + "/api/generate",
                    json={"package_id": pid, "accept_adjusted_slide_count": bool(accept)},
                )
                summary["preapproval_generation"] = {
                    "status": blocked.status_code,
                    "body": blocked.text[:2000],
                }
                if blocked.status_code != 409:
                    raise RunFailure(
                        f"Brief generation before approval returned HTTP {blocked.status_code}"
                    )
                draft = request_json(client, "GET", f"{url}/api/packages/{pid}/draft")
                save_json(directory / "draft.json", draft)
                slides = draft.get("draft", {}).get("slides") or []
                if not slides or not all(
                    slide.get("title")
                    and slide.get("bullets")
                    and all(bullet.get("text") for bullet in slide["bullets"])
                    for slide in slides
                ):
                    raise RunFailure("Brief draft is empty or incomplete")
                if draft.get("package_hash") != ready.get("package_hash") or draft.get(
                    "draft_hash"
                ) != ready.get("draft_hash"):
                    raise RunFailure("Draft hashes differ from sealed preparation")
                if budget.get("status") == "needs_input":
                    raise RunFailure(f"Preparation requires input: {budget.get('message')}")
                approved = request_json(
                    client,
                    "POST",
                    f"{url}/api/packages/{pid}/approve",
                    json={"package_hash": draft["package_hash"], "draft_hash": draft["draft_hash"]},
                )
                save_json(directory / "approval.json", approved)
                if (
                    approved.get("approved_package_hash") != draft["package_hash"]
                    or approved.get("approved_draft_hash") != draft["draft_hash"]
                ):
                    raise RunFailure("Approval did not bind the exact draft and package hashes")
                summary["draft_slides"] = len(slides)
            elif budget.get("status") == "needs_input":
                raise RunFailure(f"Preparation requires input: {budget.get('message')}")
            request_json(client, "POST", f"{url}/api/packages/{pid}/auto-generation/cancel")
            created = request_json(
                client,
                "POST",
                url + "/api/generate",
                json={"package_id": pid, "accept_adjusted_slide_count": bool(accept)},
            )
            gid = created["id"]
            ids.append(gid)
            generated = wait_job(client, url, gid, "generation", directory, timeout, poll)
            summary["steps"]["generation"] = {"id": gid, "state": generated["state"]}
            if generated["state"] not in ("needs_review", "completed"):
                raise RunFailure(f"Generation: {generated['state']}: {generated.get('error')}")
            summary["slide_counts"] = verify_artifacts(client, url, generated, directory)
            history = request_json(client, "GET", url + "/api/jobs")
            save_json(directory / "history.json", history)
            if not {pid, gid} <= {job["id"] for job in history}:
                raise RunFailure("History omitted preparation or generation")
            summary["status"] = "passed"
    except (RunFailure, httpx.HTTPError, OSError, ValueError) as exc:
        summary["status"] = "failed"
        summary["failure"] = {"type": type(exc).__name__, "message": str(exc)}
    finally:
        try:
            with httpx.Client(timeout=15, trust_env=False) as client:
                for jid in ids:
                    save_json(
                        directory / f"{jid}-diagnostics.json",
                        request_json(client, "GET", f"{url}/api/jobs/{jid}/diagnostics"),
                    )
                save_json(
                    directory / "history.json", request_json(client, "GET", url + "/api/jobs")
                )
        except (RunFailure, httpx.HTTPError, OSError) as exc:
            summary["diagnostics_error"] = str(exc)
        summary["source_unchanged"] = (
            sha256(template.read_bytes()).hexdigest() == summary["source_sha256"]
        )
        summary["images_unchanged"] = all(
            sha256(Path(path).read_bytes()).hexdigest() == expected
            for path, expected in summary["images"].items()
        )
        if not summary["source_unchanged"]:
            summary["status"] = "failed"
            summary["failure"] = {
                "type": "SourceChanged",
                "message": "Input template bytes changed",
            }
        if not summary["images_unchanged"]:
            summary["status"] = "failed"
            summary["failure"] = {"type": "ImageChanged", "message": "Input image bytes changed"}
        save_json(directory / "summary.json", summary)
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", type=Path, help="One PPTX/POTX; run sequentially for a corpus")
    parser.add_argument(
        "--output-root", type=Path, default=ROOT / "test-results/external-template-runs"
    )
    parser.add_argument(
        "--url", help="Existing HTTP server; model mode requires --allow-model-calls"
    )
    parser.add_argument("--allow-model-calls", action="store_true")
    parser.add_argument(
        "--content", type=Path, help="UTF-8 text/Markdown; default is built-in smoke content"
    )
    parser.add_argument(
        "--image", type=Path, action="append", default=[], help="PNG/JPEG upload; repeatable"
    )
    parser.add_argument("--image-presentation", choices=("plain", "device"), default="plain")
    parser.add_argument("--input-mode", choices=("content", "brief"), default="content")
    parser.add_argument("--slides", type=int, default=5, help="Requested slide count (default: 5)")
    parser.add_argument("--timeout", type=int, default=300, help="Seconds per asynchronous job")
    parser.add_argument("--poll", type=float, default=1)
    args = parser.parse_args()
    template = args.template.resolve(strict=True)
    if template.suffix.lower() not in (".pptx", ".potx") or not template.is_file():
        parser.error("Expected one existing PPTX/POTX file")
    if args.allow_model_calls and not args.url:
        parser.error("Model calls require an explicitly supplied --url")
    if args.timeout <= 0 or args.poll <= 0 or not 1 <= args.slides <= 30:
        parser.error("--timeout/--poll must be positive and --slides must be 1..30")
    images = [path.resolve(strict=True) for path in args.image]
    if images and any(not image.is_file() for image in images):
        parser.error("--image must name existing files")
    if args.image_presentation == "device" and not images:
        parser.error("--image-presentation device requires --image")
    content = args.content.read_text(encoding="utf-8") if args.content else CONTENT
    if not content.strip():
        parser.error("--content must be nonempty UTF-8 text")
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    directory = args.output_root / (
        f"{template.stem[:35]}-{args.input_mode}{args.slides}-{sha256(content.encode()).hexdigest()[:8]}-"
        f"{sha256(template.read_bytes()).hexdigest()[:10]}-{stamp}"
    )
    directory.mkdir(parents=True, exist_ok=False)
    context = (
        nullcontext((args.url.rstrip("/"), None, None))
        if args.url
        else application(directory / "server")
    )
    with context as (url, _, _):
        summary = run(
            template,
            directory,
            url=url,
            timeout=args.timeout,
            poll=args.poll,
            content=content,
            images=images,
            image_presentation=args.image_presentation,
            input_mode=args.input_mode,
            slides=args.slides,
            allow_model_calls=args.allow_model_calls,
        )
    print(json.dumps({"directory": str(directory), **summary}, ensure_ascii=False, indent=2))
    raise SystemExit(0 if summary["status"] == "passed" else 1)


if __name__ == "__main__":
    main()
