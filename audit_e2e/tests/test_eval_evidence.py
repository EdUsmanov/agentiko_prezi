import json
from pathlib import Path
from hashlib import sha256

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter

from studio.composition.office import to_pdf

from audit_e2e.controls import _FIGURE_CONTENT, build_controls
from audit_e2e.corpus import load_cases, materialize_case
from audit_e2e.evidence import audit_bundle, build_bundle


def _slide_count_bundle(directory, pptx_count, *, pdf_count=None, duplicate_pair=False):
    pdf_count = pptx_count if pdf_count is None else pdf_count
    variant_ids = ["executive", "analytical", "story"]
    variants = {}
    colors = {"executive": "173550", "analytical": "A43E35", "story": "285078"}
    if duplicate_pair:
        colors["story"] = colors["executive"]
    for variant in variant_ids:
        variant_dir = directory / "variants" / variant
        variant_dir.mkdir(parents=True)
        presentation = Presentation()
        for number in range(1, pptx_count + 1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(colors[variant])
            textbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(0.5))
            textbox.text_frame.text = f"Slide {number}"
        pptx_path = variant_dir / "deck.pptx"
        presentation.save(pptx_path)

        pdf_writer = PdfWriter()
        for _ in range(pdf_count):
            pdf_writer.add_blank_page(width=960, height=540)
        pdf_path = variant_dir / "deck.pdf"
        with pdf_path.open("wb") as stream:
            pdf_writer.write(stream)

        slides = []
        for number in range(1, pptx_count + 1):
            image_path = variant_dir / f"slide-{number}.png"
            Image.new("RGB", (16, 9), f"#{colors[variant]}").save(image_path)
            slides.append({"number": number, "image": image_path.relative_to(directory).as_posix()})
        variants[variant] = {
            "pptx": pptx_path.relative_to(directory).as_posix(),
            "pdf": pdf_path.relative_to(directory).as_posix(),
            "pptx_slide_count": pptx_count,
            "pdf_page_count": pdf_count,
            "slides": slides,
        }

    return {
        "schema_version": 1,
        "case_id": "browser-mini-range",
        "expected_variant_ids": variant_ids,
        "expected_slide_count": 5,
        "slide_contract": {
            "target": 5,
            "minimum": 3,
            "maximum": 5,
            "request_kind": "browser_preset",
            "preset": "mini",
        },
        "reference": {"status": "silver", "points": [], "requirements": []},
        "variants": variants,
        "template": {},
    }


def _visibility_case(tmp_path, *, grouped=False):
    content = "The team checks 40 requests."
    digest = sha256(content.encode()).hexdigest()
    result_dir = tmp_path / "result"
    variant_dir = result_dir / "executive"
    variant_dir.mkdir(parents=True)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if grouped:
        group = slide.shapes.add_group_shape()
        text_shape = group.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(0.8))
        text_shape.text_frame.text = content
        group.left, group.top = Inches(14), Inches(1)
    else:
        table_shape = slide.shapes.add_table(2, 2, Inches(14), Inches(1), Inches(4), Inches(1.5))
        for row_index, row in enumerate((("Team", "Requests"), ("North", "40"))):
            for column_index, value in enumerate(row):
                table_shape.table.cell(row_index, column_index).text = value
    pptx_path = variant_dir / "deck.pptx"
    presentation.save(pptx_path)
    assert to_pdf(pptx_path, variant_dir, timeout=30)

    requirements = (
        []
        if grouped
        else [
            {
                "id": "source-table",
                "kind": "table",
                "status": "gold",
                "headers": ["Team", "Requests"],
                "rows": [["North", "40"]],
            }
        ]
    )
    case = {
        "id": "grouped-text" if grouped else "off-canvas-table",
        "slides": 1,
        "variants": ["executive"],
        "content": content,
        "content_source": {"sha256": digest},
        "source_hashes": {"content": digest},
        "template": pptx_path,
        "template_source": {"source": "synthetic fixture", "format": "PPTX"},
        "template_preview_slides": [],
        "template_origin": "synthetic_analog",
        "template_fidelity_applicability": "not_applicable_synthetic_template",
        "synthetic": True,
        "reference": {
            "version": "visibility-test-1",
            "status": "gold",
            "provenance": {"review_status": "approved"},
            "points": [
                {
                    "id": "sample-size",
                    "statement": content,
                    "quote": content,
                    "required": True,
                    "origin": "user_text",
                    "numbers": [{"value": "40", "unit": "requests"}],
                }
            ],
            "requirements": requirements,
        },
    }
    bundle_dir = tmp_path / "bundle"
    bundle = build_bundle(case, result_dir, bundle_dir)
    return bundle, bundle_dir


def test_registry_is_versioned_source_anchored_and_replay_image_name_is_stable(tmp_path):
    core, extended = load_cases("core"), load_cases("extended")
    registry = json.loads(Path("audit_e2e/fixtures/cases.json").read_text(encoding="utf-8"))
    assert registry["schema_version"] == 2
    assert registry["reference_version"] == "2026-09-29.1"
    assert len(core) == 4
    assert len(extended) == 9
    assert all(case["reference"]["status"] == "silver" for case in extended)
    assert all(case["reference"]["provenance"]["review_status"] == "pending" for case in extended)
    assert all(
        point["quote"] in case["content"] for case in core for point in case["reference"]["points"]
    )
    browser = {case["id"]: case for case in core if case["interface"] == "browser"}
    assert set(browser) == {"binghamton-brief", "astate-dark-mixed"}
    assert all(
        case["slide_contract"]
        == {
            "target": 5,
            "minimum": 3,
            "maximum": 5,
            "request_kind": "browser_preset",
            "preset": "mini",
        }
        for case in browser.values()
    )
    assert all(
        case["slide_contract"]["minimum"] == case["slide_contract"]["maximum"] == case["slides"]
        for case in core
        if case["interface"] == "http"
    )

    mixed = next(case for case in core if case["id"] == "astate-dark-mixed")
    replay = materialize_case(mixed, tmp_path / "input", synthetic=True)
    assert replay["template_origin"] == "synthetic_analog"
    assert replay["images"][0].name == "dashboard-screenshot.png"
    assert "dashboard-screenshot.png" in replay["content"]
    assert replay["source_hashes"]["images"]["dashboard-screenshot.png"]

    binghamton = next(case for case in core if case["id"] == "binghamton-content")
    replay = materialize_case(binghamton, tmp_path / "binghamton", synthetic=True)
    fidelity = next(
        item for item in replay["reference"]["requirements"] if item["id"] == "template-fidelity"
    )
    assert fidelity["applicability"] == "not_applicable_synthetic_template"
    assert (
        "team decides"
        in next(
            p["statement"] for p in replay["reference"]["points"] if p["id"] == "decision"
        ).lower()
    )


def test_browser_mini_range_accepts_three_slides_and_reports_target_deviation(tmp_path):
    bundle_dir = tmp_path / "three-slides"
    bundle = _slide_count_bundle(bundle_dir, 3)

    audit = audit_bundle(bundle, bundle_dir)

    assert audit["status"] == "passed"
    assert audit["findings"] == []
    assert len(audit["diagnostics"]) == 3
    assert all(
        item["evidence"]
        == {
            "variant": variant,
            "target": 5,
            "actual": 3,
            "allowed_range": [3, 5],
            "preset": "mini",
        }
        for item, variant in zip(audit["diagnostics"], bundle["expected_variant_ids"], strict=True)
    )


@pytest.mark.parametrize("slide_count", [2, 6])
def test_browser_mini_range_rejects_counts_outside_preset(tmp_path, slide_count):
    bundle_dir = tmp_path / f"{slide_count}-slides"
    bundle = _slide_count_bundle(bundle_dir, slide_count)

    audit = audit_bundle(bundle, bundle_dir)

    assert audit["status"] == "failed"
    count_findings = [
        finding
        for finding in audit["findings"]
        if finding["message"] == "Exported slide count is outside the case contract"
    ]
    assert len(count_findings) == 3
    assert all(
        finding["evidence"]["allowed_range"] == [3, 5]
        and finding["evidence"]["actual"] == slide_count
        for finding in count_findings
    )


def test_browser_mini_range_keeps_export_and_variant_failures_hard(tmp_path):
    mismatched_dir = tmp_path / "mismatched"
    mismatch = audit_bundle(_slide_count_bundle(mismatched_dir, 3, pdf_count=2), mismatched_dir)
    assert mismatch["status"] == "failed"
    assert any("PPTX/PDF page counts differ" in item["message"] for item in mismatch["findings"])

    duplicate_dir = tmp_path / "duplicate"
    duplicate = audit_bundle(
        _slide_count_bundle(duplicate_dir, 3, duplicate_pair=True), duplicate_dir
    )
    assert duplicate["status"] == "failed"
    assert any(item["category"] == "duplicates" for item in duplicate["findings"])


def test_off_canvas_table_cannot_satisfy_visible_source_table_requirement(tmp_path):
    bundle, bundle_dir = _visibility_case(tmp_path)

    audit = audit_bundle(bundle, bundle_dir)

    assert audit["status"] == "failed"
    assert any(
        item["category"] == "hidden_text" and "source table" in item["message"].lower()
        for item in audit["findings"]
    )


def test_group_relative_text_is_inconclusive_until_slide_transform_is_known(tmp_path):
    bundle, bundle_dir = _visibility_case(tmp_path, grouped=True)

    audit = audit_bundle(bundle, bundle_dir)

    assert bundle["variants"]["executive"]["slides"][0]["text"] == ""
    assert audit["status"] == "inconclusive"
    assert any(
        item["category"] == "hidden_text" and item["severity"] == "inconclusive"
        for item in audit["findings"]
    )


def test_materialize_rejects_changed_source_and_template_before_generation(tmp_path):
    case = next(case for case in load_cases() if case["id"] == "binghamton-content")
    altered = dict(case, content=case["content"] + "altered")
    with pytest.raises(ValueError, match="content hash mismatch"):
        materialize_case(altered, tmp_path / "changed", synthetic=True)

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / case["template"]["file"]).write_bytes(b"wrong template")
    with pytest.raises(ValueError, match="template hash mismatch"):
        materialize_case(case, tmp_path / "real", corpus_root=corpus)


def test_calibration_controls_are_blind_native_exports_and_positive_controls_pass(tmp_path):
    declared_figure_text = " ".join(_FIGURE_CONTENT.values()).casefold()
    assert _FIGURE_CONTENT == {
        "title": "Pilot sample",
        "value": "73",
        "unit": "REQUESTS",
        "label": "Checked by the team",
    }
    assert not any(
        term in declared_figure_text for term in ("week", "duration", "if ", "conditional", "holds")
    )

    controls = build_controls(tmp_path / "controls")
    assert len(controls) == 13
    assert {row["category"] for row in controls} == {
        "clean",
        "paraphrase",
        "omission",
        "number",
        "negation",
        "condition",
        "table_binding",
        "units",
        "provenance",
        "hidden_text",
        "duplicates",
        "template_fidelity",
        "readability",
    }
    assert all(
        row["expected"] == ("passed" if row["category"] in {"clean", "paraphrase"} else "failed")
        for row in controls
    )

    bundles = {}
    audits = {}
    for row in controls:
        bundle_path = Path(row["bundle"])
        bundle = json.loads(bundle_path.read_text(encoding="utf-8"))
        bundles[row["category"]] = bundle
        audits[row["category"]] = audit_bundle(bundle, bundle_path.parent)
        assert bundle["case_id"] == row["id"]
        assert "category" not in bundle and "expected" not in bundle
        assert set(bundle["variants"]) == {"executive", "analytical", "story"}
        assert bundle["expected_slide_count"] == 2
        for deck in bundle["variants"].values():
            pptx = bundle_path.parent / deck["pptx"]
            pdf = bundle_path.parent / deck["pdf"]
            assert len(Presentation(pptx).slides) == 2
            assert len(PdfReader(pdf).pages) == 2
            assert all((bundle_path.parent / slide["image"]).is_file() for slide in deck["slides"])
        assert bundle["template"]["previews"]

    assert bundles["clean"]["reference"]["version"] == "control-2"
    clean_source = bundles["clean"]["source"]["text"]
    assert "| Team | Requests | Closed | Hours |" in clean_source
    assert "| North | 40 | 34 | 18 |" in clean_source
    assert "| South | 28 | 24 | 21 |" in clean_source
    assert audits["clean"]["status"] == "passed"
    assert audits["paraphrase"]["status"] == "passed"
    for category in (
        "number",
        "table_binding",
        "units",
        "hidden_text",
        "duplicates",
        "readability",
    ):
        assert category in {finding["category"] for finding in audits[category]["findings"]}
    clean_images = [
        obj
        for slide in bundles["clean"]["variants"]["executive"]["slides"]
        for obj in slide["objects"]
        if obj.get("image")
    ]
    assert clean_images
    assert clean_images[0]["geometry_inches"]["width"] > 5
    hidden_slide = bundles["hidden_text"]["variants"]["executive"]["slides"][1]
    hidden_phrase = "A privacy review is mandatory before launch."
    assert hidden_phrase not in hidden_slide["pdf_visible_text"]
    assert hidden_phrase in hidden_slide["notes"]
    assert any(
        obj.get("text") == hidden_phrase and obj.get("out_of_bounds")
        for obj in hidden_slide["objects"]
    )
    fidelity_images = [
        obj
        for slide in bundles["template_fidelity"]["variants"]["executive"]["slides"]
        for obj in slide["objects"]
        if obj.get("image")
    ]
    assert not fidelity_images
    readability_images = [
        obj
        for slide in bundles["readability"]["variants"]["executive"]["slides"]
        for obj in slide["objects"]
        if obj.get("image")
    ]
    assert (
        clean_images[0]["image"]["pixels_sha256"] == readability_images[0]["image"]["pixels_sha256"]
    )
    assert clean_images[0]["geometry_inches"]["x"] == readability_images[0]["geometry_inches"]["x"]
    assert clean_images[0]["geometry_inches"]["y"] == readability_images[0]["geometry_inches"]["y"]
    assert readability_images[0]["geometry_inches"]["width"] < 4

    control_root = tmp_path / "controls"
    result_dir = control_root / "control-01" / "result"
    template_path = control_root / "reference-template" / "source-template.pptx"
    brief = next(case for case in load_cases() if case["id"] == "binghamton-brief")
    brief["interface"], brief["slides"] = "http", 2
    brief["slide_contract"] = {
        "target": 2,
        "minimum": 2,
        "maximum": 2,
        "request_kind": "explicit_count",
    }
    brief["template_source"] = dict(brief["template"])
    brief["template_source"].update(
        source="synthetic control template", sha256=sha256(template_path.read_bytes()).hexdigest()
    )
    brief["template"] = template_path
    brief["registry_hashes"] = {"template": sha256(template_path.read_bytes()).hexdigest()}
    brief["source_hashes"] = {"content": brief["content_source"]["sha256"], "images": {}}
    brief["images_source"], brief["images"] = [], []
    brief["synthetic"], brief["template_origin"] = True, "synthetic_analog"
    brief["template_preview_slides"] = [1]
    (result_dir / "draft.json").write_text(
        json.dumps(
            {
                "package_hash": "package-hash",
                "draft_hash": "draft-hash",
                "approved": False,
                "draft": {
                    "slides": [
                        {
                            "bullets": [
                                {
                                    "text": "Add predictive routing as a possible next step.",
                                    "fact_ids": ["proposal-1"],
                                    "proposed": True,
                                },
                                {
                                    "text": "Pilot duration is twelve weeks.",
                                    "fact_ids": ["source-1"],
                                    "proposed": False,
                                },
                            ]
                        }
                    ]
                },
            }
        )
    )
    (result_dir / "approval.json").write_text(
        json.dumps(
            {
                "approved_package_hash": "package-hash",
                "approved_draft_hash": "draft-hash",
            }
        )
    )
    (result_dir / "preapproval-generation.json").write_text(json.dumps({"status": 409}))
    brief_bundle = build_bundle(brief, result_dir, control_root / "brief-bundle")
    assert brief_bundle["slide_contract"] == brief["slide_contract"]
    assert brief_bundle["slide_contract"] == brief["slide_contract"]
    assert brief_bundle["source"]["approved_proposals"] == [
        {
            "slide_number": 1,
            "text": "Add predictive routing as a possible next step.",
            "origin": "model_proposal",
            "fact_ids": ["proposal-1"],
        }
    ]
    assert brief_bundle["source"]["brief_approval"]["preapproval_generation_blocked"] is True
    assert brief_bundle["source"]["brief_approval"]["approved_draft_hash"] == "draft-hash"
    assert "quality_report" not in brief_bundle and "verdict" not in brief_bundle
