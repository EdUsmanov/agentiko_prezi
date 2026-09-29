"""Hard limits and cleanup contracts for isolated presentation evaluations."""

import asyncio
from collections import deque
from dataclasses import replace
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
import threading

import httpx
import pytest

from studio.config import Settings
from studio.providers.gateway import ModelGateway
from test_support.app_server import application
from audit_e2e.live_provider import (
    LiveProvider,
    LoopbackProxyTransport,
    evaluation_stage,
)
from audit_e2e.runtime import EvaluationDeadline, _hard_deadline, execute_case


class LocalProvider:
    def __init__(self, statuses):
        self.statuses = deque(statuses)
        self.requests = []
        owner = self

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                size = int(self.headers.get("content-length", "0"))
                body = self.rfile.read(size)
                owner.requests.append((json.loads(body), self.headers.get("authorization")))
                status = owner.statuses.popleft() if owner.statuses else 200
                if status == 200:
                    payload = {
                        "choices": [
                            {
                                "message": {"role": "assistant", "content": '{"ok":true}'},
                                "finish_reason": "stop",
                            }
                        ],
                        "usage": {"prompt_tokens": 1, "completion_tokens": 1},
                    }
                else:
                    payload = {"error": {"code": "temporary_failure"}}
                raw = json.dumps(payload).encode()
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)

            def log_message(self, *_args):
                pass

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = f"http://127.0.0.1:{self.server.server_port}/v1"

    def close(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)


def _api_settings(tmp_path, base_url, **updates):
    settings = Settings(
        data_dir=tmp_path / "data",
        mode="api",
        execution_kind="live",
        base_url=base_url,
        model_id="meta-llama/test-model",
        api_key="test-secret-never-report",
        parameters_b=27,
        open_weights=True,
        license="Apache-2.0",
        engine="native",
    )
    return replace(settings, **updates)


def test_live_proxy_counts_failed_physical_attempts_and_blocks_after_cap(tmp_path):
    upstream = LocalProvider([503, 200])
    settings = _api_settings(tmp_path, upstream.url)
    provider = LiveProvider(settings, max_requests=1, timeout=2)
    try:
        with httpx.Client(trust_env=False, timeout=3) as client:
            first = client.post(
                provider.proxy_url,
                headers={"Authorization": "Bearer test-secret-never-report"},
                json={"attempt": 1},
            )
            second = client.post(provider.proxy_url, json={"attempt": 2})
        assert first.status_code == 503
        assert second.status_code == 403
        assert len(upstream.requests) == 1
        report = provider.summary()
        assert report["model_requests"] == 1
        assert report["provider_failures"] == 1
        assert report["budget_blocked"] == 1
        assert "test-secret-never-report" not in json.dumps(report)
    finally:
        provider.shutdown()
        upstream.close()


def test_original_openrouter_branch_and_retry_records_survive_loopback_proxy(tmp_path):
    upstream = LocalProvider([503, 200])
    settings = _api_settings(
        tmp_path,
        "https://openrouter.ai/api/v1",
        thinking=True,
        thinking_token_budget=512,
        openrouter_providers=("alibaba", "deepinfra"),
        openrouter_allow_fallbacks=True,
    )
    provider = LiveProvider(
        settings,
        max_requests=2,
        timeout=5,
        request_dir=tmp_path / "provider-requests",
    )
    provider.upstream_url = upstream.url + "/chat/completions"
    gateway = ModelGateway(settings, transport=LoopbackProxyTransport(provider.proxy_url))

    async def invoke():
        try:
            with evaluation_stage("editorial"):
                return await gateway._json_request(
                    "editorial", {"slide_range": [1, 3]}, timeout=5, schema=None
                )
        finally:
            await gateway.aclose()

    try:
        assert asyncio.run(invoke()) == {"ok": True}
        body, auth = upstream.requests[-1]
        assert auth == "Bearer test-secret-never-report"
        assert body["provider"] == {
            "only": ["alibaba", "deepinfra"],
            "allow_fallbacks": True,
            "require_parameters": True,
        }
        assert body["reasoning"] == {"enabled": True, "max_tokens": 512}
        assert "chat_template_kwargs" not in body
        report = provider.summary()
        assert report["model_requests"] == 2
        assert report["provider_failures"] == 1
        assert report["provider_http_statuses"] == {"200": 1, "503": 1}
        assert [row["sequence"] for row in report["provider_requests"]] == [1, 2]
        assert all(row["stage"] == "editorial" for row in report["provider_requests"])
        assert report["generator_usage"] == {
            "input_tokens": 1,
            "output_tokens": 1,
            "total_tokens": None,
            "cached_input_tokens": None,
            "reasoning_tokens": None,
            "reported_requests": 1,
        }
        assert all(
            (tmp_path / row["request_file"]).is_file() for row in report["provider_requests"]
        )
        for row in report["provider_requests"]:
            raw = (tmp_path / row["request_file"]).read_text()
            assert "test-secret-never-report" not in raw
            assert "authorization" not in raw.casefold()
            assert '"request"' in raw and '"response"' in raw
        assert "test-secret-never-report" not in json.dumps(report)
    finally:
        provider.shutdown()
        upstream.close()


def test_live_application_settings_are_private_removed_and_server_stops(tmp_path):
    upstream = LocalProvider([200])
    settings = _api_settings(tmp_path, upstream.url)
    provider = LiveProvider(settings, max_requests=1, timeout=2)
    app_dir = tmp_path / "isolated"
    config = app_dir / "settings.json"
    try:
        with application(app_dir, settings=settings, live_provider=provider) as (url, _, _):
            assert config.exists()
            assert config.stat().st_mode & 0o777 == 0o600
            health = httpx.get(url + "/api/health", trust_env=False, timeout=2).json()
            assert health["model_mode"] == "api"
            assert "test-secret-never-report" not in json.dumps(health)
        assert not config.exists()
        with pytest.raises(httpx.ConnectError):
            httpx.get(url + "/api/health", trust_env=False, timeout=1)
    finally:
        provider.shutdown()
        upstream.close()


def test_setup_failure_removes_temporary_live_settings(tmp_path, monkeypatch):
    upstream = LocalProvider([])
    settings = _api_settings(tmp_path, upstream.url)
    provider = LiveProvider(settings, max_requests=1, timeout=2)
    app_dir = tmp_path / "isolated"
    config = app_dir / "settings.json"

    def fail_port_allocation():
        raise RuntimeError("simulated port allocation failure")

    monkeypatch.setattr("test_support.app_server.free_port", fail_port_allocation)
    try:
        with pytest.raises(RuntimeError, match="simulated port allocation failure"):
            with application(app_dir, settings=settings, live_provider=provider):
                pytest.fail("application must fail before starting the server")
        assert not config.exists()
    finally:
        provider.shutdown()
        upstream.close()


def test_deadline_interrupts_work_and_missing_replay_never_falls_back(tmp_path):
    import time

    started = time.monotonic()
    with pytest.raises(EvaluationDeadline):
        with _hard_deadline(0.05):
            time.sleep(1)
    assert time.monotonic() - started < 0.5

    template = tmp_path / "input.pptx"
    from pptx import Presentation

    Presentation().save(template)
    result = execute_case(
        {
            "id": "missing-replay",
            "input_mode": "content",
            "content": "A finite replay is required.",
            "template": template,
            "images": [],
        },
        tmp_path / "result",
        mode="replay",
        timeout=2,
    )
    assert result["model_requests"] == 0
    assert result["status"] == "inconclusive"
    assert not (tmp_path / "result" / "server" / "settings.json").exists()


def test_offline_author_preserves_long_native_table_without_oversized_caption():
    from audit_e2e.build_cassettes import authored_response
    from studio.contents.editorial_domain import validate_plan
    from studio.contents.parsing import parse_content

    rows = "\n".join(f"| Room {i} | {i + 20} |" for i in range(40))
    content = parse_content(
        "# Room inventory\n## Overview\nThe inventory records available seats.\n"
        "## Seat counts\n| Room | Seats |\n|---|---|\n" + rows
    )
    response = authored_response(
        "editorial", {"source": content.model_dump(mode="json"), "slide_range": [2, 2]}
    )
    plan = validate_plan(response, content, (2, 2), require_cover=True)
    assert plan["slides"][1]["source_table_id"] == content.tables[0].id
    assert len(content.tables[0].rows) == 40
    assert plan["slides"][1]["bullets"][0]["text"] == "Seat counts"


def test_ineffective_editorial_repair_requires_explicit_failure_capture():
    from audit_e2e.build_cassettes import authored_response
    from studio.contents.editorial_repair import apply_replacements

    slide = {
        "title": "Source fact",
        "purpose": "content",
        "bullets": [{"text": "The source records 12 days.", "evidence": [{"fact_id": "f1"}]}],
    }
    data = {"previous_plan": {"slides": [slide], "omitted": []}, "allowed_slide_indices": [1]}
    with pytest.raises(ValueError, match="Unreviewed synthetic stage"):
        authored_response("editorial_repair", data)
    patch = authored_response("editorial_repair", data, capture_failure=True)
    applied = apply_replacements(data["previous_plan"], patch, [1])
    assert applied["slides"][0]["title"] == slide["title"]
    assert applied["slides"][0]["bullets"][0]["text"] == slide["bullets"][0]["text"]
