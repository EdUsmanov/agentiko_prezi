"""Small native-deck controls for checking evaluator sensitivity and specificity."""

from hashlib import sha256
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from studio.composition.office import to_pdf

from .evidence import build_bundle


def _digest(data):
    return sha256(data).hexdigest()


_ROWS = [["North", "40", "34", "18"], ["South", "28", "24", "21"]]
_HEADERS = ["Team", "Requests", "Closed", "Hours"]
_SOURCE = "\n".join(
    [
        "Pilot duration is 12 weeks. The team checks 73 requests.",
        "Source table:",
        "| " + " | ".join(_HEADERS) + " |",
        "| " + " | ".join("---" for _ in _HEADERS) + " |",
        *["| " + " | ".join(row) + " |" for row in _ROWS],
        "Changed team membership means differences do not prove service impact.",
        "After reviewing closure share, the team chooses whether to continue for 8 more weeks if service quality holds.",
        "Model proposal: predictive routing may reduce triage time.",
        "A privacy review is mandatory before launch.",
    ]
)
_FIGURE_CONTENT = {
    "title": "Pilot sample",
    "value": "73",
    "unit": "REQUESTS",
    "label": "Checked by the team",
}
_SPEC = (
    ("clean", "passed"),
    ("paraphrase", "passed"),
    ("omission", "failed"),
    ("number", "failed"),
    ("negation", "failed"),
    ("condition", "failed"),
    ("table_binding", "failed"),
    ("units", "failed"),
    ("provenance", "failed"),
    ("hidden_text", "failed"),
    ("duplicates", "failed"),
    ("template_fidelity", "failed"),
    ("readability", "failed"),
)


def _body(category):
    if category == "paraphrase":
        return (
            "A twelve-week window is set aside for the pilot. Reviewers examine 73 requests. "
            "Since team membership changed, the time gap alone cannot establish service impact. "
            "After looking at closure share, the team may extend the trial by 8 weeks if quality remains stable. "
            "Predictive routing is only a model-suggested follow-up. "
            "A privacy review must happen before launch."
        )
    lines = [
        "Pilot duration is 12 weeks.",
        "The team checks 73 requests.",
        "Changed team membership means differences do not prove service impact.",
        "After reviewing closure share, the team chooses whether to continue for 8 more weeks if service quality holds.",
        "Model proposal: predictive routing may reduce triage time.",
        "A privacy review is mandatory before launch.",
    ]
    if category == "omission":
        lines[5] = "Privacy safeguards are planned."
    elif category == "number":
        lines[0] = "Pilot duration is 10 weeks."
    elif category == "negation":
        lines[2] = "Changed team membership means differences prove service impact."
    elif category == "condition":
        lines[3] = "After reviewing closure share, the team continues for 8 more weeks."
    elif category == "provenance":
        lines[4] = "Predictive routing will reduce triage time."
    elif category == "hidden_text":
        lines[5] = "Privacy safeguards are planned."
    return " ".join(lines)


def _table_rows(category):
    headers, rows = _HEADERS[:], [row[:] for row in _ROWS]
    if category == "table_binding":
        rows = [["North", "28", "24", "21"], ["South", "40", "34", "18"]]
    elif category == "units":
        headers[-1] = "Minutes"
    return headers, rows


def _style(slide, category, variant):
    if category == "template_fidelity":
        background, foreground, accent = "E90074", "FFFFFF", "00E5FF"
    else:
        background = {"executive": "173550", "analytical": "1A4058", "story": "193B51"}[variant]
        foreground, accent = "FFFFFF", "68C5C0"
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
    return foreground, accent


def _add_cover(prs, category, variant, cover_image):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    foreground, accent = _style(slide, category, variant)
    if category != "template_fidelity":
        size = 1.1 if category == "readability" else 5.8 if variant != "story" else 7.0
        x = {"executive": 0.75, "analytical": 6.45, "story": 2.5}[variant]
        slide.shapes.add_picture(str(cover_image), Inches(x), Inches(1.3), width=Inches(size))
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.7))
    title.text_frame.text = "Pilot decision"
    run = title.text_frame.paragraphs[0].runs[0]
    run.font.name, run.font.size, run.font.bold = "Arial", Pt(32), True
    run.font.color.rgb = RGBColor.from_string(foreground)
    band_x = {"executive": 10.5, "analytical": 6.1, "story": 0.7}[variant]
    band = slide.shapes.add_shape(1, Inches(band_x), Inches(1.25), Inches(0.18), Inches(4.4))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string(accent)
    band.line.fill.background()
    return slide


def _add_content(prs, category, variant):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    foreground, accent = _style(slide, category, variant)
    title = slide.shapes.add_textbox(Inches(0.65), Inches(0.35), Inches(12.0), Inches(0.65))
    title.text_frame.text = "Decision evidence"
    title_run = title.text_frame.paragraphs[0].runs[0]
    title_run.font.name, title_run.font.size, title_run.font.bold = "Arial", Pt(28), True
    title_run.font.color.rgb = RGBColor.from_string(foreground)
    body_geometry = {
        "executive": (0.7, 1.1, 11.9, 2.6),
        "analytical": (0.7, 1.2, 4.2, 4.9),
        "story": (0.7, 1.15, 11.9, 2.5),
    }[variant]
    body = slide.shapes.add_textbox(*(Inches(v) for v in body_geometry))
    body.text_frame.word_wrap = True
    body.text_frame.text = _body(category)
    for paragraph in body.text_frame.paragraphs:
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(18 if variant == "analytical" else 17)
        paragraph.font.color.rgb = RGBColor.from_string(foreground)
    if category == "hidden_text":
        hidden = slide.shapes.add_textbox(Inches(14), Inches(6.5), Inches(10), Inches(0.5))
        hidden.text_frame.text = "A privacy review is mandatory before launch."
        hidden._element.xpath(".//p:cNvPr")[0].set("hidden", "1")
        slide.notes_slide.notes_text_frame.text = "A privacy review is mandatory before launch."
    headers, rows = _table_rows(category)
    table_geometry = {
        "executive": (0.7, 4.0, 11.9, 2.2),
        "analytical": (5.2, 1.2, 7.4, 4.9),
        "story": (1.2, 4.0, 11.1, 2.2),
    }[variant]
    table_shape = slide.shapes.add_table(3, 4, *(Inches(v) for v in table_geometry))
    table = table_shape.table
    for col, value in enumerate(headers):
        table.cell(0, col).text = value
    for row_index, row in enumerate(rows, 1):
        for col, value in enumerate(row):
            table.cell(row_index, col).text = value
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(accent if row_index == 0 else "FFFFFF")
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Arial"
                paragraph.font.size = Pt(14)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = RGBColor.from_string("10283C")
                paragraph.alignment = PP_ALIGN.CENTER
    return slide


def _template(path):
    image_path = path.with_name("source-figure.png")
    image = Image.new("RGB", (640, 360), "#f3f6fa")
    draw = ImageDraw.Draw(image)
    fonts = Path(__file__).resolve().parents[1] / "fonts"
    title_font = ImageFont.truetype(str(fonts / "Montserrat-Bold.ttf"), 29)
    value_font = ImageFont.truetype(str(fonts / "Montserrat-Bold.ttf"), 48)
    label_font = ImageFont.truetype(str(fonts / "Montserrat-Medium.ttf"), 15)
    draw.rectangle((0, 0, 640, 70), fill="#173550")
    draw.text((24, 17), _FIGURE_CONTENT["title"], font=title_font, fill="white")
    draw.rounded_rectangle((20, 94, 620, 336), radius=12, fill="white", outline="#c8d5e2", width=2)
    draw.text((54, 112), _FIGURE_CONTENT["value"], font=value_font, fill="#173550")
    draw.text((56, 184), _FIGURE_CONTENT["unit"], font=label_font, fill="#486176")
    draw.line((56, 216, 584, 216), fill="#d6e0e9", width=2)
    draw.text((56, 238), _FIGURE_CONTENT["label"], font=label_font, fill="#486176")
    image.save(image_path)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _add_cover(prs, "clean", "executive", image_path)
    second = prs.slides.add_slide(prs.slide_layouts[6])
    second.background.fill.solid()
    second.background.fill.fore_color.rgb = RGBColor.from_string("173550")
    prs.save(path)
    return image_path


def _case(category, case_id, template_path, variants):
    content = _SOURCE
    points = [
        {
            "id": "p-duration",
            "statement": "The pilot lasts twelve weeks.",
            "quote": "Pilot duration is 12 weeks.",
            "required": True,
            "origin": "user_text",
            "numbers": [{"value": "12", "unit": "weeks"}],
        },
        {
            "id": "p-sample",
            "statement": "The team checks seventy-three requests.",
            "quote": "The team checks 73 requests.",
            "required": True,
            "origin": "user_text",
            "numbers": [{"value": "73", "unit": "requests"}],
        },
        {
            "id": "p-caveat",
            "statement": "Changed team membership prevents a causal claim from the observed time difference.",
            "quote": "Changed team membership means differences do not prove service impact.",
            "required": True,
            "origin": "user_text",
            "negation": "do not prove",
        },
        {
            "id": "p-decision",
            "statement": "The team decides after reviewing closure share, conditional on service quality.",
            "quote": "After reviewing closure share, the team chooses whether to continue for 8 more weeks if service quality holds.",
            "required": True,
            "origin": "user_text",
            "numbers": [{"value": "8", "unit": "weeks"}],
            "condition": "if service quality holds",
        },
        {
            "id": "p-proposal",
            "statement": "Predictive routing is a proposed model idea.",
            "quote": "Model proposal: predictive routing may reduce triage time.",
            "required": False,
            "origin": "model_proposal",
        },
        {
            "id": "p-privacy",
            "statement": "A privacy review is mandatory before launch.",
            "quote": "A privacy review is mandatory before launch.",
            "required": True,
            "origin": "user_text",
        },
    ]
    requirements = [
        {
            "id": "source-table",
            "kind": "table",
            "status": "gold",
            "headers": _HEADERS,
            "rows": _ROWS,
            "text": "Preserve the header units and associate each row's figures with its team.",
        },
        {
            "id": "template-fidelity",
            "kind": "template-fidelity",
            "status": "gold",
            "text": "Preserve the source template's navy and teal visual system and retain its source figure showing 73 requests checked by the team.",
        },
        {
            "id": "visual-control",
            "kind": "visual",
            "status": "gold",
            "minimum_image_width_inches": 4,
            "text": "Keep body text readable and show the source figure at least 4 inches wide so the 73-request value and its label remain legible.",
        },
    ]
    return {
        "id": case_id,
        "slides": 2,
        "variants": variants,
        "content": content,
        "content_source": {
            "file": "constructed-control-source-v2",
            "sha256": _digest(content.encode()),
        },
        "reference": {
            "version": "control-2",
            "status": "gold",
            "provenance": {"method": "constructed_control", "review_status": "approved"},
            "points": points,
            "requirements": requirements,
        },
        "template": template_path,
        "template_source": {
            "source": "constructed synthetic template",
            "format": "PPTX",
            "sha256": _digest(template_path.read_bytes()),
            "source_slides": 2,
            "layouts": 1,
        },
        "registry_hashes": {"template": _digest(template_path.read_bytes())},
        "source_hashes": {"content": _digest(content.encode()), "images": {}},
        "images": [],
        "images_source": [],
        "synthetic": True,
        "template_origin": "synthetic_analog",
        "template_fidelity_applicability": "synthetic_reference_control",
        "template_preview_slides": [1],
    }


def build_controls(directory):
    """Write thirteen blindable positive/defect bundles with real PPTX/PDF/PNG files."""
    directory = Path(directory).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    reference_dir = directory / "reference-template"
    reference_dir.mkdir(exist_ok=True)
    template_path = reference_dir / "source-template.pptx"
    cover_image = _template(template_path)
    controls = []
    for index, (category, expected) in enumerate(_SPEC, 1):
        control_id = f"control-{index:02d}"
        control_dir = directory / control_id
        result_dir = control_dir / "result"
        bundle_dir = control_dir / "bundle"
        result_dir.mkdir(parents=True, exist_ok=True)
        variants = ["executive", "analytical", "story"]
        case = _case(category, control_id, template_path, variants)
        duplicate = category == "duplicates"
        canonical_pptx = None
        for variant in variants:
            variant_dir = result_dir / variant
            variant_dir.mkdir(parents=True, exist_ok=True)
            target = variant_dir / "deck.pptx"
            if duplicate and canonical_pptx:
                target.write_bytes(canonical_pptx.read_bytes())
            else:
                prs = Presentation()
                prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
                _add_cover(prs, category, "executive" if duplicate else variant, cover_image)
                _add_content(prs, category, "executive" if duplicate else variant)
                prs.save(target)
                if duplicate:
                    canonical_pptx = target
            if not to_pdf(target, variant_dir, timeout=30):
                raise RuntimeError(
                    "Native LibreOffice PDF export is unavailable for calibration controls"
                )
        build_bundle(case, result_dir, bundle_dir)
        controls.append(
            {
                "id": control_id,
                "category": category,
                "expected": expected,
                "bundle": str((bundle_dir / "bundle.json").resolve()),
            }
        )
    return controls
