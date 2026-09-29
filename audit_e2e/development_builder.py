"""Build the owned development corpus from its source tasks and ledgers."""

from datetime import datetime
import argparse
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
import posixpath
from pathlib import Path
import re
import shutil
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
import xml.etree.ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

AUDIT_ROOT = Path(__file__).resolve().parent
DEVELOPMENT_ROOT = AUDIT_ROOT / "fixtures" / "development"
SOURCE_ROOT = DEVELOPMENT_ROOT / "sources"
LEDGER_FILE = DEVELOPMENT_ROOT / "ledgers.json"
REGISTRY_FILE = AUDIT_ROOT / "fixtures" / "development-cases.json"
DEFAULT_REVISION = "2026-09-29-development.4"

_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
_SLIDE_MASTER_REL = _R + "/slideMaster"
_SLIDE_LAYOUT_REL = _R + "/slideLayout"
_TEMPLATE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
)
_NOW = datetime(2020, 1, 1)

_SPECS = [
    {
        "id": "dev-short-brief",
        "file": "repair-pilot-4x3.potx",
        "format": "POTX",
        "geometry": "4:3",
        "layout_pattern": "full_width_content",
        "width": 10,
        "height": 7.5,
        "layouts": 4,
        "master_shapes": 2,
        "layout_shapes": 0,
        "prototype_slides": 1,
        "accent": "167D8D",
        "background": "F5F8FA",
        "dark": False,
        "features": ["layout_inheritance", "master_graphics"],
        "coverage": ["short_brief", "four_to_three", "inherited_placeholders"],
    },
    {
        "id": "dev-long-title-paragraphs",
        "file": "river-handoff-dark-wide.pptx",
        "format": "PPTX",
        "geometry": "16:9",
        "layout_pattern": "dark_sidebar",
        "width": 13.333,
        "height": 7.5,
        "layouts": 5,
        "master_shapes": 3,
        "layout_shapes": 1,
        "prototype_slides": 2,
        "accent": "E8A64A",
        "background": "172B46",
        "dark": True,
        "narrow_region": True,
        "features": ["layout_inheritance", "master_graphics", "narrow_regions", "dark_theme"],
        "coverage": ["long_title", "long_paragraphs", "dark_template", "narrow_regions"],
    },
    {
        "id": "dev-route-comparison",
        "file": "mobile-library-split.pptx",
        "format": "PPTX",
        "geometry": "16:9",
        "layout_pattern": "comparison_panels",
        "width": 13.333,
        "height": 7.5,
        "layouts": 6,
        "master_shapes": 2,
        "layout_shapes": 2,
        "prototype_slides": 1,
        "accent": "286B55",
        "background": "F6F5EF",
        "dark": False,
        "groups": 1,
        "features": ["layout_inheritance", "master_graphics", "group_shapes", "split_regions"],
        "coverage": ["route_comparison", "grouped_panels", "side_by_side_regions"],
    },
    {
        "id": "dev-project-timeline",
        "file": "canal-timeline-4x3.pptx",
        "format": "PPTX",
        "geometry": "4:3",
        "layout_pattern": "timeline_rail",
        "width": 10,
        "height": 7.5,
        "layouts": 7,
        "master_shapes": 4,
        "layout_shapes": 0,
        "prototype_slides": 2,
        "accent": "C25A37",
        "background": "FFF8F1",
        "dark": False,
        "features": ["layout_inheritance", "master_graphics", "timeline_rail"],
        "coverage": ["timeline", "four_to_three", "grouped_milestones"],
    },
    {
        "id": "dev-long-table",
        "file": "energy-long-ledger.pptx",
        "format": "PPTX",
        "geometry": "16:9",
        "layout_pattern": "ledger_table",
        "width": 13.333,
        "height": 7.5,
        "layouts": 8,
        "master_shapes": 3,
        "layout_shapes": 1,
        "prototype_slides": 2,
        "accent": "4673A1",
        "background": "F4F7FB",
        "dark": False,
        "table": {"rows": 5, "cols": 4},
        "features": ["layout_inheritance", "master_graphics", "native_tables"],
        "coverage": ["long_table", "units", "source_labels"],
    },
    {
        "id": "dev-wide-table",
        "file": "service-wide-matrix.pptx",
        "format": "PPTX",
        "geometry": "2:1",
        "layout_pattern": "wide_matrix",
        "width": 15,
        "height": 7.5,
        "layouts": 9,
        "master_shapes": 4,
        "layout_shapes": 2,
        "prototype_slides": 1,
        "accent": "7655A2",
        "background": "F7F4FA",
        "dark": False,
        "table": {"rows": 4, "cols": 8},
        "features": ["layout_inheritance", "master_graphics", "native_tables", "wide_canvas"],
        "coverage": ["wide_table", "labeled_units", "source_ids"],
    },
    {
        "id": "dev-labeled-metrics",
        "file": "workshop-source-notes.pptx",
        "format": "PPTX",
        "geometry": "16:9",
        "layout_pattern": "metric_cards",
        "width": 13.333,
        "height": 7.5,
        "layouts": 5,
        "master_shapes": 5,
        "layout_shapes": 1,
        "prototype_slides": 2,
        "accent": "B24E70",
        "background": "FFF7F9",
        "dark": False,
        "narrow_region": True,
        "features": ["layout_inheritance", "master_graphics", "narrow_regions", "source_notes"],
        "coverage": ["labeled_units", "source_ids", "denominator_limits"],
    },
    {
        "id": "dev-chart-series",
        "file": "rooftop-multiseries-chart.pptx",
        "format": "PPTX",
        "geometry": "16:9",
        "layout_pattern": "chart_with_unit_legend",
        "width": 13.333,
        "height": 7.5,
        "layouts": 10,
        "master_shapes": 2,
        "layout_shapes": 0,
        "prototype_slides": 2,
        "accent": "42805C",
        "background": "EFF8F2",
        "dark": False,
        "chart": "line",
        "features": ["layout_inheritance", "master_graphics", "native_charts", "embedded_workbook"],
        "coverage": ["chart_series", "different_units", "source_labels"],
    },
    {
        "id": "dev-screenshot-landscape",
        "file": "controlroom-capture.pptx",
        "format": "PPTX",
        "geometry": "16:9",
        "layout_pattern": "landscape_viewport",
        "width": 13.333,
        "height": 7.5,
        "layouts": 6,
        "master_shapes": 3,
        "layout_shapes": 1,
        "prototype_slides": 1,
        "accent": "2868A0",
        "background": "F1F6FA",
        "dark": False,
        "features": ["layout_inheritance", "master_graphics", "screenshot_context"],
        "image": {
            "file": "dev-screenshot-landscape.png",
            "width": 1600,
            "height": 900,
            "kind": "landscape",
        },
        "coverage": ["landscape_screenshot", "sample_labels", "image_provenance"],
    },
    {
        "id": "dev-screenshot-portrait",
        "file": "fieldcheck-portrait.potx",
        "format": "POTX",
        "geometry": "3:4",
        "layout_pattern": "portrait_device",
        "width": 7.5,
        "height": 10,
        "layouts": 7,
        "master_shapes": 4,
        "layout_shapes": 1,
        "prototype_slides": 1,
        "accent": "176C76",
        "background": "F1F8F8",
        "dark": False,
        "features": [
            "layout_inheritance",
            "master_graphics",
            "portrait_canvas",
            "screenshot_context",
        ],
        "image": {
            "file": "dev-screenshot-portrait.png",
            "width": 540,
            "height": 960,
            "kind": "portrait",
        },
        "coverage": ["portrait_screenshot", "device_fallback", "sample_labels"],
    },
    {
        "id": "dev-screenshot-square",
        "file": "catalog-square-asset.pptx",
        "format": "PPTX",
        "geometry": "4:3",
        "layout_pattern": "square_media_well",
        "width": 10,
        "height": 7.5,
        "layouts": 4,
        "master_shapes": 5,
        "layout_shapes": 0,
        "prototype_slides": 2,
        "accent": "A87432",
        "background": "FBF6ED",
        "dark": False,
        "features": ["layout_inheritance", "master_graphics", "screenshot_context"],
        "image": {
            "file": "dev-screenshot-square.png",
            "width": 960,
            "height": 960,
            "kind": "square",
        },
        "coverage": ["square_screenshot", "sample_labels", "image_provenance"],
    },
    {
        "id": "dev-narrow-region",
        "file": "archive-narrow-columns.pptx",
        "format": "PPTX",
        "geometry": "16:9",
        "layout_pattern": "three_narrow_columns",
        "width": 13.333,
        "height": 7.5,
        "layouts": 8,
        "master_shapes": 6,
        "layout_shapes": 3,
        "prototype_slides": 2,
        "accent": "7B5D4B",
        "background": "F7F4F1",
        "dark": False,
        "narrow_region": True,
        "features": [
            "layout_inheritance",
            "master_graphics",
            "narrow_regions",
            "multi_column_layout",
        ],
        "coverage": ["long_paragraphs", "narrow_regions", "exception_flow"],
    },
    {
        "id": "dev-grouped-process",
        "file": "equipment-flow-groups.pptx",
        "format": "PPTX",
        "geometry": "16:9",
        "layout_pattern": "grouped_process",
        "width": 13.333,
        "height": 7.5,
        "layouts": 9,
        "master_shapes": 3,
        "layout_shapes": 1,
        "prototype_slides": 1,
        "accent": "2C7462",
        "background": "F0F8F5",
        "dark": False,
        "groups": 3,
        "features": ["layout_inheritance", "master_graphics", "group_shapes", "connected_process"],
        "coverage": ["grouped_process", "multi_step_workflow", "source_limits"],
    },
    {
        "id": "dev-multiple-masters",
        "file": "coastal-dual-master.pptx",
        "format": "PPTX",
        "geometry": "16:9",
        "layout_pattern": "two_team_masters",
        "width": 13.333,
        "height": 7.5,
        "layouts": 7,
        "master_shapes": 4,
        "layout_shapes": 2,
        "prototype_slides": 1,
        "accent": "366A8B",
        "background": "EFF5FA",
        "dark": False,
        "second_master": True,
        "features": ["layout_inheritance", "master_graphics", "multiple_masters", "team_layouts"],
        "coverage": ["multiple_masters", "inheritance", "separate_team_content"],
    },
    {
        "id": "dev-chart-workbook-resource",
        "file": "booking-chart-workbook.pptx",
        "format": "PPTX",
        "geometry": "2:1",
        "layout_pattern": "chart_table_resource_panel",
        "width": 15,
        "height": 7.5,
        "layouts": 11,
        "master_shapes": 5,
        "layout_shapes": 3,
        "prototype_slides": 2,
        "accent": "A05A37",
        "background": "FBF4EF",
        "dark": False,
        "chart": "column",
        "table": {"rows": 3, "cols": 4},
        "embedded_picture": True,
        "features": [
            "layout_inheritance",
            "master_graphics",
            "native_charts",
            "native_tables",
            "embedded_workbook",
            "linked_media",
        ],
        "coverage": ["chart_and_workbook", "native_resources", "unit_labels"],
    },
]


def _hash(data):
    return sha256(data).hexdigest()


def _slug(value):
    return re.sub(r"[^a-z0-9-]+", "-", value.casefold()).strip("-")


def _xml_bytes(root):
    for prefix, namespace in (
        ("p", _P),
        ("a", _A),
        ("r", _R),
        ("rel", _PKG_REL),
        ("ct", _CONTENT_TYPES),
    ):
        ET.register_namespace(prefix, namespace)
    return ET.tostring(root, encoding="UTF-8", xml_declaration=True)


def _normalize_package(source_path, target_path):
    with ZipFile(source_path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    _normalize_embedded_workbooks(parts)
    if target_path.suffix.lower() == ".potx":
        content_types = ET.fromstring(parts["[Content_Types].xml"])
        main = content_types.find(
            f"{{{_CONTENT_TYPES}}}Override[@PartName='/ppt/presentation.xml']"
        )
        if main is None:
            raise ValueError("PPTX presentation content type was not found")
        main.set("ContentType", _TEMPLATE_CONTENT_TYPE)
        parts["[Content_Types].xml"] = _xml_bytes(content_types)
    _write_normalized_parts(parts, target_path)


def _write_normalized_parts(parts, target_path):
    target_path.write_bytes(_normalized_zip_bytes(parts))


def _normalized_zip_bytes(parts):
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(parts):
            info = ZipInfo(name, (2020, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            target.writestr(info, parts[name], compress_type=ZIP_DEFLATED, compresslevel=9)
    return buffer.getvalue()


def _normalize_embedded_workbooks(parts):
    """Make generated embedded XLSX parts reproducible across rebuilds."""
    for name, data in list(parts.items()):
        if not name.startswith("ppt/embeddings/") or not name.lower().endswith(".xlsx"):
            continue
        with ZipFile(BytesIO(data)) as workbook:
            workbook_parts = {
                workbook_name: workbook.read(workbook_name) for workbook_name in workbook.namelist()
            }
        core_name = "docProps/core.xml"
        if core_name in workbook_parts:
            core = workbook_parts[core_name]
            for tag in (b"created", b"modified"):
                pattern = rb"(<dcterms:" + tag + rb"[^>]*>)[^<]*(</dcterms:" + tag + rb">)"
                core = re.sub(pattern, rb"\g<1>2020-01-01T00:00:00Z\g<2>", core)
            workbook_parts[core_name] = core
        parts[name] = _normalized_zip_bytes(workbook_parts)


def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()


def _move_shape_to_tree(shape, tree):
    element = deepcopy(shape._element)
    ids = [
        int(node.get("id", "0"))
        for node in tree.iter(f"{{{_P}}}cNvPr")
        if node.get("id", "0").isdigit()
    ]
    nonvisual = element.find(f".//{{{_P}}}cNvPr")
    nonvisual.set("id", str(max(ids, default=1) + 1))
    tree.insert_element_before(element, "p:extLst")
    shape._element.getparent().remove(shape._element)


def _add_master_artwork(master, spec, prs):
    width, height = Inches(spec["width"]), Inches(spec["height"])
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = RGBColor.from_string(spec["background"])
    sample_slide = prs.slides.add_slide(prs.slide_layouts[6])
    for index in range(spec["master_shapes"]):
        if index % 3 == 0:
            left, top, shape_width, shape_height = (
                0,
                Inches(0.04 + 0.04 * index),
                width,
                Inches(0.08),
            )
        elif index % 3 == 1:
            left, top, shape_width, shape_height = (
                Inches(0.04 + 0.04 * index),
                0,
                Inches(0.1),
                height,
            )
        else:
            left = Inches(0.3 + index * 0.18)
            top = height - Inches(0.15 + index * 0.03)
            shape_width, shape_height = width - left, Inches(0.07)
        shape = sample_slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, top, shape_width, shape_height
        )
        _set_fill(shape, spec["accent"] if index == 0 else spec["background"])
        _move_shape_to_tree(shape, master._element.cSld.spTree)
    return sample_slide


def _prune_layouts(master, count):
    source_layouts = list(master.slide_layouts)
    keep_indices = [1, 6, 0, 3, 8, 2, 4, 5, 7, 9, 10][:count]
    keep = {source_layouts[index].part.partname for index in keep_indices}
    layout_id_list = master._element.sldLayoutIdLst
    for entry in list(layout_id_list):
        rel_id = entry.rId
        rel = master.part.rels[rel_id]
        if rel.target_part.partname not in keep:
            layout_id_list.remove(entry)
            master.part.drop_rel(rel_id)
    return [source_layouts[index] for index in keep_indices]


def _style_layouts(layouts, spec, sample_slide):
    title_color = "F8F5E8" if spec["dark"] else spec["accent"]
    for index, layout in enumerate(layouts):
        title = next(
            (
                placeholder
                for placeholder in layout.placeholders
                if placeholder.placeholder_format.type == PP_PLACEHOLDER.TITLE
            ),
            None,
        )
        if title:
            title.left = Inches(0.45)
            title.top = Inches(0.35)
            title.height = Inches(0.82)
            if spec.get("narrow_region"):
                title.width = Inches(4.4)
            for paragraph in title.text_frame.paragraphs:
                paragraph.font.name = "Aptos Display"
                paragraph.font.size = Pt(25 + index % 3)
                paragraph.font.color.rgb = RGBColor.from_string(title_color)
        for placeholder in layout.placeholders:
            if placeholder == title:
                continue
            if spec.get("narrow_region") and placeholder.has_text_frame:
                placeholder.left = Inches(0.55 + 0.08 * index)
                placeholder.width = Inches(3.7)
                placeholder.height = Inches(4.8)
        if index < spec["layout_shapes"]:
            shape = sample_slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.4 + 0.16 * index),
                Inches(1.45 + 0.2 * index),
                Inches(0.12),
                Inches(4.2 - 0.2 * index),
            )
            _set_fill(shape, spec["accent"])
            _move_shape_to_tree(shape, layout._element.cSld.spTree)


def _format_text(shape, text, color, size=25, bold=False):
    shape.text = text
    frame = shape.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)


def _set_table(frame, row_count, col_count):
    table = frame.table
    for row in range(row_count):
        for col in range(col_count):
            cell = table.cell(row, col)
            cell.text = "SAMPLE" if row == 0 else f"DEMO {row}-{col}"
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(8 if col_count >= 8 else 11)
                paragraph.font.color.rgb = RGBColor.from_string("FFFFFF" if row == 0 else "28384A")
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor.from_string("465D77" if row == 0 else "F0F3F6")


def _add_chart(slide, kind, accent, *, left=1.0, top=1.55, width=8.4, height=4.8):
    data = CategoryChartData()
    data.categories = ["SAMPLE 1", "SAMPLE 2", "SAMPLE 3", "SAMPLE 4"]
    data.add_series("SAMPLE A", (14, 18, 23, 20))
    data.add_series("SAMPLE B", (9, 12, 15, 19))
    chart_type = XL_CHART_TYPE.LINE_MARKERS if kind == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
    frame = slide.shapes.add_chart(
        chart_type, Inches(left), Inches(top), Inches(width), Inches(height), data
    )
    chart = frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "SAMPLE ONLY — NOT SOURCE DATA"
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.chart_style = 13
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "SAMPLE UNITS"
    return frame


def _add_groups(slide, count, accent):
    for index in range(count):
        group = slide.shapes.add_group_shape()
        x = Inches(0.8 + index * 3.3)
        card = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.1), Inches(2.4), Inches(1.6)
        )
        _set_fill(card, "FFFFFF")
        label = group.shapes.add_textbox(x + Inches(0.2), Inches(2.55), Inches(2.0), Inches(0.55))
        _format_text(label, f"SAMPLE STEP {index + 1}", accent, 14, True)
        marker = group.shapes.add_shape(
            MSO_SHAPE.OVAL, x + Inches(0.9), Inches(1.62), Inches(0.6), Inches(0.6)
        )
        _set_fill(marker, accent)


def _panel(slide, spec, left, top, width, height, label, *, fill="FFFFFF", size=14):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    _set_fill(shape, fill)
    shape.line.color.rgb = RGBColor.from_string(spec["accent"])
    text_color = "F8F5E8" if spec["dark"] else "263A4C"
    _format_text(shape, label, text_color, size, True)
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return shape


def _add_content_areas(slide, spec):
    width, height = spec["width"], spec["height"]
    accent = spec["accent"]
    pattern = spec["layout_pattern"]
    if pattern == "full_width_content":
        _panel(slide, spec, 0.8, 2.0, width - 1.6, height - 2.8, "SAMPLE • FULL-WIDTH BODY")
    elif pattern == "dark_sidebar":
        _panel(slide, spec, 0.55, 2.0, 3.1, 4.7, "SAMPLE\nNARROW SIDEBAR", fill="243E59")
        _panel(slide, spec, 3.95, 2.0, width - 4.5, 4.7, "SAMPLE\nMAIN TEXT REGION", fill="203852")
    elif pattern == "comparison_panels":
        gap, margin = 0.35, 0.8
        panel_width = (width - 2 * margin - gap) / 2
        _panel(slide, spec, margin, 2.0, panel_width, 4.2, "SAMPLE OPTION A")
        _panel(slide, spec, margin + panel_width + gap, 2.0, panel_width, 4.2, "SAMPLE OPTION B")
        _add_groups(slide, spec["groups"], accent)
    elif pattern == "timeline_rail":
        rail = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.0), Inches(0.1), Inches(4.7)
        )
        _set_fill(rail, accent)
        for index, label in enumerate(
            ("SAMPLE MILESTONE A", "SAMPLE MILESTONE B", "SAMPLE MILESTONE C")
        ):
            top = 2.0 + index * 1.45
            _panel(slide, spec, 1.35, top, 7.4, 1.0, label, size=12)
            marker = slide.shapes.add_shape(
                MSO_SHAPE.OVAL, Inches(0.78), Inches(top + 0.27), Inches(0.54), Inches(0.54)
            )
            _set_fill(marker, accent)
    elif pattern in {"ledger_table", "wide_matrix"}:
        rows, cols = spec["table"]["rows"], spec["table"]["cols"]
        left = 0.65 if pattern == "wide_matrix" else 0.8
        table_width = width - 2 * left
        frame = slide.shapes.add_table(
            rows, cols, Inches(left), Inches(2.0), Inches(table_width), Inches(3.6)
        )
        _set_table(frame, rows, cols)
        _panel(slide, spec, left, 5.85, table_width, 0.55, "SAMPLE UNITS • SOURCE LABELS", size=11)
    elif pattern == "metric_cards":
        gap, margin = 0.25, 0.7
        card_width = (width - 2 * margin - gap * 3) / 4
        for index, label in enumerate(("84 VISITS", "61 JOBS", "23 MIN", "14 RETURNS")):
            _panel(
                slide,
                spec,
                margin + index * (card_width + gap),
                2.2,
                card_width,
                2.3,
                f"SAMPLE\n{label}",
            )
        _panel(
            slide,
            spec,
            margin,
            4.9,
            width - 2 * margin,
            1.0,
            "SAMPLE SOURCE NOTES AND UNIT DEFINITIONS",
            size=12,
        )
    elif pattern == "chart_with_unit_legend":
        _add_chart(slide, spec["chart"], accent, left=0.7, top=1.9, width=8.8, height=4.8)
        _panel(slide, spec, 9.8, 2.0, 2.8, 4.5, "SAMPLE\nkg\nL\n°C\n\nDIFFERENT SCALES", size=13)
    elif pattern == "landscape_viewport":
        _panel(
            slide,
            spec,
            1.82,
            1.65,
            9.69,
            5.45,
            "SAMPLE SCREENSHOT VIEWPORT\n16:9 CONTENT AREA",
            size=16,
        )
    elif pattern == "portrait_device":
        _panel(slide, spec, 1.95, 1.8, 3.6, 6.4, "SAMPLE PHONE SCREEN\n9:16 CONTENT AREA", size=15)
        _panel(slide, spec, 0.4, 8.45, 6.7, 0.8, "SAMPLE CAPTION / SOURCE NOTE", size=11)
    elif pattern == "square_media_well":
        _panel(slide, spec, 0.65, 2.0, 4.8, 4.8, "SQUARE MEDIA AREA", size=17)
        _panel(slide, spec, 5.75, 2.0, 3.55, 4.8, "SAMPLE\nCAPTION\nSOURCE\nDATE", size=13)
    elif pattern == "three_narrow_columns":
        margin, gap = 0.45, 0.25
        col_width = (width - 2 * margin - 2 * gap) / 3
        for index, label in enumerate(("SAMPLE RECEIPT", "SAMPLE REVIEW", "SAMPLE LOCATION")):
            _panel(
                slide, spec, margin + index * (col_width + gap), 2.0, col_width, 4.7, label, size=12
            )
    elif pattern == "grouped_process":
        _add_groups(slide, spec["groups"], accent)
        for index in range(spec["groups"] - 1):
            chevron = slide.shapes.add_shape(
                MSO_SHAPE.CHEVRON,
                Inches(3.25 + index * 3.3),
                Inches(2.65),
                Inches(0.7),
                Inches(0.35),
            )
            _set_fill(chevron, accent)
        _panel(slide, spec, 0.8, 4.45, 9.0, 1.3, "SAMPLE HANDOFF AND EXCEPTION NOTES", size=12)
    elif pattern == "two_team_masters":
        _panel(slide, spec, 0.7, 2.0, 5.6, 4.5, "SAMPLE SHORE TEAM REGION", size=14)
        _panel(slide, spec, 6.9, 2.0, 5.6, 4.5, "SAMPLE WATER TEAM REGION", size=14)
    elif pattern == "chart_table_resource_panel":
        _add_chart(slide, spec["chart"], accent, left=0.55, top=1.8, width=7.6, height=4.5)
        frame = slide.shapes.add_table(3, 4, Inches(8.55), Inches(2.0), Inches(5.9), Inches(2.4))
        _set_table(frame, 3, 4)
        slide.shapes.add_picture(
            BytesIO(_sample_picture_bytes(accent)),
            Inches(11.7),
            Inches(5.1),
            width=Inches(2.4),
            height=Inches(1.35),
        )


def _sample_picture_bytes(color):
    image = Image.new("RGB", (192, 108), "#f2f5f7")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 191, 21), fill=f"#{color}")
    draw.text((10, 6), "SAMPLE ASSET", fill="white", font=ImageFont.load_default(size=10))
    draw.rectangle((15, 40, 177, 92), outline=f"#{color}", width=3)
    result = BytesIO()
    image.save(result, format="PNG", optimize=False, compress_level=9)
    return result.getvalue()


def _build_template(spec, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    raw = path.with_name(path.stem + ".building.pptx")
    final_raw = path.with_name(path.stem + ".final.pptx")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(spec["width"]), Inches(spec["height"])
    prs.core_properties.created = prs.core_properties.modified = _NOW
    prs.core_properties.title = f"Owned synthetic fixture: {spec['id']}"
    prs.core_properties.subject = "Sample-only presentation layout; no external source material"
    master = prs.slide_masters[0]
    artwork_slide = _add_master_artwork(master, spec, prs)
    original_layouts = list(master.slide_layouts)
    title_layout, blank_layout = original_layouts[1], original_layouts[6]
    layouts = _prune_layouts(master, spec["layouts"])
    _style_layouts(layouts, spec, artwork_slide)

    accent_text = "FFF8E8" if spec["dark"] else "213247"
    for slide_number in range(spec["prototype_slides"]):
        if slide_number == 0:
            slide = artwork_slide
        else:
            layout = title_layout if slide_number % 2 == 1 else blank_layout
            slide = prs.slides.add_slide(layout)
        if slide.shapes.title:
            _format_text(slide.shapes.title, "SAMPLE TEMPLATE CONTENT", accent_text, 24, True)
            if len(slide.placeholders) > 1:
                body = next(
                    (shape for shape in slide.placeholders if shape != slide.shapes.title), None
                )
                if body and body.has_text_frame:
                    _format_text(body, "SAMPLE ONLY • DO NOT USE AS SOURCE FACTS", accent_text, 17)
        if slide_number == 0:
            title_box = slide.shapes.add_textbox(
                Inches(0.7), Inches(1.35), Inches(spec["width"] - 1.4), Inches(0.6)
            )
            _format_text(title_box, "SAMPLE ONLY • NOT SOURCE DATA", spec["accent"], 17, True)
            _add_content_areas(slide, spec)
    prs.save(raw)
    if spec.get("second_master"):
        _add_second_master(raw)
        second = Presentation(raw)
        layout = second.slide_masters[1].slide_layouts[0]
        slide = second.slides.add_slide(layout)
        if slide.shapes.title:
            _format_text(slide.shapes.title, "SAMPLE SECOND MASTER", "FFFFFF", 22, True)
        text = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(6.5), Inches(0.8))
        _format_text(text, "SAMPLE TEAM LAYOUT • NOT SOURCE FACTS", spec["accent"], 18)
        second.save(final_raw)
        raw.unlink()
        _normalize_package(final_raw, path)
        final_raw.unlink()
    else:
        _normalize_package(raw, path)
        raw.unlink()
    _assert_master_layout_relationships(path)


def _assert_master_layout_relationships(path):
    path = Path(path)
    check_path = path
    with tempfile.TemporaryDirectory(prefix="template-parse-") as temporary:
        if path.suffix.lower() == ".potx":
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            content_types = ET.fromstring(parts["[Content_Types].xml"])
            main = content_types.find(
                f"{{{_CONTENT_TYPES}}}Override[@PartName='/ppt/presentation.xml']"
            )
            if main is None:
                raise ValueError("POTX presentation content type was not found")
            main.set(
                "ContentType",
                "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
            )
            parts["[Content_Types].xml"] = _xml_bytes(content_types)
            check_path = Path(temporary) / (path.stem + ".pptx")
            _write_normalized_parts(parts, check_path)
        parsed = Presentation(check_path)
        if not parsed.slide_masters:
            raise ValueError(f"Template has no slide master: {path.name}")
        for master in parsed.slide_masters:
            layouts = list(master.slide_layouts)
            if not layouts:
                raise ValueError(f"Template master has no layouts: {path.name}")
            for layout in layouts:
                if layout.part is None:
                    raise ValueError(f"Template layout relationship is dangling: {path.name}")


def _replace_relationship_target(rel, target):
    rel.set("Target", target)


def _part_target(source_part, target):
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def _relationship_part(part_name):
    folder, name = posixpath.split(part_name)
    return posixpath.join(folder, "_rels", name + ".rels")


def _add_second_master(path):
    with ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    master_path = "ppt/slideMasters/slideMaster1.xml"
    master_rels_path = _relationship_part(master_path)
    master_rels = ET.fromstring(parts[master_rels_path])
    layout_rels = [rel for rel in master_rels if rel.get("Type") == _SLIDE_LAYOUT_REL]
    next_layout = (
        max(
            [
                int(match.group(1))
                for name in parts
                if (match := re.fullmatch(r"ppt/slideLayouts/slideLayout(\d+)\.xml", name))
            ]
            + [0]
        )
        + 1
    )
    new_master_path = "ppt/slideMasters/slideMaster2.xml"
    new_master_xml = ET.fromstring(parts[master_path])
    new_master_rels = ET.fromstring(parts[master_rels_path])
    new_by_id = {rel.get("Id"): rel for rel in new_master_rels}
    layout_id_list = new_master_xml.find(f"{{{_P}}}sldLayoutIdLst")
    if layout_id_list is None:
        raise ValueError("Cannot clone a master without its layout list")
    layout_id_entries = {entry.get(f"{{{_R}}}id"): entry for entry in layout_id_list}
    maximum_layout_id = max(int(entry.get("id", "0")) for entry in layout_id_list)
    for index, rel in enumerate(layout_rels):
        old_layout_path = _part_target(master_path, rel.get("Target"))
        new_layout_path = f"ppt/slideLayouts/slideLayout{next_layout + index}.xml"
        parts[new_layout_path] = parts[old_layout_path]
        old_layout_rels_path = _relationship_part(old_layout_path)
        new_layout_rels_path = _relationship_part(new_layout_path)
        if old_layout_rels_path in parts:
            layout_rels_xml = ET.fromstring(parts[old_layout_rels_path])
            for item in layout_rels_xml:
                if item.get("Type") == _SLIDE_MASTER_REL:
                    _replace_relationship_target(item, "../slideMasters/slideMaster2.xml")
            parts[new_layout_rels_path] = _xml_bytes(layout_rels_xml)
        new_rel = new_by_id[rel.get("Id")]
        _replace_relationship_target(
            new_rel, posixpath.relpath(new_layout_path, "ppt/slideMasters")
        )
        entry = layout_id_entries.get(rel.get("Id"))
        if entry is None:
            raise ValueError("Cloned master layout relationship is missing its identifier")
        entry.set("id", str(maximum_layout_id + index + 1))
        content_types = ET.fromstring(parts["[Content_Types].xml"])
        original_override = content_types.find(
            f"{{{_CONTENT_TYPES}}}Override[@PartName='/{old_layout_path}']"
        )
        if original_override is None:
            raise ValueError(f"Missing content type for {old_layout_path}")
        content_types.append(
            ET.Element(
                f"{{{_CONTENT_TYPES}}}Override",
                {
                    "PartName": "/" + new_layout_path,
                    "ContentType": original_override.get("ContentType"),
                },
            )
        )
        parts["[Content_Types].xml"] = _xml_bytes(content_types)

    first_shape = new_master_xml.find(
        f".//{{{_P}}}sp/{{{_P}}}spPr/{{{_A}}}solidFill/{{{_A}}}srgbClr"
    )
    if first_shape is not None:
        first_shape.set("val", "4785A8")
    parts[new_master_path] = _xml_bytes(new_master_xml)
    parts[_relationship_part(new_master_path)] = _xml_bytes(new_master_rels)

    content_types = ET.fromstring(parts["[Content_Types].xml"])
    content_types.append(
        ET.Element(
            f"{{{_CONTENT_TYPES}}}Override",
            {
                "PartName": "/" + new_master_path,
                "ContentType": "application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml",
            },
        )
    )
    parts["[Content_Types].xml"] = _xml_bytes(content_types)

    presentation_rels_path = "ppt/_rels/presentation.xml.rels"
    presentation_rels = ET.fromstring(parts[presentation_rels_path])
    numeric_ids = [int(rel.get("Id", "rId0")[3:]) for rel in presentation_rels]
    new_rel_id = f"rId{max(numeric_ids) + 1}"
    presentation_rels.append(
        ET.Element(
            f"{{{_PKG_REL}}}Relationship",
            {
                "Id": new_rel_id,
                "Type": _SLIDE_MASTER_REL,
                "Target": "slideMasters/slideMaster2.xml",
            },
        )
    )
    parts[presentation_rels_path] = _xml_bytes(presentation_rels)

    presentation_xml = ET.fromstring(parts["ppt/presentation.xml"])
    master_list = presentation_xml.find(f"{{{_P}}}sldMasterIdLst")
    if master_list is None:
        raise ValueError("Presentation has no slide master list")
    master_ids = [int(entry.get("id", "0")) for entry in master_list]
    entry = ET.Element(f"{{{_P}}}sldMasterId", {"id": str(max(master_ids) + 1)})
    entry.set(f"{{{_R}}}id", new_rel_id)
    master_list.append(entry)
    parts["ppt/presentation.xml"] = _xml_bytes(presentation_xml)

    temp = path.with_name(path.stem + ".masters.pptx")
    with ZipFile(temp, "w", ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(parts):
            info = ZipInfo(name, (2020, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            target.writestr(info, parts[name], compress_type=ZIP_DEFLATED, compresslevel=9)
    temp.replace(path)


def _build_screenshot(spec, path):
    width, height = spec["width"], spec["height"]
    image = Image.new("RGB", (width, height), "#eef2f6")
    draw = ImageDraw.Draw(image)
    scale = min(width, height)
    font = ImageFont.load_default(size=max(19, scale // 20))
    small = ImageFont.load_default(size=max(13, scale // 34))
    if spec["kind"] == "landscape":
        draw.rectangle((0, 0, width, int(height * 0.15)), fill="#193d5c")
        draw.text(
            (int(width * 0.04), int(height * 0.05)),
            "SAMPLE DATA  •  SIMULATED SCREEN",
            fill="white",
            font=font,
        )
        for index, label in enumerate(("ROUTE A", "ROUTE B", "ROUTE C")):
            x = int(width * (0.04 + index * 0.31))
            draw.rounded_rectangle(
                (x, int(height * 0.21), x + int(width * 0.27), int(height * 0.42)),
                16,
                fill="white",
                outline="#b8c6d4",
                width=3,
            )
            draw.text((x + 20, int(height * 0.25)), label, fill="#41596f", font=small)
            draw.text(
                (x + 20, int(height * 0.31)), f"DEMO {18 + index * 7}", fill="#193d5c", font=font
            )
        area = (int(width * 0.04), int(height * 0.49), int(width * 0.96), int(height * 0.93))
        draw.rounded_rectangle(area, 16, fill="white", outline="#b8c6d4", width=3)
        points = [
            (int(width * x), int(height * y))
            for x, y in (
                (0.09, 0.83),
                (0.22, 0.76),
                (0.35, 0.79),
                (0.49, 0.68),
                (0.64, 0.72),
                (0.78, 0.6),
                (0.9, 0.64),
            )
        ]
        draw.line(points, fill="#c46e3c", width=max(4, scale // 170))
        for x, y in points:
            draw.ellipse((x - 7, y - 7, x + 7, y + 7), fill="#c46e3c")
        draw.text(
            (int(width * 0.06), int(height * 0.51)),
            "DECORATIVE TREND • NOT A MEASUREMENT",
            fill="#596d7e",
            font=small,
        )
    elif spec["kind"] == "portrait":
        draw.rectangle((0, 0, width, int(height * 0.11)), fill="#176c76")
        draw.text((int(width * 0.06), int(height * 0.035)), "SAMPLE ONLY", fill="white", font=font)
        draw.text(
            (int(width * 0.06), int(height * 0.13)),
            "SIMULATED CHECKLIST",
            fill="#244653",
            font=small,
        )
        for index, label in enumerate(
            ("ARRIVAL TIME", "SITE LABEL", "CHECK ITEM", "PHOTO NOTE", "SIGNATURE")
        ):
            top = int(height * (0.22 + index * 0.135))
            draw.rounded_rectangle(
                (int(width * 0.06), top, int(width * 0.94), top + int(height * 0.105)),
                12,
                fill="white",
                outline="#b9cbd0",
                width=2,
            )
            draw.rectangle(
                (
                    int(width * 0.1),
                    top + int(height * 0.03),
                    int(width * 0.17),
                    top + int(height * 0.075),
                ),
                outline="#176c76",
                width=3,
            )
            draw.text(
                (int(width * 0.22), top + int(height * 0.035)), label, fill="#47616a", font=small
            )
        draw.text(
            (int(width * 0.06), int(height * 0.93)),
            "NOT A REAL INSPECTION",
            fill="#9b4f4f",
            font=small,
        )
    else:
        draw.rounded_rectangle(
            (int(width * 0.06), int(height * 0.06), int(width * 0.94), int(height * 0.94)),
            28,
            fill="white",
            outline="#b8955f",
            width=5,
        )
        draw.rounded_rectangle(
            (int(width * 0.15), int(height * 0.14), int(width * 0.85), int(height * 0.72)),
            24,
            fill="#eadfc9",
        )
        draw.text((int(width * 0.22), int(height * 0.19)), "SAMPLE", fill="#775a33", font=font)
        draw.ellipse(
            (int(width * 0.36), int(height * 0.31), int(width * 0.66), int(height * 0.61)),
            fill="#607f70",
        )
        draw.text(
            (int(width * 0.20), int(height * 0.77)),
            "P-104 • NOT A REAL CATALOG",
            fill="#624d32",
            font=small,
        )
        draw.text(
            (int(width * 0.24), int(height * 0.84)),
            "DECORATIVE PART TILE",
            fill="#816b4e",
            font=small,
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", optimize=False, compress_level=9)


def _case_from_spec(spec, content, source_hash, ledger, fixture_root, revision):
    template_path = fixture_root / "templates" / spec["file"]
    image_sources = []
    if image_spec := spec.get("image"):
        image_path = fixture_root / "images" / image_spec["file"]
        image_sources.append(
            {
                "file": image_path.name,
                "sha256": _hash(image_path.read_bytes()),
                "source": "Owned simulated screenshot; authored for this corpus",
                "origin": "owned_synthetic_asset",
                "width_px": image_spec["width"],
                "height_px": image_spec["height"],
                "role": "layout_stress_input",
            }
        )
    slide_count = len([line for line in content.splitlines() if line.startswith("## ")])
    metadata = {
        "file": spec["file"],
        "sha256": _hash(template_path.read_bytes()),
        "source": f"audit_e2e/fixtures/development/templates/{spec['file']}",
        "format": spec["format"],
        "aspect_ratio": spec["geometry"],
        "source_slides": spec["prototype_slides"] + int(spec.get("second_master", False)),
        "layouts": spec["layouts"] * (2 if spec.get("second_master") else 1),
        "origin": "owned_native_synthetic",
        "family_id": f"owned-family:{spec['id']}",
        "related_family": spec.get("related_family", "independent_development_design"),
        "material_id": f"fixture:development:{spec['id']}:v1",
        "features": spec["features"],
        "coverage": spec["coverage"],
    }
    result = {
        "id": spec["id"],
        "suite": "development",
        "interface": "http",
        "slide_contract": {
            "target": slide_count,
            "minimum": slide_count,
            "maximum": slide_count,
            "request_kind": "explicit_count",
        },
        "template": metadata,
        "content_source": {
            "file": f"audit_e2e/fixtures/development/sources/{spec['id']}.md",
            "sha256": source_hash,
        },
        "content": content,
        "images": image_sources,
        "input_mode": "brief" if spec["id"] == "dev-short-brief" else "content",
        "image_presentation": "device" if spec["id"] == "dev-screenshot-portrait" else "plain",
        "slides": slide_count,
        "variants": ["executive", "analytical", "story"],
        "synthetic_template": {
            "kind": "owned_native",
            "asset": f"templates/{spec['file']}",
            "family_id": metadata["family_id"],
            "material_id": metadata["material_id"],
        },
        "reference": {
            "version": revision,
            "status": "silver",
            "provenance": {
                "method": "source_anchored_agent_authored",
                "review_status": "unreviewed",
            },
            "points": ledger,
            "requirements": [
                {
                    "id": "fixture-sample-exclusion",
                    "kind": "source-boundary",
                    "status": "silver",
                    "text": "Treat sample labels and values inside the owned template fixture as layout examples, not source facts.",
                }
            ],
        },
    }
    return result


def build_development_corpus(
    *,
    fixture_root=DEVELOPMENT_ROOT,
    registry_path=REGISTRY_FILE,
    source_root=SOURCE_ROOT,
    ledger_file=LEDGER_FILE,
    revision=None,
    replace=False,
):
    """Rebuild the corpus; changed fixture bytes require a new explicit revision."""
    fixture_root, registry_path = Path(fixture_root), Path(registry_path)
    source_root, ledger_file = Path(source_root), Path(ledger_file)
    revision_was_explicit = revision is not None
    revision = revision or DEFAULT_REVISION
    ledger_payload = json.loads(ledger_file.read_text(encoding="utf-8"))
    ledger_cases = ledger_payload["cases"]
    contents = {}
    for spec in _SPECS:
        case_id = spec["id"]
        source_path = source_root / f"{case_id}.md"
        content = source_path.read_text(encoding="utf-8")
        points = ledger_cases[case_id]
        if any(point["quote"] not in content for point in points):
            raise ValueError(f"Unanchored source ledger quote: {case_id}")
        if len([line for line in content.splitlines() if line.startswith("## ")]) < 5:
            raise ValueError(f"Development task needs at least five source groups: {case_id}")
        if any(len(line[3:]) > 140 for line in content.splitlines() if line.startswith("## ")):
            raise ValueError(f"Source heading exceeds 140 characters: {case_id}")
        contents[case_id] = (content, _hash(content.encode("utf-8")), points)

    with tempfile.TemporaryDirectory(prefix="development-corpus-") as temporary:
        staged_root = Path(temporary) / "development"
        staged_registry = Path(temporary) / "development-cases.json"
        (staged_root / "templates").mkdir(parents=True)
        (staged_root / "images").mkdir()
        cases = []
        for spec in _SPECS:
            _build_template(spec, staged_root / "templates" / spec["file"])
            if image_spec := spec.get("image"):
                _build_screenshot(image_spec, staged_root / "images" / image_spec["file"])
            content, source_hash, points = contents[spec["id"]]
            cases.append(_case_from_spec(spec, content, source_hash, points, staged_root, revision))
        registry = {
            "schema_version": 2,
            "reference_version": revision,
            "reference_provenance": {
                "method": "source_anchored_agent_authored",
                "review_status": "unreviewed",
            },
            "cases": cases,
        }
        staged_registry.write_text(
            json.dumps(registry, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        manifest = {
            "schema_version": 1,
            "license": "CC0-1.0; original synthetic fixtures authored for this corpus",
            "cases": [
                {
                    "id": case["id"],
                    "family_id": case["template"]["family_id"],
                    "material_id": case["template"]["material_id"],
                    "template_file": case["template"]["file"],
                    "template_sha256": case["template"]["sha256"],
                    "features": case["template"]["features"],
                    "source_file": case["content_source"]["file"],
                    "source_sha256": case["content_source"]["sha256"],
                    "images": case["images"],
                }
                for case in cases
            ],
        }
        staged_manifest = staged_root / "manifest.json"
        staged_manifest.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )

        old_registry = (
            json.loads(registry_path.read_text(encoding="utf-8"))
            if registry_path.is_file()
            else None
        )
        changed = old_registry is not None and old_registry != registry
        fixture_targets = []
        for folder in ("templates", "images"):
            for staged_path in sorted((staged_root / folder).iterdir()):
                target = fixture_root / folder / staged_path.name
                fixture_targets.append((staged_path, target))
                if target.is_file() and target.read_bytes() != staged_path.read_bytes():
                    changed = True
        manifest_target = fixture_root / "manifest.json"
        if (
            manifest_target.is_file()
            and manifest_target.read_bytes() != staged_manifest.read_bytes()
        ):
            changed = True
        if changed:
            if not replace or not revision_was_explicit:
                old_revision = (
                    old_registry.get("reference_version", "unknown") if old_registry else "unknown"
                )
                raise ValueError(
                    f"Corpus bytes differ from revision {old_revision}; pass --replace and a new --revision"
                )
            if old_registry and revision == old_registry.get("reference_version"):
                raise ValueError("Replacing corpus bytes requires a new reference revision")

        fixture_root.mkdir(parents=True, exist_ok=True)
        for staged_path, target in fixture_targets:
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(staged_path, target)
        shutil.copyfile(staged_manifest, manifest_target)
        registry_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(staged_registry, registry_path)
    return registry


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--replace", action="store_true")
    parser.add_argument("--revision")
    args = parser.parse_args()
    registry = build_development_corpus(replace=args.replace, revision=args.revision)
    print(f"Built {len(registry['cases'])} owned development cases at {REGISTRY_FILE}")


if __name__ == "__main__":
    main()
