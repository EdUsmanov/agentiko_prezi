"""Independent evidence extraction and deterministic export checks."""

from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata

from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
import pypdfium2 as pdfium

from .reporting import write_json

_NUMBER_WORDS = {
    "8": ("eight", "восемь"),
    "12": ("twelve", "двенадцать"),
    "40": ("forty", "сорок"),
}
_TEMPLATE_PREVIEW_CACHE = {}


def _hash(data):
    return sha256(data).hexdigest()


def _text(shape):
    if getattr(shape, "has_text_frame", False):
        return shape.text_frame.text.strip()
    return ""


def _hidden(shape):
    try:
        props = shape._element.xpath(".//p:cNvPr")
        return any(node.get("hidden") in {"1", "true"} for node in props)
    except (AttributeError, TypeError):
        return False


def _shape_objects(slide, width, height):
    objects = []

    def walk(shapes, grouped=False):
        for shape in shapes:
            yield shape, grouped
            if getattr(shape, "shape_type", None) == 6:
                yield from walk(shape.shapes, True)

    for shape, grouped in walk(slide.shapes):
        text = _text(shape)
        x, y, w, h = (getattr(shape, key, 0) / 914400 for key in ("left", "top", "width", "height"))
        hidden = _hidden(shape)
        out = (
            None
            if grouped
            else x < -0.01 or y < -0.01 or x + w > width + 0.02 or y + h > height + 0.02
        )
        obj = {
            "name": shape.name,
            "type": str(shape.shape_type),
            "text": text,
            "geometry_inches": {
                "x": round(x, 3),
                "y": round(y, 3),
                "width": round(w, 3),
                "height": round(h, 3),
            },
            "hidden": hidden,
            "out_of_bounds": out,
            "geometry_scope": "group_relative" if grouped else "slide",
        }
        if getattr(shape, "has_table", False):
            obj["table_rows"] = [
                [cell.text.strip() for cell in row.cells] for row in shape.table.rows
            ]
            obj["text"] = "\n".join(" | ".join(row) for row in obj["table_rows"])
        if getattr(shape, "has_chart", False):
            chart = shape.chart
            categories = []
            rows = []
            try:
                categories = [str(item.label) for item in chart.plots[0].categories]
            except (AttributeError, IndexError, TypeError):
                pass
            series = [{"name": s.name, "values": list(s.values)} for s in chart.series]
            for item in series:
                rows.extend(
                    {"category": category, "series": item["name"], "value": value}
                    for category, value in zip(categories, item["values"])
                )
            axis_titles = {}
            for key in ("category_axis", "value_axis"):
                try:
                    axis = getattr(chart, key)
                    if axis.has_title:
                        axis_titles[key] = axis.axis_title.text_frame.text.strip()
                except (AttributeError, ValueError):
                    pass
            obj["chart"] = {
                "title": chart.chart_title.text_frame.text.strip() if chart.has_title else "",
                "categories": categories,
                "axis_titles": axis_titles,
                "series": series,
                "category_series_values": rows,
            }
            obj["text"] = "\n".join(
                filter(
                    None,
                    [
                        obj["text"],
                        obj["chart"]["title"],
                        *axis_titles.values(),
                        *categories,
                        *(str(s["name"]) for s in series),
                        *(str(v) for s in series for v in s["values"]),
                    ],
                )
            )
        if getattr(shape, "shape_type", None) == 13:
            try:
                blob = shape.image.blob
                with Image.open(BytesIO(blob)) as image:
                    pixels = image.convert("RGBA")
                    pixel_hash = _hash(
                        f"{pixels.width}x{pixels.height}".encode() + pixels.tobytes()
                    )
                obj["image"] = {
                    "sha256": _hash(blob),
                    "pixels_sha256": pixel_hash,
                    "size": [pixels.width, pixels.height],
                    "content_type": shape.image.content_type,
                }
            except (AttributeError, ValueError):
                pass
        for paragraph in getattr(getattr(shape, "text_frame", None), "paragraphs", []):
            if paragraph.font.size:
                obj.setdefault("font_sizes_pt", []).append(round(paragraph.font.size.pt, 1))
            for run in paragraph.runs:
                if run.font.size:
                    obj.setdefault("font_sizes_pt", []).append(round(run.font.size.pt, 1))
        obj["visibility"] = (
            "hidden"
            if hidden
            else "uncertain_group_geometry"
            if grouped
            else "out_of_bounds"
            if out
            else "visible"
        )
        objects.append(obj)
    return objects


def _notes(slide):
    try:
        return slide.notes_slide.notes_text_frame.text.strip()
    except (AttributeError, KeyError):
        return ""


def _pdf_visible_text(page):
    """Collect page text whose origin falls inside the exported PDF page box."""
    width, height = float(page.mediabox.width), float(page.mediabox.height)
    visible = []

    def visitor(text, cm, tm, _font, _size):
        try:
            x, y = float(cm[4]) + float(tm[4]), float(cm[5]) + float(tm[5])
            if 0 <= x <= width and 0 <= y <= height:
                visible.append(text)
        except (IndexError, TypeError, ValueError):
            return

    try:
        page.extract_text(visitor_text=visitor)
    except TypeError:
        return page.extract_text() or ""
    return "".join(visible)


def _render_pdf(pdf_path, output, *, pages=None, prefix="slide"):
    document = pdfium.PdfDocument(str(pdf_path))
    count = len(document)
    selected = range(count) if pages is None else [p for p in pages if 0 <= p < count]
    paths = []
    for index in selected:
        bitmap = document[index].render(scale=1.5)
        image = bitmap.to_pil().convert("RGB")
        destination = output / f"{prefix}-{index + 1}.png"
        image.save(destination, format="PNG")
        paths.append((index + 1, destination))
    document.close()
    return count, paths


def _safe_copy(source, destination):
    source, destination = Path(source), Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return destination


def _relative(path, root):
    return Path(path).relative_to(root).as_posix()


def _template_evidence(case, bundle_dir):
    from studio.composition.office import to_pdf

    root = bundle_dir / "template"
    root.mkdir(parents=True, exist_ok=True)
    origin = case.get("template_origin", "real_external_template")
    metadata = {
        "origin": origin,
        "source_origin": "synthetic_analog" if case.get("synthetic") else "external_template",
        "source": case.get("template_source", {}).get("source"),
        "format": case.get("template_source", {}).get("format"),
        "sha256": case.get("source_hashes", {}).get(
            "template", _hash(Path(case["template"]).read_bytes())
        ),
        "registered_sha256": case.get("registry_hashes", {}).get(
            "template", case.get("template_source", {}).get("sha256")
        ),
        "source_slides": case.get("template_source", {}).get("source_slides"),
        "layouts": case.get("template_source", {}).get("layouts"),
        "fidelity_applicability": case.get(
            "template_fidelity_applicability",
            "not_applicable_synthetic_template" if case.get("synthetic") else "external_template",
        ),
        "previews": [],
        "preview_status": "unavailable",
    }
    selected = case.get("template_preview_slides", [1])[:4]
    if not selected:
        metadata["preview_status"] = "not_requested"
        return metadata
    cache_key = (_hash(Path(case["template"]).read_bytes()), tuple(selected))
    if cache_key in _TEMPLATE_PREVIEW_CACHE:
        cached = _TEMPLATE_PREVIEW_CACHE[cache_key]
        metadata["rendered_slide_count"] = cached["rendered_slide_count"]
        for number, data in cached["previews"]:
            destination = root / f"source-slide-{number}.png"
            destination.write_bytes(data)
            metadata["previews"].append(
                {"source_slide": number, "image": _relative(destination, bundle_dir)}
            )
        metadata["preview_status"] = "available" if metadata["previews"] else "unavailable"
        return metadata
    try:
        with tempfile.TemporaryDirectory(prefix="eval-template-") as temp:
            if not to_pdf(case["template"], Path(temp)):
                return metadata
            source_pdf = Path(temp) / (Path(case["template"]).stem + ".pdf")
            count, rendered = _render_pdf(
                source_pdf, root, pages=[n - 1 for n in selected], prefix="source-slide"
            )
            metadata["rendered_slide_count"] = count
            metadata["previews"] = [
                {"source_slide": number, "image": _relative(path, bundle_dir)}
                for number, path in rendered
            ]
            metadata["preview_status"] = "available" if rendered else "unavailable"
            _TEMPLATE_PREVIEW_CACHE[cache_key] = {
                "rendered_slide_count": count,
                "previews": [(number, path.read_bytes()) for number, path in rendered],
            }
    except Exception as exc:
        metadata["preview_error"] = type(exc).__name__
    return metadata


def build_bundle(case, result_dir, bundle_dir):
    """Copy final exports and extract visible PPTX/PDF/template evidence."""
    result_dir, bundle_dir = Path(result_dir), Path(bundle_dir)
    bundle_dir.mkdir(parents=True, exist_ok=True)
    source_path = bundle_dir / "input" / "source.md"
    source_path.parent.mkdir(parents=True, exist_ok=True)
    content = case["content"]
    source_path.write_text(content, encoding="utf-8")
    source = {
        "text": content,
        "sha256": case.get("source_hashes", {}).get("content", _hash(content.encode())),
        "path": _relative(source_path, bundle_dir),
        "input_mode": case.get("input_mode", "content"),
        "images": [],
    }
    for image, record in zip(case.get("images", []), case.get("images_source", [])):
        target = _safe_copy(image, bundle_dir / "input" / Path(image).name)
        with Image.open(target) as source_image:
            pixels = source_image.convert("RGBA")
            pixel_hash = _hash(f"{pixels.width}x{pixels.height}".encode() + pixels.tobytes())
            size = [pixels.width, pixels.height]
        source["images"].append(
            {
                "file": record["file"],
                "sha256": _hash(target.read_bytes()),
                "registered_sha256": record["sha256"],
                "pixels_sha256": pixel_hash,
                "size": size,
                "path": _relative(target, bundle_dir),
                "origin": case.get("image_origin"),
            }
        )
    if case.get("input_mode") == "brief":
        draft_path, approval_path = result_dir / "draft.json", result_dir / "approval.json"
        if not draft_path.is_file() or not approval_path.is_file():
            raise FileNotFoundError(
                "Brief evidence requires the saved draft and UI approval records"
            )
        draft_record = json.loads(draft_path.read_text(encoding="utf-8"))
        approval_record = json.loads(approval_path.read_text(encoding="utf-8"))
        package_hash = draft_record.get("package_hash")
        draft_hash = draft_record.get("draft_hash")
        package_approved = approval_record.get("approved_package_hash")
        draft_approved = approval_record.get("approved_draft_hash")
        if (
            not package_hash
            or not draft_hash
            or (package_hash, draft_hash) != (package_approved, draft_approved)
        ):
            raise ValueError("Saved UI approval hashes do not match the saved draft")
        proposals = []
        for slide_number, slide in enumerate(draft_record.get("draft", {}).get("slides", []), 1):
            for bullet in slide.get("bullets", []):
                if bullet.get("proposed"):
                    proposals.append(
                        {
                            "slide_number": slide_number,
                            "text": bullet.get("text", ""),
                            "origin": "model_proposal",
                            "fact_ids": bullet.get("fact_ids", []),
                        }
                    )
        source["approved_proposals"] = proposals
        preapproval_path = result_dir / "preapproval-generation.json"
        preapproval = (
            json.loads(preapproval_path.read_text(encoding="utf-8"))
            if preapproval_path.is_file()
            else {}
        )
        source["brief_approval"] = {
            "status": "approved",
            "package_hash": package_hash,
            "draft_hash": draft_hash,
            "approved_package_hash": package_approved,
            "approved_draft_hash": draft_approved,
            "preapproval_generation_blocked": preapproval.get("status") == 409,
            "preapproval_generation_http_status": preapproval.get("status"),
        }

    variants = {}
    for key in case.get("variants", []):
        variant_source = result_dir / key
        pptx_source, pdf_source = variant_source / "deck.pptx", variant_source / "deck.pdf"
        if not pptx_source.is_file() or not pdf_source.is_file():
            raise FileNotFoundError(f"Missing final exported PPTX/PDF for {case['id']}/{key}")
        pptx = _safe_copy(pptx_source, bundle_dir / "variants" / key / "deck.pptx")
        pdf = _safe_copy(pdf_source, bundle_dir / "variants" / key / "deck.pdf")
        presentation = Presentation(str(pptx))
        pdf_reader = PdfReader(str(pdf))
        text_pages = [
            (page.extract_text() or "", _pdf_visible_text(page)) for page in pdf_reader.pages
        ]
        image_dir = bundle_dir / "variants" / key
        pdf_pages, rendered = _render_pdf(pdf, image_dir)
        slides = []
        for i, slide in enumerate(presentation.slides):
            objects = _shape_objects(
                slide, presentation.slide_width / 914400, presentation.slide_height / 914400
            )
            rendered_path = next((p for page, p in rendered if page == i + 1), None)
            slides.append(
                {
                    "number": i + 1,
                    "image": _relative(rendered_path, bundle_dir) if rendered_path else None,
                    "text": "\n".join(
                        obj["text"]
                        for obj in objects
                        if obj["text"] and obj["visibility"] == "visible"
                    ),
                    "pdf_text": text_pages[i][0] if i < len(text_pages) else "",
                    "pdf_visible_text": text_pages[i][1] if i < len(text_pages) else "",
                    "notes": _notes(slide),
                    "objects": objects,
                }
            )
        variants[key] = {
            "pptx": _relative(pptx, bundle_dir),
            "pdf": _relative(pdf, bundle_dir),
            "pptx_slide_count": len(presentation.slides),
            "pdf_page_count": len(pdf_reader.pages),
            "slides": slides,
        }

    slide_contract = case.get("slide_contract") or {
        "target": case.get("slides"),
        "minimum": case.get("slides"),
        "maximum": case.get("slides"),
        "request_kind": "explicit_count",
    }
    bundle = {
        "schema_version": 1,
        "case_id": case["id"],
        "expected_variant_ids": list(case.get("variants", [])),
        "expected_slide_count": case.get("slides"),
        "slide_contract": slide_contract,
        "source": source,
        "reference": case.get(
            "reference",
            {"version": "unversioned", "status": "silver", "points": [], "requirements": []},
        ),
        "variants": variants,
        "template": _template_evidence(case, bundle_dir),
    }
    write_json(bundle_dir / "bundle.json", bundle)
    return bundle


def _norm(value):
    return " ".join(unicodedata.normalize("NFKC", str(value)).casefold().split())


def _bundle_file(bundle, bundle_dir, key):
    path = (Path(bundle_dir) / bundle[key]).resolve()
    if not path.is_relative_to(Path(bundle_dir).resolve()):
        raise ValueError("Evidence path escaped its bundle")
    return path


def _visible_text(deck):
    return "\n".join(slide.get("text", "") for slide in deck.get("slides", []))


def _number_present(text, value):
    value = str(value)
    if re.search(rf"(?<!\d){re.escape(value)}(?!\d)", text):
        return True
    return any(
        re.search(rf"(?<!\w){re.escape(word)}(?!\w)", text, re.IGNORECASE)
        for word in _NUMBER_WORDS.get(value, ())
    )


def audit_bundle(bundle, bundle_dir):
    """Run deterministic file, number, table, visibility and deck-diversity checks."""
    bundle_dir = Path(bundle_dir)
    findings = []

    def add(category, message, *, severity="failed", **evidence):
        findings.append(
            {"category": category, "severity": severity, "message": message, "evidence": evidence}
        )

    if bundle.get("schema_version") != 1:
        return {
            "status": "inconclusive",
            "findings": [
                {
                    "category": "omission",
                    "severity": "inconclusive",
                    "message": "Unsupported evidence bundle schema",
                    "evidence": {},
                }
            ],
            "diagnostics": [],
        }
    variants = bundle.get("variants", {})
    expected_variants = set(bundle.get("expected_variant_ids", []))
    expected_slides = bundle.get("expected_slide_count")
    slide_contract = bundle.get("slide_contract")
    if slide_contract is None and type(expected_slides) is int:
        slide_contract = {
            "target": expected_slides,
            "minimum": expected_slides,
            "maximum": expected_slides,
            "request_kind": "explicit_count",
        }
    valid_contract = isinstance(slide_contract, dict) and all(
        type(slide_contract.get(key)) is int for key in ("target", "minimum", "maximum")
    )
    if valid_contract:
        contract_target = slide_contract["target"]
        contract_minimum = slide_contract["minimum"]
        contract_maximum = slide_contract["maximum"]
        request_kind = slide_contract.get("request_kind")
        valid_request = (
            request_kind == "browser_preset"
            and slide_contract.get("preset") == "mini"
            and (contract_minimum, contract_maximum) == (3, 5)
        ) or (
            request_kind == "explicit_count"
            and contract_minimum == contract_target == contract_maximum
        )
        valid_contract = (
            1 <= contract_minimum <= contract_target <= contract_maximum <= 30
            and contract_target == expected_slides
            and valid_request
        )
    if not valid_contract:
        add("omission", "Invalid or missing slide-count contract")
        target_slides, minimum_slides, maximum_slides = expected_slides, 1, 30
    else:
        target_slides = slide_contract["target"]
        minimum_slides = slide_contract["minimum"]
        maximum_slides = slide_contract["maximum"]
    if expected_variants and set(variants) != expected_variants:
        add(
            "omission",
            "Exported variant ids do not match the case contract",
            expected=sorted(expected_variants),
            actual=sorted(variants),
        )
    if not variants:
        add("omission", "No exported variants in the evidence bundle")
    texts, slide_objects, all_notes = {}, {}, {}
    pixel_decks = {}
    diagnostics = []
    for name, deck in variants.items():
        try:
            pptx = _bundle_file(deck, bundle_dir, "pptx")
            pdf = _bundle_file(deck, bundle_dir, "pdf")
            presentation = Presentation(str(pptx))
            reader = PdfReader(str(pdf))
            pptx_count, pdf_count = len(presentation.slides), len(reader.pages)
            if pdf_count != pptx_count or pptx_count != deck.get("pptx_slide_count"):
                add(
                    "omission",
                    "PPTX/PDF page counts differ",
                    variant=name,
                    pptx_pages=pptx_count,
                    pdf_pages=pdf_count,
                )
            if pptx_count < minimum_slides or pptx_count > maximum_slides:
                add(
                    "omission",
                    "Exported slide count is outside the case contract",
                    variant=name,
                    target=target_slides,
                    allowed_range=[minimum_slides, maximum_slides],
                    actual=pptx_count,
                )
            elif pptx_count != target_slides:
                diagnostics.append(
                    {
                        "category": "slide_count_target",
                        "message": "Exported slide count is within the browser preset range but differs from the target.",
                        "evidence": {
                            "variant": name,
                            "target": target_slides,
                            "actual": pptx_count,
                            "allowed_range": [minimum_slides, maximum_slides],
                            "preset": slide_contract.get("preset"),
                        },
                    }
                )
            if not deck.get("slides") or any(
                not slide.get("image") for slide in deck.get("slides", [])
            ):
                add("omission", "One or more exported slides lack a rendered image", variant=name)
            objects_by_slide, notes, visible = [], [], []
            for i, slide in enumerate(presentation.slides):
                objects = _shape_objects(
                    slide, presentation.slide_width / 914400, presentation.slide_height / 914400
                )
                objects_by_slide.append(objects)
                notes.append(_notes(slide))
                visible.append(
                    "\n".join(
                        obj["text"]
                        for obj in objects
                        if obj["text"] and obj["visibility"] == "visible"
                    )
                )
            pdf_text = "\n".join(_pdf_visible_text(page) for page in reader.pages)
            texts[name] = "\n".join([*visible, pdf_text])
            slide_objects[name] = objects_by_slide
            all_notes[name] = "\n".join(notes)
            pixels = []
            for slide in deck.get("slides", []):
                path = _bundle_file({"image": slide["image"]}, bundle_dir, "image")
                with Image.open(path) as image:
                    pixels.append((image.size, image.convert("RGB").tobytes()))
            pixel_decks[name] = pixels
        except Exception as exc:
            add(
                "omission",
                "Export evidence is missing or unreadable",
                variant=name,
                error=type(exc).__name__,
            )
            texts[name] = ""
    if len(pixel_decks) > 1:
        names = list(pixel_decks)
        identical = [
            (a, b)
            for i, a in enumerate(names)
            for b in names[i + 1 :]
            if pixel_decks[a] == pixel_decks[b]
        ]
        if identical:
            add(
                "duplicates",
                "Exported variants are pixel-identical",
                pairs=[list(pair) for pair in identical],
            )

    reference = bundle.get("reference", {})
    ref_is_gold = (
        reference.get("status") == "gold"
        and reference.get("provenance", {}).get("review_status") == "approved"
    )
    for name, deck in variants.items():
        all_text = texts.get(name, "")
        visible = _norm(all_text)
        all_objects = [obj for slide in slide_objects.get(name, []) for obj in slide]
        grouped = _norm(
            "\n".join(
                obj.get("text", "")
                for obj in all_objects
                if obj.get("geometry_scope") == "group_relative"
            )
        )
        hidden = _norm(
            "\n".join(
                [
                    all_notes.get(name, ""),
                    *(
                        obj.get("text", "")
                        for obj in all_objects
                        if obj.get("hidden") or obj.get("out_of_bounds")
                    ),
                ]
            )
        )
        for point in reference.get("points", []):
            if point.get("required") is not True:
                continue
            quote = _norm(point.get("quote", ""))
            if quote and quote in grouped and quote not in visible:
                add(
                    "hidden_text",
                    "A required fact is in a grouped shape whose slide position is relative",
                    severity="inconclusive",
                    variant=name,
                    point_id=point.get("id"),
                )
            if quote and quote in hidden and quote not in visible:
                add(
                    "hidden_text",
                    "A source fact appears only in notes, hidden text, or off-slide text",
                    severity="failed" if ref_is_gold else "inconclusive",
                    variant=name,
                    point_id=point.get("id"),
                    quote=point.get("quote"),
                )
            for number in point.get("numbers", []):
                value = str(number.get("value", ""))
                if value and not _number_present(all_text, value):
                    add(
                        "number",
                        "A required source number is absent from visible exported text",
                        severity=(
                            "inconclusive"
                            if _number_present(grouped, value)
                            else "failed"
                            if ref_is_gold
                            else "inconclusive"
                        ),
                        variant=name,
                        point_id=point.get("id"),
                        expected=value,
                        unit=number.get("unit"),
                    )
        for requirement in reference.get("requirements", []):
            if requirement.get("applicability") == "not_applicable_synthetic_template":
                continue
            if requirement.get("kind") == "table":
                expected_headers = requirement.get("headers", [])
                visible_tables = [
                    obj.get("table_rows", [])
                    for slide in slide_objects.get(name, [])
                    for obj in slide
                    if obj.get("table_rows") and obj.get("visibility") == "visible"
                ]
                hidden_tables = [
                    obj
                    for slide in slide_objects.get(name, [])
                    for obj in slide
                    if obj.get("table_rows")
                    and obj.get("visibility") in {"hidden", "out_of_bounds"}
                ]
                grouped_tables = [
                    obj
                    for slide in slide_objects.get(name, [])
                    for obj in slide
                    if obj.get("table_rows")
                    and obj.get("geometry_scope") == "group_relative"
                    and not obj.get("hidden")
                ]
                charts = [
                    obj.get("chart")
                    for slide in slide_objects.get(name, [])
                    for obj in slide
                    if obj.get("chart") and obj.get("visibility") == "visible"
                ]
                if visible_tables:
                    actual = [[_norm(c) for c in row] for table in visible_tables for row in table]
                    headers_missing = (
                        bool(expected_headers)
                        and [_norm(c) for c in expected_headers] not in actual
                    )
                    rows_missing = [
                        [_norm(c) for c in row]
                        for row in requirement.get("rows", [])
                        if [_norm(c) for c in row] not in actual
                    ]
                    if headers_missing and not rows_missing:
                        add(
                            "units",
                            "A source table unit/header label is absent or changed",
                            variant=name,
                            requirement_id=requirement.get("id"),
                            expected=expected_headers[-1],
                        )
                    if rows_missing:
                        add(
                            "table_binding",
                            "An expected source table row or its label/value binding is missing",
                            variant=name,
                            requirement_id=requirement.get("id"),
                        )
                    if headers_missing and rows_missing:
                        add(
                            "table_binding",
                            "An expected source table header and row binding are missing",
                            variant=name,
                            requirement_id=requirement.get("id"),
                        )
                elif hidden_tables:
                    add(
                        "hidden_text",
                        "The source table appears only in a hidden or off-slide shape",
                        severity="failed"
                        if requirement.get("status") == "gold"
                        else "inconclusive",
                        variant=name,
                        requirement_id=requirement.get("id"),
                    )
                elif grouped_tables:
                    add(
                        "table_binding",
                        "The source table is inside a group whose slide position is uncertain",
                        severity="inconclusive",
                        variant=name,
                        requirement_id=requirement.get("id"),
                    )
                elif charts:
                    add(
                        "table_binding",
                        "Chart evidence cannot confirm every source table row binding",
                        severity="inconclusive",
                        variant=name,
                        requirement_id=requirement.get("id"),
                    )
                elif any(
                    _norm(cell) not in _norm(all_text)
                    for row in requirement.get("rows", [])
                    for cell in row
                ):
                    add(
                        "table_binding",
                        "Source table cells are absent and row binding cannot be confirmed",
                        severity="failed"
                        if requirement.get("status") == "gold"
                        else "inconclusive",
                        variant=name,
                        requirement_id=requirement.get("id"),
                    )
                else:
                    add(
                        "table_binding",
                        "Source values appear, but row associations cannot be confirmed",
                        severity="inconclusive",
                        variant=name,
                        requirement_id=requirement.get("id"),
                    )
            elif requirement.get("kind") == "image":
                for source_image in bundle.get("source", {}).get("images", []):
                    embedded = [
                        obj.get("image", {})
                        for slide in slide_objects.get(name, [])
                        for obj in slide
                        if obj.get("image")
                    ]
                    if not any(
                        item.get("pixels_sha256") == source_image.get("pixels_sha256")
                        for item in embedded
                    ):
                        add(
                            "omission",
                            "A required input image is absent or its pixels changed in the exported PPTX",
                            variant=name,
                            image=source_image.get("file"),
                            expected_pixels_sha256=source_image.get("pixels_sha256"),
                        )
            elif requirement.get("kind") == "visual":
                tiny = [
                    obj.get("name")
                    for slide in slide_objects.get(name, [])
                    for obj in slide
                    if obj.get("text") and any(size < 6 for size in obj.get("font_sizes_pt", []))
                ]
                minimum_width = requirement.get("minimum_image_width_inches")
                small_images = [
                    obj.get("name")
                    for slide in slide_objects.get(name, [])
                    for obj in slide
                    if minimum_width
                    and obj.get("image")
                    and obj.get("geometry_inches", {}).get("width", 0) < minimum_width
                ]
                if tiny or small_images:
                    add(
                        "readability",
                        "Required text or an image falls below its explicit readability threshold",
                        variant=name,
                        text_objects=tiny,
                        image_objects=small_images,
                        minimum_image_width_inches=minimum_width,
                    )

    template = bundle.get("template", {})
    if template.get("fidelity_applicability") == "external_template" and not template.get(
        "previews"
    ):
        add(
            "template_fidelity",
            "External source template previews are unavailable",
            severity="inconclusive",
            source=template.get("source"),
        )
    if any(item.get("severity") == "failed" for item in findings):
        status = "failed"
    elif findings:
        status = "inconclusive"
    else:
        status = "passed"
    return {"status": status, "findings": findings, "diagnostics": diagnostics}
