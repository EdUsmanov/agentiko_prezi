"""Seeded synthetic inputs, unrelated to organizer/customer templates."""

from datetime import datetime
from pathlib import Path
from random import Random
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image

BRIEF = "# Городской проект\nКоманда создаёт общий сервис.\nПилот длится 12 недель.\nРезультат проверяют по 40 заявкам."


def make_template(path: Path, *, seed=17, aspect="wide", columns=1, dark=False):
    random = Random(seed)
    prs = Presentation()
    width, height = (13.333, 7.5) if aspect == "wide" else (10, 7.5)
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    prs.core_properties.created = prs.core_properties.modified = datetime(2020, 1, 1)
    margin = 0.6 + random.choice([0, 0.05, 0.1])
    for index in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("172B46" if dark else "FFFFFF")
        title = slide.shapes.add_textbox(
            Inches(margin), Inches(0.5), Inches(width - 2 * margin), Inches(1.1)
        )
        fields = [(title, 36, "SYNTHETIC TEMPLATE TITLE")]
        for col in range(columns):
            gap = 0.3
            span = (width - 2 * margin - gap * (columns - 1)) / columns
            body = slide.shapes.add_textbox(
                Inches(margin + col * (span + gap)), Inches(2), Inches(span), Inches(4.5)
            )
            fields.append((body, 22, "SYNTHETIC OLD CONTENT"))
        for shape, size, text in fields:
            run = shape.text_frame.paragraphs[0].add_run()
            run.text, run.font.name, run.font.size = text, "Play", Pt(size)
            run.font.color.rgb = RGBColor.from_string("FFFFFF" if dark else "154A67")
    raw = path.with_suffix(".building.pptx")
    prs.save(raw)
    # Stable ZIP metadata permits exact input hashes across runs and machines.
    with ZipFile(raw) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if path.suffix.lower() == ".potx" and name == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.presentation.main+xml", b"presentationml.template.main+xml"
                )
            info = ZipInfo(name, (2020, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, data)
    raw.unlink()
    return path


def make_image(path: Path):
    Image.new("RGB", (160, 100), (25, 95, 120)).save(path)
    return path
