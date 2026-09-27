from zipfile import ZipFile
from pptx import Presentation
from pptx.util import Pt
import pytest
from studio.powerpoint import open_presentation
from studio.security import InputRejected, PPTX_MAIN, POTX_MAIN, presentation_content_type
from studio.template import analyze_template
from studio.analysis import template_inventory


def test_real_potx_open_preserves_source_and_design(potx, template, tmp_path):
    original = potx.read_bytes()
    with ZipFile(potx) as archive:
        assert presentation_content_type(archive) == POTX_MAIN
    with pytest.raises(ValueError):
        Presentation(potx)  # Guard against a merely renamed PPTX fixture.
    prs = open_presentation(potx)
    expected = Presentation(template)
    assert prs.part.content_type == PPTX_MAIN
    assert len(prs.slides) == len(expected.slides)
    assert len(prs.slide_layouts) == len(expected.slide_layouts)
    assert prs.slide_width == expected.slide_width
    assert prs.slides[0].shapes[0].text == expected.slides[0].shapes[0].text
    normalized = tmp_path / "normalized.pptx"
    prs.save(normalized)
    with ZipFile(potx) as source, ZipFile(normalized) as output:
        assert source.read("ppt/theme/theme1.xml") == output.read("ppt/theme/theme1.xml")
    profile = analyze_template(potx, tmp_path / "assets")
    inventory = template_inventory(potx, profile)
    assert len(inventory["slides"]) == 3
    assert potx.read_bytes() == original


def test_layout_only_potx_supported(tmp_path, template):
    prs = Presentation(template)
    for entry in list(prs.slides._sldIdLst):
        prs.part.drop_rel(entry.rId)
        prs.slides._sldIdLst.remove(entry)
    for master in prs.slide_masters:
        for surface in [master, *master.slide_layouts]:
            for shape in surface.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "Play"
                            run.font.size = Pt(20)
    path = tmp_path / "empty.pptx"
    prs.save(path)
    target = tmp_path / "layout-only.potx"
    with ZipFile(path) as source, ZipFile(target, "w") as output:
        for entry in source.infolist():
            data = source.read(entry)
            if entry.filename == "[Content_Types].xml":
                data = data.replace(PPTX_MAIN.encode(), POTX_MAIN.encode())
            output.writestr(entry, data)
    profile = analyze_template(target, tmp_path / "layout-assets")
    assert profile.slide_count == 0 and profile.layout_count > 0
    assert any(p.title_zone and p.body_zones for p in profile.patterns)
    assert template_inventory(target, profile)["slides"] == []
    from studio.models import PreparedPackage, Constraints
    from studio.content import parse_content
    from studio.planner import extractive_plans
    from studio.composer import compose_variant
    from studio.render import render_pptx

    package = PreparedPackage(
        id="test",
        created_at="test",
        template=profile,
        content=parse_content("# Проект\nКоманда работает с заявками."),
        constraints=Constraints(slides=1),
        manifest={},
    )
    scenes = compose_variant(extractive_plans(package).variants[0], package)
    output = tmp_path / "from-layout-only.pptx"
    render_pptx(scenes, profile, target, output)
    assert len(Presentation(output).slides) == 1


@pytest.mark.parametrize(
    "main_type",
    [
        "application/vnd.ms-powerpoint.template.macroEnabled.main+xml",
        "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    ],
)
def test_disguised_unsupported_main_type_rejected(potx, tmp_path, main_type):
    target = tmp_path / "disguised.potx"
    with ZipFile(potx) as source, ZipFile(target, "w") as output:
        for entry in source.infolist():
            raw = source.read(entry)
            if entry.filename == "[Content_Types].xml":
                raw = raw.replace(POTX_MAIN.encode(), main_type.encode())
            output.writestr(entry, raw)
    with pytest.raises(InputRejected, match="PPTX и POTX"):
        open_presentation(target)


def test_index_potx_preserves_original_container(potx, tmp_path):
    from studio.examples import index_examples
    from studio.config import Settings

    settings = Settings(data_dir=tmp_path / "index")
    indexed = index_examples([potx], settings)
    assert indexed[0]["name"] == "template.potx"
    stored = settings.data_dir / "references" / indexed[0]["id"] / "input.pptx"
    assert stored.read_bytes() == potx.read_bytes()
    assert len(open_presentation(stored).slides) == 3
