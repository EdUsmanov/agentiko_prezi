"""Marker measurement and repair hints must agree with editable PPTX output."""

from copy import deepcopy
from copy import copy
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.util import Pt

from studio.config import ROOT
from studio.editorial_domain import EditorialPlan
from studio.editorial_patch_validation import shortening_contracts, validate_contracts
from studio.fonts import text_width, wrap_text
from studio.models import Box, Element, Fact, Finding, SlideScene
from studio.pptx_text import set_text
from studio.repair_policy import scene_fit_feedback, shortening_target
from studio.text_composer import fact_elements


@pytest.fixture
def profile():
    return SimpleNamespace(
        font="Montserrat",
        font_file=str(ROOT / "fonts/Montserrat-Regular.ttf"),
        font_roles={},
        font_assets=[],
        body_size=18,
        font_sizes=[12, 14, 15, 16, 18],
        foreground="#222222",
    )


@pytest.mark.parametrize("explicit_list", [False, True])
@pytest.mark.parametrize("text", ["Материалы для вашей команды", "Materials for your team"])
def test_single_paragraph_uses_full_width_but_explicit_list_keeps_marker(
    profile, text, explicit_list
):
    width = text_width(text, profile.font_file, 16) + 1
    box = Box(x=20, y=30, w=width, h=23)
    fact = Fact(id="summary-1-1", text=text, list_item=explicit_list)
    element = fact_elements([fact], box, profile, profile.foreground)[0]
    assert element.bullet is explicit_list
    assert element.text == text and element.source_ids == [fact.id]
    assert (element.size >= 16) is not explicit_list
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Pt(box.x), Pt(box.y), Pt(box.w), Pt(box.h))
    set_text(shape.text_frame, element.text, element, profile)
    assert bool(shape._element.xpath(".//a:buChar")) is explicit_list
    assert not shape._element.xpath(".//a:br")


def test_implicit_list_and_source_or_heading_keep_their_own_marker_policy(profile):
    texts = ["First ordinary item", "Source: verified material", "Section:", "Second item"]
    facts = [Fact(id=f"f{i}", text=t) for i, t in enumerate(texts)]
    elements = fact_elements(facts, Box(x=0, y=0, w=400, h=200), profile, "#222222")
    assert [e.bullet for e in elements] == [True, False, False, True]
    assert [e.text for e in elements] == texts
    assert all(e.size == 18 for e in elements)
    assert all(a.box.y + a.box.h <= b.box.y for a, b in zip(elements, elements[1:]))


def test_grouped_sentences_remain_one_unbulleted_paragraph(profile):
    facts = [Fact(id="f1", text="First sentence.", line=1), Fact(id="f2", text="Next.", line=1)]
    elements = fact_elements(facts, Box(x=0, y=0, w=400, h=100), profile, "#222222")
    assert len(elements) == 1 and not elements[0].bullet
    assert elements[0].text == "First sentence. Next."
    assert elements[0].source_ids == ["f1", "f2"]


def test_bold_authored_field_matches_native_line_breaks(profile):
    bold_path = str(ROOT / "fonts/Montserrat-Bold.ttf")
    profile.font_assets = [{"id": "bold", "requested": "Montserrat Bold", "path": bold_path}]
    profile.model_copy = lambda: copy(profile)
    text = "A bold paragraph near the field edge"
    box = Box(x=20, y=30, w=text_width(text, bold_path, 18) * 0.97, h=22.5)
    element = fact_elements(
        [Fact(id="f1", text=text)],
        box,
        profile,
        "#222222",
        field_style={"family": "Montserrat", "bold": True, "size": 18},
    )[0]
    assert element.bold and element.size >= 16 and not element.bullet
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Pt(box.x), Pt(box.y), Pt(box.w), Pt(box.h))
    set_text(shape.text_frame, element.text, element, profile)
    assert not shape._element.xpath(".//a:br")
    assert shape.text_frame.paragraphs[0].runs[0].font.bold


@pytest.mark.parametrize("bullet", [False, True])
@pytest.mark.parametrize("bold", [False, True])
def test_shortening_hint_uses_font_and_actual_marker(profile, bullet, bold):
    text = "A longer explanation needs a smaller field"
    element = Element(
        kind="text",
        text=text,
        box=Box(x=0, y=0, w=170, h=20),
        size=12,
        bullet=bullet,
        bold=bold,
        source_ids=["summary-1-1"],
    )
    target = shortening_target(element, profile)
    width = (170 - (16 * 1.4 if bullet else 0)) * (0.94 if bold else 1)
    assert 0 < target < len(text)
    assert len(wrap_text(text[:target], profile.font_file, 16, width)) == 1
    assert text_width(text[:target], profile.font_file, 16) <= width


def feedback_for(profile, elements, bad_index=0):
    scene = SlideScene(
        title="Topic", background="#FFFFFF", source_ids=[], layout="columns", elements=elements
    )
    row = scene_fit_feedback(
        scene,
        [Finding(code="readability", severity="warning", message="Small", element=bad_index)],
        1,
        {"summary-1-1", "summary-1-2"},
        profile=profile,
    )
    return {"repair_issues": row["repair_issues"], "fields": [row]}


def plan(texts):
    return EditorialPlan.model_validate(
        {
            "slides": [
                {
                    "title": "Topic",
                    "purpose": "content",
                    "bullets": [
                        {"text": t, "evidence": [{"fact_id": f"f{i}"}]}
                        for i, t in enumerate(texts, 1)
                    ],
                }
            ]
        }
    ).model_dump()


def test_single_field_contract_cannot_offer_whole_slide_budget(profile):
    text = "A long description of the service"
    element = Element(
        kind="text", text=text, box=Box(x=0, y=0, w=180, h=20), size=12, source_ids=["summary-1-1"]
    )
    feedback = feedback_for(profile, [element])
    contract = shortening_contracts(plan([text]), [1], feedback, 500)[0]
    target = feedback["fields"][0]["fields"][0]["target_max_characters"]
    assert contract["target_characters_per_bullet"] == target < len(text)
    assert contract["field_character_targets"][0]["target_max_characters"] == target
    assert contract["fixed_title"] == "Topic"


def test_local_target_does_not_shorten_a_fitting_neighbor(profile):
    texts = ["Keep this verified and already fitting explanation unchanged.", "Too much detail."]
    elements = [
        Element(
            kind="text",
            text=t,
            box=Box(x=0, y=i * 80, w=120, h=20),
            size=12,
            source_ids=[f"summary-1-{i + 1}"],
        )
        for i, t in enumerate(texts)
    ]
    previous = plan(texts)
    contracts = shortening_contracts(previous, [1], feedback_for(profile, elements, 1), 500)
    contract = contracts[0]
    assert contract["target_characters_per_bullet"] is None
    assert contract["field_character_targets"][0]["fact_ids"] == ["summary-1-2"]
    assert contract["fixed_bullets"][0]["index"] == 1
    changed = deepcopy(previous)
    changed["slides"][0]["bullets"][1]["text"] = "Detail."
    validate_contracts(changed, contracts)
    changed["slides"][0]["bullets"][0]["text"] = "Changed."
    with pytest.raises(ValueError, match="unaffected bullet"):
        validate_contracts(changed, contracts)


def test_title_target_never_becomes_body_budget(profile):
    previous = plan(["Already fits."])
    element = Element(
        kind="text", text="Topic", role="title", size=12, box=Box(x=0, y=0, w=30, h=20)
    )
    contract = shortening_contracts(previous, [1], feedback_for(profile, [element]), 500)[0]
    assert contract["target_characters_per_bullet"] is None
    assert contract["field_character_targets"][0]["role"] == "title"
    assert contract["fixed_bullets"][0]["index"] == 1


def test_non_geometric_repair_keeps_existing_budget():
    previous = plan(["Some text."])
    assert (
        shortening_contracts(
            previous,
            [1],
            {
                "repair_issues": [
                    {"slide": 1, "code": "unsupported_number", "action": "revise_content"}
                ]
            },
            500,
        )
        == []
    )
