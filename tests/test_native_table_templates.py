"""Synthetic OOXML only: source cell words must never enter a reusable blueprint."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from studio.models import Box, Element, SlideScene, TemplateProfile
from studio.templates.native_template import native_patterns
from studio.templates.table_templates import table_template
from studio.composition.table_style import (
    apply_table_style,
    table_widths,
    table_heights,
    table_cell_style,
)
from studio.composition.render import render_pptx, render_pdf, render_html, clean_base
from studio.config import ROOT
from types import SimpleNamespace


def _source():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "OLD PRIVATE TITLE"
    shape = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(8), Inches(2))
    table = shape.table
    table.columns[0].width = Inches(5)
    table.columns[1].width = Inches(3)
    table.rows[0].height = Inches(0.7)
    table.rows[1].height = Inches(1.3)
    for ri, row in enumerate(table.rows):
        for ci in range(2):
            cell = table.cell(ri, ci)
            cell.text = f"OLD PRIVATE {ri} {ci}"
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("17365D" if ri == 0 else "E7EDF4")
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(15 if ri == 0 else 12)
            run.font.bold = ri == 0
            run.font.color.rgb = RGBColor.from_string("FFFFFF" if ri == 0 else "15202B")
    return prs


def _profile(source, pattern):
    return TemplateProfile(
        sha256="synthetic",
        name="Synthetic",
        width=720,
        height=540,
        slide_count=1,
        master_count=1,
        layout_count=len(Presentation(source).slide_layouts),
        object_count=2,
        placeholder_count=1,
        fonts=["Play"],
        font="Play",
        font_file=str(ROOT / "fonts" / "Play-Regular.ttf"),
        font_sizes=[12, 15],
        title_size=24,
        body_size=12,
        colors=["#15202B", "#FFFFFF"],
        background="#FFFFFF",
        foreground="#15202B",
        accent="#17365D",
        margin=36,
        patterns=[pattern],
        source_kind="pptx",
    )


def test_table_only_slide_becomes_typed_reusable_pattern_without_source_words():
    pattern = next(p for p in native_patterns(_source()) if p.source_slide == 1)
    assert pattern.role == "table"
    template = pattern.table_template
    assert template is not None
    assert template.column_widths == [360, 216]
    assert template.cells[0][0].fill == "#17365D"
    assert template.cells[1][1].color == "#15202B"
    assert "OLD PRIVATE" not in template.model_dump_json()


def test_compatible_new_content_uses_geometry_and_repeats_body_style():
    pattern = next(p for p in native_patterns(_source()) if p.source_slide == 1)
    element = Element(
        kind="table",
        box=Box(x=0, y=0, w=300, h=100),
        size=13,
        rows=[["New", "Header"], ["One", "1"], ["Two", "2"]],
    )
    apply_table_style(element, pattern)
    assert element.table_template is not pattern.table_template
    assert element.box == pattern.table_template.box
    assert table_widths(element, "unused") == [360, 216]
    assert round(sum(table_heights(element, [360, 216], "unused"))) == round(element.box.h)
    assert table_cell_style(element, 2, 1).fill == "#E7EDF4"
    element.table_template.cells = element.table_template.cells[:1]
    assert table_cell_style(element, 2, 1).fill == "#17365D"
    incompatible = Element(
        kind="table", box=Box(x=0, y=0, w=300, h=100), size=13, rows=[["A", "B", "C"]]
    )
    apply_table_style(incompatible, pattern)
    assert incompatible.table_template is None


def test_layout_table_is_inherited_by_slide_pattern(tmp_path):
    prs = Presentation()
    layout = prs.slide_layouts[5]
    donor = prs.slides.add_slide(prs.slide_layouts[6])
    shape = donor.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(8), Inches(2))
    table = shape.table
    table.cell(0, 0).text = "OLD PRIVATE LAYOUT"
    layout.shapes._spTree.insert_element_before(shape._element, "p:extLst")
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "New title"
    pattern = next(p for p in native_patterns(prs) if p.source_slide == 2)
    assert pattern.table_template is not None
    assert pattern.role == "table"
    assert "OLD PRIVATE" not in pattern.table_template.model_dump_json()
    source = tmp_path / "layout-table.pptx"
    prs.save(source)
    cleaned = clean_base(source, SimpleNamespace(background_source="", assets=[]))
    assert not any(sh.has_table for sh in cleaned.slide_layouts[5].shapes)


def test_native_pptx_exports_new_editable_cells_without_old_words(tmp_path):
    source = tmp_path / "synthetic.pptx"
    _source().save(source)
    pattern = next(p for p in native_patterns(Presentation(source)) if p.source_slide == 1)
    profile = _profile(source, pattern)
    element = Element(
        kind="table",
        box=Box(x=0, y=0, w=200, h=80),
        size=12,
        rows=[["Fresh heading", "Fresh metric"], ["Alpha", "10"], ["Beta", "20"]],
        color="#15202B",
    )
    apply_table_style(element, pattern)
    scene = SlideScene(
        title="Fresh",
        background="#FFFFFF",
        elements=[element],
        source_ids=[],
        layout="table",
        pattern_id=pattern.id,
    )
    output = tmp_path / "out.pptx"
    render_pptx([scene], profile, source, output, verify_text=False)
    result = Presentation(output)
    tables = [s.table for s in result.slides[0].shapes if s.has_table]
    assert len(tables) == 1
    assert tables[0].cell(0, 0).text == "Fresh heading"
    assert tables[0].cell(2, 1).text == "20"
    assert "OLD PRIVATE" not in " ".join(cell.text for row in tables[0].rows for cell in row.cells)
    assert tables[0].columns[0].width > tables[0].columns[1].width
    pdf = tmp_path / "out.pdf"
    html = tmp_path / "out.html"
    render_pdf([scene], profile, pdf)
    render_html([scene], profile, html)
    assert pdf.stat().st_size > 500
    markup = html.read_text()
    assert "Fresh heading" in markup and "OLD PRIVATE" not in markup
    assert "rgba(23,54,93,1.0)" in markup and "rgba(231,237,244,1.0)" in markup


def test_one_row_source_expands_and_empty_rows_render_safely(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Old title"
    shape = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(8), Inches(1))
    shape.table.cell(0, 0).text = "OLD PRIVATE"
    source = tmp_path / "one-row.pptx"
    prs.save(source)
    pattern = next(p for p in native_patterns(Presentation(source)) if p.source_slide == 1)
    element = Element(
        kind="table",
        box=Box(x=0, y=0, w=100, h=50),
        size=12,
        rows=[["New", "Header"], ["New", "Body"]],
        color="#15202B",
    )
    apply_table_style(element, pattern)
    assert table_cell_style(element, 1, 0) == table_cell_style(element, 0, 0)
    scene = SlideScene(
        title="New",
        background="#FFFFFF",
        elements=[element],
        source_ids=[],
        layout="table",
        pattern_id=pattern.id,
    )
    output = tmp_path / "expanded.pptx"
    render_pptx([scene], _profile(source, pattern), source, output, verify_text=False)
    assert Presentation(output).slides[0].shapes[-1].table.cell(1, 0).text == "New"
    empty = Element(kind="table", box=Box(x=0, y=0, w=100, h=50), size=12, rows=[])
    apply_table_style(empty, pattern)
    scene.elements = [empty]
    render_pptx(
        [scene], _profile(source, pattern), source, tmp_path / "empty.pptx", verify_text=False
    )


def test_merged_source_table_is_not_misrepresented_as_rectangular_blueprint():
    prs = _source()
    shape = next(sh for sh in prs.slides[0].shapes if sh.has_table)
    shape.table.cell(0, 0).merge(shape.table.cell(0, 1))
    assert table_template(shape) is None
