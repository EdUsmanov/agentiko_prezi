"""Branded empty covers can use BODY placeholders for both text fields."""

import pytest
from pptx import Presentation
from pptx.util import Pt

from studio.templates.native_style import native_styles
from studio.templates.native_template import native_patterns


def _patterns(tmp_path, *, name="Title Navy Flame", size=60, overlap=False):
    prs = Presentation()
    layout = prs.slide_layouts[0]
    layout._element.attrib.pop("type", None)
    layout._element.cSld.set("name", name)
    for index, shape in enumerate(list(layout.placeholders)[:2]):
        shape._element.nvSpPr.nvPr.ph.set("type", "body")
        shape.left, shape.width = Pt(60), Pt(500)
        shape.top = Pt(140 if index == 0 else 200 if overlap else 300)
        shape.height = Pt(130 if index == 0 else 40)
        shape.text_frame.paragraphs[0].font.size = Pt(size if index == 0 else 32)
    prs.slides.add_slide(prs.slide_layouts[6])  # First source slide is a logo page.
    slide = prs.slides.add_slide(layout)
    assert all(not shape.text for shape in slide.placeholders)
    path = tmp_path / "body-cover.pptx"
    prs.save(path)
    return native_patterns(Presentation(path), native_styles(path))


def test_empty_body_cover_uses_inherited_typographic_hierarchy(tmp_path):
    patterns = _patterns(tmp_path)
    cover = next(p for p in patterns if p.source_slide == 2)
    assert cover.role == "cover"
    assert cover.title_size == 60
    assert len(cover.body_zones) == 1
    assert {field["role"] for field in cover.fields} == {"title", "body"}
    assert any(p.role == "cover" and p.source_slide == 0 for p in patterns)


@pytest.mark.parametrize(
    "options",
    [{"name": "Title and Content"}, {"size": 32}, {"overlap": True}],
)
def test_ambiguous_body_pair_is_not_promoted_to_cover(tmp_path, options):
    assert not any(p.source_slide == 2 for p in _patterns(tmp_path, **options))
