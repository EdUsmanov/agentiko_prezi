import asyncio
from types import SimpleNamespace
from zipfile import ZipFile
import pytest
from studio.models import Fact,SlidePlan,Pattern,Box,TableData
from studio.content import parse_content
from studio.contracts import compatible,apply_meanings
from studio.document import structure_document
from studio.storyboard import prepare_storyboard
from studio.planner import extractive_plans,validate_plans,assign_compositions
from studio.composer import compose_variant
from studio.audit import audit_scenes


def test_plain_headings_and_visual_directives_are_not_body(prepared):
    _,_,p=prepared
    p.content=parse_content('Проект\nОбзор проекта.\nДинамика\n| Неделя | Просмотры |\n|---|---|\n| 1 | 10 |\n| 2 | 20 |\nРекомендуемый вид: линейный график.')
    class Gateway:
        settings=SimpleNamespace(mode='api')
        async def json_request(self,*args,**kwargs):
            return {'blocks':[{'fact_id':f'f{i}','kind':k} for i,k in enumerate(
                ['heading','body','heading','body','visualization'],1)]}
    asyncio.run(structure_document(p,Gateway()))
    assert [f.id for f in p.content.facts]==['f2','f4']
    assert p.content.tables[0].visualization=='line'
    assert p.content.tables[0].section=='Динамика'
    assert len(p.analysis['document_structure']['source_blocks'])==5


def test_invalid_document_roles_preserve_assertions_without_failing_analysis(prepared):
    _,_,p=prepared
    p.content=parse_content('Название\nСущественное утверждение.')
    original=[f.model_dump() for f in p.content.facts]
    class Gateway:
        settings=SimpleNamespace(mode='api')
        async def json_request(self,*args,**kwargs):
            return {'blocks':[{'fact_id':f'f{i}','kind':'heading'} for i in (1,2)]}
    asyncio.run(structure_document(p,Gateway()))
    assert [f.text for f in p.content.facts]==[original[1]['text']]
    assert p.content.headings[-1]['text']==original[0]['text']
    assert p.analysis['document_structure']['status']=='degraded'
    assert p.analysis['document_fallback_blocks']==['f2']


def test_semantic_service_and_reference_cannot_be_dividers(prepared):
    _,_,p=prepared
    pattern=p.template.patterns[0]
    pattern.role='divider'
    apply_meanings(p.template,{'patterns':[{'pattern_id':pattern.id,'purpose':'reference','reusable':True}]})
    assert not pattern.reusable and pattern.role!='divider'
    assert not compatible(pattern,SlidePlan(title='Раздел',fact_ids=[],layout='divider'))


def test_title_fact_is_visible_once_and_still_covered(prepared):
    _,_,p=prepared
    p.content=parse_content('Один исходный тезис.')
    p.constraints.slides=1
    plans=extractive_plans(p)
    for v in plans.variants:
        v.slides[0].title=p.content.facts[0].text
        scenes=compose_variant(v,p)
        assert sum(e.text==v.slides[0].title for e in scenes[0].elements)==1
        assert any(e.role=='title' and e.source_ids==['f1'] for e in scenes[0].elements)
        assert not [f for f in audit_scenes(scenes,p) if f.code in ('coverage','duplicate_title','empty_slide')]


def test_storyboard_reserves_cover_and_dividers_before_planner(prepared):
    _,_,p=prepared
    p.constraints.slides=8
    p.content=parse_content('# Продукт\n## Контекст\nПервое утверждение.\nВторое утверждение.\nТретье утверждение.\nЧетвёртое утверждение.\n## Результаты\nПервый результат.\nВторой результат.\nТретий результат.\nЧетвёртый результат.')
    pattern=p.template.patterns[0].model_copy(deep=True)
    pattern.id='real-cover';pattern.role=pattern.purpose='cover'
    p.template.patterns.append(pattern)
    # Remove geometry-only divider guesses to exercise the explicit derivative.
    p.template.patterns=[x for x in p.template.patterns if x.role!='divider']
    prepare_storyboard(p)
    plans=validate_plans(extractive_plans(p),p)
    assert p.analysis['derived_divider']['authored_divider_found'] is False
    for v in plans.variants:
        assert len(v.slides)==8 and v.slides[0].purpose=='cover'
        assert v.slides[0].fact_ids==[]
        assert any(s.layout=='divider' for s in v.slides)
        assert [f for s in v.slides for f in s.fact_ids]==[f.id for f in p.content.facts]
    plans.variants[0].slides[0].pattern_id=p.template.patterns[0].id
    with pytest.raises(ValueError,match='назначение|Назначение'):
        validate_plans(plans,p)


@pytest.mark.parametrize('kind',['bar','column','line','pie'])
def test_real_editable_chart_type_and_data(prepared,tmp_path,kind):
    from studio.charts import make_chart
    from studio.render import render_pptx
    from studio.models import SlideScene
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    _,store,p=prepared
    table=TableData(id='t1',headers=['Неделя','Значение'],rows=[['Первая','10'],['Вторая','20']])
    plan=SlidePlan(title='Данные',fact_ids=['f1'],layout='chart',table_id='t1',chart_type=kind)
    e=make_chart(table,plan,Box(x=80,y=130,w=650,h=320),p.template,p.template.foreground,['f1'])
    scene=SlideScene(title='Данные',background=p.template.background,elements=[e],source_ids=['f1'],layout='chart')
    path=tmp_path/(kind+'.pptx')
    render_pptx([scene],p.template,store.directory(p.id)/'input.pptx',path,verify_text=False)
    chart=next(s.chart for s in Presentation(path).slides[0].shapes if s.has_chart)
    expected={'bar':XL_CHART_TYPE.BAR_CLUSTERED,'column':XL_CHART_TYPE.COLUMN_CLUSTERED,'line':XL_CHART_TYPE.LINE_MARKERS,'pie':XL_CHART_TYPE.PIE}
    assert chart.chart_type==expected[kind]
    assert list(chart.series[0].values)==[10,20]
    assert chart.plots[0].data_labels.position == (
        XL_LABEL_POSITION.ABOVE if kind=='line' else XL_LABEL_POSITION.OUTSIDE_END)
    with ZipFile(path) as z:
        assert any(n.endswith('.xlsx') for n in z.namelist())


def test_refinement_rejects_unrelated_or_wrong_purpose_edits(prepared):
    from studio.refinement import apply_edits,LayoutEdit
    _,_,p=prepared
    plans=assign_compositions(extractive_plans(p),p)
    decks={v.key:compose_variant(v,p) for v in plans.variants}
    with pytest.raises(ValueError,match='allowlist'):
        apply_edits(p,plans,decks,[LayoutEdit(variant='executive',slide=1,pattern_id='../../secret')],{})


def test_multi_series_chart_preserves_all_columns(prepared):
    from studio.charts import make_chart
    _,_,p=prepared
    table=TableData(id='t1',headers=['Период','Автор','Сообщество'],rows=[['Первый','12','18'],['Второй','15','46']])
    e=make_chart(table,SlidePlan(title='Сравнение',fact_ids=['f1'],chart_type='column'),
        Box(x=0,y=0,w=600,h=300),p.template,p.template.foreground,['f1'])
    assert e.series_names==['Автор','Сообщество'] and e.series_values==[[12,15],[18,46]]


@pytest.mark.parametrize('improves',[True,False])
@pytest.mark.parametrize('model_edits',[True,False])
def test_refinement_adopts_only_reviewed_improvement(prepared,tmp_path,monkeypatch,improves,model_edits):
    import studio.refinement as refinement
    _,_,p=prepared
    p.content=parse_content('# Изменение\n| Период | Значение |\n|---|---|\n| Первый | 10 |\n| Второй | 20 |')
    p.content.tables[0].visualization='line';p.constraints.slides=1
    plans=assign_compositions(extractive_plans(p),p)
    decks={v.key:compose_variant(v,p) for v in plans.variants}
    results=[{'key':v.key,'title':v.title,'slides':1,'rendering':{'native_render':True},'findings':[]} for v in plans.variants]
    finding={'variant':'executive','slide':1,'code':'readability','severity':'warning','message':'Тесные подписи оси'}
    visual={'status':'completed','findings':[finding],'checked':3,'total':3}
    def render(scenes,profile,source,folder):
        folder.mkdir(parents=True,exist_ok=True)
        (folder/'marker.txt').write_text('revised')
        return {'native_render':True}
    async def review(*args,**kwargs):
        return {'status':'completed','findings':[] if improves else [finding],'checked':1,'total':1}
    monkeypatch.setattr(refinement,'render_variant',render)
    monkeypatch.setattr(refinement,'audit_export',lambda *a:[])
    monkeypatch.setattr(refinement,'review_visuals',review)
    class Gateway:
        settings=SimpleNamespace(mode='api')
        async def json_request(self,*args,**kwargs):
            return {'edits':[{'variant':'executive','slide':1,'operation':'readable_chart','pattern_id':None}] if model_edits else []}
    out=asyncio.run(refinement.refine(p,plans,decks,results,visual,tmp_path,tmp_path/'source.pptx',Gateway(),60,lambda _:None))
    assert out[-1]['accepted'] is improves
    assert (tmp_path/'executive'/'marker.txt').exists() is improves
    assert out[0].variants[0].slides[0].chart_style==('readable' if improves else 'standard')


def test_refinement_reserves_deadline_without_calling_model(prepared,tmp_path):
    from studio.refinement import refine
    _,_,p=prepared
    class Gateway:
        settings=SimpleNamespace(mode='api')
        async def json_request(self,*args,**kwargs): raise AssertionError('No budget for model')
    visual={'status':'completed','findings':[{'variant':'executive','slide':1}]}
    out=asyncio.run(refine(p,None,{},[],visual,tmp_path,None,Gateway(),20,lambda _:None))
    assert out[-1]['status']=='not_run'


def test_template_timeout_retries_only_the_same_batch(prepared):
    from studio.analysis import analyze_meaning
    _,_,p=prepared
    pattern=p.template.patterns[0]
    inventory={'patterns':[{'id':pattern.id,'source_slide':1}],'width':960,'height':540}
    class Gateway:
        settings=SimpleNamespace(mode='api')
        requests=0
        async def json_request(self,*args,**kwargs):
            self.requests+=1
            if self.requests==1: raise TimeoutError()
            return {'patterns':[{'pattern_id':pattern.id,'roles':['context'],'density':'low','purpose':'content','reusable':True}]}
    gateway=Gateway()
    result=asyncio.run(analyze_meaning(inventory,gateway))
    assert result['status']=='completed' and gateway.requests==2


def test_incomplete_template_induction_cannot_reach_planner(prepared,monkeypatch):
    import studio.analysis as analysis
    _,store,p=prepared
    async def failed(*args,**kwargs):
        return {'status':'failed','method':'text_and_geometry','patterns':[]}
    monkeypatch.setattr(analysis,'analyze_meaning',failed)
    gateway=SimpleNamespace(settings=SimpleNamespace(mode='api',model_id='test'))
    with pytest.raises(ValueError,match='непроверенному каталогу'):
        asyncio.run(analysis.prepare_intelligence(p,store.directory(p.id)/'input.pptx',gateway,lambda *_:None))


def test_title_only_cover_is_available_to_design(prepared):
    from studio.deeppresenter import CompositionEnvironment
    _,_,p=prepared
    cover=p.template.patterns[0].model_copy(deep=True)
    cover.id='cover-without-subtitle';cover.role=cover.purpose='cover';cover.body_zones=[]
    p.template.patterns.append(cover)
    prepare_storyboard(p)
    plans=extractive_plans(p)
    for v in plans.variants:
        v.slides[0].pattern_id=cover.id
    validate_plans(plans,p)
    assert cover.id in CompositionEnvironment(p,plans).allowed
