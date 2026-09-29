from pathlib import Path
from types import SimpleNamespace

import pytest

from studio.composition.variant_body import compose_variant_body
from studio.models import Box, ContentModel, Fact, SlidePlan, Pattern, VariantPlan


FONT = str(Path(__file__).resolve().parents[1] / "fonts" / "Montserrat-Regular.ttf")


def package(facts, units=()):
    return SimpleNamespace(
        content=ContentModel(title="Example", facts=facts),
        analysis={"archetypes": {"units": list(units)}},
    )


PROFILE = SimpleNamespace(
    font="Montserrat",
    font_file=FONT,
    font_roles={},
    font_assets=[],
    font_sizes=[12, 16, 20, 24],
    body_size=20,
    height=540,
    accent="#336699",
    foreground="#222222",
)
BOX = Box(x=50, y=120, w=900, h=320)


def test_executive_lead_and_grounded_analytical_comparison():
    facts = [
        Fact(id="a", text="Alpha costs 100.", section="Alpha"),
        Fact(id="b", text="Beta costs 120.", section="Beta"),
    ]
    unit = {
        "purpose": "comparison",
        "fact_ids": ["a", "b"],
        "slots": [
            {"role": "entity", "fact_id": "a", "quote": "Alpha"},
            {"role": "entity", "fact_id": "b", "quote": "Beta"},
        ],
    }
    p = package(facts, [unit])
    slide = SlidePlan(title="Price", fact_ids=["a", "b"], purpose="comparison")
    lead = compose_variant_body(slide, p, "executive", facts, BOX, PROFILE, PROFILE.foreground)
    table = compose_variant_body(slide, p, "analytical", facts, BOX, PROFILE, PROFILE.foreground)
    assert lead and [e.kind for e in lead] == ["text", "text"]
    assert lead[0].box.y < lead[1].box.y
    assert [e.source_ids for e in lead] == [["a"], ["b"]]
    assert table and len(table) == 1 and table[0].kind == "table"
    assert table[0].rows == [["Alpha", "Beta"], [facts[0].text, facts[1].text]]
    assert table[0].source_ids == ["a", "b"]


def test_no_unreviewed_comparison_table_or_unearned_sequence():
    facts = [Fact(id="a", text="Alpha costs 100."), Fact(id="b", text="Beta costs 120.")]
    p = package(facts)
    slide = SlidePlan(title="Price", fact_ids=["a", "b"], purpose="comparison")
    columns = compose_variant_body(slide, p, "analytical", facts, BOX, PROFILE, PROFILE.foreground)
    story = compose_variant_body(slide, p, "story", facts, BOX, PROFILE, PROFILE.foreground)
    assert columns and [e.kind for e in columns] == ["text", "text"]
    assert columns[0].box.x < columns[1].box.x
    assert story and len(story) == 1 and story[0].source_ids == ["a", "b"]
    assert all(e.kind != "table" for e in columns + story)


def test_fitted_executive_lead_stays_at_least_as_large_as_its_details():
    facts = [
        Fact(id="lead", text="The team reviews every incoming request before assigning an owner."),
        Fact(id="detail", text="Owners receive a daily summary."),
    ]
    elements = compose_variant_body(
        SlidePlan(title="Requests", fact_ids=[f.id for f in facts]),
        package(facts),
        "executive",
        facts,
        Box(x=50, y=120, w=300, h=260),
        PROFILE,
        PROFILE.foreground,
    )
    assert elements and elements[0].bold
    assert all(16 <= e.size <= elements[0].size for e in elements)
    assert [e.text for e in elements] == [f.text for f in facts]


def test_story_uses_only_source_reviewed_steps():
    facts = [Fact(id="a", text="Collect requests."), Fact(id="b", text="Review requests.")]
    unit = {
        "purpose": "process",
        "fact_ids": ["a", "b"],
        "slots": [
            {"role": "step", "fact_id": "a", "quote": facts[0].text},
            {"role": "step", "fact_id": "b", "quote": facts[1].text},
        ],
    }
    p = package(facts, [unit])
    slide = SlidePlan(title="Workflow", fact_ids=["a", "b"], purpose="process")
    executive = compose_variant_body(slide, p, "executive", facts, BOX, PROFILE, PROFILE.foreground)
    analytical = compose_variant_body(
        slide, p, "analytical", facts, BOX, PROFILE, PROFILE.foreground
    )
    story = compose_variant_body(slide, p, "story", facts, BOX, PROFILE, PROFILE.foreground)
    assert executive and analytical and story
    assert [e.source_ids for e in executive] == [["a"], ["b"]]
    assert [e.source_ids for e in analytical] == [["a"], ["b"]]
    assert executive[0].box.y < executive[1].box.y
    assert analytical[0].box.x < analytical[1].box.x
    assert len(story) == 1 and story[0].source_ids == ["a", "b"]
    assert story[0].text == "Collect requests.\nReview requests."


def test_story_retains_reviewed_timeline_label():
    facts = [Fact(id="a", text="Pilot opened in 2025."), Fact(id="b", text="Review began in 2026.")]
    unit = {
        "purpose": "timeline",
        "fact_ids": ["a", "b"],
        "slots": [
            {"role": "date", "fact_id": "a", "quote": "2025"},
            {"role": "date", "fact_id": "b", "quote": "2026"},
        ],
    }
    slide = SlidePlan(title="Milestones", fact_ids=["a", "b"], purpose="timeline")
    story = compose_variant_body(
        slide, package(facts, [unit]), "story", facts, BOX, PROFILE, PROFILE.foreground
    )
    assert story and len(story) == 1
    assert story[0].text == "Pilot opened in 2025.\nReview began in 2026."
    assert story[0].source_ids == ["a", "b"]


def test_single_native_field_keeps_bounds_and_authored_color(prepared):
    from studio.composition.composer import compose_variant
    from studio.contents.planner import assign_compositions, extractive_plans
    from studio.contents.parsing import parse_content

    _, _, p = prepared  # The fixture creates a synthetic PPTX in a temporary directory.
    p.content = parse_content("# Topic\nOne fact.\nSecond fact.\nThird fact.")
    p.constraints.slides = 1
    zone = Box(x=50, y=150, w=700, h=300)
    p.template.patterns = [
        Pattern(
            id="native-one",
            source_slide=1,
            source_layout="Native",
            role="statement",
            text_zones=[],
            title_zone=Box(x=50, y=35, w=600, h=60),
            body_zones=[zone],
            title_foreground="#154A67",
            zone_foregrounds=["#154A67"],
        )
    ]
    plans = assign_compositions(extractive_plans(p), p)
    for variant in plans.variants:
        scene = compose_variant(variant, p)[0]
        body = [e for e in scene.elements if e.role == "body"]
        assert body and all(
            zone.x <= e.box.x
            and e.box.x + e.box.w <= zone.x + zone.w
            and zone.y <= e.box.y
            and e.box.y + e.box.h <= zone.y + zone.h
            and e.color == "#154A67"
            for e in body
        )


def test_reviewed_comparison_uses_native_table_blueprint(prepared, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from studio.composition.composer import compose_native
    from studio.templates.native_template import native_patterns

    _, _, p = prepared
    source = Presentation()
    slide_shape = source.slides.add_slide(source.slide_layouts[5])
    slide_shape.shapes.title.text = "OLD PRIVATE TITLE"
    table = slide_shape.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(8), Inches(2)).table
    table.columns[0].width = Inches(5)
    table.columns[1].width = Inches(3)
    table.cell(0, 0).text = "OLD PRIVATE A"
    table.cell(0, 1).text = "OLD PRIVATE B"
    native = next(x for x in native_patterns(source) if x.source_slide == 1)
    assert native.table_template and len(native.table_template.column_widths) == 2
    p.template.patterns = [native]
    p.content = ContentModel(
        title="Comparison",
        facts=[
            Fact(id="a", text="Alpha costs 100.", section="Alpha"),
            Fact(id="b", text="Beta costs 120.", section="Beta"),
        ],
    )
    p.analysis = {
        "archetypes": {
            "units": [
                {
                    "purpose": "comparison",
                    "fact_ids": ["a", "b"],
                    "slots": [
                        {"role": "entity", "fact_id": "a", "quote": "Alpha"},
                        {"role": "entity", "fact_id": "b", "quote": "Beta"},
                    ],
                }
            ]
        }
    }
    plan = SlidePlan(
        title="Comparison", fact_ids=["a", "b"], purpose="comparison", pattern_id=native.id
    )
    scene = compose_native(plan, p, 0, "analytical")
    new_table = next(e for e in scene.elements if e.kind == "table")
    assert new_table.table_template is not None
    assert new_table.box == native.table_template.box
    assert new_table.table_template.column_widths == native.table_template.column_widths
    assert new_table.rows == [["Alpha", "Beta"], ["Alpha costs 100.", "Beta costs 120."]]
    assert "OLD PRIVATE" not in str(new_table.model_dump())
    from studio.composition.render import render_pptx

    input_file, output_file = tmp_path / "source.pptx", tmp_path / "composed.pptx"
    source.save(input_file)
    render_pptx([scene], p.template, input_file, output_file, verify_text=False)
    rendered = Presentation(output_file)
    cells = [shape.table for shape in rendered.slides[0].shapes if shape.has_table]
    assert len(cells) == 1
    assert cells[0].cell(0, 0).text == "Alpha"
    assert cells[0].cell(1, 1).text == "Beta costs 120."
    assert "OLD PRIVATE" not in " ".join(cell.text for row in cells[0].rows for cell in row.cells)
    ordinary = native.model_copy(deep=True)
    ordinary.id = "ordinary-body"
    ordinary.table_template = None
    ordinary.role = "statement"
    p.template.patterns = [ordinary, native]
    plan.pattern_id = None
    selected = compose_native(plan, p, 0, "analytical")
    assert selected.pattern_id == native.id
    narrow = native.model_copy(deep=True)
    narrow.id = "narrow-source-table"
    narrow.table_template.column_widths = [20, 556]
    p.template.patterns = [ordinary, narrow]
    selected = compose_native(plan, p, 0, "analytical")
    assert selected.pattern_id == ordinary.id


def test_oversized_sample_fonts_keep_three_real_exported_structures(prepared, template, tmp_path):
    from pptx import Presentation
    from studio.composition.composer import compose_variant
    from studio.composition.render import render_pptx
    from studio.checks.export_audit import audit_export

    _, _, p = prepared
    p.analysis = {}
    p.constraints.slides = 1
    p.template.body_size = 48
    p.template.font_sizes = [44, 48, 52]
    p.template.patterns = [
        Pattern(
            id="dense-body",
            source_slide=1,
            source_layout="Synthetic",
            role="statement",
            text_zones=[],
            title_zone=Box(x=50, y=35, w=600, h=60),
            body_zones=[Box(x=50, y=150, w=600, h=170)],
            zone_foregrounds=["#154A67"],
        )
    ]
    facts = [
        Fact(
            id="a",
            text="The team collects requests from several departments and records each request in one shared queue.",
        ),
        Fact(
            id="b",
            text="Reviewers check the submitted details and return incomplete requests for clarification.",
        ),
        Fact(
            id="c",
            text="Owners record the final outcome and make the status visible to the requesting department.",
        ),
    ]
    p.content = ContentModel(title="Workflow", facts=facts)
    scenes = {}
    for key in ("executive", "analytical", "story"):
        variant = VariantPlan(
            key=key,
            title=key,
            slides=[
                SlidePlan(
                    title="Workflow",
                    fact_ids=[f.id for f in facts],
                    layout="columns",
                    purpose="content",
                    pattern_id="dense-body",
                )
            ],
        )
        scene = compose_variant(variant, p)[0]
        scenes[key] = scene
        content = [e for e in scene.elements if e.kind == "text" and e.source_ids]
        assert content and all(e.size >= 16 for e in content)
        assert {fid for e in content for fid in e.source_ids} == {"a", "b", "c"}
        output = tmp_path / f"{key}.pptx"
        render_pptx([scene], p.template, template, output, verify_text=False)
        exported = Presentation(output)
        rendered_text = "\n".join(
            shape.text for shape in exported.slides[0].shapes if shape.has_text_frame
        )
        import re

        normalized = re.sub(r"\s+", " ", rendered_text)
        assert all(f.text in normalized for f in facts), (key, rendered_text)
        assert not any(f.code == "fact_coverage" for f in audit_export(output, variant, p))
    executive = [e for e in scenes["executive"].elements if e.kind == "text" and e.source_ids]
    analytical = [e for e in scenes["analytical"].elements if e.kind == "text" and e.source_ids]
    story = [e for e in scenes["story"].elements if e.kind == "text" and e.source_ids]
    assert len(story) == 1 and len(executive) == len(analytical) == 3
    assert executive[0].box.y < executive[1].box.y
    assert analytical[0].box.x < analytical[1].box.x


@pytest.mark.parametrize("key", ["executive", "analytical", "story"])
def test_variant_prefers_safe_evidence_structure_among_authored_patterns(
    prepared, template, tmp_path, key
):
    from studio.composition.composer import compose_native
    from studio.composition.render import render_pptx
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    _, _, p = prepared
    p.analysis = {}
    p.content = ContentModel(
        title="Overview",
        facts=[
            Fact(id="a", text="The team records each request in the shared queue."),
            Fact(id="b", text="Reviewers check the details before assigning an owner."),
            Fact(id="c", text="The owner records the outcome for the requester."),
        ],
    )
    one = Pattern(
        id="one-field",
        source_slide=1,
        source_layout="Authored one",
        role="statement",
        text_zones=[],
        title_zone=Box(x=50, y=35, w=600, h=60),
        body_zones=[Box(x=50, y=150, w=600, h=210)],
        graphic_count=1,
    )
    two = Pattern(
        id="two-fields",
        source_slide=2,
        source_layout="Authored two",
        role="split",
        text_zones=[],
        title_zone=Box(x=50, y=35, w=600, h=60),
        body_zones=[Box(x=50, y=150, w=290, h=210), Box(x=360, y=150, w=290, h=210)],
    )
    p.template.patterns = [two, one]
    plan = SlidePlan(title="Overview", fact_ids=["a", "b", "c"], purpose="content", layout="split")
    scene = compose_native(plan, p, 0, key)
    assert scene.pattern_id == "one-field"
    body = [e for e in scene.elements if e.kind == "text" and e.role == "body"]
    if key == "story":
        assert [e.source_ids for e in body] == [["a", "b", "c"]]
    else:
        assert [e.source_ids for e in body] == [["a"], ["b"], ["c"]]
        if key == "executive":
            assert body[0].box.y < body[1].box.y and body[0].bold
        else:
            assert body[0].box.x < body[1].box.x < body[2].box.x
    decorated = Presentation(template)
    decorated.slides[0].shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.1), Inches(0.1), Inches(0.3), Inches(0.3)
    )
    source, output = tmp_path / "decorated.pptx", tmp_path / f"{key}.pptx"
    decorated.save(source)
    p.template.background_source = ""
    render_pptx([scene], p.template, source, output, verify_text=False)
    exported = Presentation(output)
    assert any(
        shape.shape_type == 1 and shape.left == Inches(0.1) and shape.top == Inches(0.1)
        for shape in exported.slides[0].shapes
    )
    one.safe_text_zone = {"field_checks": [{"status": "unknown"}]}
    scene = compose_native(plan, p, 0, key)
    assert scene.pattern_id == "two-fields"


def test_reviewed_groups_do_not_filter_story_single_field_or_explicit_recompose(prepared):
    from studio.composition.composer import compose_native

    _, _, p = prepared
    p.content = ContentModel(
        title="Comparison",
        facts=[
            Fact(id="a", text="Alpha has a shared request queue."),
            Fact(id="b", text="Beta routes requests to its owners."),
        ],
    )
    p.analysis = {
        "editorial": {
            "bindings": [
                {
                    "fact_ids": ["a", "b"],
                    "purpose": "comparison",
                    "groups": [
                        {"label": "Alpha", "fact_ids": ["a"]},
                        {"label": "Beta", "fact_ids": ["b"]},
                    ],
                }
            ]
        }
    }
    one = Pattern(
        id="reviewed-one",
        source_slide=1,
        source_layout="One",
        role="statement",
        text_zones=[],
        title_zone=Box(x=50, y=35, w=600, h=60),
        body_zones=[Box(x=50, y=150, w=600, h=210)],
    )
    two = Pattern(
        id="reviewed-two",
        source_slide=2,
        source_layout="Two",
        role="split",
        text_zones=[],
        title_zone=Box(x=50, y=35, w=600, h=60),
        body_zones=[Box(x=50, y=150, w=290, h=210), Box(x=360, y=150, w=290, h=210)],
    )
    p.template.patterns = [two, one]
    plan = SlidePlan(title="Comparison", fact_ids=["a", "b"], purpose="comparison", layout="split")
    selected = compose_native(plan, p, 0, "story")
    assert selected.pattern_id == one.id
    assert [e.source_ids for e in selected.elements if e.kind == "text" and e.role == "body"] == [
        ["a", "b"]
    ]
    restored = SlidePlan.model_validate({**plan.model_dump(), "pattern_id": selected.pattern_id})
    recomposed = compose_native(restored, p, 0, "story")
    assert recomposed.pattern_id == one.id
    assert [e.source_ids for e in recomposed.elements if e.kind == "text" and e.role == "body"] == [
        ["a", "b"]
    ]
