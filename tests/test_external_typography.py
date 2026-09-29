"""A PPTX may specify every font size in its layout/master, not its runs."""

from zipfile import ZipFile

import pytest
from pptx import Presentation

from studio.templates.native_style import native_styles
from studio.templates.parsing import analyze_template


@pytest.mark.parametrize("filled", [False, True])
def test_inherited_title_and_body_sizes_are_used_without_direct_run_sizes(tmp_path, filled):
    source = tmp_path / "inherited.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if filled:
        slide.shapes.title.text = "Inherited title"
        slide.placeholders[1].text = "Inherited subtitle"
    prs.save(source)
    with ZipFile(source) as archive:
        xml = archive.read("ppt/slides/slide1.xml")
    assert b" sz=" not in xml  # No direct size to collect from a run.

    styles = native_styles(source)
    title = styles[("ppt/slides/slide1.xml", slide.shapes.title.shape_id)]["size"]
    subtitle = styles[("ppt/slides/slide1.xml", slide.placeholders[1].shape_id)]["size"]
    assert title == 44 and subtitle == 32  # Defined by the Office master.

    profile = analyze_template(source, tmp_path / "artifacts", allow_download=False)
    assert profile.font_sizes == [32, 44]
    assert profile.title_size == title
    assert profile.body_size == subtitle
    assert any(name.casefold().startswith("calibri") for name in profile.fonts)
