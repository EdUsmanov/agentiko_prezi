from pptx import Presentation
from studio.config import ROOT
from studio.embedded_fonts import check_glyphs
from studio.fonts import font_runs, symbol_font, text_width, pdf_font
from studio.render import set_text
from studio.models import Element, Box
from reportlab.pdfbase import pdfmetrics
import pytest


def test_arrow_has_explicit_symbol_fallback_not_a_replaced_body_font(prepared):
    _, _, package = prepared
    path = str(ROOT / "fonts/Play-Regular.ttf")
    text = "Вход → Выход"
    check_glyphs(path, text)
    assert list(font_runs(text, path)) == [("Вход ", path), ("→", symbol_font()), (" Выход", path)]
    expected = pdfmetrics.stringWidth("Вход  Выход", pdf_font(path), 20) + pdfmetrics.stringWidth(
        "→", pdf_font(symbol_font()), 20
    )
    assert text_width(text, path, 20) == pytest.approx(expected)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, 1000000, 1000000)
    element = Element(kind="text", box=Box(x=0, y=0, w=400, h=80), size=20, color="#000000")
    set_text(shape.text_frame, text, element, package.template)
    assert shape.text == text
    assert [(r.text, r.font.name) for r in shape.text_frame.paragraphs[0].runs] == [
        ("Вход ", "Play"),
        ("→", "Montserrat"),
        (" Выход", "Play"),
    ]


def test_missing_non_symbol_is_still_rejected():
    with pytest.raises(ValueError, match="U\\+"):
        check_glyphs(str(ROOT / "fonts/Play-Regular.ttf"), "漢")
