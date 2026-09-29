"""Final PPTX evidence is independent of scene labels and bookkeeping."""

from types import SimpleNamespace

from pptx import Presentation
from pptx.util import Pt
import pytest

from studio.checks.export_audit import content_scenes
from studio.models import ContentModel, Fact, SlidePlan, TableData
from studio.contents.editorial_domain import EditorialPlan, editorial_schema
from studio.contents.editorial_repair import EditorialPatch
from studio.contents.editorial_outline import SlideBatch


def extracted(rows, facts, *, source_table=None):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(len(rows), len(rows[0]), Pt(20), Pt(40), Pt(600), Pt(300)).table
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            table.cell(ri, ci).text = value
    package = SimpleNamespace(
        content=ContentModel(
            title="Synthetic", facts=facts, tables=[source_table] if source_table else []
        ),
        template=SimpleNamespace(background="#FFFFFF"),
    )
    plan = SlidePlan(
        title="Synthetic",
        fact_ids=[f.id for f in facts],
        table_id=source_table.id if source_table else None,
    )
    return content_scenes(prs, SimpleNamespace(slides=[plan]), package)[0]


def test_qualitative_table_is_proven_by_actual_cell_text():
    facts = [
        Fact(id="a", text="Requests arrive by email."),
        Fact(id="b", text="A reviewer checks each request."),
    ]
    scene = extracted([["Intake", "Review"], [facts[0].text, facts[1].text]], facts)
    assert len(scene.elements) == 1
    assert scene.elements[0].source_ids == ["a", "b"]
    assert scene.elements[0].rows[1] == [f.text for f in facts]
    empty = extracted([["Intake", "Review"], ["Unrelated", "Unrelated"]], facts)
    assert empty.elements == []


def test_native_table_requires_actual_source_cells():
    source = TableData(id="t", headers=["Topic", "Count"], rows=[["A", "12"]])
    facts = [Fact(id="row", text="A has count 12.", source="t")]
    assert extracted([source.headers] + source.rows, facts, source_table=source).elements[
        0
    ].source_ids == ["row"]
    assert extracted([source.headers, ["A", "120"]], facts, source_table=source).elements == []


@pytest.mark.parametrize("model", [EditorialPlan, EditorialPatch, SlideBatch])
def test_model_requests_require_explicit_grounded_grouping(model):
    schema = editorial_schema(model)
    assert "group" in schema["$defs"]["Claim"]["required"]
    assert "purpose" in schema["$defs"]["EditorialSlide"]["required"]
    assert "quote" not in schema["$defs"]["Citation"]["properties"]


def test_publication_rechecks_exported_structure_instead_of_scene_metadata(tmp_path):
    from studio.models import Plans, VariantPlan
    from studio.generation.audit import audit_diversity
    from studio.generation.results import DiversityResult

    facts = [
        Fact(id="a", text="Alpha accepts requests."),
        Fact(id="b", text="Beta checks requests."),
    ]
    package = SimpleNamespace(
        content=ContentModel(title="Synthetic", facts=facts),
        template=SimpleNamespace(background="#FFFFFF", patterns=[]),
    )
    variants = []
    for i, key in enumerate(("executive", "analytical", "story")):
        variant = VariantPlan(
            key=key, title=key, slides=[SlidePlan(title="Synthetic", fact_ids=["a", "b"])]
        )
        variants.append(variant)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # Different sizes and page position used to produce three signatures.
        for j, fact in enumerate(facts):
            slide.shapes.add_textbox(
                Pt(20 + i * 5), Pt(50 + j * 120), Pt(400 - i * 40), Pt(90 - i * 5)
            ).text = fact.text
        (tmp_path / key).mkdir()
        prs.save(tmp_path / key / "deck.pptx")
    stale = DiversityResult(
        policy="stale",
        distinct=3,
        expected=3,
        verified=True,
        adjustments=[],
        signatures={},
        findings=[],
    )
    result = audit_diversity(
        Plans(variants=variants), {v.key: [] for v in variants}, package, tmp_path, stale
    )
    assert not result.verified and result.distinct == 1
    assert result.method == "exported_source_linked_structure"
    assert result.findings
