"""Real numeric whitespace, totals-only charts and sparse font scales."""

from copy import deepcopy
from types import SimpleNamespace

import pytest
from pptx import Presentation

from studio.charts import chart_projection, make_chart, render_chart, table_series
from studio.config import ROOT
from studio.content import numeric_column, parse_content
from studio.fonts import text_width
from studio.models import Box, Fact, SlidePlan, TableData
from studio.text_composer import fact_elements, text_element
from studio.template_geometry import walk_shapes


@pytest.fixture
def profile():
    return SimpleNamespace(
        font="Montserrat",
        font_file=str(ROOT / "fonts/Montserrat-Regular.ttf"),
        font_roles={},
        font_assets=[],
        body_size=20,
        title_size=32,
        font_sizes=[10, 12, 14, 20, 32],
        foreground="#222222",
        background="#FFFFFF",
        accent="#336699",
        colors=["#336699", "#992255"],
    )


@pytest.mark.parametrize("space", [" ", "\u00a0", "\u202f", "\u2009"])
@pytest.mark.parametrize("suffix", ["", "%", "₽", "млн"])
def test_grouping_whitespace_works_in_analysis_and_chart_without_rewriting_cells(space, suffix):
    value = f"1{space}234,5{suffix}"
    source = parse_content(f"# Данные\n| Период | Значение |\n|---|---|\n| А | {value} |")
    table = source.tables[0]
    before = table.model_dump()
    assert numeric_column(table) == (1, [1234.5], suffix)
    assert table_series(table) == [[1234.5]]
    assert table.model_dump() == before and table.rows[0][1] == value


@pytest.mark.parametrize("minus", ["-", "−"])
def test_negative_grouped_value_keeps_sign(minus):
    table = TableData(
        id="t", headers=["Категория", "Изменение"], rows=[["А", f"{minus}1\u202f234,5"]]
    )
    assert table_series(table) == [[-1234.5]]
    # Legacy proportional bars support only non-negative values.
    assert numeric_column(table) is None


@pytest.mark.parametrize(
    "value",
    ["1,234.50", "not a number", "9" * 400, "NaN", "Infinity", "1 2", "12\u00a034"],
)
def test_invalid_or_nonfinite_values_remain_unplottable(value):
    table = TableData(id="t", headers=["Category", "Value"], rows=[["A", value]])
    assert table_series(table) is None and numeric_column(table) is None


@pytest.mark.parametrize("chart_type", ["bar", "column", "line", "pie", "column_stacked"])
@pytest.mark.parametrize("label", ["Итого", "Всего", "Total"])
def test_totals_alone_remain_editable_chart_categories(profile, tmp_path, chart_type, label):
    table = TableData(id="t", headers=["Категория", "Количество"], rows=[[label, "82"]])
    before = table.model_dump()
    projected, captions = chart_projection(table)
    assert projected.rows == table.rows and not captions
    element = make_chart(
        table,
        SlidePlan(title="Итог", fact_ids=["f1"], chart_type=chart_type),
        Box(x=20, y=20, w=850, h=440),
        profile,
        profile.foreground,
        ["f1"],
    )
    assert element.kind == "chart" and element.labels == [label]
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_chart(slide, element, profile)
    path = tmp_path / "total.pptx"
    prs.save(path)
    actual = next(
        s.chart for s, _ in walk_shapes(Presentation(path).slides[0].shapes) if s.has_chart
    )
    assert [str(c.label) for c in actual.plots[0].categories] == [label]
    assert [list(s.values) for s in actual.series] == [[82.0]]
    assert table.model_dump() == before


def test_detail_rows_still_exclude_aggregate_and_retain_mixed_units():
    table = TableData(
        id="t",
        headers=["Период", "Количество", "Доля"],
        rows=[["А", "20", "40%"], ["Б", "30", "60%"], ["Итого", "50", "100%"]],
    )
    projected, captions = chart_projection(table, compact_captions=True)
    assert projected.rows == [["А", "20"], ["Б", "30"]]
    assert "40%" in " ".join(captions) and "60%" in " ".join(captions)
    assert "Итого" in " ".join(captions) and "100%" in " ".join(captions)


def test_totals_only_mixed_units_keep_unplotted_values():
    table = TableData(
        id="t", headers=["Период", "Количество", "Доля"], rows=[["Итого", "50", "100%"]]
    )
    for compact in [False, True]:
        projected, captions = chart_projection(table, compact_captions=compact)
        assert projected.rows == [["Итого", "50"]]
        assert "100%" in " ".join(captions)


def test_empty_series_never_creates_a_native_chart(profile):
    table = TableData(id="t", headers=["Category", "Value"], rows=[])
    assert table_series(table) is None
    element = make_chart(
        table,
        SlidePlan(title="Empty", fact_ids=["f1"], chart_type="column"),
        Box(x=0, y=0, w=600, h=300),
        profile,
        profile.foreground,
        ["f1"],
    )
    assert element.kind == "table" and element.rows == [table.headers]


@pytest.mark.parametrize(
    "body_size,scale", [(12, [10, 12]), (14, [10, 14, 20]), (20, [12, 15, 20])]
)
def test_readable_floor_is_available_even_when_missing_from_template_scale(
    profile, body_size, scale
):
    profile.body_size = body_size
    profile.font_sizes = scale
    text = "Short evidence"
    box = Box(x=0, y=0, w=text_width(text, profile.font_file, 16) + 1, h=20.5)
    before = deepcopy(vars(profile))
    element = fact_elements([Fact(id="f1", text=text)], box, profile, profile.foreground)[0]
    assert element.size == 16 and element.text == text and element.source_ids == ["f1"]
    assert vars(profile) == before


@pytest.mark.parametrize("role,floor", [("title", 18), ("body", 16), ("subheading", 16)])
def test_direct_text_cannot_be_capped_below_readability_by_small_source_style(profile, role, floor):
    profile.body_size = 12
    profile.font_sizes = [10, 12]
    element = text_element(
        "Evidence",
        Box(x=0, y=0, w=400, h=50),
        profile,
        role=role,
        size=12,
        source_ids=["f1"],
    )
    assert element.size == floor


def test_decorative_footer_size_and_impossible_containers_are_not_hidden(profile):
    profile.body_size = 12
    profile.font_sizes = [10, 12]
    footer = text_element("Page", Box(x=0, y=0, w=100, h=20), profile, role="footer", size=10)
    assert footer.size == 10
    crowded = fact_elements(
        [Fact(id="f1", text="Evidence")], Box(x=0, y=0, w=300, h=8), profile, "#222222"
    )[0]
    assert crowded.box.h <= 8 and crowded.size < 16
