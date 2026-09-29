from types import SimpleNamespace
import asyncio
import pytest
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from studio.templates.native_style import native_styles
from studio.templates.native_template import native_patterns
from studio.contents.parsing import parse_content
from studio.models import Pattern, Box
from studio.contents.sections import prepare_sections, add_dividers
from studio.contents.planner import extractive_plans, assign_compositions, validate_plans
from studio.composition.composer import compose_variant
from studio.checks.audit import audit_scenes


def test_effective_shape_colors_and_empty_layout(template, tmp_path):
    prs = Presentation(template)
    layout = prs.slide_layouts[1]
    for shape in layout.placeholders:
        if shape.has_text_frame:
            shape.text = ""
            shape.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string("C5522E")
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Native title"
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string("17384A")
    path = tmp_path / "style.pptx"
    prs.save(path)
    patterns = native_patterns(prs, native_styles(path))
    assert next(p for p in patterns if p.id == "native-slide-4").title_foreground == "#17384A"
    assert next(p for p in patterns if p.id == "native-layout-0-1").title_foreground == "#C5522E"


def test_small_brand_label_is_not_main_heading(template):
    prs = Presentation(template)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def text(value, y, w, h, size):
        shape = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(w), Inches(h))
        shape.text = value
        shape.text_frame.paragraphs[0].font.size = Pt(size)
        return shape

    text("Brand", 0.2, 3, 0.25, 8)
    title = text("Large heading", 1, 7, 1, 36)
    text("Body", 2.2, 7, 2, 18)
    p = next(p for p in native_patterns(prs) if p.source_slide == 4)
    assert p.title_zone.y >= title.top / 12700
    assert p.title_size == 36 and len(p.body_zones) == 1


def test_variant_bodies_stay_inside_native_zone_and_preserve_colors(prepared):
    _, _, p = prepared
    p.content = parse_content("# Topic\nOne fact.\nSecond fact.\nThird fact.")
    p.constraints.slides = 1
    pattern = Pattern(
        id="native-one",
        source_slide=1,
        source_layout="Native",
        role="statement",
        text_zones=[],
        title_zone=Box(x=50, y=35, w=600, h=60),
        body_zones=[Box(x=50, y=150, w=700, h=300)],
        title_foreground="#154A67",
        zone_foregrounds=["#154A67"],
    )
    p.template.patterns = [pattern]
    plan = assign_compositions(extractive_plans(p), p)
    for v in plan.variants:
        scene = compose_variant(v, p)[0]
        body = [e for e in scene.elements if e.role == "body"]
        assert body
        assert all(
            e.box.x >= 50
            and e.box.x + e.box.w <= 750
            and e.box.y >= 150
            and e.box.y + e.box.h <= 450
            for e in body
        )
        assert {fid for e in body for fid in e.source_ids} == set(v.slides[0].fact_ids)
        assert all(e.color == "#154A67" for e in scene.elements if e.kind == "text")


def test_bullet_uses_text_color_and_size(prepared):
    from studio.models import Element
    from studio.composition.render import set_text

    _, _, p = prepared
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    element = Element(
        kind="text",
        box=Box(x=72, y=72, w=360, h=72),
        text="Bullet",
        font=p.template.font,
        size=20,
        color="#FFFFFF",
        bullet=True,
    )
    set_text(shape.text_frame, element.text, element, p.template)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    from lxml import etree

    root = etree.fromstring(shape._element.xml)
    assert root.xpath("//a:buClr/a:srgbClr/@val", namespaces=ns) == ["FFFFFF"]
    assert root.xpath("//a:defRPr/@sz", namespaces=ns) == ["2000"]


def test_overlapping_and_off_canvas_native_layouts_are_clipped(template):
    prs = Presentation(template)
    layout = prs.slide_layouts[1]
    title, body = layout.placeholders[0], layout.placeholders[1]
    for shape in (title, body):
        shape.left = Inches(0.5)
        shape.width = Inches(9)
    title.top = Inches(0.5)
    title.height = Inches(2)
    body.top = Inches(2)
    body.height = Inches(7)
    pattern = next(p for p in native_patterns(prs) if p.id == "native-layout-0-1")
    assert pattern.title_zone.y + pattern.title_zone.h < pattern.body_zones[0].y
    assert pattern.body_zones[0].y + pattern.body_zones[0].h < prs.slide_height / 12700


def test_numbered_outline_gets_semantic_dividers_without_fact_loss(prepared):
    _, _, p = prepared
    p.content = parse_content(
        "# Project\n"
        + "\n".join(f"## Слайд {i}. Тема {i}\nКороткий факт {i}." for i in range(1, 7))
    )
    p.constraints.slides = 6
    p.template.patterns.append(
        Pattern(
            id="divider",
            role="divider",
            source_slide=1,
            source_layout="Section",
            title_zone=Box(x=80, y=180, w=650, h=100),
            text_zones=[],
            title_size=32,
        )
    )

    class Gateway:
        settings = SimpleNamespace(mode="api")

        async def json_request(self, name, payload, **kwargs):
            assert name == "sections"
            return {
                "chapters": [
                    {"title": "Контекст", "sections": [0, 1, 2]},
                    {"title": "Результаты", "sections": [3, 4, 5]},
                ]
            }

    asyncio.run(prepare_sections(p, Gateway()))
    plans = add_dividers(assign_compositions(extractive_plans(p), p), p)
    validate_plans(plans, p)
    expected = [f.id for f in p.content.facts]
    for v in plans.variants:
        assert len(v.slides) == 6
        assert [s.title for s in v.slides if s.layout == "divider"] == ["Результаты"]
        assert [f for s in v.slides for f in s.fact_ids] == expected
        assert not [f for f in audit_scenes(compose_variant(v, p), p) if f.severity == "error"]
    broken = plans.model_copy(deep=True)
    broken.variants[0].slides[-1].fact_ids.reverse()
    # A swapped source position, even with all facts present, is invalid.
    v = broken.variants[0]
    v.slides[0], v.slides[1] = v.slides[1], v.slides[0]
    with pytest.raises(ValueError, match="порядок"):
        validate_plans(broken, p)


def test_bad_model_partition_is_not_used(prepared):
    _, _, p = prepared
    p.content = parse_content(
        "# Project\n" + "\n".join(f"## Слайд {i}. Тема\nФакт {i}." for i in range(1, 7))
    )
    p.template.patterns.append(
        Pattern(
            id="divider",
            role="divider",
            source_slide=1,
            source_layout="Section",
            title_zone=Box(x=80, y=180, w=650, h=100),
            text_zones=[],
        )
    )

    class Gateway:
        settings = SimpleNamespace(mode="api")

        async def json_request(self, *args, **kwargs):
            return {
                "chapters": [
                    {"title": "A", "sections": [0, 1, 3]},
                    {"title": "B", "sections": [2, 4, 5]},
                ]
            }

    asyncio.run(prepare_sections(p, Gateway()))
    assert "section_groups" not in p.analysis
    assert p.analysis["section_grouping"]["status"] == "failed"
