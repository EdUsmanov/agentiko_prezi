import asyncio
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from studio.contents.parsing import parse_content
from studio.models import ContextualAudit
from studio.contents.planner import plan, extractive_plans, validate_plans
from studio.composition.composer import compose_variant
from studio.composition.render import render_pptx
from studio.templates.parsing import analyze_template


def test_markdown_rules_not_facts():
    content = parse_content(
        "# Проект\n---\n## Слайд 1. **Контекст**\n**Первый факт**\n***\n___\nСлайд 2. Итог\nВторой `факт`"
    )
    assert [f.text for f in content.facts] == ["Первый факт", "Второй факт"]
    assert [f.section for f in content.facts] == ["Слайд 1. Контекст", "Слайд 2. Итог"]


def test_outline_preserved_and_not_sorted_by_numbers(prepared):
    _, _, package = prepared
    package.content = parse_content(
        "# Проект\n## Слайд 1. Контекст\nПервый факт\nВторой факт\n## Слайд 2. Итог\nОбъём 2026\nДетали 75"
    )
    package.constraints.slides = 2
    result = validate_plans(extractive_plans(package), package)
    for variant in result.variants:
        assert [s.fact_ids for s in variant.slides] == [["f1", "f2"], ["f3", "f4"]]
        assert [s.title for s in variant.slides] == ["Контекст", "Итог"]
    result.variants[0].slides.reverse()
    with pytest.raises(ValueError, match="порядок"):
        validate_plans(result, package)


def test_semantic_repair_is_bounded_and_explained(prepared):
    _, _, package = prepared
    valid = extractive_plans(package).model_dump()
    invalid = deepcopy(valid)
    invalid["variants"][0]["slides"][0]["fact_ids"] = ["unknown"]

    class Gateway:
        settings = SimpleNamespace(mode="api")
        calls = []
        requests = []

        async def json_request(self, stage, payload, **kwargs):
            self.requests.append(payload)
            return invalid if len(self.requests) == 1 else valid

    gateway = Gateway()
    plans, warning = asyncio.run(plan(package, gateway, 10))
    assert warning is None and plans
    assert len(gateway.requests) == 2
    assert "несуществующие факты" in gateway.requests[1]["validation_errors"][0]["message"]
    assert gateway.calls[0]["status"] == "rejected"


def test_semantic_repair_failure_reports_safe_reason(prepared):
    _, _, package = prepared
    valid = extractive_plans(package).model_dump()
    valid["variants"][0]["slides"][0]["fact_ids"] = ["unknown"]

    class Gateway:
        settings = SimpleNamespace(mode="api")
        calls = []
        attempts = 0

        async def json_request(self, *args, **kwargs):
            self.attempts += 1
            return valid

    gateway = Gateway()
    _, warning = asyncio.run(plan(package, gateway, 10))
    assert gateway.attempts == 2
    assert "несуществующие факты" in warning


def test_provider_exception_does_not_leak_payload(prepared):
    _, _, package = prepared

    class Gateway:
        settings = SimpleNamespace(mode="api")

        async def json_request(self, *args, **kwargs):
            raise ValueError("secret-api-token PRIVATE CONTENT")

    _, warning = asyncio.run(plan(package, Gateway(), 10))
    assert "secret" not in warning and "PRIVATE" not in warning


def test_critic_has_required_schema():
    assert ContextualAudit.model_json_schema()["required"] == ["findings"]
    with pytest.raises(ValueError):
        ContextualAudit.model_validate({})
    assert ContextualAudit.model_validate({"findings": []}).findings == []


def test_master_artwork_preserved_without_old_copy(template, tmp_path, prepared):
    prs = Presentation(template)
    prs.slide_masters[0].background.fill.solid()
    prs.slide_masters[0].background.fill.fore_color.rgb = RGBColor.from_string("154A67")
    sh = prs.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(5), Pt(5), Pt(12), Pt(12))
    sh.name = "brand-marker"
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    prs.slide_masters[0].shapes._spTree.append(deepcopy(sh._element))
    source = tmp_path / "source.pptx"
    prs.save(source)
    _, _, package = prepared
    package.template = analyze_template(source, tmp_path)
    scenes = compose_variant(extractive_plans(package).variants[0], package)
    output = tmp_path / "native.pptx"
    render_pptx(scenes, package.template, source, output)
    deck = Presentation(output)
    assert str(deck.slide_masters[0].background.fill.fore_color.rgb) == "154A67"
    assert any(sh.name == "brand-marker" for sh in deck.slide_masters[0].shapes)
    assert all(s.strategy == "native_template" for s in scenes)
    with ZipFile(output) as z:
        assert not any(b"OLD PRIVATE" in z.read(n) for n in z.namelist() if n.endswith(".xml"))


def test_native_zones_and_background_layer_are_not_rasterized_in_pptx(prepared, tmp_path):
    _, store, package = prepared
    scenes = compose_variant(extractive_plans(package).variants[0], package)
    for scene in scenes:
        pattern = next(p for p in package.template.patterns if p.id == scene.pattern_id)
        title = next(e for e in scene.elements if e.role == "title")
        assert title.box == pattern.title_zone
        assert not any(e.text.startswith("F1") for e in scene.elements)
    output = tmp_path / "native.pptx"
    render_pptx(scenes, package.template, store.directory(package.id) / "input.pptx", output)
    deck = Presentation(output)
    assert not any(
        hasattr(sh, "image") and sh.width == deck.slide_width
        for s in deck.slides
        for sh in s.shapes
    )


def test_cached_template_layer_tampering_is_rejected(prepared):
    from studio.pipeline import load_package

    _, store, package = prepared
    paths = [p.background_image for p in package.template.patterns if p.background_image]
    if not paths:
        pytest.skip("LibreOffice unavailable")
    Path(paths[0]).write_bytes(b"tampered")
    with pytest.raises(ValueError, match="Фоновый слой"):
        load_package(store, package.id)


def test_office_does_not_receive_api_credentials(tmp_path, monkeypatch):
    from studio.composition import office

    monkeypatch.setenv("LLM_API_KEY", "DO-NOT-SEND")
    monkeypatch.setattr(office, "executable", lambda: "/fake/soffice")
    seen = {}

    def run(args, **kwargs):
        seen.update(kwargs["env"])
        (tmp_path / "deck.pdf").write_bytes(b"%PDF" + b" " * 200)
        assert Path(kwargs["env"]["FONTCONFIG_FILE"]).is_file()
        assert any("UserInstallation=" in a for a in args)
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(office.subprocess, "run", run)
    assert office.to_pdf(tmp_path / "deck.pptx", tmp_path, font_file=tmp_path / "font.ttf")
    assert "LLM_API_KEY" not in seen


def test_deadline_kills_dedicated_worker_group(monkeypatch):
    from studio.jobs import runtime as app

    if app.os.name != "posix":
        pytest.skip("POSIX process groups")
    calls = []
    monkeypatch.setattr(app.os, "killpg", lambda pid, sig: calls.append((pid, sig)))
    app.kill_worker(SimpleNamespace(pid=12345))
    assert calls == [(12345, app.signal.SIGKILL)]


@pytest.mark.parametrize(
    "content",
    ["# Проект\n" + "\n".join(f"## Этап {i}\nИсходный факт {i}." for i in range(1, 13))],
    ids=["roomy"],
)
def test_variants_differ_in_visible_content_geometry(prepared):
    _, _, package = prepared
    from studio.checks.diversity import ensure_diversity, geometry_signature

    decks = {v.key: compose_variant(v, package) for v in extractive_plans(package).variants}
    assert ensure_diversity(decks, package)["verified"]
    signatures = []
    for scenes in decks.values():
        signatures.append(geometry_signature(scenes))
    assert len(set(signatures)) == 3
