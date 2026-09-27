from copy import deepcopy
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from studio.native_template import title_bounds, native_patterns
from studio.export_audit import slide_text, inspect_content
from studio.models import Box, Element, SlideScene, TableData, VariantPlan, SlidePlan


def test_inherited_artwork_limits_title_and_persisted_contract():
    prs = Presentation()
    prs.slide_width = Pt(960)
    prs.slide_height = Pt(540)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    title.left = Pt(50)
    title.top = Pt(220)
    title.width = Pt(800)
    title.height = Pt(140)
    subtitle = slide.placeholders[1]
    subtitle.left = Pt(50)
    subtitle.top = Pt(400)
    subtitle.width = Pt(380)
    subtitle.height = Pt(50)
    art = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(500), Pt(200), Pt(200), Pt(250))
    art.fill.solid()
    art.fill.fore_color.rgb = RGBColor(0, 100, 255)
    slide.slide_layout.shapes._spTree.append(deepcopy(art._element))
    art._element.getparent().remove(art._element)
    pattern = next(p for p in native_patterns(prs) if p.source_slide == 1)
    field = next(f for f in pattern.fields if f["role"] == "title")
    assert pattern.title_zone.x + pattern.title_zone.w == 490
    assert field["box"] == pattern.title_zone.model_dump()


def test_transparent_field_and_full_background_do_not_clip_title():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Pt(960), Pt(540))
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    title = slide.shapes.add_textbox(Pt(50), Pt(230), Pt(800), Pt(140))
    title.text = "Title"
    slide.shapes.add_textbox(Pt(500), Pt(200), Pt(200), Pt(250))
    assert title_bounds(title, slide, slide.slide_layout).w > 780


def test_exported_native_chart_values_reach_text_review(prepared):
    _, _, package = prepared
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Pt(40), Pt(10), Pt(500), Pt(50))
    title.text = "Динамика"
    data = CategoryChartData()
    data.categories = ["Апрель", "Май", "Июнь"]
    data.add_series("Заявки, шт.", [40, 55, 70])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(40), Pt(70), Pt(500), Pt(300), data)
    assert "Май — 55" in slide_text(slide)
    package.content.tables = [
        TableData(
            id="t1",
            headers=["Месяц", "Заявки, шт."],
            rows=[["Апрель", "40"], ["Май", "55"], ["Июнь", "70"]],
        )
    ]
    variant = VariantPlan(
        key="executive",
        title="Данные",
        slides=[
            SlidePlan(
                title="Динамика", fact_ids=[], table_id="t1", layout="chart", chart_type="column"
            )
        ],
    )
    actual, findings = inspect_content(prs, variant, package)
    assert "Май — 55" in actual[0]["actual_text"]
    assert not findings


def test_table_header_fit_measures_bold_face():
    from studio.fonts import resolve_font, text_width, table_cell_fits

    regular = resolve_font("Montserrat")
    bold = resolve_font("Montserrat Bold")
    assert regular and bold
    text = "Проверено"
    size = 17
    regular_width = text_width(text, regular, size)
    bold_width = text_width(text, bold, size)
    assert bold_width > regular_width
    width = (regular_width + bold_width) / 2 / 0.94
    assert table_cell_fits(text, regular, size, width, 100, bold=False)
    assert not table_cell_fits(text, regular, size, width, 100, bold=True)


def test_native_table_padding_is_applied_once(prepared, tmp_path):
    from studio.render import render_pptx

    _, store, package = prepared
    element = Element(
        kind="table",
        box=Box(x=50, y=120, w=360, h=210),
        rows=[["Сценарий", "Проверено"], ["Поиск", "20"]],
        font=package.template.font,
        size=17,
        color="#000000",
        fill="#A0DFE0",
    )
    title = Element(
        kind="text",
        box=Box(x=50, y=20, w=600, h=70),
        text="Таблица",
        font=package.template.font,
        size=25,
        color="#000000",
        role="title",
    )
    scene = SlideScene(
        title="Таблица",
        background="#FFFFFF",
        elements=[title, element],
        source_ids=[],
        layout="table",
    )
    target = tmp_path / "padding.pptx"
    render_pptx([scene], package.template, store.directory(package.id) / "input.pptx", target)
    table = next(s.table for s in Presentation(target).slides[0].shapes if s.has_table)
    for row in table.rows:
        for cell in row.cells:
            assert cell.margin_left == Pt(8) and cell.margin_top == Pt(6)
            assert all(
                getattr(cell.text_frame, "margin_" + side) == 0
                for side in ("left", "right", "top", "bottom")
            )
