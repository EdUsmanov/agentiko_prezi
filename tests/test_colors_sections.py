from zipfile import ZipFile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from studio.colors import (
    agent_color_context,
    color_schemes,
    extract_colors,
    resolve_rendered_schemes,
)
from studio.template import analyze_template
from studio.models import Box, Pattern, Element, SlideScene
from studio.artwork import safe_body_zone
from studio.planner import extractive_plans, validate_plans, assign_compositions
from studio.sections import add_dividers
from studio.composer import compose_variant
from studio.content import parse_content
from studio.render import render_pptx, render_html
from studio.audit import audit_scenes, repair_scenes


def test_color_roles_table_evidence_and_potx(template, potx, tmp_path):
    roles, summary = extract_colors(potx, tmp_path / "report")
    assert "#154A67" in roles["text.other"] and "#FFFFFF" in roles["background"]
    assert summary["slides"] == 3 and summary["visual_accuracy_verified"] is False
    prs = Presentation(template)
    table = prs.slides[0].shapes.add_table(2, 2, Inches(1), Inches(4), Inches(5), Inches(1)).table
    for ri in range(2):
        for ci in range(2):
            table.cell(ri, ci).fill.solid()
            table.cell(ri, ci).fill.fore_color.rgb = RGBColor.from_string(
                "123456" if ri == 0 else "AABBCC"
            )
    path = tmp_path / "colors.pptx"
    prs.save(path)
    roles, summary = extract_colors(path, tmp_path / "report2")
    assert summary["table_styles"]["1"]["header"] == {"color": "#123456", "opacity": 1}
    assert summary["table_styles"]["1"]["body"] == {"color": "#AABBCC", "opacity": 1}


def test_color_schemes_keep_text_with_its_source_background():
    def slide(number, background, text):
        return {
            "number": number,
            "uses": [
                {"role": "background", "color": background, "opacity": 1},
                {"role": "text.other", "color": text, "opacity": 1},
            ],
        }

    schemes, mapping = color_schemes(
        {
            "slides": [
                slide(1, "#EF3124", "#FFFFFF"),
                slide(2, "#EBEBEB", "#000000"),
                slide(3, "#EF3124", "#FFFFFF"),
            ]
        }
    )
    assert [(row["background"], row["foreground"], row["slides"]) for row in schemes] == [
        ("#EF3124", "#FFFFFF", [1, 3]),
        ("#EBEBEB", "#000000", [2]),
    ]
    assert mapping == {"1": "scheme-1", "2": "scheme-2", "3": "scheme-1"}

    # Card copy may be more frequent than the slide headline; it belongs to
    # its own local surface and must not redefine the canvas scheme.
    report = {
        "slides": [
            {
                "number": 13,
                "uses": [
                    {"role": "background", "color": "#EF3124", "opacity": 1},
                    *[{"role": "text.other", "color": "#000000", "opacity": 1} for _ in range(8)],
                    {"role": "text.other", "color": "#FFFFFF", "opacity": 1},
                ],
            }
        ]
    }
    schemes, _ = color_schemes(report, {13: "#FFFFFF"})
    assert schemes[0]["foreground"] == "#FFFFFF"
    assert "#000000" in schemes[0]["text_colors"]


def test_large_authored_white_text_stays_white_on_brand_red(template, tmp_path):
    from studio.template_adaptation import adapt_native_text_fields

    profile = analyze_template(template, tmp_path / "profile")
    pattern = profile.patterns[0]
    profile.patterns = [pattern]
    profile.width, profile.height = 720, 405
    pattern.source_slide = 1
    pattern.reusable = True
    pattern.title_zone = Box(x=30, y=20, w=600, h=80)
    pattern.body_zones = [Box(x=30, y=120, w=600, h=80), Box(x=30, y=230, w=600, h=60)]
    pattern.title_background = "#EF3124"
    pattern.zone_backgrounds = ["#EF3124", "#EF3124"]
    pattern.title_foreground = "#FFFFFF"
    pattern.zone_foregrounds = ["#FFFFFF", "#FFFFFF"]
    profile.colors = ["#EF3124", "#FFFFFF", "#000000"]
    image = tmp_path / "red.png"
    Image.new("RGB", (720, 405), "#EF3124").save(image)
    pattern.background_image = str(image)
    pattern.fields = [
        {"role": "title", "index": 0, "font_size": 60, "style": {"bold": True}},
        {"role": "body", "index": 0, "font_size": 28.5, "style": {"bold": False}},
        {"role": "body", "index": 1, "font_size": 16, "style": {"bold": False}},
    ]
    adapt_native_text_fields(profile)
    assert pattern.title_foreground == "#FFFFFF"
    assert pattern.zone_foregrounds == ["#FFFFFF", "#000000"]


def test_agent_color_context_keeps_canvas_and_card_pairs(template, tmp_path):
    profile = analyze_template(template, tmp_path / "profile")
    pattern = profile.patterns[0]
    profile.patterns = [pattern]
    pattern.color_scheme_id = "scheme-red"
    pattern.title_background = "#EF3124"
    pattern.title_foreground = "#FFFFFF"
    pattern.body_zones = [Box(x=10, y=100, w=300, h=100)]
    pattern.zone_backgrounds = ["#FFFFFF"]
    pattern.zone_foregrounds = ["#000000"]
    profile.color_schemes = [
        {"id": "scheme-red", "background": "#EF3124", "foreground": "#FFFFFF", "slides": [1]}
    ]
    context = agent_color_context(profile)
    assert context["source_schemes"][0]["text"] == "#FFFFFF"
    assert context["patterns"][pattern.id] == {
        "scheme_id": "scheme-red",
        "title": {"background": "#EF3124", "text": "#FFFFFF"},
        "body": [{"background": "#FFFFFF", "text": "#000000"}],
    }


def test_rendered_schemes_resolve_dominant_raster_background(template, tmp_path):
    profile = analyze_template(template, tmp_path / "profile")
    base = profile.patterns[0]
    profile.patterns = []
    for number, background, foreground in (
        (1, "#152848", "#EEF2F5"),
        (2, "#1B2F52", "#EEF2F5"),
        (3, "#F5F1EA", "#111827"),
    ):
        pattern = base.model_copy(deep=True)
        pattern.id = f"slide-{number}"
        pattern.source_slide = number
        pattern.title_background = background
        pattern.title_foreground = foreground
        profile.patterns.append(pattern)
    original = profile.color_schemes
    resolve_rendered_schemes(profile)
    assert profile.color_analysis["static_schemes"] == original
    assert len(profile.color_schemes) == 2
    assert profile.color_schemes[0]["slides"] == [1, 2]
    assert profile.background == "#152848"
    assert profile.foreground == "#EEF2F5"
    assert profile.patterns[0].color_scheme_id == profile.patterns[1].color_scheme_id


def test_transparent_table_native_and_html(prepared, tmp_path):
    _, store, p = prepared
    e = Element(
        kind="table",
        box=Box(x=70, y=150, w=600, h=100),
        rows=[["Name", "Value"], ["A", "10"]],
        font=p.template.font,
        size=16,
        color=p.template.foreground,
        fill=p.template.accent,
        fill_opacity=0.4,
    )
    scene = SlideScene(
        title="Table", background=p.template.background, elements=[e], source_ids=[], layout="table"
    )
    target = tmp_path / "table.pptx"
    render_pptx(
        [scene], p.template, store.directory(p.id) / "input.pptx", target, verify_text=False
    )
    with ZipFile(target) as z:
        from lxml import etree

        root = etree.fromstring(z.read("ppt/slides/slide1.xml"))
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        assert root.xpath("count(//a:tr[2]/a:tc/a:tcPr/a:noFill)", namespaces=ns) == 2
        assert root.xpath("//a:tr[1]/a:tc/a:tcPr/a:solidFill//a:alpha/@val", namespaces=ns) == [
            "40000",
            "40000",
        ]
        assert root.xpath("count(//a:tr[2]/a:tc/a:tcPr/a:lnB/a:solidFill)", namespaces=ns) == 2
    render_html([scene], p.template, tmp_path / "table.html")
    html = (tmp_path / "table.html").read_text()
    assert "transparent" in html and ",0.4)" in html


def test_sections_keep_budget_facts_and_native_artwork(prepared):
    _, _, p = prepared
    p.content = parse_content(
        "# Проект\n## Контекст\nПервый факт.\nВторой факт.\nТретий факт.\n## Результаты\nРезультат один.\nРезультат два.\nРезультат три."
    )
    p.constraints.slides = 6
    p.template.patterns.append(
        Pattern(
            id="divider",
            role="divider",
            source_slide=1,
            source_layout="Section",
            title_zone=Box(x=80, y=180, w=650, h=100),
            text_zones=[],
            title_size=32,
        )
    )
    plans = add_dividers(assign_compositions(extractive_plans(p), p), p)
    validate_plans(plans, p)
    for variant in plans.variants:
        assert len(variant.slides) == 6
        assert any(s.layout == "divider" for s in variant.slides)
        assert {f for s in variant.slides for f in s.fact_ids} == {f.id for f in p.content.facts}
        scenes = compose_variant(variant, p)
        assert not [f for f in audit_scenes(scenes, p) if f.severity == "error"]
    from studio.deeppresenter import CompositionEnvironment

    env = CompositionEnvironment(p, plans)
    assert "divider" in env.allowed


def test_native_title_only_layout_retained(template):
    from studio.native_template import native_patterns

    prs = Presentation(template)
    divider = prs.slides.add_slide(prs.slide_layouts[5])
    divider.shapes.title.text = "Раздел"
    assert any(p.source_slide == 4 and p.role == "divider" for p in native_patterns(prs))


def test_safe_area_avoids_bottom_decoration():
    image = Image.new("RGB", (400, 200), "white")
    # A pixel fixture, not a generated presentation asset.
    image.paste((30, 40, 50), (0, 150, 400, 200))
    box = safe_body_zone(image, Box(x=0, y=0, w=400, h=200), scale=1)
    assert box.y + box.h <= 150
    assert box.w == 400


def test_table_repair_does_not_split_long_header_word(prepared):
    from studio.fonts import table_cell_fits, element_font

    _, _, p = prepared
    e = Element(
        kind="table",
        box=Box(x=80, y=150, w=400, h=250),
        rows=[["Официальные ролики", "Фанатские публикации"], ["12", "18"]],
        font=p.template.font,
        size=32,
        color=p.template.foreground,
    )
    s = SlideScene(
        title="Таблица",
        background=p.template.background,
        elements=[e],
        source_ids=[],
        layout="table",
    )
    repair_scenes([s], p)
    assert e.size < 32
    assert table_cell_fits(e.rows[0][0], element_font(p.template, e)[1], e.size, 184, 113, True)
