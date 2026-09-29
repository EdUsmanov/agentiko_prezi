"""Integration checks for the supplied font extraction kit and Studio adapter."""

import hashlib
import json
import asyncio
import base64
import struct
import shutil
from types import SimpleNamespace
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.util import Pt
from fontTools import subset
from fontTools.ttLib import TTFont

from studio.config import ROOT
from studio.templates.font_coverage import ensure_text_coverage
from studio.templates.font_extraction import (
    _apply_missing_font_fallbacks,
    _kit,
    extract_template_fonts,
)
from studio.templates.fonts import check_glyphs, role_font
from studio.templates.parsing import analyze_template


def test_bundle_model_and_role_adapter(template, tmp_path):
    prs = Presentation(template)
    for slide in prs.slides:
        for run in slide.shapes[0].text_frame.paragraphs[0].runs:
            run.font.name = "Montserrat Medium"
    source = tmp_path / "roles.pptx"
    prs.save(source)

    profile = analyze_template(source, tmp_path / "analysis")
    model = json.loads((tmp_path / "analysis/font-model.json").read_text())
    assert model["schemaVersion"] == 2
    assert len(model["slides"]) == 3
    assert "OLD PRIVATE" not in json.dumps(model)
    assert role_font(profile, "title")[0] == "Montserrat Medium"
    assert role_font(profile, "body")[0] == "Play"
    assert not profile.missing_fonts
    for asset in model["fontAssets"]:
        path = tmp_path / "analysis" / asset["path"]
        assert path.is_file()
        assert hashlib.sha256(path.read_bytes()).hexdigest() == asset["sha256"]


def test_missing_face_is_downloaded_or_substituted_without_blocking(tmp_path, monkeypatch):
    from studio.templates import font_extraction

    monkeypatch.setattr(font_extraction, "ROOT", tmp_path)
    bundled = tmp_path / "fonts"
    bundled.mkdir()
    shutil.copy2(ROOT / "fonts/Montserrat-Regular.ttf", bundled / "Montserrat-Regular.ttf")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, Pt(400), Pt(80))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "Missing face"
    run.font.name = "Unobtainable Test Family"
    run.font.size = Pt(32)
    source = tmp_path / "missing.pptx"
    prs.save(source)

    report = extract_template_fonts(source, tmp_path / "analysis", allow_download=False)
    assert report["unresolved"][0]["requested"] == "Unobtainable Test Family"
    assert not report["unresolved"][0]["required_for_generation"]
    assert report["unresolved"][0]["substituted_by"] == "Montserrat"
    assert report["primary"]["requested"] == "Montserrat"
    assert report["replacements"][0]["reason"] == "missing_font"
    model = json.loads((tmp_path / "analysis/font-model.json").read_text())
    assert model["slides"][0]["elements"]["title"][0]["status"] == "missing"


def test_missing_table_face_falls_back_to_resolved_body_font(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, Pt(400), Pt(80))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "Main text"
    run.font.name = "Play"
    table = slide.shapes.add_table(1, 1, 0, Pt(100), Pt(400), Pt(100)).table
    cell_run = table.cell(0, 0).text_frame.paragraphs[0].add_run()
    cell_run.text = "Table text"
    cell_run.font.name = "Unobtainable Table Family"
    source = tmp_path / "missing-table.pptx"
    prs.save(source)

    report = extract_template_fonts(source, tmp_path / "analysis", allow_download=False)

    table_face = next(a for a in report["assets"] if a["id"] == report["roles"]["table"])
    assert table_face["requested"] == "Play"
    assert "Unobtainable Table Family" in table_face["template_aliases"]
    assert report["primary"]["requested"] == "Play"
    assert report["replacements"][0]["fallback_font"] == "Play"


def test_missing_bold_face_uses_bold_variant_of_fallback_family():
    regular = {
        "id": "regular",
        "requested": "Available",
        "family": "Available",
        "weight": 400,
        "style": "normal",
    }
    bold = {**regular, "id": "bold", "requested": "Available Bold", "weight": 700}
    model = SimpleNamespace(
        unresolved=[SimpleNamespace(family="Unavailable", weight=700, style="normal")]
    )
    _, replacements, fallback = _apply_missing_font_fallbacks(
        model,
        [regular, bold],
        {"body": regular["id"]},
        {"body": ("Available", 400, "normal")},
        allow_download=True,
    )
    assert fallback[("Unavailable", 700, "normal")]["requested"] == "Available Bold"
    assert replacements[0]["style_changed"] is False


def test_exact_local_font_is_used_without_network(template, tmp_path):
    report = extract_template_fonts(template, tmp_path / "analysis", allow_download=False)
    assert report["primary"]["requested"] == "Play"
    assert report["primary"]["status"] == "installed"
    assert (
        hashlib.sha256(Path(report["primary"]["path"]).read_bytes()).hexdigest()
        == report["primary"]["sha256"]
    )


def test_custom_data_directory_font_is_found(template, tmp_path, monkeypatch):
    from studio.templates import font_extraction

    monkeypatch.setattr(font_extraction, "ROOT", tmp_path / "empty")
    local = tmp_path / "data/local-fonts"
    local.mkdir(parents=True)
    shutil.copy2(ROOT / "fonts/Play-Regular.ttf", local / "play.ttf")
    report = extract_template_fonts(template, tmp_path / "data/jobs/example", allow_download=False)
    assert report["primary"]["requested"] == "Play"
    assert not report["unresolved"]


def test_embedded_medium_uses_binary_weight(template, tmp_path):
    raw = (ROOT / "fonts/Montserrat-Medium.ttf").read_bytes()
    prs = Presentation(template)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Montserrat Medium"
    source = tmp_path / "styled.pptx"
    prs.save(source)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    p = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    r = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    presentation = ET.fromstring(parts["ppt/presentation.xml"])
    entry = ET.SubElement(ET.SubElement(presentation, p + "embeddedFontLst"), p + "embeddedFont")
    ET.SubElement(entry, p + "font", {"typeface": "Montserrat Medium"})
    ET.SubElement(entry, p + "regular", {r + "id": "rIdFontTest"})
    parts["ppt/presentation.xml"] = ET.tostring(presentation)
    relations = ET.fromstring(parts["ppt/_rels/presentation.xml.rels"])
    ET.SubElement(
        relations,
        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
        {
            "Id": "rIdFontTest",
            "Type": r[1:-1] + "/font",
            "Target": "fonts/font1.fntdata",
        },
    )
    parts["ppt/_rels/presentation.xml.rels"] = ET.tostring(relations)
    parts["ppt/fonts/font1.fntdata"] = raw
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    report = extract_template_fonts(source, tmp_path / "analysis", allow_download=False)
    assert not report["unresolved"]
    assert report["primary"]["status"] == "embedded"
    assert report["primary"]["weight"] == 500
    assert Path(report["primary"]["path"]).read_bytes() == raw


def test_bundled_eot_decoder_round_trip():
    raw = (ROOT / "fonts/Play-Regular.ttf").read_bytes()
    header = bytearray(96)
    struct.pack_into("<4I", header, 0, len(header) + len(raw), len(raw), 0x10000, 0)
    struct.pack_into("<H", header, 34, 0x504C)
    decoded = asyncio.run(_kit("decoder").EmbeddedFontDecoder().decode_eot(bytes(header) + raw))
    assert decoded["embeddable"]
    assert base64.b64decode(decoded["data"]) == raw


def test_cyrillic_fallback_keeps_reference_model_intact(template, tmp_path):
    profile = analyze_template(template, tmp_path / "analysis")
    reference_model = (tmp_path / "analysis/font-model.json").read_bytes()
    latin_only = tmp_path / "latin-only.ttf"
    font = TTFont(ROOT / "fonts/Play-Regular.ttf")
    subsetter = subset.Subsetter()
    subsetter.populate(unicodes=range(32, 127))
    subsetter.subset(font)
    font.save(latin_only)
    font.close()
    primary = next(a for a in profile.font_assets if a["path"] == profile.font_file)
    primary["path"] = str(latin_only)
    primary["requested"] = "Fixture Sans"
    primary["family"] = "Fixture Sans"
    profile.font_file = str(latin_only)
    profile.font = "Fixture Sans"
    profile.missing_fonts = [
        {
            "requested": "Unavailable Bold",
            "weight": 700,
            "style": "normal",
            "required_for_generation": False,
            "substituted_by": "Fixture Sans",
        }
    ]
    profile.font_replacements = [
        {
            "scope": "font",
            "template_font": "Unavailable Bold",
            "fallback_font": "Fixture Sans",
            "reason": "missing_font",
            "style_changed": True,
        }
    ]

    records = ensure_text_coverage(profile, "Привет, мир")

    assert [(r["template_font"], r["fallback_font"]) for r in records] == [
        ("Unavailable Bold", "Montserrat"),
        ("Fixture Sans", "Montserrat"),
    ]
    assert profile.missing_fonts[0]["substituted_by"] == "Montserrat"
    assert profile.font == "Montserrat"
    assert role_font(profile, "body")[0] == "Montserrat"
    check_glyphs(profile.font_file, "Привет, мир")
    assert (tmp_path / "analysis/font-model.json").read_bytes() == reference_model
