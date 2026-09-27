"""No-wrap PowerPoint text must be audited using its saved line breaks."""

from types import SimpleNamespace
import pytest
from pptx import Presentation
from pptx.util import Pt
from studio.config import ROOT
from studio.export_audit import geometry
from studio.fonts import text_width


@pytest.fixture
def profile():
    regular = str(ROOT / "fonts/Montserrat-Regular.ttf")
    bold = str(ROOT / "fonts/Montserrat-Bold.ttf")
    return SimpleNamespace(
        width=1440,
        height=810,
        body_size=20,
        font="Montserrat",
        font_file=regular,
        font_roles={"body": "r"},
        font_assets=[
            {"id": "r", "path": regular, "requested": "Montserrat"},
            {"id": "b", "path": bold, "requested": "Montserrat Bold"},
        ],
    )


def textbox(text, width, height, size=24, bold=False, wrap=False):
    prs = Presentation()
    prs.slide_width = Pt(1440)
    prs.slide_height = Pt(810)
    shape = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Pt(20), Pt(20), Pt(width), Pt(height)
    )
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = wrap
    p = frame.paragraphs[0]
    p.text = text
    p.line_spacing = Pt(size * 1.25)
    p.font.name = "Montserrat"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.space_before = p.space_after = Pt(0)
    for run in p.runs:
        run.font.name = "Montserrat"
        run.font.size = Pt(size)
        run.font.bold = bold
    return prs, shape


@pytest.mark.parametrize(
    "text,width,height,size",
    [
        ("Six Stages\vof Hitler's\vLife", 385.3068, 310.5953, 70),
        ("Nazi Vote Share\v1928–1932", 496.125, 166.5, 58),
    ],
)
def test_saved_headline_lines_fit_without_artificial_rewrapping(
    profile, text, width, height, size, tmp_path
):
    prs, shape = textbox(text, width, height, size, bold=True)
    path = tmp_path / "saved.pptx"
    prs.save(path)
    prs = Presentation(path)
    before = prs.slides[0].shapes[0]._element.xml
    findings, repairs = geometry(prs, profile, repair=True)
    assert findings == [] and repairs == []
    assert prs.slides[0].shapes[0]._element.xml == before


def test_no_wrap_horizontal_overflow_is_not_hidden_by_hypothetical_wrapping(profile):
    prs, _ = textbox("Wide text that must not wrap", 100, 500)
    assert any(f["code"] == "pptx_text_overflow" for f in geometry(prs, profile)[0])


@pytest.mark.parametrize("wrap", [False, True])
def test_explicit_breaks_cannot_be_collapsed_to_hide_vertical_overflow(profile, wrap):
    prs, _ = textbox("A\vB\vC", 400, 40, wrap=wrap)
    assert any(f["code"] == "pptx_text_overflow" for f in geometry(prs, profile)[0])


def test_each_run_uses_its_own_bold_metrics(profile):
    prefix = "Normal "
    tail = "WWWWWWWW"
    regular, bold = [a["path"] for a in profile.font_assets]
    underestimated = text_width(prefix + tail, regular, 24)
    actual = text_width(prefix, regular, 24) + text_width(tail, bold, 24)
    assert actual > underestimated + 2
    prs, shape = textbox(prefix, (underestimated + actual) / 2, 100)
    p = shape.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = tail
    r.font.name = "Montserrat"
    r.font.size = Pt(24)
    r.font.bold = True
    assert any(f["code"] == "pptx_text_overflow" for f in geometry(prs, profile)[0])


def test_bullet_margin_and_frame_padding_reduce_real_available_width(profile):
    text = "A short label"
    width = text_width(text, profile.font_file, 24) + 5
    prs, shape = textbox(text, width, 100)
    assert not geometry(prs, profile)[0]
    shape.text_frame.margin_right = Pt(10)
    assert any(f["code"] == "pptx_text_overflow" for f in geometry(prs, profile)[0])
    shape.text_frame.margin_right = 0
    props = shape.text_frame.paragraphs[0]._p.get_or_add_pPr()
    props.set("marL", str(Pt(20)))
    props.set("indent", str(-Pt(20)))
    assert any(f["code"] == "pptx_text_overflow" for f in geometry(prs, profile)[0])


def test_repair_still_fixes_genuine_no_wrap_overflow_without_moving_or_erasing_text(profile):
    prs, shape = textbox("Wide text", 95, 100)
    before = (shape.text, shape.left, shape.top, shape.width, shape.height)
    assert geometry(prs, profile)[0]
    findings, repairs = geometry(prs, profile, repair=True)
    assert not findings and repairs
    assert before == (shape.text, shape.left, shape.top, shape.width, shape.height)
    assert not geometry(prs, profile)[0]
