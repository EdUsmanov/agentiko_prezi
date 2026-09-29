import asyncio
from dataclasses import replace
import time

import pytest
from fastapi.testclient import TestClient

from studio.app import create_app
from studio.checks.audit import audit_scenes, repair_scenes
from studio.composition.composer import compose_variant
from studio.config import Settings
from studio.contents.parsing import parse_content, numeric_column
from studio.models import Box, Element, TableData
from studio.contents.planner import extractive_plans, validate_plans, assign_compositions
from studio.composition.render import primitives
from studio.jobs.store import Store


def test_visible_coverage_not_scene_metadata(prepared):
    _, _, package = prepared
    scenes = compose_variant(extractive_plans(package).variants[0], package)
    for scene in scenes:
        scene.elements = [e for e in scene.elements if not e.source_ids]
    assert any(f.code == "coverage" for f in audit_scenes(scenes, package))


def test_missing_table_link_is_recovered_without_losing_values(prepared):
    _, _, package = prepared
    package.content = parse_content(
        "# Проект\n| Канал | Заявки |\n|---|---|\n| А | 200 |\n| Б | 300 |"
    )
    package.constraints.slides = 1
    plans = extractive_plans(package)
    for variant in plans.variants:
        variant.slides[0].table_id = None
        variant.slides[0].layout = "statement"
    plans = assign_compositions(validate_plans(plans, package), package)
    for variant in plans.variants:
        assert variant.slides[0].table_id == package.content.tables[0].id
        scenes = compose_variant(variant, package)
        assert any(e.kind in ("table", "chart") for s in scenes for e in s.elements)
        assert not any(f.code == "coverage" for f in audit_scenes(scenes, package))


def test_large_chart_number_and_zero_are_not_misrepresented(prepared):
    _, _, package = prepared
    chart = Element(
        kind="chart",
        box=Box(x=0, y=0, w=600, h=300),
        labels=["А", "Б"],
        values=[1000001, 0],
        value_labels=["1 000 001", "0"],
        size=20,
        font=package.template.font,
        color=package.template.foreground,
        fill=package.template.accent,
    )
    items = primitives(chart, package.template)
    assert any(e.text == "1 000 001" for e in items)
    assert len([e for e in items if e.kind == "rect"]) == 1


def test_chart_label_overflow_repairs_to_full_table(prepared):
    _, _, package = prepared
    package.content = parse_content(
        "# Проект\n| Канал | Заявки |\n|---|---|\n| А | 1000001 |\n| Б | 0 |"
    )
    package.constraints.slides = 1
    plans = assign_compositions(extractive_plans(package), package)
    scenes = compose_variant(plans.variants[0], package)
    chart = next(e for e in scenes[0].elements if e.kind == "chart")
    chart.box.h = 20
    assert any(f.code == "chart_overflow" for f in audit_scenes(scenes, package))
    repairs = repair_scenes(scenes, package)
    assert any(f.code == "chart_to_table" for f in repairs)
    assert (
        next(e for e in scenes[0].elements if e.kind == "table").rows
        == [package.content.tables[0].headers] + package.content.tables[0].rows
    )


def test_infinite_numeric_values_do_not_become_charts():
    assert numeric_column(TableData(id="t", headers=["А", "Б"], rows=[["x", "9" * 400]])) is None


def test_restart_recovers_jobs_beyond_history_pagination(tmp_path):
    store = Store(tmp_path)
    old = store.create("generation")
    for _ in range(31):
        job = store.create("preparation")
        store.update(job["id"], "ready")
    store.recover()
    assert store.get(old["id"])["state"] == "failed"
    assert store.create("generation")["state"] == "accepted"


@pytest.mark.parametrize("length", ["broken", "-1"])
def test_invalid_content_length_returns_400(tmp_path, length):
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        result = client.post("/api/generate", content=b"{}", headers={"content-length": length})
        assert result.status_code == 400


def test_deeppresenter_failure_never_silently_falls_back(prepared, monkeypatch):
    import studio.providers.deeppresenter as runtime
    from studio.pipeline import generate

    settings, store, package = prepared

    # Retain offline prepared plan so the test does not need network at all.
    async def fail(*args, **kwargs):
        raise ValueError("PRIVATE PAYLOAD")

    monkeypatch.setattr(runtime, "design", fail)
    monkeypatch.setattr(
        "studio.pipeline.ModelGateway",
        lambda settings: type("G", (), {"settings": settings, "calls": [], "usage": []})(),
    )
    job = store.create("generation", {"package_id": package.id, "deadline_at": time.time() + 300})
    with pytest.raises(ValueError, match="Подмена движка не выполнялась") as caught:
        asyncio.run(generate(store, job["id"], replace(settings, engine="deeppresenter")))
    assert "PRIVATE PAYLOAD" not in str(caught.value)
    assert not (store.directory(job["id"]) / "presentations.zip").exists()


def test_revision_preserves_tables_and_model_provenance(prepared):
    from studio.models import Fact
    from studio.pipeline import prepare, load_package
    import shutil

    settings, store, package = prepared
    content = parse_content("# Проект\n| Канал | Значение |\n|---|---|\n| А | 20 |\n| Б | 30 |")
    content.facts.append(
        Fact(
            id="draft1",
            text="Проектный тезис — требует проверки: организовать единый приём заявок.",
            source="model_proposal",
        )
    )
    package.constraints.slides = 3
    package.constraints.count_mode = "maximum"
    job = store.create("preparation")
    shutil.copyfile(
        store.directory(package.id) / "input.pptx", store.directory(job["id"]) / "input.pptx"
    )
    prepare(
        store,
        job["id"],
        "",
        "",
        "Сделай заголовки короче",
        None,
        settings,
        content,
        package.constraints,
    )
    revised = load_package(store, job["id"])
    assert revised.content.model_dump() == content.model_dump()
    assert revised.constraints.count_mode == "maximum" and revised.constraints.slides == 3
    assert revised.manifest["content_hash_format"] == "canonical_json"


def test_later_count_instruction_wins():
    from studio.contents.parsing import parse_constraints

    result = parse_constraints(None, "", "Ровно 10 слайдов. Уточнение: до 5 слайдов.")
    assert (result.slides, result.count_mode) == (5, "maximum")


def test_brief_reexpansion_keeps_unique_fact_ids(prepared):
    from studio.contents.author import expand_brief
    from studio.models import Fact
    from types import SimpleNamespace

    _, _, package = prepared
    package.content.facts = [
        Fact(id="draft1", text="Предложение из прошлого запуска", source="model_proposal")
    ]
    package.constraints.slides = 2

    class Gateway:
        settings = SimpleNamespace(mode="api")

        async def json_request(self, *args, **kwargs):
            return {
                "proposals": [
                    "Организовать единый интерфейс для обработки внутренних заявок сотрудников."
                ]
            }

    result, warning = asyncio.run(expand_brief(package, Gateway(), 10))
    assert warning is None
    assert [f.id for f in result.content.facts] == ["draft1", "draft2"]


def test_mixed_placeholder_title_and_card_pattern(tmp_path):
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.enum.shapes import MSO_SHAPE
    from studio.templates.native_template import native_patterns, source_slide

    prs = Presentation()
    prs.slide_width = Pt(960)
    prs.slide_height = Pt(540)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.left = Pt(30)
    title.top = Pt(30)
    title.width = Pt(600)
    title.height = Pt(80)
    title.text = "OLD TITLE"
    for x in (30, 330, 630):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(140), Pt(270), Pt(280))
        sh.text = "OLD CARD TEXT"
        sh.text_frame.margin_top = Pt(60)
    from io import BytesIO
    from PIL import Image

    picture = BytesIO()
    Image.new("RGB", (20, 20), "red").save(picture, format="PNG")
    picture.seek(0)
    group = slide.shapes.add_group_shape()
    group.shapes.add_picture(picture, Pt(800), Pt(500), Pt(30), Pt(30))
    patterns = native_patterns(prs)
    pattern = next(p for p in patterns if p.source_slide == 1)
    assert len(pattern.body_zones) == 3
    assert all(b.y == 200 for b in pattern.body_zones)
    assert all(b.h < 221 for b in pattern.body_zones)
    # Retain card geometry, but not old copy or an unapproved grouped photo.
    prs._studio_sources = list(prs.slides)
    prs._studio_brand_hashes = set()
    generated = source_slide(prs, pattern)
    assert len(generated.shapes) == 3
    assert all(not sh.text for sh in generated.shapes)


def test_picture_artwork_reduces_body_safe_area(tmp_path):
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.enum.shapes import MSO_SHAPE
    from studio.templates.native_template import native_patterns
    from io import BytesIO
    from PIL import Image, ImageDraw

    prs = Presentation()
    prs.slide_width = Pt(960)
    prs.slide_height = Pt(540)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Заголовок"
    picture = Image.new("RGB", (400, 300), "white")
    draw = ImageDraw.Draw(picture)
    draw.rectangle((0, 210, 400, 300), fill="#0077ff")
    raw = BytesIO()
    picture.save(raw, format="PNG")
    raw.seek(0)
    slide.shapes.add_picture(raw, Pt(40), Pt(140), Pt(400), Pt(300))
    body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(60), Pt(160), Pt(360), Pt(260))
    body.text = "Исходный текст"
    body.fill.background()
    body.line.fill.background()
    pattern = next(p for p in native_patterns(prs) if p.source_slide == 1)
    assert pattern.body_zones[0].h < 210


def test_critic_cannot_attribute_other_titles_or_facts(prepared):
    from studio.checks.review_grounding import grounded_review

    _, _, package = prepared
    plans = extractive_plans(package)
    slide = plans.variants[0].slides[0]
    finding = {
        "variant": "executive",
        "slide": 1,
        "title_quote": slide.title,
        "fact_ids": slide.fact_ids,
        "code": "test",
        "severity": "warning",
        "message": "Test finding",
    }
    assert grounded_review({"findings": [finding]}, plans)["findings"]
    for changed in ({"title_quote": "Invented title"}, {"fact_ids": ["unknown"]}, {"slide": 30}):
        with pytest.raises(ValueError):
            grounded_review({"findings": [{**finding, **changed}]}, plans)


def test_layout_scoring_accounts_for_chart_fallback_table(prepared):
    _, _, package = prepared
    package.content = parse_content(
        "# Данные\n| Канал | Заявки |\n|---|---|\n| Почта | 180 |\n| Мессенджеры | 120 |\n| Портал | 90 |"
    )
    package.constraints.slides = 1
    source = next(p for p in package.template.patterns if p.title_zone and p.body_zones)
    small = source.model_copy(deep=True)
    small.id = "small-chart"
    small.body_zones = [Box(x=50, y=150, w=300, h=40)]
    large = source.model_copy(deep=True)
    large.id = "large-chart"
    large.body_zones = [Box(x=50, y=150, w=700, h=300)]
    package.template.patterns = [small, large]
    plans = assign_compositions(extractive_plans(package), package)
    scenes = compose_variant(plans.variants[0], package)
    assert scenes[0].pattern_id == "large-chart"
    repair_scenes(scenes, package)
    assert not any(f.severity == "error" for f in audit_scenes(scenes, package))


def test_semantic_plan_cannot_pin_untested_physical_layout(prepared):
    from studio.contents.planner import planning_schema

    _, _, package = prepared
    plans = extractive_plans(package)
    for variant in plans.variants:
        for slide in variant.slides:
            slide.pattern_id = next(p.id for p in package.template.patterns if p.title_zone)
    baseline = assign_compositions(plans, package)
    assert all(s.pattern_id is None for v in baseline.variants for s in v.slides)
    assert planning_schema(package)["$defs"]["SlidePlan"]["properties"]["pattern_id"] == {
        "type": "null"
    }


def test_dense_slide_does_not_veto_other_safe_diversity(prepared, monkeypatch):
    import studio.checks.diversity as diversity
    from studio.models import Finding

    _, _, package = prepared
    scenes = compose_variant(extractive_plans(package).variants[0], package)
    # The shared template produces one-line body boxes. Give the other slides
    # real wrapping space: the quality guard must reject overflow even when
    # this test's synthetic dense-slide error detector accepts a candidate.
    for scene in scenes[1:]:
        for element in scene.elements:
            if element.kind == "text" and element.role == "body":
                element.box.h = 100
    baseline_width = next(e.box.w for e in scenes[0].elements if e.source_ids and e.role != "title")
    # Make only the first slide too dense to shrink; the others still have room.
    scenes[0].title = "DENSE"

    def audit(items, package):
        if any(
            s.title == "DENSE"
            and any(
                e.source_ids and e.role != "title" and e.box.w < baseline_width for e in s.elements
            )
            for s in items
        ):
            return [
                Finding(
                    code="text_overflow", severity="error", message="Dense slide overflow", slide=1
                )
            ]
        return []

    monkeypatch.setattr(diversity, "audit_scenes", audit)
    package.template.patterns = []
    dense_candidate = [s.model_copy(deep=True) for s in scenes]
    next(e for e in dense_candidate[0].elements if e.source_ids and e.role != "title").box.w *= 0.84
    assert not diversity.preserves_quality(scenes, dense_candidate, package)
    safe_candidate = [s.model_copy(deep=True) for s in scenes]
    next(e for e in safe_candidate[1].elements if e.source_ids and e.role != "title").box.w *= 0.84
    assert diversity.preserves_quality(scenes, safe_candidate, package)

    decks = {
        k: [s.model_copy(deep=True) for s in scenes] for k in ("executive", "analytical", "story")
    }
    report = diversity.ensure_diversity(decks, package)
    assert not report["verified"] and report["findings"]
    assert report["adjustments"] == []
    assert all(deck == scenes for deck in decks.values())
