"""Geometry regressions built from synthetic templates, without user samples."""

from copy import deepcopy
from types import SimpleNamespace
import pytest
from studio.models import Box, Element, SlideScene
from studio.template import analyze_template
from studio.template_adaptation import derive_safe_cover_patterns
from studio.table_style import column_widths
from studio.fonts import element_font, table_cell_fits
from studio.render import render_pptx, render_html


@pytest.mark.parametrize("width,height", [(720, 405), (960, 540), (800, 600)])
def test_cover_uses_only_verified_region(template, tmp_path, width, height):
    p = analyze_template(template, tmp_path / "profile")
    p.width = width
    p.height = height
    pattern = p.patterns[0]
    p.patterns = [pattern]
    pattern.role = pattern.purpose = "cover"
    pattern.reusable = True
    pattern.title_zone = Box(x=200, y=120, w=300, h=15)
    pattern.body_zones = [Box(x=200, y=160, w=300, h=14)]
    p.colors = list(dict.fromkeys(p.colors + ["#FFFFFF"]))
    pattern.safe_text_zone = {
        "box": [48, 48, 352, 672],
        "minimum_contrast": 12,
        "text_color": "white",
        "coordinate_space": {"width": 1280, "height": 720},
    }
    original = pattern.model_copy(deep=True)
    records = derive_safe_cover_patterns(p)
    assert len(records) == 1 and pattern == original
    adapted = p.patterns[-1]
    region = Box.model_validate(records[0]["region"])
    for zone in [adapted.title_zone, *adapted.body_zones]:
        assert region.x <= zone.x and region.y <= zone.y
        assert zone.x + zone.w <= region.x + region.w + 0.01
        assert zone.y + zone.h <= region.y + region.h + 0.01
    assert adapted.title_zone.y + adapted.title_zone.h < adapted.body_zones[0].y
    assert adapted.title_size >= 18 and adapted.background_image == original.background_image
    assert derive_safe_cover_patterns(p) == []


@pytest.mark.parametrize(
    "fault", ["no_region", "contrast", "missing_contrast", "unknown_color", "outside"]
)
def test_cover_cannot_invent_a_safe_region(template, tmp_path, fault):
    p = analyze_template(template, tmp_path / "profile")
    pattern = p.patterns[0]
    p.patterns = [pattern]
    pattern.role = pattern.purpose = "cover"
    pattern.reusable = True
    pattern.title_zone.h = 12
    p.colors = list(dict.fromkeys(p.colors + ["#FFFFFF"]))
    report = {
        "box": [40, 40, 340, 480],
        "minimum_contrast": 8,
        "text_color": "white",
        "coordinate_space": {"width": 960, "height": 540},
    }
    if fault == "no_region":
        report.pop("box")
    elif fault == "contrast":
        report["minimum_contrast"] = 2
    elif fault == "missing_contrast":
        report["minimum_contrast"] = None
    elif fault == "unknown_color":
        report["text_color"] = "unknown"
    else:
        report["box"] = [40, 40, 1200, 480]
    pattern.safe_text_zone = report
    assert derive_safe_cover_patterns(p) == []


def test_table_columns_share_geometry_across_audit_and_exports(template, tmp_path):
    p = analyze_template(template, tmp_path / "profile")
    rows = [
        ["Description", "Count"],
        ["A considerably longer activity label", "7"],
        ["Another complete source label", "18"],
    ]
    e = Element(
        kind="table",
        box=Box(x=50, y=100, w=650, h=180),
        rows=deepcopy(rows),
        font=p.font,
        size=16,
        color=p.foreground,
        fill=p.accent,
        source_ids=["f1"],
    )
    widths = column_widths(rows, e.box.w, element_font(p, e)[1], 16)
    assert sum(widths) == pytest.approx(e.box.w) and widths[0] > widths[1] * 2
    assert all(
        table_cell_fits(cell, element_font(p, e)[1], 16, widths[ci] - 16, 48, ri == 0)
        for ri, row in enumerate(rows)
        for ci, cell in enumerate(row)
    )
    title = Element(
        kind="text",
        box=Box(x=50, y=25, w=600, h=50),
        text="Summary",
        font=p.font,
        size=24,
        color=p.foreground,
        role="title",
    )
    scene = SlideScene(
        title="Summary",
        background=p.background,
        elements=[title, e],
        source_ids=["f1"],
        layout="table",
    )
    from studio.audit import audit_scenes, repair_scenes
    from studio.models import Constraints, ContentModel, Fact

    package = SimpleNamespace(
        template=p,
        images=[],
        analysis={},
        constraints=Constraints(slides=1),
        content=ContentModel(title="Summary", facts=[Fact(id="f1", text="Source")]),
    )
    repair_scenes([scene], package)
    assert e.size == 16 and e.rows == rows
    assert not any(
        f.code in ("table_overflow", "readability") for f in audit_scenes([scene], package)
    )
    render_pptx([scene], p, template, tmp_path / "result.pptx")
    render_html([scene], p, tmp_path / "result.html")
    from pptx import Presentation

    table = next(
        s.table for s in Presentation(tmp_path / "result.pptx").slides[0].shapes if s.has_table
    )
    assert [c.width / 12700 for c in table.columns] == pytest.approx(widths, abs=0.001)
    assert [[c.text for c in row.cells] for row in table.rows] == rows
    html = (tmp_path / "result.html").read_text()
    assert "<colgroup>" in html
    for row in rows:
        for value in row:
            assert value in html


def test_data_layout_retries_geometry_before_rewriting_source(template, tmp_path):
    from studio.content import parse_content
    from studio.models import Constraints, SlidePlan, VariantPlan
    from studio.composer import compose_slide, _compose_slide
    from studio.audit import audit_scenes, repair_scenes

    p = analyze_template(template, tmp_path / "profile")
    p.assets = []
    pattern = p.patterns[0]
    p.patterns = [pattern]
    pattern.purpose = "content"
    pattern.role = "content"
    pattern.reusable = True
    pattern.title_zone = Box(x=30, y=20, w=700, h=60)
    pattern.body_zones = [Box(x=30, y=110, w=250, h=150)]
    pattern.text_zones = [pattern.title_zone, *pattern.body_zones]
    pattern.fields = [
        {"role": "title", "index": 0, "shape_id": 1, "box": pattern.title_zone.model_dump()},
        {"role": "body", "index": 0, "shape_id": 2, "box": pattern.body_zones[0].model_dump()},
    ]
    content = parse_content(
        "# Results\n\n| Activity description | Count |\n|---|---|\n| Quality verification work | 17 |\n| Requirements discovery | 28 |\n| Integration maintenance | 31 |"
    )
    table = content.tables[0]
    ids = [f.id for f in content.facts if f.source == table.id]
    package = SimpleNamespace(
        template=p,
        content=content,
        original_content=content,
        images=[],
        analysis={},
        constraints=Constraints(slides=1),
    )
    variant = VariantPlan(
        key="executive",
        title="Results",
        slides=[SlidePlan(title="Results", fact_ids=ids, table_id=table.id, layout="table")],
    )
    native = _compose_slide(variant, package, 0)
    repair_scenes([native], package)
    assert native.pattern_id == pattern.id
    assert any(f.code in ("readability", "table_overflow") for f in audit_scenes([native], package))
    scene = compose_slide(variant, package, 0)
    repair_scenes([scene], package)
    assert scene.pattern_id is None and scene.strategy == "token_composition"
    assert not any(
        f.code in ("readability", "table_overflow", "overlap")
        for f in audit_scenes([scene], package)
    )
    actual = next(e for e in scene.elements if e.kind == "table")
    assert actual.rows == [table.headers] + table.rows and actual.source_ids == ids
    variant.slides[0].pattern_id = pattern.id
    assert compose_slide(variant, package, 0).pattern_id == pattern.id


@pytest.mark.parametrize("width", [300, 420, 640])
def test_chart_caption_space_is_checked_before_native_export(template, tmp_path, width):
    from studio.models import TableData, SlidePlan
    from studio.charts import make_chart, chart_caption_layout, render_chart
    from studio.render import chart_fits
    from pptx import Presentation

    p = analyze_template(template, tmp_path / "profile")
    table = TableData(
        id="t",
        headers=["Period", "Planned", "Delivered", "Completion"],
        rows=[[str(i), str(20 + i), str(10 + i), f"{50 + i}%"] for i in range(1, 7)],
    )
    rows = [table.headers] + table.rows
    caption, height = chart_caption_layout(rows, width, p)
    assert all(f"{i} — {50 + i}%" in caption for i in range(1, 7))
    assert "Completion (Period)" in caption
    element = make_chart(
        table,
        SlidePlan(title="Delivery", fact_ids=["f1"], chart_type="line"),
        Box(x=20, y=20, w=width, h=160 + height + 12),
        p,
        p.foreground,
        ["f1"],
    )
    assert chart_fits(element, p)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = render_chart(slide, element, p)
    assert list(chart.series[0].values) == [21, 22, 23, 24, 25, 26]
    assert list(chart.series[1].values) == [11, 12, 13, 14, 15, 16]
    visible = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
    for i in range(1, 7):
        assert f"{50 + i}%" in visible
    from studio.export_audit import inspect_content
    from studio.models import VariantPlan, ContentModel

    variant = VariantPlan(
        key="executive",
        title="Delivery",
        slides=[
            SlidePlan(
                title="Delivery", fact_ids=["f1"], table_id="t", layout="chart", chart_type="line"
            )
        ],
    )
    package = SimpleNamespace(
        content=ContentModel(title="Delivery", facts=[], tables=[table]), images=[]
    )
    assert not any(
        f["code"] == "chart_supplement" for f in inspect_content(prs, variant, package)[1]
    )
    caption_shape = next(s for s in slide.shapes if s.has_text_frame and "51%" in s.text)
    caption_shape.text = caption_shape.text.replace("51%", "99%")
    assert any(f["code"] == "chart_supplement" for f in inspect_content(prs, variant, package)[1])
    element.box.h -= 1
    assert not chart_fits(element, p)
    with pytest.raises(ValueError, match="не помещаются"):
        render_chart(slide, element, p)
