import base64
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import pytest
from pptx import Presentation
from pptx.util import Pt
from studio.composition.pictures import A, P, R, SVG, is_picture, embedded_picture_blob
from studio.composition.powerpoint import open_presentation
from studio.security import PPTX_MAIN, POTX_MAIN
from studio.templates.parsing import analyze_template, walk_shapes
from studio.templates.template_analysis import template_inventory
from studio.composition.render import render_pptx
from studio.composition.composer import compose_variant
from studio.contents.planner import extractive_plans

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jTQAAAABJRU5ErkJggg=="
)
VECTOR = (
    b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><path d="M0 0L10 10"/></svg>'
)


def picture_input(template, tmp_path, mode, extension):
    prs = Presentation(template)
    for slide in prs.slides:
        slide.shapes.add_picture(BytesIO(PNG), Pt(70), Pt(100), Pt(20), Pt(20))
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(BytesIO(PNG), Pt(100), Pt(100), Pt(20), Pt(20))
    seed = tmp_path / "seed.pptx"
    prs.save(seed)
    with ZipFile(seed) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    for index in range(1, 4):
        name = f"ppt/slides/slide{index}.xml"
        root = ET.fromstring(parts[name])
        relname = f"ppt/slides/_rels/slide{index}.xml.rels"
        rels = ET.fromstring(parts[relname])
        for pic in root.iter(P + "pic"):
            blip = pic.find(P + "blipFill/" + A + "blip")
            rid = blip.attrib.pop(R + "embed")
            rel = next(r for r in rels if r.get("Id") == rid)
            if mode == "svg":
                ext = ET.SubElement(
                    ET.SubElement(blip, A + "extLst"), A + "ext", {"uri": "test-svg"}
                )
                ET.SubElement(ext, SVG + "svgBlip", {R + "embed": rid})
                rel.set("Target", "../media/vector.svg")
            elif mode == "linked":
                blip.set(R + "link", rid)
                rel.set("TargetMode", "External")
                rel.set("Target", "https://never-fetch.invalid/image.png")
        parts[name] = ET.tostring(root)
        parts[relname] = ET.tostring(rels)
    if mode == "svg":
        parts["ppt/media/vector.svg"] = VECTOR
        ct = ET.fromstring(parts["[Content_Types].xml"])
        ET.SubElement(
            ct,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
            {"Extension": "svg", "ContentType": "image/svg+xml"},
        )
        parts["[Content_Types].xml"] = ET.tostring(ct)
    if extension == "potx":
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
            PPTX_MAIN.encode(), POTX_MAIN.encode()
        )
    source = tmp_path / f"picture-{mode}.{extension}"
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return source


@pytest.mark.parametrize("mode", ["svg", "linked", "empty"])
@pytest.mark.parametrize("extension", ["pptx", "potx"])
def test_non_raster_pictures_analyze_and_export(template, prepared, tmp_path, mode, extension):
    source = picture_input(template, tmp_path, mode, extension)
    original = source.read_bytes()
    prs = open_presentation(source)
    pictures = [sh for sh, _ in walk_shapes(prs.slides[0].shapes) if is_picture(sh)]
    assert len(pictures) == 2
    with pytest.raises(ValueError, match="no embedded image"):
        pictures[0].image
    assert embedded_picture_blob(pictures[0]) == (VECTOR if mode == "svg" else None)
    profile = analyze_template(source, tmp_path / "analysis")
    inventory = template_inventory(source, profile)
    assert inventory["slides"][0]["pictures"] == 2
    assert inventory["slides"][0]["pictures_without_embedded_data"] == (0 if mode == "svg" else 2)
    assert any("без доступного встроенного" in w for w in profile.warnings) == (mode != "svg")
    _, _, package = prepared
    package.template = profile
    scenes = compose_variant(extractive_plans(package).variants[0], package)
    output = tmp_path / "output.pptx"
    render_pptx(scenes, profile, source, output)
    assert len(Presentation(output).slides) == 5
    assert source.read_bytes() == original
    with ZipFile(output) as archive:
        assert not any(
            b"never-fetch.invalid" in archive.read(n)
            for n in archive.namelist()
            if n.endswith(".rels")
        )


def test_regular_embedded_picture_still_readable(template):
    prs = Presentation(template)
    picture = prs.slides[0].shapes.add_picture(BytesIO(PNG), Pt(70), Pt(100), Pt(20), Pt(20))
    assert is_picture(picture) and embedded_picture_blob(picture) == PNG
