from io import BytesIO

import pytest
from fontTools.ttLib import TTFont
from pptx import Presentation

from studio.config import ROOT, Settings
from studio.embedded_fonts import check_glyphs
from studio.font_fallback import content_text, ensure_text_fonts
from studio.font_disclosure import exported_substitutions, warnings
from studio.fonts import element_font, role_font
from studio.models import Box, SlideScene
from studio.composer import text_element
from studio.security import InputRejected
from studio.template import analyze_template
from studio.template_cache import TemplateCache
from tests.test_fonts import embed


@pytest.fixture
def latin_template(tmp_path, template):
    # Reproduce a genuinely incomplete embedded face, without a private template.
    with TTFont(ROOT / "fonts/Play-Regular.ttf") as font:
        for table in font["cmap"].tables:
            if table.isUnicode():
                table.cmap = {key: value for key, value in table.cmap.items() if key < 128}
        stream = BytesIO()
        font.save(stream)
    source = tmp_path / "latin-only.pptx"
    embed(template, source, stream.getvalue())
    profile = analyze_template(source, tmp_path / "analysis")
    return source, profile


def test_complete_exact_face_preferred_over_different_family(latin_template):
    _, profile = latin_template
    ensure_text_fonts(profile, "Новый русский текст")
    assert role_font(profile)[0] == "Play"
    assert profile.font_file == str(ROOT / "fonts/Play-Regular.ttf")
    assert profile.font_substitutions[0]["fallback_font"] == "Play"
    check_glyphs(profile.font_file, "Новый русский текст")


def test_cyrillic_replacement_drives_field_measurement_and_native_export(
    latin_template, tmp_path, monkeypatch
):
    from studio import font_fallback
    from studio.render import render_pptx

    monkeypatch.setattr(font_fallback, "resolve_font", lambda name: "")
    source, profile = latin_template
    ensure_text_fonts(profile, "Анализ проекта → Результат", tmp_path / "analysis")
    assert role_font(profile, "title")[0] == "Montserrat"
    assert role_font(profile, "body")[0] == "Montserrat"
    assert profile.font == "Montserrat"
    element = text_element(
        "Анализ проекта",
        Box(x=50, y=50, w=500, h=100),
        profile,
        role="title",
        size=28,
        field_style={"family": "Play"},
    )
    assert element.font == "Montserrat"
    assert element_font(profile, element)[1] == str(ROOT / "fonts/Montserrat-Regular.ttf")
    scene = SlideScene(
        title=element.text,
        background=profile.background,
        elements=[element],
        source_ids=[],
        layout="statement",
    )
    target = tmp_path / "result.pptx"
    render_pptx([scene], profile, source, target)
    actual = Presentation(target)
    runs = [
        r
        for shape in actual.slides[0].shapes
        if shape.has_text_frame
        for p in shape.text_frame.paragraphs
        for r in p.runs
    ]
    assert "".join(r.text for r in runs) == element.text
    assert {r.font.name for r in runs} == {"Montserrat"}
    records = exported_substitutions(actual, profile)
    assert records[0]["template_font"] == "Play"
    assert "автоматически использован" in warnings(records)[0]
    assert "path" not in records[0]
    import json

    report = json.loads((tmp_path / "analysis/font-model.json").read_text())
    assert report["assets"][0]["requested"] == "Play"  # Source evidence retained.
    assert report["generation"]["assets"][0]["requested"] == "Montserrat"
    assert str(ROOT) not in json.dumps(report)


def test_no_replacement_for_supported_input_or_symbol_fallback(latin_template):
    _, profile = latin_template
    before = profile.model_dump()
    ensure_text_fonts(profile, "English → text")
    assert profile.model_dump() == before


def test_unsupported_character_still_blocks_without_partial_changes(latin_template, monkeypatch):
    from studio import font_fallback

    monkeypatch.setattr(font_fallback, "resolve_font", lambda name: "")
    _, profile = latin_template
    before = profile.model_dump()
    with pytest.raises(InputRejected, match="автоматическая замена тоже"):
        ensure_text_fonts(profile, "Текст 漢")
    assert profile.model_dump() == before


def test_substitution_is_idempotent_and_not_cached_for_other_text(
    latin_template, tmp_path, monkeypatch
):
    from studio import font_fallback

    monkeypatch.setattr(font_fallback, "resolve_font", lambda name: "")
    _, profile = latin_template
    ensure_text_fonts(profile, "Русский текст")
    once = profile.model_dump()
    ensure_text_fonts(profile, "Русский текст")
    assert profile.model_dump() == once
    cache = TemplateCache(Settings(data_dir=tmp_path / "data"))
    assert not cache.save(
        profile, tmp_path / "analysis", {"template_semantics": {"status": "completed"}}
    )
    assert not cache.location(profile.sha256).exists()


def test_prepare_reaches_ready_with_disclosed_replacement(
    latin_template, content, tmp_path, monkeypatch
):
    from studio import font_fallback
    from studio.pipeline import prepare, load_package
    from studio.store import Store
    import shutil

    monkeypatch.setattr(font_fallback, "resolve_font", lambda name: "")
    source, _ = latin_template
    store = Store(tmp_path / "data")
    job = store.create("preparation")
    shutil.copyfile(source, store.directory(job["id"]) / "input.pptx")
    prepare(store, job["id"], content, "", "", 5)
    done = store.get(job["id"])
    assert done["state"] == "ready", done
    assert any("автоматически будет использован" in w for w in done["warnings"])
    package = load_package(store, job["id"])
    assert package.manifest["font"]["name"] == "Montserrat"
    assert package.manifest["font_substitutions"][0]["scope"] == "font"
    check_glyphs(package.template.font_file, content_text(package.content, package.prepared_plans))


def test_model_added_cyrillic_updates_frozen_font_manifest(latin_template, tmp_path, monkeypatch):
    from studio import font_fallback, pipeline
    from studio.store import Store
    import shutil

    monkeypatch.setattr(font_fallback, "resolve_font", lambda name: "")

    async def model_text(package, path, gateway, progress):
        assert not package.template.font_substitutions
        package.content.title = "Заголовок"
        package.content.facts[0].text = "Проверка кириллицы"
        package.analysis = {"warnings": []}
        return package

    monkeypatch.setattr(pipeline, "prepare_intelligence", model_text)
    source, _ = latin_template
    store = Store(tmp_path / "data")
    job = store.create("preparation")
    shutil.copyfile(source, store.directory(job["id"]) / "input.pptx")
    pipeline.prepare(store, job["id"], "# Title\nEnglish source text.", "", "", 3)
    assert store.get(job["id"])["state"] == "ready", store.get(job["id"])
    package = pipeline.load_package(store, job["id"])
    assert package.template.font == "Montserrat"
    assert package.manifest["font"]["name"] == "Montserrat"
    assert package.manifest["font"]["sha256"] == package.template.font_origin["sha256"]
