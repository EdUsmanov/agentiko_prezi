"""A first-slide cover can use authored text boxes instead of placeholders."""

from pptx import Presentation
from pptx.util import Inches, Pt

from studio.templates.native_template import native_patterns


def _authored_slide(prs, layout_index):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
    for shape in list(slide.placeholders):
        slide.shapes._spTree.remove(shape._element)
    for text, top, size in (("Headline", 1.8, 32), ("Subhead", 2.7, 20)):
        shape = slide.shapes.add_textbox(Inches(1.8), Inches(top), Inches(5), Inches(0.6))
        shape.text_frame.paragraphs[0].text = text
        shape.text_frame.paragraphs[0].font.size = Pt(size)
    return slide


def _role(prs, slide_number):
    return next(p.role for p in native_patterns(prs) if p.source_slide == slide_number)


def test_first_authored_title_layout_is_cover_without_placeholder_or_type():
    prs = Presentation()
    layout = prs.slide_layouts[0]
    layout._element.attrib.pop("type", None)  # Adelphi: only cSld name identifies Title Slide.
    _authored_slide(prs, 0)

    assert _role(prs, 1) == "cover"


def test_ooxml_title_type_recognizes_renamed_cover_layout():
    prs = Presentation()
    prs.slide_layouts[0]._element.cSld.set("name", "Branded Opening")
    _authored_slide(prs, 0)

    assert _role(prs, 1) == "cover"


def test_regular_statement_on_title_and_content_layout_stays_statement():
    prs = Presentation()
    _authored_slide(prs, 1)

    assert _role(prs, 1) == "statement"


def test_later_authored_title_layout_stays_statement():
    prs = Presentation()
    _authored_slide(prs, 1)
    _authored_slide(prs, 0)

    assert _role(prs, 2) == "statement"
