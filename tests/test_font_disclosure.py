from pptx import Presentation
from pptx.util import Pt
from studio.config import ROOT
from studio.font_disclosure import substitutions, exported_substitutions, warnings, preparation_substitutions
from studio.export_audit import repair_symbols


def test_disclosure_is_only_for_replaced_symbols():
    path = str(ROOT/'fonts/Play-Regular.ttf')
    records = substitutions('Вход → Выход →', 'Play', path)
    assert len(records) == 1 and records[0]['codepoint'] == 'U+2192'
    assert records[0]['fallback_font'] == 'Montserrat'
    assert 'только символа' in warnings(records)[0]
    assert substitutions('Обычный текст', 'Play', path) == []
    assert substitutions('→', 'Montserrat', str(ROOT/'fonts/Montserrat-Regular.ttf')) == []


def test_actual_pptx_disclosure_reads_runs_not_source_predictions(prepared):
    _, _, p = prepared
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, Pt(400), Pt(200))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = 'Вход → Выход'
    run.font.name = 'Play'
    assert exported_substitutions(prs, p.template) == []
    repair_symbols(prs, p.template)
    records = exported_substitutions(prs, p.template)
    assert len(records) == 1 and records[0]['template_font'] == 'Play'
    assert 'использован Montserrat' in warnings(records)[0]
    assert all('path' not in r for r in records)


def test_preparation_warns_after_model_titles_are_finalized(prepared):
    _, _, p = prepared
    p.prepared_plans.variants[0].slides[0].title = 'Вход → Выход'
    records = preparation_substitutions(p)
    assert any(r['symbol'] == '→' for r in records)
    assert 'будет использован' in warnings(records, planned=True)[0]

