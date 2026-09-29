from io import BytesIO
from types import SimpleNamespace

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from studio.models import Box
from studio.templates.template_resources import (
    _candidate_shapes,
    _screen_is_empty,
    discover_resources,
)


def test_native_icon_candidate_uses_top_level_shape_id(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    icon = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(2), Inches(2), Inches(0.5), Inches(0.5))
    path = tmp_path / "source.pptx"
    prs.save(path)
    profile = SimpleNamespace(
        width=720, height=540, patterns=[SimpleNamespace(source_slide=1)], assets=[]
    )
    candidates = _candidate_shapes(path, profile)
    assert [(row[1], row[2]) for row in candidates] == [(1, icon.shape_id)]


def test_icon_crop_rejects_old_overlapping_text(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(2), Inches(2), Inches(0.5), Inches(0.5))
    old = slide.shapes.add_textbox(Inches(2), Inches(2), Inches(0.5), Inches(0.5))
    old.text = "PRIVATE OLD CONTENT"
    path = tmp_path / "source.pptx"
    prs.save(path)
    profile = SimpleNamespace(
        width=720, height=540, patterns=[SimpleNamespace(source_slide=1)], assets=[]
    )
    assert _candidate_shapes(path, profile) == []


def test_opaque_raster_screenshot_is_not_an_icon_candidate(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    image = Image.new("RGB", (120, 120), "#3498db")
    blob = BytesIO()
    image.save(blob, format="PNG")
    slide.shapes.add_picture(
        BytesIO(blob.getvalue()), Inches(2), Inches(2), Inches(0.7), Inches(0.7)
    )
    path = tmp_path / "source.pptx"
    prs.save(path)
    profile = SimpleNamespace(
        width=720, height=540, patterns=[SimpleNamespace(source_slide=1)], assets=[]
    )
    assert _candidate_shapes(path, profile) == []


def test_device_requires_open_group_screen():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    solid = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(1), Inches(4), Inches(3))
    assert not _screen_is_empty(solid, Box(x=170, y=100, w=160, h=100))


def test_resource_model_rejects_unknown_id(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(2), Inches(2), Inches(0.5), Inches(0.5))
    path = tmp_path / "source.pptx"
    prs.save(path)
    profile = SimpleNamespace(
        width=720, height=540, patterns=[SimpleNamespace(source_slide=1)], assets=[]
    )
    image = Image.new("RGB", (720, 540), "white")
    out = BytesIO()
    image.save(out, format="PNG")

    class Gateway:
        settings = SimpleNamespace(mode="api")

        async def json_request(self, *args, **kwargs):
            return {"choices": [{"id": "forged", "kind": "icon", "confidence": 1}]}

    import asyncio
    import pytest

    with pytest.raises(ValueError, match="неизвестный"):
        asyncio.run(discover_resources(path, profile, Gateway(), {1: out.getvalue()}))


def test_resource_catalogue_keeps_native_shape_reference_and_preview(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    icon = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(2), Inches(2), Inches(0.5), Inches(0.5))
    path = tmp_path / "source.pptx"
    prs.save(path)
    profile = SimpleNamespace(
        width=720, height=540, patterns=[SimpleNamespace(source_slide=1)], assets=[]
    )
    image = Image.new("RGB", (720, 540), "white")
    out = BytesIO()
    image.save(out, format="PNG")

    class Gateway:
        settings = SimpleNamespace(mode="api")

        async def json_request(self, stage, payload, **kwargs):
            assert stage == "template_resources" and kwargs["images"][0].startswith(b"\x89PNG")
            return {
                "choices": [
                    {
                        "id": payload["candidates"][0]["id"],
                        "kind": "icon",
                        "description": "Значок безопасности",
                        "tags": ["безопасность"],
                        "confidence": 0.9,
                    }
                ]
            }

    import asyncio

    resources, report = asyncio.run(
        discover_resources(path, profile, Gateway(), {1: out.getvalue()})
    )
    assert report["selected"] == 1
    assert resources[0].shape_ids == [icon.shape_id]
    assert resources[0].tags == ["безопасность"]
    assert resources[0].preview_path.startswith(str(tmp_path))


def test_open_device_frame_preview_has_transparent_screen(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    for x, y, w, h in [(2, 1, 4, 0.08), (2, 3.92, 4, 0.08), (2, 1, 0.08, 3), (5.92, 1, 0.08, 3)]:
        group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    path = tmp_path / "source.pptx"
    prs.save(path)
    profile = SimpleNamespace(
        width=720, height=540, patterns=[SimpleNamespace(source_slide=1)], assets=[]
    )
    image = Image.new("RGB", (720, 540), "white")
    out = BytesIO()
    image.save(out, format="PNG")

    class Gateway:
        settings = SimpleNamespace(mode="api")

        async def json_request(self, stage, payload, **kwargs):
            return {
                "choices": [
                    {
                        "id": payload["candidates"][0]["id"],
                        "kind": "device_frame",
                        "description": "Рамка браузера",
                        "tags": ["браузер"],
                        "confidence": 0.95,
                        "screen_box": [50, 50, 950, 950],
                    }
                ]
            }

    import asyncio

    resources, _ = asyncio.run(discover_resources(path, profile, Gateway(), {1: out.getvalue()}))
    assert resources[0].shape_ids == [group.shape_id]
    assert resources[0].screen_box is not None
    with Image.open(resources[0].preview_path) as cutout:
        assert cutout.mode == "RGBA"
        assert cutout.getpixel((cutout.width // 2, cutout.height // 2))[3] == 0
