import asyncio
import json
from dataclasses import replace
from types import SimpleNamespace
import httpx
import pytest

from studio.content import parse_content
from studio.document import visualization
from studio.models import TableData, SlidePlan
from studio.charts import chart_projection
from studio.quality_gate import require_publishable


def test_table_markdown_is_not_numeric_data():
    content=parse_content('# Данные\n| Спринт | Задачи |\n|---|---|\n| **Всего** | **82** |')
    assert content.tables[0].rows==[['Всего','82']]


@pytest.mark.parametrize('text,kind',[
    ('Данные для столбчатого графика:', 'column'),
    ('Тип графика: круговой', 'pie'),
    ('Тип графика: горизонтальные столбцы', 'bar'),
    ('Тип графика: таблица', 'table'),
    ('Данные для накопительной столбчатой диаграммы:','column_stacked'),
    ('Данные для двух линий на одном графике:','line'),
    ('Данные для горизонтальных столбцов «было / стало»:','bar'),
    ('Данные для парных столбцов «план / факт»:','column')])
def test_presentation_hints_are_metadata(text,kind):
    assert visualization(text)==kind


@pytest.mark.parametrize('text',[
    'Данные для четырёх карточек:',
    'Четыре цифры для финального слайда:',
    '*Показатели разных типов: отдельные карточки, без сравнения высоты столбцов.*'])
def test_metric_hints_are_not_visible_body(text):
    assert visualization(text)=='metrics'


def test_metric_cards_are_editable_and_preserve_every_cell(prepared,tmp_path):
    from pptx import Presentation
    from pptx.util import Pt
    from studio.models import Element,Box,VariantPlan
    from studio.metrics import metric_elements,render_metric_cards
    from studio.export_audit import inspect_content
    _,_,package=prepared
    table=TableData(id='t',headers=['Результат','Значение'],
        rows=[['Задачи','200'],['Дефекты','152'],['План','100%'],['Команда','4']],visualization='metrics')
    element=Element(kind='table',box=Box(x=50,y=100,w=500,h=360),rows=[table.headers]+table.rows,
        font=package.template.font,color=package.template.foreground,fill=package.template.accent,source_ids=['f1'])
    assert len([e for e in metric_elements(element,package.template) if e.role=='metric_value'])==4
    prs=Presentation();prs.slide_width=Pt(960);prs.slide_height=Pt(540)
    slide=prs.slides.add_slide(prs.slide_layouts[6]);render_metric_cards(slide,element,package.template)
    package.content.tables=[table]
    variant=VariantPlan(key='executive',title='Карточки',slides=[SlidePlan(title='Результат',fact_ids=['f1'],table_id='t',layout='table')])
    package.images=[]
    _,findings=inspect_content(prs,variant,package)
    assert not [f for f in findings if f['severity']=='error']
    assert not any(s.has_table for s in slide.shapes)
    assert all(any(cell in s.text for s in slide.shapes if s.has_text_frame) for row in [table.headers]+table.rows for cell in row)
    with pytest.raises(ValueError,match='не помещаются'):
        metric_elements(element.model_copy(update={'box':Box(x=0,y=0,w=150,h=60)}),package.template)


def test_chart_projection_preserves_mixed_units_and_totals():
    table=TableData(id='t',headers=['Спринт','План','Факт','Выполнение'],
        rows=[['1','30','18','60%'],['2','32','24','75%'],['Всего','62','42','67.7%']])
    projected,caption=chart_projection(table)
    assert projected.headers==['Спринт','План','Факт']
    assert projected.rows==[['1','30','18'],['2','32','24']]
    assert '60%' in caption[0] and '75%' in caption[1]
    assert 'Всего' in caption[2] and '67.7%' in caption[2]
    assert table.rows[-1][0]=='Всего'


def test_chart_fallback_keeps_original_table(prepared):
    from studio.charts import make_chart
    from studio.models import Box
    _,_,package=prepared
    table=TableData(id='t',headers=['Спринт','План','Факт','Выполнение'],
        rows=[['1','30','18','60%'],['Всего','30','18','60%']])
    # A pie cannot represent two count series. Its fallback must retain the
    # total and percentages excluded from the comparable chart projection.
    element=make_chart(table,SlidePlan(title='Данные',fact_ids=['f1'],chart_type='pie'),
        Box(x=10,y=10,w=600,h=400),package.template,package.template.foreground,['f1'])
    assert element.kind=='table'
    assert element.rows==[table.headers]+table.rows


def test_known_error_is_not_published_as_needs_review():
    with pytest.raises(ValueError,match='не опубликованы'):
        require_publishable({'errors':1,'variants':[{'key':'story','findings':[
            {'severity':'error','code':'text_overflow','slide':2}]}]})
    require_publishable({'errors':0,'visual_audit':{'findings':[{'severity':'warning'}]}})


def test_explicit_slide_structure_does_not_call_role_model(prepared):
    from studio.document import structure_document
    from studio.archetypes import analyze_content_archetypes
    _,_,package=prepared
    package.content=parse_content('## Слайд 1. Проект\nЗадача команды.\n'
        '## Слайд 2. Динамика\nДанные для линейного графика:\n'
        '| Спринт | Задачи |\n|---|---|\n| 1 | 10 |\n| 2 | 20 |')
    class Gateway:
        settings=SimpleNamespace(mode='api')
        async def json_request(self,*args,**kwargs):
            pytest.fail('Explicit slide/data structure must not be reclassified by a model')
    asyncio.run(structure_document(package,Gateway()))
    report=asyncio.run(analyze_content_archetypes(package,Gateway()))
    assert report['method']=='explicit_data_storyboard'
    assert [f for unit in report['units'] for f in unit['fact_ids']]==[f.id for f in package.content.facts]
    assert package.content.tables[0].visualization=='line'


def test_explicit_author_plan_is_validated_without_rewriting(prepared):
    from studio.planner import extractive_plans,plan
    _,_,package=prepared
    original=extractive_plans(package)
    package.analysis.update(storyboard=[s.model_dump() for s in original.variants[0].slides],
        archetypes={'method':'explicit_data_storyboard'})
    class Gateway:
        settings=SimpleNamespace(mode='api')
        async def json_request(self,*args,**kwargs):
            pytest.fail('The complete authored scenario must not be rewritten')
    plans,warning=asyncio.run(plan(package,Gateway(),600))
    assert warning is None
    assert package.analysis['planning_method']=='explicit_author_storyboard'
    for variant in plans.variants:
        assert [s.fact_ids for s in variant.slides]==[s.fact_ids for s in original.variants[0].slides]


def test_thinking_budget_uses_provider_contract_not_reasoning_effort(tmp_path):
    from studio.config import Settings
    from studio.gateway import ModelGateway
    settings=Settings(data_dir=tmp_path,mode='api',model_id='qwen3.8-27b',base_url='https://test.invalid/v1',
        api_key='test',open_weights=True,parameters_b=27,license='Apache-2.0',thinking=True,thinking_token_budget=2048)
    def answer(request):
        body=json.loads(request.content)
        assert body['chat_template_kwargs']=={'enable_thinking':True,'thinking_token_budget':2048}
        assert 'reasoning_effort' not in body
        assert body['max_tokens']==12048
        return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{
            'content':'{"ok":true}','reasoning_content':'This is not final JSON'}}]})
    gateway=ModelGateway(settings,transport=httpx.MockTransport(answer))
    assert asyncio.run(gateway.json_request('planner',{}))=={'ok':True}


def test_visual_control_disables_thinking_but_keeps_multimodal_model(tmp_path):
    from studio.config import Settings
    from studio.gateway import ModelGateway
    settings=Settings(data_dir=tmp_path,mode='api',model_id='qwen3.8-27b',base_url='https://test.invalid/v1',
        api_key='test',open_weights=True,parameters_b=27,license='Apache-2.0',thinking=True)
    def answer(request):
        body=json.loads(request.content)
        assert body['model']=='qwen3.8-27b'
        assert body['chat_template_kwargs']=={'enable_thinking':False}
        return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':'{"ok":true}'}}]})
    gateway=ModelGateway(settings,transport=httpx.MockTransport(answer))
    assert asyncio.run(gateway.json_request('visual_critic',{}))=={'ok':True}
