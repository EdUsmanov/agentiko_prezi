from pptx import Presentation
from pptx.util import Pt
from studio.export_audit import contains_text, repair_symbols, geometry, slide_text


def test_fact_matching_does_not_accept_numeric_prefix():
    assert contains_text('Доход 1 млн.','Доход 1 млн.')
    assert not contains_text('Рост 15','Рост 1')
    assert not contains_text('Рост 1.5','Рост 1')
    assert not contains_text('Рост 1,5','Рост 1')
    assert contains_text('**Рост**\n 1','Рост 1')


def test_symbol_repair_keeps_text_order_and_style(prepared):
    _,_,package=prepared
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    shape=slide.shapes.add_textbox(0,0,Pt(400),Pt(200))
    p=shape.text_frame.paragraphs[0]
    r=p.add_run();r.text='А → Б';r.font.name='Play';r.font.bold=True
    r=p.add_run();r.text=' конец';r.font.name='Play'
    assert repair_symbols(prs,package.template)==1
    assert slide_text(slide)=='А → Б конец'
    assert [r.font.name for r in p.runs]==['Play','Montserrat','Play','Play']
    assert p.runs[1].font.bold


def test_actual_geometry_repair_changes_no_text_or_box(prepared):
    _,_,package=prepared
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    shape=slide.shapes.add_textbox(Pt(30),Pt(30),Pt(220),Pt(50))
    # This case relies on auto-wrapping; new text boxes default to no-wrap.
    shape.text_frame.word_wrap=True
    r=shape.text_frame.paragraphs[0].add_run();r.text='Сначала принять заявку, затем проверить данные.'
    r.font.name='Play';r.font.size=Pt(28)
    before=(shape.left,shape.top,shape.width,shape.height,shape.text)
    findings,repairs=geometry(prs,package.template,repair=True)
    assert repairs and not findings
    assert before==(shape.left,shape.top,shape.width,shape.height,shape.text)
    assert not geometry(prs,package.template)[0]

def test_content_geometry_ignores_text_and_ids_but_detects_position():
    from studio.export_audit import content_geometry_signature
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    shape=slide.shapes.add_textbox(Pt(10),Pt(20),Pt(200),Pt(100))
    shape.text='Body'
    original=content_geometry_signature(prs)
    shape.text='Different words';shape.name='Different ID'
    assert content_geometry_signature(prs)==original
    shape.left=Pt(90)
    assert content_geometry_signature(prs)!=original
