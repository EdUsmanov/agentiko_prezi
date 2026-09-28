from copy import deepcopy
import pytest
from studio.models import Fact, SlidePlan, Box
from studio.composer import compose
from studio.semantic_bindings import labeled_facts, missing_editorial_labels, object_contract


@pytest.mark.parametrize("native", [False, True])
def test_generic_layout_keeps_editorial_actor_names(prepared, native):
    _, _, package = prepared
    package.content.facts = [
        Fact(id="a", text="Формирует план."),
        Fact(id="b", text="Проверяет файл."),
    ]
    package.content.tables = []
    package.analysis = {
        "editorial": {
            "provenance": [{"fact_id": "a", "group": "ИИ"}, {"fact_id": "b", "group": "Код"}],
            "bindings": [
                {
                    "purpose": "structure",
                    "fact_ids": ["a", "b"],
                    "groups": [
                        {"label": "ИИ", "fact_ids": ["a"], "parent": None},
                        {"label": "Код", "fact_ids": ["b"], "parent": None},
                    ],
                }
            ],
        }
    }
    pattern = package.template.patterns[0]
    pattern.body_zones = [Box(x=50, y=130, w=800, h=300)]
    pattern.heading_zones = []
    pattern.purpose = "content"
    before = deepcopy(package.content)
    slide = SlidePlan(
        title="Архитектура",
        fact_ids=["a", "b"],
        purpose="structure",
        pattern_id=pattern.id if native else "token:auto",
    )
    scene = compose(slide, package, 0, "executive")
    assert not missing_editorial_labels(scene, package)
    assert package.content == before
    assert any("ИИ. Формирует план." in e.text for e in scene.elements)
    assert any("Код. Проверяет файл." in e.text for e in scene.elements)
    # A renderer regression must be detected despite unchanged fact IDs.
    for element in scene.elements:
        element.text = element.text.replace("ИИ. ", "")
    assert missing_editorial_labels(scene, package) == ["ИИ"]
    from studio.audit import audit_scenes

    assert any(
        f.code == "semantic_label_missing" and f.severity == "error"
        for f in audit_scenes([scene], package)
    )
    if native:
        contract = object_contract(slide, package, pattern)
        text = "\n".join(p for f in contract["fields"].values() for p in f["paragraphs"])
        assert "ИИ. Формирует план." in text and "Код. Проверяет файл." in text


def test_inline_labels_do_not_duplicate_existing_owner_or_mutate_source():
    from types import SimpleNamespace

    package = SimpleNamespace(
        analysis={"editorial": {"provenance": [{"fact_id": "a", "group": "ИИ"}]}}
    )
    fact = Fact(id="a", text="ИИ формирует план.")
    assert labeled_facts([fact], package)[0].text == fact.text
    package.analysis["editorial"]["provenance"][0]["group"] = "Инженер"
    assert labeled_facts([fact], package)[0].text == "Инженер. ИИ формирует план."
    assert fact.text == "ИИ формирует план."


def test_native_export_audit_detects_missing_owner_with_intact_claim_text():
    from types import SimpleNamespace
    from pptx import Presentation
    from pptx.util import Pt
    from studio.models import VariantPlan
    from studio.export_audit import inspect_content

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Pt(20), Pt(20), Pt(600), Pt(50))
    title.text = "Архитектура"
    body = slide.shapes.add_textbox(Pt(20), Pt(100), Pt(600), Pt(100))
    body.text = "Формирует план."
    package = SimpleNamespace(
        images=[], analysis={"editorial": {"provenance": [{"fact_id": "a", "group": "ИИ"}]}}
    )
    variant = VariantPlan(
        key="executive", title="Тест", slides=[SlidePlan(title="Архитектура", fact_ids=["a"])]
    )
    _, findings = inspect_content(prs, variant, package)
    assert any(f["code"] == "semantic_label_missing" and f["severity"] == "error" for f in findings)
    body.text = "ИИ. Формирует план."
    assert not inspect_content(prs, variant, package)[1]
