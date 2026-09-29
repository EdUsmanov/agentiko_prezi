from pptx import Presentation
from pptx.util import Pt
from studio.templates.parsing import analyze_template
from studio.templates.field_style import field_style
from studio.composition.composer import text_element
from studio.templates.fonts import element_font
from studio.composition.render import set_text
from studio.models import Box


def test_local_field_font_survives_global_role_and_native_export(template, tmp_path):
    prs = Presentation(template)
    title = prs.slides[0].shapes[0]
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Montserrat"
            run.font.bold = True
            run.font.size = Pt(34)
    source = tmp_path / "mixed.pptx"
    prs.save(source)
    profile = analyze_template(source, tmp_path / "analysis")
    pattern = next(p for p in profile.patterns if p.id == "native-slide-1")
    style = field_style(pattern, "title")
    assert style["family"] == "Montserrat" and style["bold"] is True
    assert style["size"] == 34
    element = text_element(
        "Заголовок", Box(x=20, y=20, w=700, h=100), profile, "title", 34, field_style=style
    )
    assert element_font(profile, element)[0] == "Montserrat Bold"
    output = Presentation()
    slide = output.slides.add_slide(output.slide_layouts[6])
    box = slide.shapes.add_textbox(Pt(20), Pt(20), Pt(700), Pt(100))
    set_text(box.text_frame, element.text, element, profile)
    target = tmp_path / "export.pptx"
    output.save(target)
    run = Presentation(target).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Montserrat" and run.font.bold
    assert run.font.size.pt == 34


def test_missing_field_face_is_explicit_not_claimed_exact(template, tmp_path):
    profile = analyze_template(template, tmp_path / "analysis")
    element = text_element(
        "Текст",
        Box(x=20, y=20, w=700, h=100),
        profile,
        field_style={"family": "Unavailable Test Font", "size": 20},
    )
    assert element.field_style["unresolved_font"] is True
    assert element_font(profile, element)[0] != "Unavailable Test Font"


def test_subheading_fits_shallow_native_heading_zone(template, tmp_path):
    from studio.templates.fonts import wrap_text

    profile = analyze_template(template, tmp_path / "analysis")
    profile.body_size = 48
    profile.font_sizes = [48]
    element = text_element("Выбор", Box(x=70, y=135, w=390, h=58), profile, "subheading", size=48)
    assert element.size < 48
    assert (
        len(wrap_text(element.text, element_font(profile, element)[1], element.size, 390))
        * element.size
        * 1.25
        <= element.box.h
    )


def test_table_compaction_uses_final_readable_size_and_keeps_field_origin(template, tmp_path):
    from studio.models import Element
    from studio.composition.table_style import compact_table

    profile = analyze_template(template, tmp_path / "analysis")
    table = Element(
        kind="table",
        box=Box(x=150, y=100, w=650, h=360),
        size=24,
        rows=[["Сценарий", "Проверено"], ["Загрузка", "20"], ["Поиск", "18"]],
    )
    compact_table(table, profile)
    assert (table.box.x, table.box.y, table.box.w) == (150, 100, 650)
    assert 100 < table.box.h < 200 and table.size == 24
