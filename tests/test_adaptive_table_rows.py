"""Independent row sizing across synthetic templates and actual native exports."""

from copy import deepcopy
from types import SimpleNamespace
import pytest
from pptx import Presentation
from studio.config import ROOT
from studio.models import (
    PreparationControl,
    Box,
    Element,
    SlideScene,
    Constraints,
    ContentModel,
    Fact,
)
from studio.templates.parsing import analyze_template
from studio.composition.table_style import column_widths, row_heights
from studio.templates.fonts import element_font, table_cell_fits
from studio.checks.audit import audit_scenes, repair_scenes
from studio.composition.render import render_pptx, render_html, render_pdf


@pytest.mark.parametrize(
    "font,width,rows_count",
    [("Play", 510, 6), ("Montserrat", 560, 6), ("Play", 640, 9), ("Montserrat", 700, 8)],
)
def test_multiline_header_does_not_shrink_numeric_rows(template, tmp_path, font, width, rows_count):
    profile = analyze_template(template, tmp_path / "profile")
    profile.patterns = []
    profile.font_roles = {}
    profile.font_assets = []
    profile.font = font
    profile.fonts = [font]
    profile.font_file = str(ROOT / f"fonts/{font}-Regular.ttf")
    profile.body_size = 22
    profile.font_sizes = [14, 22, 32]  # No authored 16pt step.
    rows = [["Период", "Запланировано задач", "Выполнено задач", "Процент выполнения"]]
    rows += [[str(i), str(20 + i), str(10 + i), f"{50 + i}%"] for i in range(rows_count)]
    element = Element(
        kind="table",
        box=Box(x=40, y=100, w=width, h=rows_count * 36 + 74),
        rows=deepcopy(rows),
        font=font,
        size=22,
        color=profile.foreground,
        fill=profile.accent,
        source_ids=["f1"],
    )
    scene = SlideScene(
        title="Progress",
        background=profile.background,
        elements=[
            Element(
                kind="text",
                box=Box(x=40, y=20, w=600, h=50),
                text="Progress",
                font=font,
                size=24,
                color=profile.foreground,
                role="title",
            ),
            element,
        ],
        source_ids=["f1"],
        layout="table",
    )
    package = SimpleNamespace(
        template=profile,
        images=[],
        analysis={},
        control=PreparationControl(),
        constraints=Constraints(slides=1),
        content=ContentModel(title="Progress", facts=[Fact(id="f1", text="Data")]),
    )
    repair_scenes([scene], package)
    assert element.size >= 16 and element.rows == rows
    defects = {"readability", "table_overflow", "out_of_bounds"}
    assert not [f for f in audit_scenes([scene], package) if f.code in defects]
    path = element_font(profile, element)[1]
    widths = column_widths(rows, width, path, element.size)
    heights = row_heights(rows, widths, path, element.size, element.box.h)
    assert heights[0] >= heights[1] and sum(heights) == pytest.approx(element.box.h)
    if width < 600:
        assert heights[0] > heights[1]
    assert all(
        table_cell_fits(cell, path, element.size, widths[ci] - 16, heights[ri] - 12, ri == 0)
        for ri, row in enumerate(rows)
        for ci, cell in enumerate(row)
    )
    render_pptx([scene], profile, template, tmp_path / "output.pptx")
    render_html([scene], profile, tmp_path / "output.html")
    render_pdf([scene], profile, tmp_path / "output.pdf")
    table = next(
        s.table for s in Presentation(tmp_path / "output.pptx").slides[0].shapes if s.has_table
    )
    assert [r.height / 12700 for r in table.rows] == pytest.approx(heights, abs=0.002)
    assert [[" ".join(c.text.split()) for c in r.cells] for r in table.rows] == rows
    assert all(
        run.font.size.pt >= 16
        for r in table.rows
        for c in r.cells
        for p in c.text_frame.paragraphs
        for run in p.runs
    )
    html = (tmp_path / "output.html").read_text()
    for height in heights:
        assert f"height:{height}px" in html
    import pypdfium2 as pdfium

    with pdfium.PdfDocument(str(tmp_path / "output.pdf")) as pdf:
        page = pdf[0]
        textpage = page.get_textpage()
        text = textpage.get_text_range()
        assert "Процент" in text and f"{50 + rows_count - 1}%" in text
        textpage.close()
        page.close()
    from studio.checks.export_audit import geometry

    assert not any(
        f["code"] == "pptx_table_overflow"
        for f in geometry(Presentation(tmp_path / "output.pptx"), profile)[0]
    )


def test_tall_body_rows_and_real_overflow_remain_visible(template, tmp_path):
    profile = analyze_template(template, tmp_path / "profile")
    rows = [
        ["Item", "Value"],
        ["Short", "1"],
        ["A multi-line description of independently measured table rows", "2"],
    ]
    path = profile.font_file
    size = 16
    width = 360
    widths = column_widths(rows, width, path, size)
    natural = row_heights(rows, widths, path, size)
    assert natural[2] > natural[1]
    cramped = row_heights(rows, widths, path, size, height=70)
    assert sum(cramped) == pytest.approx(70)
    assert any(
        not table_cell_fits(cell, path, size, widths[ci] - 16, cramped[ri] - 12, ri == 0)
        for ri, row in enumerate(rows)
        for ci, cell in enumerate(row)
    )
    e = Element(
        kind="table",
        box=Box(x=40, y=100, w=width, h=70),
        rows=rows,
        font=profile.font,
        size=16,
        color=profile.foreground,
        source_ids=["f1"],
    )
    scene = SlideScene(
        title="Dense",
        background=profile.background,
        elements=[e],
        source_ids=["f1"],
        layout="table",
    )
    package = SimpleNamespace(
        template=profile,
        images=[],
        analysis={},
        control=PreparationControl(),
        constraints=Constraints(slides=1),
        content=ContentModel(title="Dense", facts=[Fact(id="f1", text="Data")]),
    )
    assert any(f.code == "table_overflow" for f in audit_scenes([scene], package))
    e.size = 14
    assert any(f.code == "readability" and f.element == 0 for f in audit_scenes([scene], package))
