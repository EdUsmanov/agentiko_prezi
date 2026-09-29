from copy import deepcopy

import pytest
from pptx import Presentation

from studio.composition.chart_layout import bar_layout
from studio.composition.charts import make_chart, render_chart
from studio.models import Box, Fact, Plans, SlidePlan, TableData, VariantPlan
from studio.composition.render import chart_fits
from studio.templates.parsing import analyze_template


@pytest.fixture
def profile(template, tmp_path):
    return analyze_template(template, tmp_path / "profile")


def chart(profile, width=630, height=320, count=4, negative=False):
    table = TableData(
        id="metrics",
        headers=["Время обработки запросов, рабочие часы", "До изменений", "После изменений"],
        rows=[
            [
                f"Проверка результата обработки запроса {i}",
                str(-10 - i if negative else 10 + i),
                str(3 + i),
            ]
            for i in range(count)
        ],
    )
    return make_chart(
        table,
        SlidePlan(title="Результаты", fact_ids=["data"], chart_type="bar"),
        Box(x=20, y=30, w=width, h=height),
        profile,
        profile.foreground,
        ["data"],
    )


def test_narrow_native_chart_rejected_before_export(profile):
    element = chart(profile, 315, 184)
    assert not chart_fits(element, profile)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    with pytest.raises(ValueError, match="Подписи диаграммы не помещаются"):
        render_chart(slide, element, profile)


@pytest.mark.parametrize("width,height,count", [(420, 340, 4), (640, 380, 6), (850, 520, 8)])
@pytest.mark.parametrize("negative", [False, True])
def test_measured_native_export_preserves_categories_values_and_heading(
    profile, tmp_path, width, height, count, negative
):
    element = chart(profile, width, height, count, negative)
    before = element.model_dump()
    layout = bar_layout(element, profile)
    assert layout.fits
    assert chart_fits(element, profile)
    assert layout.height / count >= max(
        len(x.splitlines()) * layout.label_size * 1.25 for x in layout.labels
    )
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_chart(slide, element, profile)
    target = tmp_path / "chart.pptx"
    prs.save(target)
    reopened = Presentation(target)
    exported = next(s.chart for s in reopened.slides[0].shapes if s.has_chart)
    assert [" ".join(str(c.label).split()) for c in exported.plots[0].categories] == element.labels
    assert [list(s.values) for s in exported.series] == element.series_values
    assert not exported.category_axis.has_title
    assert element.category_title in " ".join(
        s.text.replace("\n", " ") for s in reopened.slides[0].shapes if s.has_text_frame
    )
    assert element.model_dump() == before


def test_count_and_font_size_affect_capacity_without_losing_labels(profile):
    element = chart(profile, 630, 300)
    assert chart_fits(element, profile)
    crowded = chart(profile, 630, 300, count=20)
    assert not chart_fits(crowded, profile)
    unbreakable = element.model_copy(deep=True)
    unbreakable.labels[0] = "UnbreakableCategoryIdentifier" * 8
    assert not chart_fits(unbreakable, profile)
    assert bar_layout(unbreakable, profile).labels[0] == unbreakable.labels[0]


def test_repair_reselects_roomier_layout_without_changing_story(prepared, monkeypatch):
    from studio.composition.composer import compose_slide
    from studio.checks.refinement import apply_edits
    from studio.composition.layout_edits import LayoutEdit

    _, _, package = prepared
    package.template.width = 720
    package.template.height = 405
    package.template.margin = 30
    package.template.title_size = 20
    package.template.body_size = 16
    pattern = next(
        p for p in package.template.patterns if p.source_slide and p.body_zones
    ).model_copy(deep=True)
    pattern.purpose = "content"
    pattern.role = "statement"
    pattern.title_zone = Box(x=30, y=30, w=640, h=40)
    pattern.body_zones = [Box(x=350, y=100, w=315, h=215)]
    body_field = next(f for f in pattern.fields if f["role"] == "body")
    body_field["box"] = pattern.body_zones[0].model_dump()
    pattern.safe_text_zone = {}
    package.template.patterns = [pattern]
    element = chart(package.template)
    for index, row in enumerate(element.rows[1:]):
        row[0] = f"Обработка заявки отдела {index}"
    table = TableData(
        id="metrics", headers=element.rows[0], rows=element.rows[1:], visualization="bar"
    )
    package.content.facts = [
        Fact(id="data", text="Данные измерений", source="metrics"),
        Fact(id="claim", text="Результат улучшился.", source="user_text"),
    ]
    package.content.tables = [table]
    package.original_content = package.content.model_copy(deep=True)
    package.constraints.slides = 1
    package.analysis = {}
    plans = Plans(
        variants=[
            VariantPlan(
                key=key,
                title=key,
                slides=[
                    SlidePlan(
                        title="Результаты",
                        fact_ids=["data", "claim"],
                        table_id="metrics",
                        layout="chart",
                        chart_type="bar",
                        purpose="content",
                        pattern_id=pattern.id,
                    )
                ],
            )
            for key in ("executive", "analytical", "story")
        ]
    )
    decks = {v.key: [compose_slide(v, package, 0)] for v in plans.variants}
    # A repeated narrow assignment from Design must also be repaired before export.
    from studio.providers.deeppresenter import CompositionEnvironment, Assignment
    from studio.contents.planner import assign_compositions

    environment = CompositionEnvironment(package, assign_compositions(plans, package))
    environment.compose(
        [Assignment(variant=v.key, slide=1, pattern_id=pattern.id) for v in plans.variants]
    )
    for variant in environment.plans.variants:
        assert variant.slides[0].pattern_id == "token:auto"

    old = deepcopy(plans.model_dump())
    trial, changed = apply_edits(
        package,
        plans,
        decks,
        [LayoutEdit(variant="executive", slide=1, operation="readable_chart")],
        {("executive", 1): ["@readable_chart"]},
    )
    scene = changed["executive"][0]
    repaired = next(e for e in scene.elements if e.kind == "chart")
    original = next(e for e in decks["executive"][0].elements if e.kind == "chart")
    assert repaired.box.w * repaired.box.h > original.box.w * original.box.h
    assert chart_fits(repaired, package.template)
    assert repaired.rows == original.rows
    assert repaired.series_values == original.series_values
    assert scene.source_ids == decks["executive"][0].source_ids
    assert scene.title == decks["executive"][0].title
    assert plans.model_dump() == old
    assert trial.variants[0].slides[0].chart_style == "readable"

    from studio.composition import composer

    original_compose = composer._compose_slide

    def no_token_layout(variant, package, index, image_groups=None):
        if variant.slides[index].pattern_id == "token:auto":
            raise ValueError("No token space")
        return original_compose(variant, package, index, image_groups)

    monkeypatch.setattr(composer, "_compose_slide", no_token_layout)
    unconstrained = assign_compositions(plans, package).variants[0]
    retained = composer.compose_slide(unconstrained, package, 0)
    assert retained.pattern_id == pattern.id
