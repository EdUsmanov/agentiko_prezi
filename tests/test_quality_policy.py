import json
from pathlib import Path
from types import SimpleNamespace
import pytest
from studio.checks.quality import candidate_regressions, quality_report, meaningful_diversity
from studio.models import Box, Element, SlideScene


def quality_profile(patterns=None):
    from studio.config import ROOT

    return SimpleNamespace(
        patterns=patterns or [],
        width=960,
        height=540,
        body_size=20,
        font="Play",
        font_file=str(ROOT / "fonts/Play-Regular.ttf"),
        font_roles={},
        font_assets=[],
        foreground="#222222",
        colors=["#222222", "#FFFFFF"],
    )


def manifest():
    return {
        "variants": [{"key": "executive", "findings": []}],
        "visual_audit": {"status": "completed", "checked": 1, "total": 1, "findings": []},
        "contextual_audit": {"status": "completed", "findings": []},
        "composition_diversity": {"verified": True},
        "errors": 0,
    }


@pytest.mark.parametrize("source", ["variant", "contextual_audit", "visual_audit"])
def test_warning_from_every_source_reaches_status_and_report(source):
    value = manifest()
    warning = {"code": "contrast", "severity": "warning", "message": "Low contrast", "slide": 1}
    (value["variants"][0] if source == "variant" else value[source])["findings"] = [warning]
    report = quality_report(value)
    assert report["status"] == "needs_review" and report["warnings"] == 1
    assert report["findings"][0]["message"] == "Low contrast"


def test_incomplete_review_and_error_require_review():
    value = manifest()
    value["visual_audit"]["checked"] = 0
    assert quality_report(value)["status"] == "needs_review"
    value["variants"][0]["findings"] = [
        {"severity": "error", "code": "missing_fact", "message": "Missing"}
    ]
    report = quality_report(value)
    assert report["status"] == "needs_review"
    assert report["errors"] == 1
    assert any(f["code"] == "missing_fact" for f in report["findings"])


def scene():
    return SlideScene(
        title="Title",
        background="#FFFFFF",
        source_ids=["f"],
        layout="table",
        purpose="metrics",
        elements=[
            Element(
                kind="text", role="title", size=40, text="Title", box=Box(x=20, y=10, w=800, h=70)
            ),
            Element(
                kind="table",
                size=20,
                source_ids=["f"],
                box=Box(x=20, y=100, w=500, h=220),
                rows=[["a"], ["b"]],
            ),
        ],
    )


def test_title_without_fact_id_and_data_area_are_protected():
    old = scene()
    new = old.model_copy(deep=True)
    new.elements[0].size = 12
    new.elements[1].box.h = 100
    package = SimpleNamespace(template=quality_profile())
    codes = {f["code"] for f in candidate_regressions([old], [new], package, audit=lambda *_: [])}
    assert codes == {"readability_regression", "data_area_regression"}


def test_translation_does_not_count_as_new_composition():
    originals = [scene() for _ in range(3)]
    second = [s.model_copy(deep=True) for s in originals]
    profile = SimpleNamespace(width=960, height=540)
    for s in second:
        s.elements[1].box.x += 2
    assert not meaningful_diversity({"a": originals, "b": second}, profile)["verified"]
    second[0].elements[1].box.x += 180
    assert not meaningful_diversity({"a": originals, "b": second}, profile)["verified"]
    second[1].elements[1].box.x += 180
    assert not meaningful_diversity({"a": originals, "b": second}, profile)["verified"]


@pytest.fixture
def saved_european():
    from studio.models import PreparedPackage, Plans

    path = Path(
        "/Users/edward/Documents/ChatGPT/Хакатон - ВК/template-benchmark/results-2026-09-26/08-main/result.json"
    )
    if not path.exists():
        pytest.skip("Local benchmark fixture not available")
    result = json.loads(path.read_text())
    generation = Path(result["manifest"]).parent
    package = PreparedPackage.model_validate_json(
        (generation.parent / result["preparation_id"] / "package.json").read_text()
    )
    plans = Plans.model_validate_json((generation / "plans.json").read_text())
    decks = {
        v.key: [
            SlideScene.model_validate(s)
            for s in json.loads(
                (generation / "refinement-1/.originals" / v.key / "slides.json").read_text()
            )
        ]
        for v in plans.variants
    }
    for v in plans.variants:
        for p, s in zip(v.slides, decks[v.key]):
            p.pattern_id = s.pattern_id
    return package, plans, decks


def test_european_repair_rejects_real_20_to_12_candidate(saved_european):
    from studio.checks.refinement import apply_edits, LayoutEdit

    package, plans, decks = saved_european
    with pytest.raises(ValueError, match="quality"):
        apply_edits(
            package,
            plans,
            decks,
            [LayoutEdit(variant="executive", slide=5, pattern_id="native-slide-17")],
            {("executive", 5): ["native-slide-17"]},
        )
    assert [e.size for e in decks["executive"][4].elements if e.role == "body"] == [20]


def test_european_design_reverts_same_real_candidate(saved_european):
    from studio.providers.deeppresenter import CompositionEnvironment, Assignment

    package, plans, _ = saved_european
    env = CompositionEnvironment(package, plans)
    assignments = [Assignment.model_validate(row) for row in env.baseline_assignments]
    next(
        a for a in assignments if a.variant == "executive" and a.slide == 5
    ).pattern_id = "native-slide-17"
    env.compose(assignments)
    assert env.plans.variants[0].slides[4].pattern_id != "native-slide-17"


def test_exported_geometry_ignores_titles_and_rejects_one_point_motion():
    from pptx import Presentation
    from pptx.util import Pt
    from studio.models import Fact, SlidePlan, VariantPlan
    from studio.checks.export_audit import content_scenes

    facts = [Fact(id=f"f{i}", text=f"Evidence {i}") for i in range(3)]
    package = SimpleNamespace(
        content=SimpleNamespace(facts=facts),
        template=SimpleNamespace(width=960, height=540, background="#FFFFFF"),
    )
    variant = VariantPlan(
        key="executive",
        title="Test",
        slides=[
            SlidePlan(title=f"Title {i}", fact_ids=[f"f{i}"], purpose="content") for i in range(3)
        ],
    )
    prs = Presentation()
    for plan, fact in zip(variant.slides, facts):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Pt(20), Pt(20), Pt(800), Pt(80)).text = plan.title
        slide.shapes.add_textbox(Pt(20), Pt(120), Pt(400), Pt(180)).text = fact.text
    before = content_scenes(prs, variant, package)
    for slide in prs.slides:
        slide.shapes[0].left += Pt(200)
        slide.shapes[1].left += Pt(1)
    after = content_scenes(prs, variant, package)
    assert not meaningful_diversity({"a": before, "b": after}, package.template)["verified"]
    for slide in list(prs.slides)[:2]:
        slide.shapes[1].left += Pt(180)
    assert not meaningful_diversity(
        {"a": before, "b": content_scenes(prs, variant, package)}, package.template
    )["verified"]


def test_diversity_does_not_rewrite_slides_just_to_move_content(monkeypatch):
    from studio.models import Pattern, Fact
    from studio.checks import diversity
    from studio.composition import composer

    title = Box(x=20, y=10, w=800, h=60)
    left = Box(x=20, y=120, w=350, h=200)
    right = Box(x=500, y=120, w=350, h=200)
    patterns = [
        Pattern(
            id=key,
            source_slide=1,
            source_layout=key,
            role="statement",
            purpose="content",
            title_zone=title,
            body_zones=[box],
            text_zones=[title, box],
        )
        for key, box in [("left", left), ("right", right)]
    ]
    scenes = [
        SlideScene(
            title=f"Title {i}",
            background="#FFFFFF",
            source_ids=[f"f{i}"],
            layout="statement",
            purpose="content",
            pattern_id="left",
            elements=[
                Element(kind="text", role="title", text=f"Title {i}", size=32, box=title),
                Element(kind="text", text=f"Evidence {i}", size=20, source_ids=[f"f{i}"], box=left),
            ],
        )
        for i in range(2)
    ]
    package = SimpleNamespace(
        template=quality_profile(patterns),
        content=SimpleNamespace(
            tables=[], facts=[Fact(id=f"f{i}", text=f"Evidence {i}") for i in range(2)]
        ),
    )

    def compose(plan, package, index, key):
        candidate = scenes[index].model_copy(deep=True)
        candidate.pattern_id = plan.pattern_id
        candidate.elements[1].box = right.model_copy()
        return candidate

    monkeypatch.setattr(composer, "compose", compose)
    monkeypatch.setattr(diversity, "audit_scenes", lambda *_: [])
    decks = {
        "executive": [s.model_copy(deep=True) for s in scenes],
        "analytical": [s.model_copy(deep=True) for s in scenes],
    }
    report = diversity.ensure_diversity(decks, package)
    assert not report["verified"]
    assert not report["adjustments"]
    assert decks["analytical"] == scenes
