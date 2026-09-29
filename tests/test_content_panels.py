from types import SimpleNamespace

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
import pytest

from studio.models import Box, Element, Pattern, SlideScene
from studio.scene_regions import unused_body_regions
from studio.scene_quality import scene_quality_findings


@pytest.fixture
def panel(tmp_path):
    def build(box=Box(x=360, y=0, w=360, h=405), color="#262626", layout=False):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Pt(720), Pt(405)
        source_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(source_layout)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, *[Pt(v) for v in (box.x, box.y, box.w, box.h)]
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color[1:])
        shape.line.fill.background()
        if layout:
            # Layout artwork is inherited even when the pattern has no source slide.
            source_layout.shapes._spTree.insert_element_before(shape._element, "p:extLst")
        path = tmp_path / "background.pptx"
        prs.save(path)
        image = tmp_path / "background.png"
        canvas = Image.new("RGB", (720, 405), "white")
        ImageDraw.Draw(canvas).rectangle((box.x, box.y, box.x + box.w, box.y + box.h), fill=color)
        canvas.save(image)
        body = Box(x=30, y=100, w=290, h=240)
        pattern = Pattern(
            id="arbitrary",
            source_slide=0 if layout else 1,
            source_layout="Arbitrary name",
            master_index=0,
            layout_index=6,
            role="statement",
            purpose="content",
            text_zones=[body],
            body_zones=[body],
            background_image=str(image),
        )
        profile = SimpleNamespace(
            width=720, height=405, patterns=[pattern], background_source=str(path)
        )
        scene = SlideScene(
            title="Title",
            background="#FFFFFF",
            layout="statement",
            purpose="content",
            pattern_id=pattern.id,
            source_ids=["f"],
            elements=[
                Element(
                    kind="text", role="body", text="Content", size=20, source_ids=["f"], box=body
                ),
            ],
        )
        return SimpleNamespace(template=profile), scene, box, image

    return build


@pytest.mark.parametrize("layout", [False, True])
@pytest.mark.parametrize("color", ["#262626", "#FFD900"])
def test_empty_panel_reaches_selection_and_absolute_audit(panel, layout, color):
    package, scene, _, _ = panel(color=color, layout=layout)
    assert unused_body_regions(scene, package) == 1
    assert "unused_template_regions" in {f.code for f in scene_quality_findings([scene], package)}


@pytest.mark.parametrize("kind", ["text", "chart", "table", "image", "title"])
def test_filled_panel_and_title_plate_are_not_empty(panel, kind):
    package, scene, zone, _ = panel()
    scene.elements.append(
        Element(
            kind="text" if kind == "title" else kind,
            role="title" if kind == "title" else "body",
            box=Box(x=zone.x + 30, y=zone.y + 60, w=zone.w - 60, h=140),
            source_ids=[] if kind in ("image", "title") else ["data"],
            image_id="uploaded" if kind == "image" else "",
        )
    )
    assert unused_body_regions(scene, package) == 0


@pytest.mark.parametrize(
    "box",
    [
        Box(x=0, y=0, w=720, h=405),  # Full background.
        Box(x=0, y=375, w=720, h=30),  # Footer stripe.
        Box(x=620, y=10, w=70, h=50),  # Small decoration/logo.
    ],
)
def test_background_and_small_artwork_are_preserved(panel, box):
    package, scene, _, _ = panel(box=box)
    assert unused_body_regions(scene, package) == 0


def test_photo_over_panel_is_not_treated_as_blank_and_source_is_immutable(panel):
    package, scene, zone, image = panel()
    source = package.template.background_source
    from pathlib import Path

    before = Path(source).read_bytes()
    canvas = Image.open(image)
    draw = ImageDraw.Draw(canvas)
    for x in range(int(zone.x), int(zone.x + zone.w), 15):
        draw.rectangle((x, 0, x + 7, 405), fill="#1177FF")
    canvas.save(image)
    assert unused_body_regions(scene, package) == 0
    assert Path(source).read_bytes() == before


@pytest.mark.parametrize("purpose", ["cover", "divider"])
def test_cover_and_divider_whitespace_is_intentional(panel, purpose):
    package, scene, _, _ = panel()
    scene.purpose = purpose
    assert unused_body_regions(scene, package) == 0


def test_panel_occupied_by_template_background_alone_still_counts(panel):
    package, scene, zone, _ = panel()
    scene.elements.append(Element(kind="image", role="template_background", box=zone))
    assert unused_body_regions(scene, package) == 1


def test_background_donor_checks_artwork_without_requiring_original_text_slots(panel):
    package, scene, _, _ = panel()
    scene.background_pattern_id, scene.pattern_id = scene.pattern_id, None
    package.template.patterns[0].body_zones.append(Box(x=10, y=10, w=100, h=50))
    assert unused_body_regions(scene, package) == 1


def test_same_color_panel_is_not_a_visible_empty_slot(panel):
    package, scene, _, _ = panel(color="#FFFFFF")
    assert unused_body_regions(scene, package) == 0


@pytest.mark.parametrize("missing", [True, False])
def test_unavailable_optional_source_does_not_break_audit(panel, missing):
    from pathlib import Path

    package, scene, _, _ = panel()
    source = Path(package.template.background_source)
    if missing:
        source.unlink()
    else:
        source.write_bytes(b"")
    assert unused_body_regions(scene, package) == 0
