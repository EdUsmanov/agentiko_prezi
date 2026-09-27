from pathlib import Path
from zipfile import ZipFile
import httpx
import pytest
from pptx import Presentation
from pptx.util import Pt
from fastapi.testclient import TestClient
from studio.config import ROOT, Settings
from studio.template import analyze_template
from studio.fonts import role_font, resolve_font
from studio import font_manifest, open_fonts
from studio.app import create_app
from tests.test_api import wait_job


def test_roles_are_distinct_and_report_contains_no_source_text(template, tmp_path):
    prs = Presentation(template)
    for slide in prs.slides:
        for run in slide.shapes[0].text_frame.paragraphs[0].runs:
            run.font.name = "Montserrat Medium"
    source = tmp_path / "roles.pptx"
    prs.save(source)
    profile = analyze_template(source, tmp_path / "analysis")
    assert role_font(profile, "title")[0] == "Montserrat Medium"
    assert role_font(profile, "body")[0] == "Play"
    assert not profile.missing_fonts
    report = (tmp_path / "analysis/font-model.json").read_text()
    assert "OLD PRIVATE" not in report and str(ROOT) not in report
    assert all(a["sha256"] and a["bytes"] for a in profile.font_assets)
    from studio.composer import text_element
    from studio.models import Box

    assert (
        text_element("Заголовок", Box(x=1, y=1, w=300, h=50), profile, "title").font
        == "Montserrat Medium"
    )


def test_theme_placeholder_inheritance(template, tmp_path):
    from lxml import etree
    from studio.font_manifest import A

    prs = Presentation(template)
    for entry in list(prs.slides._sldIdLst):
        prs.part.drop_rel(entry.rId)
        prs.slides._sldIdLst.remove(entry)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Название"
    slide.placeholders[1].text = "Содержание"
    for master in prs.slide_masters:
        rel = next(r for r in master.part.rels.values() if r.reltype.endswith("/theme"))
        root = etree.fromstring(rel.target_part.blob)
        root.find(".//" + A + "majorFont/" + A + "latin").set("typeface", "Montserrat Medium")
        root.find(".//" + A + "minorFont/" + A + "latin").set("typeface", "Play")
        rel.target_part._blob = etree.tostring(root)
    source = tmp_path / "theme.pptx"
    prs.save(source)
    report = font_manifest.build_font_manifest(prs, source, tmp_path / "theme-report")
    assert not report["unresolved"]
    assert {a["requested"] for a in report["assets"]} == {"Play", "Montserrat Medium"}


def test_missing_font_can_resume_without_reupload(template, content, tmp_path, monkeypatch):
    original = font_manifest.resolve_font
    monkeypatch.setattr(font_manifest, "resolve_font", lambda name: "")
    with TestClient(create_app(Settings(data_dir=tmp_path / "api"))) as client:
        response = client.post(
            "/api/prepare",
            data={"text": content, "slides": 5},
            files={"template": ("test.pptx", template.read_bytes())},
        )
        jid = response.json()["id"]
        waiting = wait_job(client, jid)
        assert waiting["state"] == "waiting_fonts"
        assert waiting["missing_fonts"][0]["requested"] == "Play"
        assert client.post("/api/generate", json={"package_id": jid}).status_code == 409
        assert client.get(f"/api/jobs/{jid}/files/font-model.json").status_code == 200
        assert client.get(f"/api/jobs/{jid}/files/font-resume.json").status_code == 409
        monkeypatch.setattr(font_manifest, "resolve_font", original)
        assert client.post(f"/api/packages/{jid}/retry-fonts").status_code == 202
        assert wait_job(client, jid)["state"] == "ready"
        assert client.post(f"/api/packages/{jid}/retry-fonts").status_code == 409


def test_download_is_pinned_and_validated(tmp_path, monkeypatch):
    raw = (ROOT / "fonts/Play-Regular.ttf").read_bytes()
    metadata = b'license: "OFL"\nfonts {\n name: "Play"\n full_name: "Play Regular"\n filename: "Play-Regular.ttf"\n}\n'
    seen = []

    def respond(request):
        seen.append(str(request.url))
        name = request.url.path.rsplit("/", 1)[-1]
        return httpx.Response(
            200,
            content={
                "METADATA.pb": metadata,
                "OFL.txt": b"SIL OPEN FONT LICENSE test",
                "Play-Regular.ttf": raw,
            }[name],
        )

    real_client = httpx.Client
    monkeypatch.setattr(
        open_fonts.httpx,
        "Client",
        lambda **kw: real_client(transport=httpx.MockTransport(respond), **kw),
    )
    monkeypatch.setattr(open_fonts, "CACHE", tmp_path / "download")
    file, issue = open_fonts.download_face("Play")
    assert not issue and Path(file).read_bytes() == raw
    assert Path(file).with_suffix(".json").exists()
    assert all(url.startswith(open_fonts.BASE + "play/") for url in seen)
    assert (Path(file).parent / "OFL.txt").exists()
    before = len(seen)
    assert not open_fonts.download_face("https://evil.example/font")[0]
    assert not open_fonts.download_face("../../secret")[0]
    assert len(seen) == before


@pytest.mark.parametrize("failure", ["redirect", "oversize", "bad_font", "mismatch"])
def test_download_failures_leave_no_font(tmp_path, monkeypatch, failure):
    _raw = (ROOT / "fonts/Play-Regular.ttf").read_bytes()

    def respond(request):
        name = request.url.path.rsplit("/", 1)[-1]
        if name == "METADATA.pb":
            if failure == "redirect":
                return httpx.Response(302, headers={"Location": "http://127.0.0.1/private"})
            if failure == "oversize":
                return httpx.Response(200, content=b"x" * 256001)
            return httpx.Response(
                200,
                content=b'license: "OFL"\nfonts {\n full_name: "Play Regular"\n filename: "Play-Regular.ttf"\n}\n',
            )
        if name == "OFL.txt":
            return httpx.Response(200, content=b"SIL OPEN FONT LICENSE test")
        return httpx.Response(
            200,
            content=b"not a font"
            if failure == "bad_font"
            else (ROOT / "fonts/Montserrat-Regular.ttf").read_bytes(),
        )

    real_client = httpx.Client
    monkeypatch.setattr(
        open_fonts.httpx,
        "Client",
        lambda **kw: real_client(transport=httpx.MockTransport(respond), **kw),
    )
    monkeypatch.setattr(open_fonts, "CACHE", tmp_path / "download")
    file, issue = open_fonts.download_face("Play")
    assert not file and issue
    assert not list(tmp_path.rglob("*.ttf"))


def test_all_role_font_hashes_checked(prepared):
    from studio.pipeline import load_package
    from studio.security import digest

    settings, store, package = prepared
    package.template.font_assets.append(
        {"id": "bad", "path": "/nonexistent-font.ttf", "sha256": "bad"}
    )
    raw = package.model_dump_json()
    (store.directory(package.id) / "package.json").write_text(raw)
    store.update(package.id, package_hash=digest(raw.encode()))
    with pytest.raises(ValueError, match="Начертание"):
        load_package(store, package.id)


def test_automatic_download_enabled_only_by_setting(template, tmp_path, monkeypatch):
    original = resolve_font("Play")
    monkeypatch.setattr(font_manifest, "resolve_font", lambda name: "")
    calls = []
    monkeypatch.setattr(
        open_fonts, "download_face", lambda name: (calls.append(name) or original, "")
    )
    profile = analyze_template(template, tmp_path / "offline")
    assert profile.missing_fonts and calls == []
    profile = analyze_template(template, tmp_path / "online", allow_download=True)
    assert not profile.missing_fonts and calls == ["Play"]
    assert profile.font_assets[0]["status"] == "downloaded"


def test_optional_face_does_not_trigger_download(template, tmp_path, monkeypatch):
    prs = Presentation(template)
    box = prs.slides[0].shapes.add_textbox(Pt(100), Pt(300), Pt(200), Pt(50))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "x"
    run.font.name = "Calibri Light"
    run.font.bold = True
    source = tmp_path / "optional-face.pptx"
    prs.save(source)
    play = resolve_font("Play")
    original_resolve = font_manifest.resolve_font
    monkeypatch.setattr(
        font_manifest,
        "resolve_font",
        lambda name: (
            play
            if name == "Play"
            else original_resolve(name)
            if name != "Calibri Light Bold"
            else ""
        ),
    )
    calls = []
    monkeypatch.setattr(
        open_fonts, "download_face", lambda name: (calls.append(name) or "", "unexpected")
    )
    profile = analyze_template(source, tmp_path / "analysis", allow_download=True)
    assert calls == []
    unresolved = next(
        item for item in profile.missing_fonts if item["requested"] == "Calibri Light Bold"
    )
    assert unresolved["required_for_generation"] is False
    assert all("Calibri Light Bold" not in warning for warning in profile.warnings)


def test_font_manifest_is_validated_before_atomic_publication(template, tmp_path):
    prs = Presentation(template)
    source = tmp_path / "source.pptx"
    prs.save(source)
    report = font_manifest.build_font_manifest(prs, source, tmp_path / "analysis")
    model = tmp_path / "analysis/font-model.json"
    assert model.is_file()
    assert not list(model.parent.glob(".font-model.json.*"))
    report["assets"][0]["sha256"] = "0" * 64
    with pytest.raises(ValueError, match="SHA-256"):
        font_manifest._validate_manifest(report)


@pytest.mark.skipif(
    not (ROOT / "data/local-fonts/aptos/Aptos.ttf").is_file(),
    reason="Optional local Microsoft package",
)
def test_installed_aptos_roles_and_cyrillic(template, tmp_path):
    from studio.embedded_fonts import check_glyphs
    from studio.models import Element, Box, SlideScene
    from studio.render import render_variant

    prs = Presentation(template)
    for slide in prs.slides:
        for i, shape in enumerate(slide.shapes):
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Aptos Display" if i == 0 else "Aptos"
    source = tmp_path / "aptos.pptx"
    prs.save(source)
    profile = analyze_template(source, tmp_path / "analysis")
    assert not profile.missing_fonts
    assert role_font(profile, "title")[0] == "Aptos Display"
    assert role_font(profile, "body")[0] == "Aptos"
    for asset in profile.font_assets:
        check_glyphs(asset["path"], "Проверка русского текста — 2026 •")
        assert not asset["redistributable"]
    scene = SlideScene(
        title="Проверка Aptos",
        background=profile.background,
        elements=[
            Element(
                kind="text",
                box=Box(x=50, y=50, w=600, h=100),
                text="Проверка Aptos Display",
                role="title",
                font="Aptos Display",
                size=32,
                color=profile.foreground,
            )
        ],
        source_ids=[],
        layout="statement",
    )
    render_variant([scene], profile, source, tmp_path / "rendered")
    assert "data:font" not in (tmp_path / "rendered/deck.html").read_text()


def test_local_only_font_bytes_never_exported(template, tmp_path):
    # Use an ordinary test font under local-only policy, independent of Aptos installation.
    from studio.models import Element, Box, SlideScene
    from studio.render import render_variant
    import pypdfium2 as pdfium

    profile = analyze_template(template, tmp_path / "analysis")
    for asset in profile.font_assets:
        asset["redistributable"] = False
    scene = SlideScene(
        title="Тест",
        background=profile.background,
        elements=[
            Element(
                kind="text",
                box=Box(x=50, y=50, w=400, h=100),
                text="Редактируемый текст",
                font=profile.font,
                size=24,
                color=profile.foreground,
            )
        ],
        source_ids=[],
        layout="statement",
    )
    target = tmp_path / "out"
    render_variant([scene], profile, template, target)
    html = (target / "deck.html").read_text()
    assert "data:image/png" in html and "data:font" not in html
    with ZipFile(target / "deck.pptx") as z:
        assert not any(n.startswith("ppt/fonts/") for n in z.namelist())
    assert any(
        s.has_text_frame and "Редактируемый" in s.text
        for s in Presentation(target / "deck.pptx").slides[0].shapes
    )
    with pdfium.PdfDocument(str(target / "deck.pdf")) as pdf:
        assert len(pdf) == 1
        page = pdf[0]
        text = page.get_textpage()
        assert text.get_text_range() == ""
        text.close()
        page.close()
