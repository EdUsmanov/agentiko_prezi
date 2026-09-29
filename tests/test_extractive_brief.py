"""A direct-fact brief stays reviewable and frozen through the real HTTP worker."""

from io import BytesIO
from copy import deepcopy
import time

import httpx
from pptx import Presentation

from test_support.app_server import application
from test_support.inputs import make_template


def wait_job(client, url, jid):
    until = time.monotonic() + 120
    while time.monotonic() < until:
        job = client.get(f"{url}/api/jobs/{jid}").json()
        if job["state"] not in ("accepted", "running"):
            return job
        time.sleep(0.2)
    raise AssertionError(f"Job {jid} timed out")


def test_extractive_brief_needs_exact_approval_and_keeps_copy(tmp_path):
    source = (
        "# Пилот\nКоманда открывает единый сервис заявок. Руководитель проверяет результат пилота."
    )
    template = make_template(tmp_path / "brief-template.pptx")
    with application(tmp_path / "app") as (url, _, _):
        with httpx.Client(timeout=30, trust_env=False) as client:
            with template.open("rb") as file:
                uploaded = client.post(
                    f"{url}/api/prepare",
                    data={"text": source, "slides": "2", "input_mode": "brief"},
                    files={"template": (template.name, file, "application/octet-stream")},
                )
            assert uploaded.status_code == 202, uploaded.text
            ready = wait_job(client, url, uploaded.json()["id"])
            assert ready["state"] == "ready", ready
            pid = ready["id"]
            blocked = client.post(
                f"{url}/api/generate",
                json={"package_id": pid, "accept_adjusted_slide_count": True},
            )
            assert blocked.status_code == 409, blocked.text
            review = client.get(f"{url}/api/packages/{pid}/draft")
            assert review.status_code == 200, review.text
            draft = review.json()
            assert draft["package_hash"] == ready["package_hash"]
            assert draft["draft_hash"] == ready["draft_hash"]
            bullets = [bullet for slide in draft["draft"]["slides"] for bullet in slide["bullets"]]
            assert len(bullets) == 2
            assert all(bullet["fact_ids"] and not bullet["proposed"] for bullet in bullets)
            approved = client.post(
                f"{url}/api/packages/{pid}/approve",
                json={"package_hash": draft["package_hash"], "draft_hash": draft["draft_hash"]},
            )
            assert approved.status_code == 200, approved.text
            assert approved.json()["approved_draft_hash"] == draft["draft_hash"]
            assert (
                client.post(f"{url}/api/packages/{pid}/auto-generation/cancel").status_code == 200
            )
            created = client.post(
                f"{url}/api/generate",
                json={"package_id": pid, "accept_adjusted_slide_count": True},
            )
            assert created.status_code == 202, created.text
            done = wait_job(client, url, created.json()["id"])
            assert done["state"] in ("needs_review", "completed"), done
            assert done["quality_report"]["errors"] == 0
            output = client.get(f"{url}/api/jobs/{done['id']}/files/executive/deck.pptx")
            assert output.status_code == 200
            deck = Presentation(BytesIO(output.content))
            text = "\n".join(
                shape.text
                for slide in deck.slides
                for shape in slide.shapes
                if shape.has_text_frame
            )
            assert all(bullet["text"] in text for bullet in bullets)

            edited = deepcopy(draft["draft"])
            edited["slides"][0]["bullets"][0]["text"] = "Команда запустила единый сервис заявок."
            revision = client.post(
                f"{url}/api/packages/{pid}/draft",
                json={"package_hash": draft["package_hash"], "draft": edited},
            )
            assert revision.status_code == 202, revision.text
            child = wait_job(client, url, revision.json()["id"])
            assert child["state"] == "ready", child
            assert child["id"] != pid
            assert child.get("approved_draft_hash") is None
            assert child.get("approved_package_hash") is None
            blocked_child = client.post(
                f"{url}/api/generate",
                json={"package_id": child["id"], "accept_adjusted_slide_count": True},
            )
            assert blocked_child.status_code == 409, blocked_child.text
            stale_approval = client.post(
                f"{url}/api/packages/{child['id']}/approve",
                json={"package_hash": draft["package_hash"], "draft_hash": draft["draft_hash"]},
            )
            assert stale_approval.status_code == 409, stale_approval.text
            current = client.get(f"{url}/api/packages/{child['id']}/draft")
            assert current.status_code == 200, current.text
            current = current.json()
            assert current["package_hash"] != draft["package_hash"]
            assert current["draft_hash"] != draft["draft_hash"]
            assert (
                current["draft"]["slides"][0]["bullets"][0]["text"]
                == edited["slides"][0]["bullets"][0]["text"]
            )
