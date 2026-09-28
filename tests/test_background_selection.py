from pathlib import Path
from types import SimpleNamespace

from PIL import Image, ImageDraw
import pytest

from studio.background_selection import apply_background, background_is_safe, background_candidates
from studio.models import Box, Element, Pattern, SlideScene


@pytest.fixture
def background(tmp_path):
    path = tmp_path / "clean.png"
    image = Image.new("RGB", (720, 405), "white")
    ImageDraw.Draw(image).rectangle((600, 250, 719, 404), fill="blue")
    image.save(path)
    source = tmp_path / "clean.pptx"
    source.touch()
    pattern = Pattern(
        id="donor",
        source_layout="content",
        role="content",
        purpose="content",
        text_zones=[],
        background="#FFFFFF",
        foreground="#000000",
        background_image=str(path),
        source_slide=1,
    )
    package = SimpleNamespace(
        template=SimpleNamespace(
            patterns=[pattern],
            width=720,
            height=405,
            background_source=str(source),
            foreground="#FFFFFF",
            colors=["#FFFFFF", "#000000"],
        )
    )
    scene = SlideScene(
        title="Title",
        background="#000000",
        source_ids=["f"],
        layout="split",
        purpose="content",
        elements=[
            Element(
                kind="text",
                role="title",
                text="Title",
                size=32,
                color="#FFFFFF",
                box=Box(x=40, y=40, w=400, h=60),
            ),
            Element(
                kind="text",
                text="Evidence",
                size=20,
                color="#FFFFFF",
                source_ids=["f"],
                box=Box(x=40, y=150, w=400, h=150),
            ),
        ],
    )
    return package, scene


def test_background_changes_paint_without_changing_content_or_geometry(background):
    package, scene = background
    before = scene.model_dump()
    result = apply_background(scene, package, "donor")
    assert scene.model_dump() == before
    assert result.background_pattern_id == "donor" and result.pattern_id is None
    assert result.elements[0].role == "template_background"
    for old, new in zip(scene.elements, result.elements[1:]):
        assert old.model_dump(exclude={"color", "background_hint"}) == new.model_dump(
            exclude={"color", "background_hint"}
        )
        assert new.color == "#000000" and new.background_hint == "#FFFFFF"
    assert background_is_safe(SlideScene.model_validate_json(result.model_dump_json()), package)


@pytest.mark.parametrize("change", ["artwork", "outside", "upload", "cover", "unknown", "native"])
def test_unsafe_or_inapplicable_background_is_rejected(background, change):
    package, scene = background
    pid = "donor"
    if change == "artwork":
        scene.elements[1].box.x = 300
    elif change == "outside":
        scene.elements[1].box.y = 350
    elif change == "upload":
        scene.elements[1].image_id = "uploaded"
    elif change == "cover":
        scene.purpose = "cover"
    elif change == "unknown":
        pid = "missing"
    elif change == "native":
        scene.pattern_id = "native"
    assert apply_background(scene, package, pid) is None


def test_no_background_without_sanitized_source(background):
    package, scene = background
    package.template.background_source = ""
    assert not background_candidates(package)
    assert apply_background(scene, package, "donor") is None


def test_later_geometry_or_paint_change_blocks_audit_and_export(background, tmp_path):
    from studio.scene_quality import scene_quality_findings
    from studio.render import render_pptx
    from studio.repair_policy import FIT_CODES

    package, scene = background
    result = apply_background(scene, package, "donor")
    result.elements[-1].box.x = 300
    assert not background_is_safe(result, package)
    assert any(
        f.code == "background_conflict" and f.severity == "error"
        for f in scene_quality_findings([result], package)
    )
    assert "background_conflict" in FIT_CODES
    with pytest.raises(ValueError, match="фон"):
        render_pptx([result], package.template, None, tmp_path / "out.pptx")
    assert not (tmp_path / "out.pptx").exists()
    result = apply_background(scene, package, "donor")
    result.elements[-1].color = "#FFFFFF"
    assert not background_is_safe(result, package)


def test_background_color_does_not_inflate_composition_diversity(background):
    from studio.quality import meaningful_diversity
    from studio.background_diversity import background_report
    from studio.stage_results import BackgroundDiversity

    package, scene = background
    changed = apply_background(scene, package, "donor")
    assert not meaningful_diversity({"a": [scene], "b": [changed]}, package.template)["verified"]
    report = BackgroundDiversity.model_validate(
        background_report([scene, changed], package.template)
    )
    assert report.background_colors == {"#000000": 1, "#FFFFFF": 1}


def test_flat_master_and_token_background_count_as_one_family(background, tmp_path):
    from studio.background_diversity import background_report

    package, scene = background
    path = tmp_path / "white.png"
    Image.new("RGB", (720, 405), "white").save(path)
    package.template.patterns[0].background_image = str(path)
    generic = scene.model_copy(update={"background": "#FFFFFF"})
    selected = apply_background(scene, package, "donor")
    assert background_report([generic, selected], package.template)["unique_backgrounds"] == 1


def test_native_export_and_plan_validation_use_clean_source(prepared, tmp_path):
    from pptx import Presentation
    from studio.planner import validate_plans, planning_schema
    from studio.render import render_pptx

    _, store, package = prepared
    pattern = next(p for p in background_candidates(package) if p.source_slide)
    scene = SlideScene(
        title="New title",
        background=package.template.background,
        elements=[
            Element(
                kind="text",
                role="title",
                text="New title",
                size=28,
                color=package.template.foreground,
                box=Box(x=50, y=40, w=500, h=60),
            )
        ],
        source_ids=[],
        layout="statement",
        purpose="content",
    )
    selected = apply_background(scene, package, pattern.id)
    assert selected is not None
    target = tmp_path / "selected.pptx"
    render_pptx([selected], package.template, store.directory(package.id) / "input.pptx", target)
    prs = Presentation(target)
    text = " ".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert "New title" in text and "OLD PRIVATE" not in text
    assert not any(s.shape_type == 13 for slide in prs.slides for s in slide.shapes)
    assert (
        "background_pattern_id" not in planning_schema(package)["$defs"]["SlidePlan"]["properties"]
    )
    plans = package.prepared_plans.model_copy(deep=True)
    plan = plans.variants[0].slides[1]
    plan.pattern_id = "token:auto"
    plan.background_pattern_id = "missing"
    with pytest.raises(ValueError, match="фон"):
        validate_plans(plans, package)
    plan.background_pattern_id = pattern.id
    validate_plans(plans, package)
    plan.pattern_id = pattern.id
    with pytest.raises(ValueError, match="фон"):
        validate_plans(plans, package)


def test_roomy_master_remains_eligible_when_authored_layout_is_too_small(prepared):
    from studio.composer import compose_variant
    from studio.contracts import candidates
    from studio.models import VariantPlan, SlidePlan
    from studio.audit import audit_scenes

    _, _, package = prepared
    package.analysis["editorial"] = {"status": "completed"}
    package.constraints.summarize = True
    original = next(p for p in package.template.patterns if p.source_slide and p.body_zones)
    authored = original.model_copy(deep=True)
    authored.id = "small-authored"
    authored.purpose = "content"
    authored.role = "content"
    authored.body_zones = [Box(x=50, y=160, w=60, h=25)]
    master = authored.model_copy(deep=True)
    master.id = "roomy-master"
    master.source_slide = 0
    master.fields = []
    master.body_zones = [Box(x=50, y=160, w=800, h=260)]
    package.template.patterns = [authored, master]
    slide = SlidePlan(
        title="Evidence",
        fact_ids=[package.content.facts[-1].id],
        layout="statement",
        purpose="content",
    )
    assert [p.id for p in candidates(package, slide)] == [authored.id]
    assert master in candidates(package, slide, prefer_specialized=False)
    variant = VariantPlan(key="executive", title="Test", slides=[slide])
    scene = compose_variant(variant, package)[0]
    assert scene.pattern_id == master.id
    assert not [f for f in audit_scenes([scene], package) if f.severity == "error" and f.slide == 1]


def test_visual_review_compares_actual_background_without_imposing_donor_geometry(
    background, tmp_path
):
    from studio.visual import visual_payload

    package, scene = background
    package.prepared_plans = None
    selected = apply_background(scene, package, "donor")
    folder = tmp_path / "executive"
    folder.mkdir()
    (folder / "slide-1.png").write_bytes(
        Path(package.template.patterns[0].background_image).read_bytes()
    )
    result = visual_payload([1], [selected.model_dump()], {"key": "executive"}, package, tmp_path)
    assert result.payload["reference_images"][0]["source_slide"] == 1
    assert len(result.images) == 2
    assert result.payload["slides"][0]["field_safety"] == []
    assert result.payload["slides"][0]["authored_title_position"] is False


@pytest.mark.parametrize("kind", ["chart", "table"])
def test_background_variety_cannot_trade_data_height_for_width(background, kind):
    from studio.background_selection import preserves_data_space

    _, scene = background
    scene.elements[-1].kind = kind
    alternative = scene.model_copy(deep=True)
    alternative.elements[-1].box.h -= 20
    alternative.elements[-1].box.w *= 2
    assert not preserves_data_space(scene, alternative)
    alternative.elements[-1].box = scene.elements[-1].box.model_copy()
    assert preserves_data_space(scene, alternative)
    alternative.elements[-1].box.w -= 20
    assert not preserves_data_space(scene, alternative)
