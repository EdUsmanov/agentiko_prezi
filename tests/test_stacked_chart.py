import pytest
from pptx import Presentation
from pptx.util import Pt

from studio.charts import make_chart, render_chart
from studio.export_audit import geometry, slide_text
from studio.models import Box, SlidePlan, TableData
from studio.stacked_chart import stacked_layout, value_label_positions
from studio.template import analyze_template


@pytest.fixture
def profile(template, tmp_path):
    p = analyze_template(template, tmp_path / "profile")
    p.width = 720
    p.height = 405
    return p


def example(profile, width=460, height=300, total=True):
    rows = [["A", "160", "240"], ["B", "80", "32"], ["C", "48", "32"], ["D", "32", "16"]]
    if total:
        rows.append(["Total", "320", "320"])
    table = TableData(
        id="data", headers=["Work category", "Earlier period", "Later period"], rows=rows
    )
    return make_chart(
        table,
        SlidePlan(title="Hours", fact_ids=["f"], chart_type="column_stacked"),
        Box(x=20, y=40, w=width, h=height),
        profile,
        profile.foreground,
        ["f"],
    )


@pytest.mark.parametrize("width,height", [(420, 300), (500, 330), (640, 340)])
def test_export_keeps_totals_values_and_contained_annotations(profile, tmp_path, width, height):
    e = example(profile, width, height)
    before = e.model_dump()
    layout = stacked_layout(e, profile)
    assert layout["fits"] and layout["height"] >= 110 and layout["legend_size"] < 16
    prs = Presentation()
    prs.slide_width = Pt(720)
    prs.slide_height = Pt(405)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = render_chart(slide, e, profile)
    assert [list(s.values) for s in chart.series] == e.series_values
    target = tmp_path / "stacked.pptx"
    prs.save(target)
    actual = Presentation(target)
    assert geometry(actual, profile)[0] == []
    text = slide_text(actual.slides[0])
    assert "Earlier period — Total: 320" in text and "Later period — Total: 320" in text
    assert e.model_dump() == before


def test_labels_on_tiny_segments_are_separated_and_bounded():
    ideal, placed = value_label_positions([32, 16, 1], 0, 400, 180)
    assert ideal != placed
    assert min(placed) >= 10 and max(placed) <= 170
    assert all(b - a >= 20 for a, b in zip(sorted(placed), sorted(placed)[1:]))


def test_small_or_unbreakable_chart_remains_rejected(profile):
    e = example(profile, height=160)
    assert not stacked_layout(e, profile)["fits"]
    e = example(profile)
    e.labels[0] = "LongUnbreakableIdentifier" * 10
    assert not stacked_layout(e, profile)["fits"]
    e.series_values[0].pop()
    assert not stacked_layout(e, profile)["fits"]


def test_chart_annotations_do_not_exempt_unrelated_text_or_label_overlap(profile):
    prs = Presentation()
    prs.slide_width = Pt(720)
    prs.slide_height = Pt(405)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_chart(slide, example(profile), profile)
    group = slide.shapes[0]
    labels = [s for s in group.shapes if s.has_text_frame and s.text == "16"]
    assert labels
    box = labels[0]
    other = group.shapes.add_textbox(box.left, box.top, box.width, box.height)
    other.text = "unrelated"
    assert any(f["code"] == "pptx_content_overlap" for f in geometry(prs, profile)[0])
    other.text = "16"
    assert any(f["code"] == "pptx_content_overlap" for f in geometry(prs, profile)[0])
    other.left = Pt(719)
    assert any(f["code"] == "pptx_out_of_bounds" for f in geometry(prs, profile)[0])


def test_percent_values_keep_units_in_visible_labels(profile):
    e = example(profile, width=640, height=340)
    e.value_labels = [str(v) + "%" for v in e.values]
    prs = Presentation()
    prs.slide_width = Pt(720)
    prs.slide_height = Pt(405)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_chart(slide, e, profile)
    assert "16%" in slide_text(slide)
    assert not geometry(prs, profile)[0]


@pytest.mark.parametrize("width,height", [(720, 405), (960, 540)])
def test_scene_uses_height_without_crossing_title_footer_or_prose(profile, width, height):
    from types import SimpleNamespace
    from studio.models import Constraints, ContentModel, Element, Fact, SlideScene
    from studio.stacked_chart import adapt_stacked_scene
    from studio.audit import audit_scenes
    from studio.repair_policy import FIT_CODES

    profile.width, profile.height = width, height
    e = example(profile, width=width * 0.5, height=160)
    e.box.x, e.box.y = 40, 100
    e.source_ids = ["data"]
    table = TableData(id="t", headers=e.rows[0], rows=e.rows[1:], visualization="column_stacked")
    package = SimpleNamespace(
        template=profile,
        content=ContentModel(
            title="Report",
            facts=[
                Fact(id="data", text="Measurements", source="t"),
                Fact(id="body", text="Measured progress across periods.", source="user_text"),
            ],
            tables=[table],
        ),
        constraints=Constraints(slides=1),
        analysis={},
        images=[],
    )
    scene = SlideScene(
        title="Report",
        background=profile.background,
        source_ids=["data", "body"],
        layout="chart",
        elements=[
            Element(
                kind="text",
                role="title",
                text="Report",
                size=22,
                font=profile.font,
                color=profile.foreground,
                box=Box(x=40, y=30, w=width - 80, h=30),
            ),
            e,
            Element(
                kind="text",
                text="Measured progress across periods.",
                source_ids=["body"],
                size=16,
                font=profile.font,
                color=profile.foreground,
                box=Box(x=width * 0.6, y=100, w=width * 0.3, h=100),
            ),
            Element(
                kind="text",
                role="footer",
                color=profile.foreground,
                text="1",
                size=10,
                font=profile.font,
                box=Box(x=width - 50, y=height - 25, w=25, h=15),
            ),
        ],
    )
    before = scene.model_dump()
    result = adapt_stacked_scene(scene, package)
    data = next(e for e in result.elements if e.kind == "chart")
    assert data.box.h > 160 and stacked_layout(data, profile)["fits"]
    assert data.box.y >= 68 and data.box.y + data.box.h <= height - 33
    assert data.series_values == e.series_values
    assert scene.model_dump() == before
    assert not [f for f in audit_scenes([result], package) if f.code in FIT_CODES]


def test_compact_legend_must_retain_each_source_total(profile):
    from types import SimpleNamespace
    from studio.export_audit import inspect_content
    from studio.models import ContentModel, VariantPlan
    from studio.template_geometry import walk_shapes

    e = example(profile)
    table = TableData(id="t", headers=e.rows[0], rows=e.rows[1:])
    package = SimpleNamespace(
        content=ContentModel(title="Hours", facts=[], tables=[table]), images=[]
    )
    variant = VariantPlan(
        key="executive",
        title="Result",
        slides=[
            SlidePlan(
                title="Hours",
                fact_ids=[],
                table_id="t",
                layout="chart",
                chart_type="column_stacked",
            )
        ],
    )
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_chart(slide, e, profile)
    assert not [
        f for f in inspect_content(prs, variant, package)[1] if f["code"].startswith("chart_")
    ]
    legend = next(
        s
        for s, _ in walk_shapes(slide.shapes)
        if s.has_text_frame and "Earlier period — Total: 320" in s.text
    )
    legend.text = legend.text.replace("320", "321")
    assert any(f["code"] == "chart_supplement" for f in inspect_content(prs, variant, package)[1])
