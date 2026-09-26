from copy import deepcopy
from types import SimpleNamespace
import pytest
from studio.content import parse_content
from studio.editorial import validate_plan,apply_plan,EditorialPlan
from studio.editorial_repair import apply_replacements,validation_targets
from studio.editorial_tables import validate_headers
from studio.models import TableData


def claim(text,fid='f1'):
    return {'title':text,'bullets':[{'text':text,'evidence':[{'fact_id':fid}]}]}


def test_patch_keeps_every_accepted_neighbour_and_rejects_wrong_scope():
    previous=EditorialPlan.model_validate({'slides':[claim('Один'),claim('Два'),claim('Три')]}).model_dump()
    snapshot=deepcopy(previous)
    patched=apply_replacements(previous,{'replacements':[{'slide':2,'content':claim('Исправлено')}]},[2])
    assert previous==snapshot and patched['slides'][0]==previous['slides'][0] and patched['slides'][2]==previous['slides'][2]
    assert len(patched['slides'])==3
    for indices in ([1],[2,2],[2,3]):
        with pytest.raises(ValueError):apply_replacements(previous,{'replacements':[{'slide':i,'content':claim('Новое')} for i in indices]},[2])


def test_cover_is_required_first_and_included_in_count():
    content=parse_content('Исторический обзор. Последствия войны.')
    raw={'slides':[claim('Исторический обзор.'),claim('Последствия войны.','f2')]}
    with pytest.raises(ValueError,match='s1'):validate_plan(raw,content,(2,2),require_cover=True)
    raw['slides'][0]['purpose']='cover'
    result=validate_plan(raw,content,(2,2),require_cover=True)
    assert len(result['slides'])==2 and result['slides'][0]['purpose']=='cover'
    raw['slides'][1]['purpose']='cover'
    with pytest.raises(ValueError,match='s2'):validate_plan(raw,content,(2,2),require_cover=True)


def test_source_columns_preserve_rows_and_separate_chart_units():
    content=parse_content('| Выборы | Голоса, % | Места |\n|---|---|---|\n| Май 1928 | 2,6 | 12 |\n| Сентябрь 1930 | 18,3 | 107 |\n| Июль 1932 | 37,3 | 230 |\n| Ноябрь 1932 | 33,1 | 196 |')
    table=content.tables[0];fid=next(f.id for f in content.facts if f.source==table.id)
    raw={'slides':[dict(claim('Доля голосов',fid),source_table_id=table.id,source_columns=[0,1],chart_type='line'),dict(claim('Места в парламенте',fid),source_table_id=table.id,source_columns=[0,2],chart_type='column')]}
    plan=validate_plan(raw,content,(2,2));p=SimpleNamespace(content=content,original_content=content.model_copy(deep=True),analysis={})
    apply_plan(p,plan,{},(2,2))
    assert len({t.id for t in p.content.tables})==2
    assert [t.visualization for t in p.content.tables]==['line','column']
    assert p.content.tables[0].rows==[[r[0],r[1]] for r in table.rows]
    assert p.content.tables[1].rows==[[r[0],r[2]] for r in table.rows]
    assert p.original_content.tables[0]==table
    with pytest.raises(ValueError,match='every source table column'):validate_plan({'slides':raw['slides'][:1]},content,(1,1))


def test_header_recovery_cannot_invent_units():
    broken=[TableData(id='t1',headers=['ВыборыГолоса, %Места','',''],rows=[['Май','2,6','12']])]
    assert validate_headers({'tables':[{'table_id':'t1','headers':['Выборы','Голоса, %','Места']}]},broken)
    with pytest.raises(ValueError):validate_headers({'tables':[{'table_id':'t1','headers':['Выборы','Голоса, %','Места, %']}]},broken)
    assert broken[0].headers[1]==''


def test_target_selection_and_uniform_region_artwork_guard(tmp_path):
    from PIL import Image,ImageDraw
    from studio.models import Box
    from studio.template_adaptation import uniform_region
    assert validation_targets(ValueError('s2b1 number; s4 title'),5)==[2,4]
    path=tmp_path/'background.png';im=Image.new('RGB',(400,300),'blue');im.save(path)
    box=Box(x=0,y=0,w=400,h=300)
    assert uniform_region(path,box,400,300)
    ImageDraw.Draw(im).rectangle((30,30,100,100),fill='pink');im.save(path)
    assert not uniform_region(path,box,400,300)


def test_compact_list_and_large_year_labels_are_separate_fields():
    from pptx import Presentation
    from pptx.util import Pt
    from studio.native_template import native_patterns
    prs=Presentation();prs.slide_width=Pt(720);prs.slide_height=Pt(405)
    for count in (6,3):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        title=slide.shapes.add_textbox(Pt(20),Pt(30),Pt(230),Pt(80));title.text='Хронология'
        title.text_frame.paragraphs[0].runs[0].font.size=Pt(40)
        for i in range(count):
            if count==3:
                heading=slide.shapes.add_textbox(Pt(320),Pt(30+i*110),Pt(350),Pt(40));heading.text=str(1920+i)
                heading.text_frame.paragraphs[0].runs[0].font.size=Pt(32)
            body=slide.shapes.add_textbox(Pt(320),Pt((30+i*50) if count==6 else (75+i*110)),Pt(350),Pt(22 if count==6 else 40))
            body.text='Событие';body.text_frame.paragraphs[0].runs[0].font.size=Pt(14)
    patterns=native_patterns(prs)
    six=next(p for p in patterns if p.source_slide==1)
    three=next(p for p in patterns if p.source_slide==2)
    assert len(six.body_zones)==6
    assert len(three.body_zones)==3 and all(three.heading_zones)


def test_duplicate_classifications_require_identical_meaning():
    from studio.analysis import normalize_meanings
    one={'pattern_id':'native-slide-3','roles':['context'],'density':'low','purpose':'agenda'}
    assert len(normalize_meanings({'patterns':[one,one]}).patterns)==1
    with pytest.raises(ValueError,match='Conflicting'):normalize_meanings({'patterns':[one,{**one,'purpose':'timeline'}]})


def test_chart_band_expansion_preserves_title_and_artwork(tmp_path):
    from PIL import Image,ImageDraw
    from studio.models import Pattern,Box
    from studio.template_adaptation import derive_data_patterns,uniform_region
    path=tmp_path/'art.png';im=Image.new('RGB',(720,405),'blue')
    ImageDraw.Draw(im).rectangle((0,0,300,405),fill='pink');im.save(path)
    zones=[Box(x=462,y=120+i*44,w=130,h=20) for i in range(6)]
    title=Box(x=40,y=110,w=280,h=80)
    pattern=Pattern(id='list',source_slide=3,source_layout='list',role='content',purpose='agenda',
        text_zones=[title]+zones,title_zone=title,body_zones=zones,background_image=str(path),
        fields=[{'role':'title','index':0,'shape_id':1,'box':title.model_dump()}]+[
            {'role':'body','index':i,'shape_id':i+2,'box':box.model_dump()} for i,box in enumerate(zones)])
    p=SimpleNamespace(patterns=[pattern],width=720,height=405,margin=40)
    additions=derive_data_patterns(p);assert len(additions)==1
    derived=p.patterns[-1];region=derived.body_zones[0]
    assert region.w>=260 and region.x>=336 and region.x+region.w<=680
    assert uniform_region(path,region,720,405) and pattern.body_zones==zones
    assert len([f for f in derived.fields if f['role']=='body'])==1


def test_explicit_chart_requests_are_output_requirements():
    from studio.editorial_repair import requested_chart_types
    assert requested_chart_types('Покажи долю линейным графиком, а места отдельной столбчатой диаграммой.')==['line','column']
    assert requested_chart_types('Исторический обзор')==[]
    assert requested_chart_types('Use a line chart and a column chart')==['line','column']


def test_existing_plan_can_receive_scoped_quality_feedback(monkeypatch):
    import asyncio
    from studio import narrative
    from studio.editorial import prepare_editorial
    from studio.models import Constraints
    monkeypatch.setattr(narrative,'narrative_storyboard',lambda p:p.analysis.update(slide_budget={'status':'adjusted'}))
    source=parse_content('Пилот ускоряет обработку. Экономия не гарантирована.')
    first=dict(claim('Пилот'),purpose='cover')
    second=dict(claim('Экономия не гарантирована.','f2'),purpose='content')
    initial={'slides':[first,second]}
    package=SimpleNamespace(content=source,original_content=None,template=SimpleNamespace(patterns=[]),
        constraints=Constraints(slides=2,count_mode='exact',summarize=True),analysis={})
    class Gateway:
        settings=SimpleNamespace()
        async def json_request(self,stage,payload,**kwargs):
            assert stage!='editorial'
            if stage=='editorial_repair':
                assert payload['allowed_slide_indices']==[2]
                return {'replacements':[{'slide':2,'content':dict(claim('Экономический эффект не обещан.','f2'),purpose='content')}]}
            return {'claims':[{'claim_id':c['claim_id'],'supported':True,'meaning_preserved':True} for c in payload['claims']],
                'narrative_coherent':True}
    assert asyncio.run(prepare_editorial(package,Gateway(),starting_plan=initial,quality_feedback=[{'slide':2,'message':'Уточнить формулировку'}]))
    final=package.analysis['editorial']['plan']
    assert final['slides'][0]['bullets'][0]['text']=='Пилот'
    assert final['slides'][1]['bullets'][0]['text']=='Экономический эффект не обещан.'
    assert package.analysis['editorial']['repair_history'][0]['unchanged_slides']==[1]


def test_derived_header_moves_only_to_proven_empty_band(tmp_path):
    from PIL import Image,ImageDraw
    from studio.models import Pattern,Box
    from studio.template_adaptation import reposition_derived_titles,uniform_region
    path=tmp_path/'art.png';im=Image.new('RGB',(720,405),'blue')
    ImageDraw.Draw(im).rectangle((0,30,100,350),fill='pink');im.save(path)
    title=Box(x=40,y=120,w=250,h=70);body=Box(x=350,y=140,w=320,h=220)
    pattern=Pattern(id='data-source',source_slide=1,source_layout='list',role='content',
        title_zone=title,body_zones=[body],text_zones=[title,body],background_image=str(path),
        fields=[{'role':'title','index':0,'shape_id':1,'box':title.model_dump()},
            {'role':'body','index':0,'shape_id':2,'box':body.model_dump()}])
    profile=SimpleNamespace(patterns=[pattern],width=720,height=405,margin=40)
    assert reposition_derived_titles(profile)
    assert pattern.title_zone.y+pattern.title_zone.h<body.y
    assert uniform_region(path,pattern.title_zone,720,405)
    assert pattern.fields[0]['box']==pattern.title_zone.model_dump()
    assert pattern.safe_text_zone['field_checks'][0]['status']=='safe'


def test_readable_native_chart_wraps_dates_by_available_width(tmp_path):
    from pptx import Presentation
    from studio.charts import render_chart
    from studio.models import Element,Box
    # Use a bundled exact font, including on clean Linux installations.
    from studio.template import analyze_template
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.util import Pt
    sample=slide.shapes.add_textbox(Pt(20),Pt(20),Pt(500),Pt(80));sample.text='Пример'
    sample.text_frame.paragraphs[0].runs[0].font.name='Play'
    sample.text_frame.paragraphs[0].runs[0].font.size=Pt(16)
    source=tmp_path/'template.pptx';prs.save(source)
    profile=analyze_template(source,tmp_path,allow_download=False)
    element=Element(kind='chart',box=Box(x=200,y=100,w=360,h=200),font=profile.font,size=16,
        color='#000000',fill='#213FFF',chart_type='line',chart_style='readable',
        labels=['Май 1928','Сентябрь 1930','Июль 1932','Ноябрь 1932'],
        series_names=['Голоса, %'],series_values=[[2.6,18.3,37.3,33.1]],values=[2.6,18.3,37.3,33.1])
    chart=render_chart(slide,element,profile)
    assert [c.label for c in chart.plots[0].categories]==['Май\n1928','Сентябрь\n1930','Июль\n1932','Ноябрь\n1932']
    assert 'rot="0"' in chart.category_axis._element.xml


def test_export_accepts_wrapped_categories_but_rejects_changed_data():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Pt
    from studio.export_audit import inspect_content
    from studio.models import SlidePlan
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    data=CategoryChartData();data.categories=['Май\n1928','Сентябрь\n1930'];data.add_series('Места',[12,107])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Pt(20),Pt(20),Pt(400),Pt(250),data)
    table=TableData(id='t1',headers=['Выборы','Места'],rows=[['Май 1928','12'],['Сентябрь 1930','107']])
    package=SimpleNamespace(content=SimpleNamespace(tables=[table]),images=[])
    variant=SimpleNamespace(slides=[SlidePlan(title='Места',fact_ids=[],layout='chart',chart_type='column',table_id='t1')])
    def errors():return [f['code'] for f in inspect_content(prs,variant,package)[1]]
    assert 'chart_values' not in errors()
    table.rows[1][1]='108'
    assert 'chart_values' in errors()
    table.rows[1]=['Сентябрь 1931','107']
    assert 'chart_values' in errors()



def test_stacked_chart_requirement_and_editorial_plan_preserve_all_series():
    from studio.editorial_repair import requested_chart_types
    assert requested_chart_types('Накопительная столбчатая диаграмма')==['column_stacked']
    assert requested_chart_types('Столбчатая диаграмма с накоплением')==['column_stacked']
    assert requested_chart_types('A stacked column chart and a column chart')==['column_stacked','column']
    content=parse_content('| Спринт | Ильф | Вадим | Остальные |\n|---|---|---|---|\n| 1 | 7 | 7 | 4 |\n| 2 | 10 | 9 | 5 |')
    table=content.tables[0];fid=next(f.id for f in content.facts if f.source==table.id)
    raw={'slides':[dict(claim('Распределение задач',fid),source_table_id=table.id,chart_type='column_stacked')]}
    plan=validate_plan(raw,content,(1,1));p=SimpleNamespace(content=content,original_content=content.model_copy(deep=True),analysis={})
    apply_plan(p,plan,{},(1,1))
    assert p.content.tables[0].visualization=='column_stacked'
    assert p.content.tables[0].rows==table.rows and p.content.tables[0].headers==table.headers
