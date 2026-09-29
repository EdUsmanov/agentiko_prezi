"""Readability checks must inspect the same data objects as final generation."""

from copy import deepcopy
from types import SimpleNamespace
import random

import pytest
from studio.config import ROOT
from studio.contents.parsing import parse_content
from studio.contents.editorial import apply_plan, validate_plan, EditorialPlan
from studio.contents.editorial_repair import EditorialPatch, apply_replacements
from studio.models import Constraints, PreparationControl, VariantPlan
from studio.contents.narrative import narrative_storyboard
from studio.templates.parsing import analyze_template
from studio.composition.composer import compose_slide
from studio.checks.audit import repair_scenes


def package_with_data(template, tmp_path, seed):
    profile = analyze_template(template, tmp_path / "profile")
    rng = random.Random(seed)
    profile.width = rng.choice([960, 1080, 1280])
    profile.height = rng.choice([600, 680, 720])
    profile.margin = 38
    profile.body_size = rng.choice([18, 20, 22])
    profile.title_size = 32
    profile.font_sizes = [12, 14, 16, 18, 20, 22, 32]
    font = rng.choice(["Play", "Montserrat"])
    profile.font = font
    profile.fonts = [font]
    profile.font_file = str(ROOT / f"fonts/{font}-Regular.ttf")
    profile.font_roles = {}
    profile.font_assets = []
    profile.patterns = []
    profile.assets = []
    source = parse_content(
        "Delivery became more predictable.\n\n| Period | Planned | Delivered | Completion |\n|---|---|---|---|\n| A | 30 | 18 | 60% |\n| B | 40 | 36 | 90% |\n| C | 50 | 50 | 100% |"
    )
    raw = {
        "slides": [
            {
                "title": "Delivery progress",
                "purpose": "trend",
                "bullets": [
                    {"text": source.facts[0].text, "evidence": [{"fact_id": source.facts[0].id}]}
                ],
                "source_table_id": source.tables[0].id,
                "chart_type": "line",
                "relationship": "time",
            }
        ]
    }
    raw = validate_plan(raw, source, (1, 1), require_cover=False)
    package = SimpleNamespace(
        template=profile,
        content=source,
        original_content=source,
        images=[],
        constraints=Constraints(slides=1, count_mode="exact", summarize=True, include_cover=False),
        analysis={},
        control=PreparationControl(),
    )
    apply_plan(package, raw, {}, (1, 1))
    return package


@pytest.mark.parametrize("seed", [17, 29, 43, 71])
def test_mixed_units_chart_probe_matches_export_on_varied_templates(template, tmp_path, seed):
    package = package_with_data(template, tmp_path, seed)
    cells = deepcopy(package.content.tables[0].rows)
    narrative_storyboard(package)
    assert package.control.slide_budget.status == "adjusted", package.control.slide_budget
    variant = VariantPlan(key="executive", title="Output", slides=package.analysis["storyboard"])
    scene = compose_slide(variant, package, 0)
    repair_scenes([scene], package)
    chart = next(e for e in scene.elements if e.kind == "chart")
    assert chart.chart_type == "line" and chart.size >= 16
    assert chart.series_values == [[30, 40, 50], [18, 36, 50]]
    assert chart.rows[1:] == cells  # Percentage series retained for visible captions.
    assert not any(e.kind == "table" for e in scene.elements)
    from studio.composition.charts import render_chart
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    actual = render_chart(slide, chart, package.template)
    assert len(actual.series) == 2
    assert [list(s.values) for s in actual.series] == chart.series_values
    captions = "\n".join(sh.text for sh in slide.shapes if sh.has_text_frame)
    for value in ("60%", "90%", "100%"):
        assert value in captions


def test_real_data_capacity_failure_remains_blocking(template, tmp_path):
    package = package_with_data(template, tmp_path, 17)
    package.template.width = 400
    package.template.height = 240
    narrative_storyboard(package)
    assert package.control.slide_budget.status == "needs_input"
    assert package.control.slide_budget.fit_issues


def patch_case():
    plan = EditorialPlan.model_validate(
        {
            "slides": [
                {
                    "title": "Original",
                    "bullets": [{"text": "A supported claim.", "evidence": [{"fact_id": "f1"}]}],
                }
            ]
        }
    ).model_dump()
    changed = deepcopy(plan["slides"][0])
    changed["title"] = "Updated"
    return plan, {"slide": 1, "content": changed}


def test_bare_replacement_array_is_only_wrapped_and_fully_validated():
    previous, row = patch_case()
    assert (
        EditorialPatch.model_validate([row]).model_dump()
        == EditorialPatch.model_validate({"replacements": [row]}).model_dump()
    )
    assert apply_replacements(previous, [row], [1])["slides"][0]["title"] == "Updated"
    assert previous["slides"][0]["title"] == "Original"


@pytest.mark.parametrize(
    "case", ["empty", "duplicate", "wrong_index", "missing_content", "extra_field"]
)
def test_array_normalization_cannot_bypass_patch_contract(case):
    previous, row = patch_case()
    raw = [row]
    if case == "empty":
        raw = []
    elif case == "duplicate":
        raw = [row, row]
    elif case == "wrong_index":
        row["slide"] = 2
    elif case == "missing_content":
        del row["content"]
    else:
        row["unexpected"] = "ignored?"
    with pytest.raises(ValueError):
        apply_replacements(previous, raw, [1])


def test_numeric_repair_hints_find_sources_without_silencing_validation():
    from studio.contents.editorial_patch_validation import numeric_evidence_hints

    source = parse_content(
        "The pilot lasted 11 weeks.\nThe team had 7 people.\nSeven unrelated devices lasted 7 days."
    )
    plan = EditorialPlan.model_validate(
        {
            "slides": [
                {
                    "title": "Pilot",
                    "bullets": [
                        {"text": "7 people worked for 11 weeks.", "evidence": [{"fact_id": "f1"}]}
                    ],
                }
            ]
        }
    ).model_dump()
    hint = numeric_evidence_hints(plan, [1], source)[0]
    assert hint["unsupported_numbers"] == ["7"] and hint["current_fact_ids"] == ["f1"]
    assert {row["fact_id"] for row in hint["candidate_evidence"]} == {"f2", "f3"}
    with pytest.raises(ValueError, match="Unsupported number"):
        validate_plan(plan, source, (1, 1))
    # The hint itself never adds citations or treats a digit match as semantic proof.
    assert plan["slides"][0]["bullets"][0]["evidence"] == [{"fact_id": "f1", "quote": None}]
    plan["slides"][0]["bullets"][0]["evidence"].append({"fact_id": "f2"})
    assert numeric_evidence_hints(plan, [1], source) == []
    validate_plan(plan, source, (1, 1))
    plan["slides"][0]["bullets"][0]["text"] = "999 people."
    assert numeric_evidence_hints(plan, [1], source)[0]["candidate_evidence"] == []
    with pytest.raises(ValueError, match="Unsupported number"):
        validate_plan(plan, source, (1, 1))


@pytest.mark.parametrize("seed", [17, 29, 43, 71])
def test_table_and_explanation_reallocate_space_without_losing_data(template, tmp_path, seed):
    from studio.checks.audit import audit_scenes
    from studio.checks.repair_policy import FIT_CODES

    package = package_with_data(template, tmp_path, seed)
    package.template.width = 720
    package.template.height = 405
    package.template.body_size = 16
    package.template.title_size = 20
    table = package.content.tables[0]
    table.visualization = "table"
    table.headers = ["Период", "Запланировано задач", "Выполнено задач", "Выполнение плана"]
    table.rows = [[str(i), "40", "36", "90%"] for i in range(1, 7)]
    expected = deepcopy([table.headers] + table.rows)
    narrative_storyboard(package)
    assert package.control.slide_budget.status == "adjusted"
    variant = VariantPlan(key="executive", title="Output", slides=package.analysis["storyboard"])
    scene = compose_slide(variant, package, 0)
    repair_scenes([scene], package)
    actual = next(e for e in scene.elements if e.kind == "table")
    assert actual.rows == expected
    assert actual.size >= 16
    assert any(e.text == "Delivery became more predictable." for e in scene.elements)
    assert not [f for f in audit_scenes([scene], package) if f.code in FIT_CODES]


def test_dense_table_remains_blocked_without_dropping_rows(template, tmp_path):
    package = package_with_data(template, tmp_path, 17)
    table = package.content.tables[0]
    table.visualization = "table"
    table.rows *= 20
    expected = deepcopy(table.rows)
    narrative_storyboard(package)
    assert package.control.slide_budget.status == "needs_input"
    assert table.rows == expected
