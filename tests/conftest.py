import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from studio.config import Settings
from studio.jobs.store import Store
from studio.pipeline import prepare, load_package


@pytest.fixture
def template(tmp_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        for y, size, text in [(0.5, 32, "OLD PRIVATE TITLE"), (2, 20, "OLD PRIVATE CONTENT")]:
            shape = slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(10), Inches(1))
            r = shape.text_frame.paragraphs[0].add_run()
            r.text = text
            r.font.name = "Play"
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor.from_string("154A67")
    path = tmp_path / "unknown.pptx"
    prs.save(path)
    return path


@pytest.fixture
def content():
    return "# Проект\n" + "\n".join(
        f"## Этап {i}\nПодразделение {i} обрабатывает заявки через единый интерфейс."
        for i in range(1, 13)
    )


@pytest.fixture
def potx(tmp_path, template):
    from zipfile import ZipFile
    from studio.security import PPTX_MAIN, POTX_MAIN

    target = tmp_path / "template.potx"
    with ZipFile(template) as source, ZipFile(target, "w") as output:
        for entry in source.infolist():
            data = source.read(entry)
            if entry.filename == "[Content_Types].xml":
                data = data.replace(PPTX_MAIN.encode(), POTX_MAIN.encode())
            output.writestr(entry, data)
    return target


@pytest.fixture
def prepared(tmp_path, template, content):
    import shutil

    settings = Settings(data_dir=tmp_path / "data")
    store = Store(settings.data_dir)
    job = store.create("preparation", {"template_name": "unknown.pptx"})
    shutil.copyfile(template, store.directory(job["id"]) / "input.pptx")
    prepare(store, job["id"], content, "Команда", "", 5)
    assert store.get(job["id"])["state"] == "ready", store.get(job["id"])
    return settings, store, load_package(store, job["id"])


@pytest.fixture
def preparation_worker(monkeypatch):
    """Exercise HTTP persisted requests while replacing only process execution."""
    import json
    from studio.models import ContentModel, Constraints
    from studio.jobs.runtime import JobRuntime

    def install(operation):
        async def supervise(runtime, job):
            payload = json.loads((runtime.store.directory(job["id"]) / "request.json").read_text())
            operation(
                runtime.store,
                job["id"],
                payload["text"],
                payload["audience"],
                payload["instructions"],
                payload["slides"],
                runtime.settings,
                ContentModel.model_validate(payload["content_model"])
                if payload.get("content_model")
                else None,
                Constraints.model_validate(payload["base_constraints"])
                if payload.get("base_constraints")
                else None,
            )

        monkeypatch.setattr(JobRuntime, "_supervise", supervise)

    return install
