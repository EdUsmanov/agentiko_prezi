from pptx import Presentation
from pptx.util import Pt
from studio.native_template import native_patterns


def test_field_font_inherits_from_its_own_layout():
    prs=Presentation()
    layout=prs.slide_layouts[1]
    layout.placeholders[1].text_frame.paragraphs[0].font.size=Pt(18)
    # The old field-collection loop retained this unrelated final layout.
    prs.slide_layouts[-1].placeholders[1].text_frame.paragraphs[0].font.size=Pt(42)
    slide=prs.slides.add_slide(layout)
    slide.shapes.title.text='Заголовок'
    slide.placeholders[1].text='Текст'
    pattern=next(p for p in native_patterns(prs) if p.source_slide==1)
    assert next(f for f in pattern.fields if f['role']=='body')['font_size']==18
