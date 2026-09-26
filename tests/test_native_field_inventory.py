from pptx import Presentation
from pptx.util import Pt
from studio.native_template import native_patterns


def test_labeled_illustration_is_a_physical_evidence_field_not_body_text():
    from pptx.enum.shapes import MSO_SHAPE
    prs=Presentation();prs.slide_width=Pt(960);prs.slide_height=Pt(540)
    slide=prs.slides.add_slide(prs.slide_layouts[5])
    title=slide.shapes.title
    title.left=Pt(30);title.top=Pt(30);title.width=Pt(800);title.height=Pt(60)
    body=slide.shapes.add_textbox(Pt(30),Pt(140),Pt(300),Pt(270));body.text='Текст'
    visual=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(370),Pt(140),Pt(470),Pt(298))
    visual.text='Иллюстрация'
    pattern=next(p for p in native_patterns(prs) if p.source_slide==1)
    assert len(pattern.body_zones)==1 and len(pattern.image_zones)==1
    field=next(f for f in pattern.fields if f['role']=='image')
    assert field['shape_id']==visual.shape_id and field['evidence_placeholder']
    assert not any(f['shape_id']==visual.shape_id and f['role']!='image' for f in pattern.fields)


def test_four_authored_cards_are_not_misread_as_title_only():
    prs=Presentation();prs.slide_width=Pt(720);prs.slide_height=Pt(405)
    slide=prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.left=Pt(20);slide.shapes.title.top=Pt(20)
    slide.shapes.title.width=Pt(650);slide.shapes.title.height=Pt(45)
    for i in range(4):
        heading=slide.shapes.add_textbox(Pt(20+i*172),Pt(95),Pt(150),Pt(20));heading.text='Heading'
        body=slide.shapes.add_textbox(Pt(20+i*172),Pt(125),Pt(150),Pt(140));body.text='Description'
    pattern=next(p for p in native_patterns(prs) if p.source_slide==1)
    assert len(pattern.body_zones)==4
    assert len([f for f in pattern.fields if f['role']=='body'])==4
    assert all(pattern.heading_zones)


def test_short_cover_subtitle_does_not_remove_the_entire_layout():
    prs=Presentation();prs.slide_width=Pt(720);prs.slide_height=Pt(405)
    slide=prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.left=Pt(40);slide.shapes.title.top=Pt(50)
    slide.shapes.title.width=Pt(550);slide.shapes.title.height=Pt(60)
    subtitle=slide.placeholders[1];subtitle.left=Pt(40);subtitle.top=Pt(140)
    subtitle.width=Pt(550);subtitle.height=Pt(23)
    pattern=next(p for p in native_patterns(prs) if p.source_slide==1)
    assert pattern.role=='cover' and len(pattern.body_zones)==1


def test_scaled_group_fields_use_absolute_slide_coordinates():
    prs=Presentation();prs.slide_width=Pt(720);prs.slide_height=Pt(405)
    slide=prs.slides.add_slide(prs.slide_layouts[5])
    title=slide.shapes.title;title.left=Pt(20);title.top=Pt(20);title.width=Pt(650);title.height=Pt(45)
    group=slide.shapes.add_group_shape()
    for i in range(3):
        body=group.shapes.add_textbox(Pt(20+i*200),Pt(120),Pt(170),Pt(130));body.text='Описание шага'
    group.left=Pt(40);group.top=Pt(140);group.width=Pt(600)
    before=slide._element.xml
    pattern=next(p for p in native_patterns(prs) if p.source_slide==1)
    assert len(pattern.body_zones)==3
    assert all(z.x>=40 and z.y>=140 for z in pattern.body_zones)
    assert all(f['shape_id'] in {s.shape_id for s in group.shapes} for f in pattern.fields if f['role']=='body')
    assert slide._element.xml==before


def test_verified_graphic_path_reorders_all_field_bindings():
    from studio.contracts import apply_meanings
    from types import SimpleNamespace
    from studio.models import Pattern,Box
    zones=[Box(x=i*100,y=80,w=80,h=80) for i in range(3)]
    p=Pattern(id='route',source_slide=1,source_layout='route',text_zones=zones,body_zones=zones,
        role='columns',fields=[{'role':'body','index':i,'shape_id':i+10} for i in range(3)])
    apply_meanings(SimpleNamespace(patterns=[p]),{'patterns':[{'pattern_id':'route','purpose':'process',
        'body_order':[12,10,11],'graphic_flow_confirmed':True}]})
    assert p.graphic_order_verified and [z.x for z in p.body_zones]==[200,0,100]
    assert [f['shape_id'] for f in sorted(p.fields,key=lambda f:f['index'])]==[12,10,11]


def test_vector_paths_survive_cleaning_without_sample_labels(tmp_path):
    from studio.portable_templates import extract_backgrounds
    from studio.powerpoint import open_presentation
    from studio.models import TemplateProfile
    from types import SimpleNamespace
    from pptx.enum.shapes import MSO_SHAPE
    prs=Presentation();prs.slide_width=Pt(720);prs.slide_height=Pt(405)
    slide=prs.slides.add_slide(prs.slide_layouts[5]);slide.shapes.title.text='Template title'
    slide.shapes.title.left=Pt(20);slide.shapes.title.top=Pt(20);slide.shapes.title.width=Pt(650);slide.shapes.title.height=Pt(40)
    group=slide.shapes.add_group_shape()
    for i in range(2):
        shape=group.shapes.add_shape(MSO_SHAPE.CHEVRON,Pt(30+i*310),Pt(140),Pt(270),Pt(140))
        shape.text='Old confidential sample'
    source=tmp_path/'source.pptx';prs.save(source)
    patterns=[p for p in native_patterns(prs) if p.source_slide==1]
    for p in patterns:
        p.purpose='process';p.graphic_order_verified=True;p.graphic_kind='sequence'
        p.graphic_shape_ids=[s.shape_id for s in group.shapes]
    profile=SimpleNamespace(patterns=patterns,background_source='')
    extract_backgrounds(profile,source,tmp_path)
    cleaned=open_presentation(profile.background_source)
    assert not any((node.text or '').strip() for node in cleaned.slides[0]._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'))
    assert len(cleaned.slides[0]._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom'))>=2
