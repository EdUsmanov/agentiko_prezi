from types import SimpleNamespace

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from studio.templates.fonts import symbol_font

from studio.composition.image_composer import attach_template_resources
from studio.models import Box, Element, SlideScene, TemplateResource
from studio.composition.native_surface import place_template_resource
from studio.composition.render import render_html, render_pdf, render_pptx
from studio.security import digest


def test_native_resource_copies_editable_shape_with_fresh_id():
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[6])
    original = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(10), Pt(10), Pt(30), Pt(30))
    target = Presentation()
    output = target.slides.add_slide(target.slide_layouts[6])
    output.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(1), Pt(1), Pt(5), Pt(5))
    resource = TemplateResource(
        id="icon-a",
        kind="icon",
        source_slide=1,
        shape_ids=[original.shape_id],
        box=Box(x=10, y=10, w=30, h=30),
        preview_path="unused.png",
    )
    element = Element(kind="image", box=Box(x=100, y=100, w=45, h=45), resource_id="icon-a")
    place_template_resource(output, element, SimpleNamespace(resources=[resource]), source)
    assert len(output.shapes) == 2
    copied = output.shapes[-1]
    assert copied.shape_id != output.shapes[0].shape_id
    assert abs(copied.left / Pt(1) - 100) < 0.1
    assert abs(copied.width / Pt(1) - 45) < 0.1


def test_device_frame_wraps_uploaded_image_without_replacing_its_source(tmp_path):
    preview = tmp_path / "frame.png"
    Image.new("RGBA", (100, 180), (0, 0, 0, 0)).save(preview)
    frame = TemplateResource(
        id="frame",
        kind="device_frame",
        source_slide=1,
        shape_ids=[5],
        box=Box(x=0, y=0, w=100, h=180),
        screen_box=Box(x=10, y=20, w=80, h=140),
        preview_path=str(preview),
    )
    image = Element(
        kind="image",
        box=Box(x=100, y=100, w=150, h=300),
        image_id="uploaded",
        image_path="original.png",
        role="user_image",
    )
    second = image.model_copy(deep=True)
    second.box.x += 200
    scene = SlideScene(
        title="Слайд",
        background="#FFFFFF",
        elements=[image, second],
        source_ids=[],
        layout="split",
    )
    package = SimpleNamespace(
        template=SimpleNamespace(resources=[frame]),
        images=[SimpleNamespace(id="uploaded", presentation="device", width=600, height=900)],
    )
    attach_template_resources(scene, package, SimpleNamespace(title="Слайд", fact_ids=[]))
    assert image.image_path == "original.png"
    assert image.box.w < scene.elements[-1].box.w
    assert scene.elements[-1].resource_id == "frame"
    assert sum(e.resource_id == "frame" for e in scene.elements) == 2


def test_device_export_keeps_picture_and_editable_frame_separate(tmp_path):
    source = Presentation()
    original = source.slides.add_slide(source.slide_layouts[6])
    frame_shape = original.shapes.add_group_shape()
    for x, y, w, h in [(0, 0, 100, 6), (0, 174, 100, 6), (0, 0, 6, 180), (94, 0, 6, 180)]:
        frame_shape.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    source_path = tmp_path / "source.pptx"
    source.save(source_path)
    screenshot = tmp_path / "screenshot.png"
    Image.new("RGB", (600, 900), "#5279aa").save(screenshot)
    cutout = tmp_path / "frame.png"
    border = Image.new("RGBA", (100, 180), (0, 0, 0, 0))
    for rectangle in ((0, 0, 100, 6), (0, 174, 100, 180), (0, 0, 6, 180), (94, 0, 100, 180)):
        from PIL import ImageDraw

        ImageDraw.Draw(border).rectangle(rectangle, fill="#202020")
    border.save(cutout)
    resource = TemplateResource(
        id="frame",
        kind="device_frame",
        source_slide=1,
        shape_ids=[frame_shape.shape_id],
        box=Box(x=0, y=0, w=100, h=180),
        screen_box=Box(x=6, y=6, w=88, h=168),
        preview_path=str(cutout),
    )
    scene = SlideScene(
        title="Screenshot",
        background="#FFFFFF",
        layout="split",
        source_ids=[],
        elements=[
            Element(
                kind="image",
                box=Box(x=120, y=50, w=88, h=168),
                image_id="uploaded",
                image_path=str(screenshot),
                role="user_image",
            ),
            Element(
                kind="image",
                box=Box(x=114, y=44, w=100, h=180),
                resource_id="frame",
                image_path=str(cutout),
                role="template_resource",
                field_style={"paired_image_id": "uploaded"},
            ),
        ],
    )
    profile = SimpleNamespace(
        background_source="",
        resource_source=str(source_path),
        resources=[resource],
        sha256=digest(source_path.read_bytes()),
        assets=[],
        patterns=[],
        layout_index=6,
        font_file=symbol_font(),
        font_assets=[],
        font_roles={},
        font="Montserrat",
        width=300,
        height=270,
        foreground="#202020",
    )
    deck_path = tmp_path / "result.pptx"
    render_pptx([scene], profile, source_path, deck_path, verify_text=False)
    output = Presentation(deck_path)
    assert len(output.slides[0].shapes) == 2
    assert output.slides[0].shapes[0].shape_type.name == "PICTURE"
    assert output.slides[0].shapes[1].shape_type.name == "GROUP"
    html_path, pdf_path = tmp_path / "result.html", tmp_path / "result.pdf"
    render_html([scene], profile, html_path)
    render_pdf([scene], profile, pdf_path)
    assert html_path.read_text().count("data:image/png;base64,") == 2
    assert pdf_path.stat().st_size > 0


def test_brief_notes_use_approved_content_boundary(prepared):
    from studio.composition.composer import compose_variant
    from studio.contents.planner import extractive_plans

    _, _, package = prepared
    package.original_content.facts[0].text += " https://removed.example/private"
    package.input_mode = "brief"
    scenes = compose_variant(extractive_plans(package).variants[0], package)
    assert all("removed.example" not in scene.notes for scene in scenes)


def test_icon_placement_preserves_source_aspect_and_requires_free_space(tmp_path):
    preview = tmp_path / "icon.png"
    Image.new("RGBA", (80, 40), (0, 0, 0, 0)).save(preview)
    icon = TemplateResource(
        id="cloud",
        kind="icon",
        source_slide=1,
        shape_ids=[1],
        box=Box(x=10, y=10, w=80, h=40),
        preview_path=str(preview),
        tags=["облако"],
        confidence=0.9,
    )
    pattern = SimpleNamespace(id="body", body_zones=[Box(x=10, y=50, w=400, h=250)])
    package = SimpleNamespace(
        template=SimpleNamespace(resources=[icon], patterns=[pattern]),
        images=[],
        content=SimpleNamespace(facts=[]),
    )
    scene = SlideScene(
        title="Облако",
        pattern_id="body",
        background="#FFFFFF",
        elements=[],
        source_ids=[],
        layout="split",
    )
    attach_template_resources(scene, package, SimpleNamespace(title="Облако", fact_ids=[]))
    assert len(scene.elements) == 1
    assert scene.elements[0].box.w / scene.elements[0].box.h == 2
    occupied = scene.model_copy(deep=True)
    occupied.elements = [Element(kind="text", box=pattern.body_zones[0], text="Материал")]
    attach_template_resources(occupied, package, SimpleNamespace(title="Облако", fact_ids=[]))
    assert len(occupied.elements) == 1
    assert not occupied.elements[0].resource_id
