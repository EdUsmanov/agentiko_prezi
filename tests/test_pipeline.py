import asyncio
import json
import time
from zipfile import ZipFile
import pytest
from pptx import Presentation
from pypdf import PdfReader
from studio.pipeline import generate,load_package
from studio.planner import extractive_plans,validate_plans
from studio.content import parse_content,numeric_column

def test_three_variants_exports_and_coverage(prepared):
    settings,store,package=prepared
    job=store.create("generation",{"package_id":package.id,"deadline_at":time.time()+300})
    asyncio.run(generate(store,job["id"],settings))
    result=store.get(job["id"])
    assert result["state"]=="completed",result
    assert result["elapsed_seconds"]<300
    assert len(result["variants"])==3
    root=store.directory(job["id"])
    plans=json.loads((root/"plans.json").read_text())
    expected={f.id for f in package.content.facts}
    signatures=[]
    for variant in plans["variants"]:
        assert len(variant["slides"])==5
        assert {i for s in variant["slides"] for i in s["fact_ids"]}==expected
        signatures.append([(s["fact_ids"],s["layout"]) for s in variant["slides"]])
        folder=root/variant["key"]
        deck=Presentation(folder/"deck.pptx")
        assert len(deck.slides)==5
        assert all(any(sh.has_text_frame for sh in slide.shapes) for slide in deck.slides)
        with ZipFile(folder/"deck.pptx") as z:
            assert not any(b"OLD PRIVATE" in z.read(n) for n in z.namelist() if n.endswith(".xml"))
        html=(folder/"deck.html").read_text()
        assert "<script" not in html
        assert "<div" in html
        pdf=PdfReader(folder/"deck.pdf")
        assert len(pdf.pages)==5
        assert any("Подразделение" in p.extract_text() for p in pdf.pages)
    assert len({json.dumps(s) for s in signatures})==3
    with ZipFile(root/"presentations.zip") as z:
        assert sum(n.endswith(".pptx") for n in z.namelist())==3

def test_package_is_immutable(prepared):
    _,store,package=prepared
    path=store.directory(package.id)/"package.json"
    path.write_text(path.read_text()+" ")
    with pytest.raises(ValueError,match="изменён"):
        load_package(store,package.id)

def test_unknown_fact_rejected(prepared):
    _,_,package=prepared
    plans=extractive_plans(package)
    plans.variants[0].slides[0].fact_ids=["fake"]
    with pytest.raises(ValueError):validate_plans(plans,package)

def test_new_number_rejected(prepared):
    _,_,package=prepared
    plans=extractive_plans(package)
    plans.variants[0].slides[0].title="Эффект составил 999999 процентов"
    with pytest.raises(ValueError):validate_plans(plans,package)

def test_comparable_numbers_only():
    content=parse_content("# Данные\n| Канал | Заявки |\n|---|---|\n| A | 20 |\n| B | 30 |")
    assert numeric_column(content.tables[0])[1]==[20,30]
    content.tables[0].rows[1][1]="30%"
    assert numeric_column(content.tables[0]) is None

def test_generation_reservation_and_no_late_success(prepared):
    _,store,package=prepared
    job=store.create("generation",{"package_id":package.id})
    with pytest.raises(ValueError):store.create("generation")
    store.update(job["id"],"timed_out")
    store.update(job["id"],"completed")
    assert store.get(job["id"])["state"]=="timed_out"

def test_no_default_padding(prepared):
    _,_,package=prepared
    package.content.facts=package.content.facts[:1]
    plans=extractive_plans(package)
    assert all(len(v.slides)==1 for v in plans.variants)

def test_one_slide_table_still_has_three_variants(prepared):
    from studio.composer import compose_variant
    _,_,package=prepared
    package.content=parse_content('# Данные\n| Канал | Объём |\n|---|---|\n| A | 20 |\n| B | 30 |')
    package.constraints.slides=1
    plans=validate_plans(extractive_plans(package),package)
    scenes=[compose_variant(v,package) for v in plans.variants]
    assert len({json.dumps([s.model_dump() for s in scene],sort_keys=True) for scene in scenes})==3

def test_schema_constrains_count_and_source_ids(prepared):
    from studio.planner import planning_schema
    _,_,package=prepared
    schema=planning_schema(package)
    slides=schema['$defs']['VariantPlan']['properties']['slides']
    assert slides['minItems']==slides['maxItems']==5
    facts=schema['$defs']['SlidePlan']['properties']['fact_ids']
    assert facts['items']['enum']==[f.id for f in package.content.facts]
    assert facts['minItems']==1
    assert schema['$defs']['SlidePlan']['properties']['table_id']['enum']==[None]

def test_model_timeout_is_needs_review_not_green_success(prepared,monkeypatch):
    from dataclasses import replace
    from studio import pipeline
    settings,store,package=prepared
    settings=replace(settings,mode='api')
    class FailingGateway:
        def __init__(self,settings):
            self.settings=settings;self.usage=[];self.calls=[]
        async def json_request(self,*args,**kwargs):
            raise TimeoutError()
    monkeypatch.setattr(pipeline,'ModelGateway',FailingGateway)
    job=store.create('generation',{'package_id':package.id,'deadline_at':time.time()+300})
    asyncio.run(generate(store,job['id'],settings))
    done=store.get(job['id'])
    assert done['state']=='needs_review'
    assert done['model_degraded'] is True
    assert done['planning_source']=='extractive'
    assert done['errors']==0
