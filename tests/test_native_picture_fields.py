"""Filled picture placeholders retain image geometry without text-style inspection."""

from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.util import Pt

from studio.templates.native_template import inherited_title_size, native_patterns


def test_filled_picture_placeholder_is_not_a_text_field():
    prs = Presentation()
    layout = prs.slide_layouts[8]  # Picture with caption.
    slide = prs.slides.add_slide(layout)
    slide.shapes.title._element.nvSpPr.nvPr.ph.set("type", "ctrTitle")
    slide.shapes.title.text = "Project image"
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)
    raw = BytesIO()
    Image.new("RGB", (120, 80), "blue").save(raw, format="PNG")
    raw.seek(0)
    picture = slide.placeholders[1].insert_picture(raw)

    assert inherited_title_size(picture, layout) == 0
    assert inherited_title_size(slide.shapes.title, layout) == 32
    patterns = native_patterns(prs)
    source = next(p for p in patterns if p.source_slide == 1)
    image = next(field for field in source.fields if field["role"] == "image")
    assert image["shape_id"] == picture.shape_id
    assert image["font_size"] == 0
    assert len(source.image_zones) == 1
    assert not picture.has_text_frame
