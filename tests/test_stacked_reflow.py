from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.util import Pt

from studio.checks.audit import audit_scenes, repair_scenes
from studio.checks.export_audit import geometry, slide_text
from studio.checks.repair_policy import FIT_CODES
from studio.composition.charts import make_chart, render_chart
from studio.composition.stacked_adaptation import adapt_stacked_scene
from studio.composition.stacked_chart import stacked_layout
from studio.models import (
    Box,
    Constraints,
    ContentModel,
    Element,
    Fact,
    PreparationControl,
    SlidePlan,
    SlideScene,
    TableData,
)
from studio.templates.parsing import analyze_template


@pytest.fixture
def case(template, tmp_path):
    profile = analyze_template(template, tmp_path / "profile")
    profile.width, profile.height, profile.margin = 720, 405, 40
    table = TableData(
        id="hours",
        visualization="column_stacked",
        headers=["Куда ушли командные часы", "Спринт 1", "Спринт 6"],
        rows=[
            ["Запланированная работа", "160", "240"],
            ["Переделки", "80", "32"],
            ["Совещания", "48", "32"],
            ["Аварийные исправления", "32", "16"],
            ["Всего", "320", "320"],
        ],
    )
    claims = [
        "План. Доля запланированной работы выросла с 50% до 75%.",
        "Экономия. На переделки и аварийные исправления уходит на 64 часа меньше за спринт.",
        "Итог. Те же четверо, те же часы, заметно меньше хаоса.",
    ]
    facts = [Fact(id="data", text="Исходные измерения", source=table.id)] + [
        Fact(id=f"claim-{i}", text=text, source="user_text") for i, text in enumerate(claims)
    ]
    package = SimpleNamespace(
        template=profile,
        content=ContentModel(title="Результат", facts=facts, tables=[table]),
        constraints=Constraints(slides=1),
        analysis={},
        control=PreparationControl(),
        images=[],
    )
    chart = make_chart(
        table,
        SlidePlan(title="Результат", fact_ids=["data"], chart_type="column_stacked"),
        Box(x=40, y=93, w=397, h=272),
        profile,
        profile.foreground,
        ["data"],
    )
    scene = SlideScene(
        title="Результат",
        background=profile.background,
        source_ids=[f.id for f in facts],
        layout="chart",
        strategy="token_composition",
        elements=[
            Element(
                kind="text",
                role="title",
                text="Результат",
                size=24,
                font=profile.font,
                color=profile.foreground,
                box=Box(x=40, y=40, w=640, h=31),
            ),
            chart,
            *[
                Element(
                    kind="text",
                    role="body",
                    text=text,
                    size=16.9,
                    bullet=True,
                    source_ids=[f"claim-{i}"],
                    font=profile.font,
                    color=profile.foreground,
                    box=Box(x=455, y=93 + i * 90, w=225, h=85),
                )
                for i, text in enumerate(claims)
            ],
            Element(
                kind="text",
                role="footer",
                text="1",
                size=10,
                font=profile.font,
                color=profile.foreground,
                box=Box(x=645, y=381, w=35, h=15),
            ),
        ],
    )
    return package, scene


def test_dense_stacked_chart_reflows_without_table_downgrade(case, tmp_path):
    package, scene = case
    original = scene.model_dump()
    assert not stacked_layout(scene.elements[1], package.template)["fits"]
    result = adapt_stacked_scene(scene, package)
    chart = next(e for e in result.elements if e.kind == "chart")
    assert stacked_layout(chart, package.template)["fits"]
    assert chart.box.w == 640
    assert chart.series_values == scene.elements[1].series_values
    assert chart.rows == scene.elements[1].rows
    assert [e.text for e in result.elements] == [e.text for e in scene.elements]
    assert [e.source_ids for e in result.elements] == [e.source_ids for e in scene.elements]
    assert all(
        e.box.y >= chart.box.y + chart.box.h
        for e in result.elements
        if e.kind == "text" and e.role == "body"
    )
    assert not repair_scenes([result], package)
    assert not [f for f in audit_scenes([result], package) if f.code in FIT_CODES]
    assert scene.model_dump() == original
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(720), Pt(405)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rendered = render_chart(slide, chart, package.template)
    assert [list(s.values) for s in rendered.series] == chart.series_values
    path = tmp_path / "chart.pptx"
    prs.save(path)
    actual = Presentation(path)
    assert not geometry(actual, package.template)[0]
    text = slide_text(actual.slides[0])
    assert "Спринт 1 — Всего: 320" in text and "Спринт 6 — Всего: 320" in text


def test_reflow_does_not_accept_unreadable_prose_or_rewrite_native_fields(case):
    package, scene = case
    native = scene.model_copy(
        update={"pattern_id": "authored", "strategy": "native_template"}, deep=True
    )
    assert adapt_stacked_scene(native, package) is native
    scene.elements[2].text *= 30
    before = scene.model_dump()
    assert adapt_stacked_scene(scene, package) is scene
    assert scene.model_dump() == before


def test_compact_legend_keeps_entries_inside_the_chart(case):
    package, scene = case
    chart = scene.elements[1].model_copy(deep=True)
    chart.box.w, chart.box.h = 640, 210
    layout = stacked_layout(chart, package.template)
    assert layout["fits"]
    assert layout["label_size"] >= 12 and layout["legend_size"] >= 12
    assert len(layout["legend"]) == len(chart.series_names)
    for cell in layout["legend_cells"]:
        assert 0 <= cell["x"] < cell["x"] + cell["w"] <= chart.box.w
        assert 0 <= cell["y"] < cell["y"] + cell["h"] <= layout["legend_h"]
