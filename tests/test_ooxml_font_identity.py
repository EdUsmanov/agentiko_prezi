from types import SimpleNamespace
import pytest
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Pt
from studio.config import ROOT
from studio.font_identity import ooxml_face, apply_ooxml_font
from studio.fonts import resolve_font
from studio.models import Box, Element
from studio.render import set_text
from studio.export_audit import _face


def test_bold_face_is_family_plus_flag_not_a_fictitious_family():
    path = str(ROOT / "fonts/Montserrat-Bold.ttf")
    assert ooxml_face(path) == ("Montserrat", True, False)
    prs = Presentation()
    font = (
        prs.slides.add_slide(prs.slide_layouts[6])
        .shapes.add_textbox(0, 0, 1, 1)
        .text_frame.paragraphs[0]
        .font
    )
    apply_ooxml_font(font, path)
    assert font.name == "Montserrat" and font.bold and not font.italic


def test_legacy_family_preserves_light_and_display_faces(tmp_path):
    path = tmp_path / "light.ttf"
    with TTFont(ROOT / "fonts/Play-Regular.ttf") as font:
        platforms = {(n.platformID, n.platEncID, n.langID) for n in font["name"].names}
        for platform, encoding, language in platforms:
            for key, value in ((1, "Test Light"), (2, "Regular"), (16, "Test"), (17, "Light")):
                font["name"].setName(value, key, platform, encoding, language)
        font.save(path)
    assert ooxml_face(path) == ("Test Light", False, False)


def test_full_face_request_exports_canonical_runs_and_uses_same_audit_metrics():
    path = str(ROOT / "fonts/Montserrat-Bold.ttf")
    profile = SimpleNamespace(
        font_roles={"title": "bold"},
        font_assets=[{"id": "bold", "requested": "Montserrat Bold", "path": path}],
        font="Montserrat",
        font_file=str(ROOT / "fonts/Montserrat-Regular.ttf"),
    )
    prs = Presentation()
    shape = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Pt(500), Pt(100))
    element = Element(
        kind="text",
        box=Box(x=0, y=0, w=500, h=100),
        text="Синтетический пилот",
        role="title",
        size=30,
        color="#000000",
    )
    set_text(shape.text_frame, element.text, element, profile)
    run = shape.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Montserrat" and run.font.bold
    assert _face(profile, run.font.name, bold=run.font.bold) == path


def test_calibri_bold_native_render_does_not_fall_back(tmp_path):
    path = resolve_font("Calibri Bold")
    if not path:
        pytest.skip("optional local Calibri integration test")
    from studio.office import executable, to_pdf

    if not executable():
        pytest.skip("LibreOffice unavailable")
    from pypdf import PdfReader

    profile = SimpleNamespace(
        font_roles={"title": "bold"},
        font_assets=[{"id": "bold", "requested": "Calibri Bold", "path": path}],
        font="Calibri Bold",
        font_file=path,
    )
    prs = Presentation()
    shape = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Pt(20), Pt(20), Pt(470), Pt(184)
    )
    element = Element(
        kind="text",
        box=Box(x=20, y=20, w=470, h=184),
        text="Синтетический пилот",
        role="title",
        size=66,
        bold=True,
        color="#000000",
    )
    set_text(shape.text_frame, element.text, element, profile)
    prs.save(tmp_path / "deck.pptx")
    to_pdf(
        tmp_path / "deck.pptx",
        tmp_path,
        font_files=[path, str(ROOT / "fonts/Montserrat-Regular.ttf")],
    )
    # Use the project's declared PDF dependency, including on clean installs.
    spans = []

    def collect(text, _cm, _tm, font, _size):
        face = str((font or {}).get("/BaseFont", "")).lstrip("/").split("+")[-1]
        spans.extend((line.strip(), face) for line in text.splitlines() if line.strip())

    PdfReader(tmp_path / "deck.pdf").pages[0].extract_text(visitor_text=collect)
    assert [text for text, _ in spans] == ["Синтетический", "пилот"]
    assert all(face == "Calibri-Bold" for _, face in spans)
