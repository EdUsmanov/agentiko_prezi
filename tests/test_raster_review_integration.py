import asyncio
import json
from io import BytesIO
from types import SimpleNamespace

from PIL import Image, ImageDraw
from pptx import Presentation

from studio.templates.portable_templates import extract_backgrounds
from studio.checks.raster_review import review_template_rasters


def test_vl_raster_regions_reach_background_extractor(tmp_path):
    image = Image.new("RGB", (800, 450), "#F3F4F5")
    draw = ImageDraw.Draw(image)
    draw.rectangle((260, 150, 540, 300), outline="#E7E8E9", width=3)
    payload = BytesIO()
    image.save(payload, format="PNG")
    picture = tmp_path / "background.png"
    picture.write_bytes(payload.getvalue())

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(
        str(picture), 0, 0, presentation.slide_width, presentation.slide_height
    )
    source = tmp_path / "input.pptx"
    presentation.save(source)

    class Gateway:
        settings = SimpleNamespace(
            model_id="vision-test", base_url="https://provider.test/v1", data_dir=tmp_path
        )
        calls = 0

        async def json_request(self, prompt_name, payload, **kwargs):
            assert prompt_name == "background_raster"
            assert kwargs["images"][0].startswith(b"\x89PNG")
            self.calls += 1
            return {
                "regions": [{"x1": 300, "y1": 300, "x2": 700, "y2": 700, "reason": "sample card"}]
                if self.calls == 1
                else []
            }

    gateway = Gateway()
    review = asyncio.run(review_template_rasters(source, gateway, tmp_path))
    assert gateway.calls == 2
    assert len(review["assets"]) == 1
    assert len(review["regions"]) == 1

    profile = SimpleNamespace(patterns=[], background_source="")
    model = extract_backgrounds(profile, source, tmp_path)
    assert model["rasterReview"]["status"] == "completed"
    assert model["rasterRegions"] == review["regions"]
    assert json.loads((tmp_path / "background-model.json").read_text())["rasterRegions"]
    assert (tmp_path / "template-layers/background-source.pptx").is_file()
