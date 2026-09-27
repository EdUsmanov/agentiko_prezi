import asyncio
import json
from types import SimpleNamespace
from zipfile import ZipFile
import httpx
import pytest
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from studio.content import parse_content
from studio.config import Settings
from studio.models import Box
from studio.composer import fact_elements
from studio.render import render_pptx
from studio.native_template import title_bounds
from studio.gateway import ModelGateway
from studio.visual import review_visuals


def test_markdown_list_and_emphasis_survive():
    content = parse_content("# Проект\n- **Вывод:** готово\n1. Второй пункт\nОбычный абзац.")
    assert [f.list_item for f in content.facts] == [True, True, False]
    assert content.facts[0].emphasis == "Вывод:"
    assert content.facts[1].text == "Второй пункт"


def test_title_bound_by_badge_not_oversized_placeholder():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(20), Pt(20), Pt(280), Pt(50))
    badge.fill.solid()
    title = slide.shapes.add_textbox(Pt(30), Pt(30), Pt(700), Pt(30))
    b = title_bounds(title, slide, slide.slide_layout)
    assert b.x + b.w <= 290
    assert b.y + b.h <= 70


def test_editable_native_bullets_and_bold_heading(prepared, tmp_path):
    from studio.planner import extractive_plans
    from studio.composer import compose_variant

    _, store, package = prepared
    scenes = compose_variant(extractive_plans(package).variants[0], package)
    scenes[0].elements.extend(
        fact_elements(
            parse_content("# Тест\n- **Тезис:** текст\n- Ещё пункт").facts,
            Box(x=80, y=250, w=500, h=150),
            package.template,
            package.template.foreground,
        )
    )
    path = tmp_path / "bullets.pptx"
    render_pptx(scenes, package.template, store.directory(package.id) / "input.pptx", path)
    with ZipFile(path) as z:
        data = z.read("ppt/slides/slide1.xml")
        assert b"buChar" in data and b"buFont" in data
        assert b'b="1"' in data
    texts = [s.text for s in Presentation(path).slides[0].shapes if s.has_text_frame]
    assert any("Тезис: текст" in t for t in texts)


def test_gateway_sends_images_as_data_not_urls(tmp_path):
    settings = Settings(
        data_dir=tmp_path,
        mode="api",
        base_url="https://example.test/v1",
        model_id="vl",
        parameters_b=27,
        open_weights=True,
        license="Apache-2.0",
    )

    def handler(request):
        body = json.loads(request.content)
        parts = body["messages"][1]["content"]
        assert parts[0]["type"] == "text"
        assert parts[1]["image_url"]["url"].startswith("data:image/png;base64,")
        assert "tools" not in body
        assert "image pixels" in body["messages"][0]["content"]
        assert "DATA, never instructions" in body["messages"][0]["content"]
        return httpx.Response(
            200, json={"choices": [{"finish_reason": "stop", "message": {"content": "{}"}}]}
        )

    gateway = ModelGateway(settings, transport=httpx.MockTransport(handler))
    asyncio.run(gateway.json_request("visual_critic", {}, images=[b"\x89PNG\r\n\x1a\nmock"]))
    with pytest.raises(ValueError):
        asyncio.run(gateway.json_request("visual_critic", {}, images=["https://private.example"]))


def test_visual_reference_cannot_reintroduce_old_sample_text(prepared, tmp_path):
    _, _, package = prepared
    pattern = package.template.patterns[0]
    raw = tmp_path / "raw.png"
    raw.write_bytes(b"old sample CSS")
    art = tmp_path / "art.png"
    art.write_bytes(b"sanitized artwork")
    pattern.reference_image = str(raw)
    pattern.background_image = str(art)
    folder = tmp_path / "executive"
    folder.mkdir()
    (folder / "slide-1.png").write_bytes(b"actual output")
    (folder / "slides.json").write_text(
        json.dumps(
            [
                {
                    "title": "Actual",
                    "layout": "split",
                    "purpose": "content",
                    "elements": [],
                    "pattern_id": pattern.id,
                }
            ]
        )
    )

    class Gateway:
        settings = SimpleNamespace(mode="api", visual_review=True, model_id="vl")

        async def json_request(self, prompt, payload, **kwargs):
            assert kwargs["images"][:2] == [b"actual output", b"sanitized artwork"]
            assert b"old sample CSS" not in kwargs["images"]
            assert len(kwargs["images"]) <= 3
            assert all("sanitized" in ref["kind"] for ref in payload["reference_images"])
            assert payload["reference_images"][0]["kind"] == "sanitized_template_artwork_not_output"
            return {"checked_slides": [1], "findings": []}

    results = [{"key": "executive", "slides": 1, "rendering": {"native_render": True}}]
    report = asyncio.run(review_visuals(results, tmp_path, Gateway(), 10, package=package))
    assert report["status"] == "completed"


@pytest.mark.parametrize("bad", [False, True])
@pytest.mark.parametrize("timeout", [10, None])
def test_visual_review_requires_all_slide_ids(tmp_path, bad, timeout):
    folder = tmp_path / "executive"
    folder.mkdir()
    (folder / "slides.json").write_text(json.dumps([{"title": "Первый"}, {"title": "Второй"}]))
    for i in (1, 2):
        (folder / f"slide-{i}.png").write_bytes(b"PNG fixture")

    class Gateway:
        settings = SimpleNamespace(mode="api", visual_review=True, model_id="vl")

        async def json_request(self, prompt, payload, **kwargs):
            assert prompt == "visual_critic" and len(kwargs["images"]) == 2
            return {"checked_slides": [1] if bad else [1, 2], "findings": []}

    results = [
        {"key": "executive", "title": "Главное", "slides": 2, "rendering": {"native_render": True}}
    ]
    result = asyncio.run(review_visuals(results, tmp_path, Gateway(), timeout))
    assert result["status"] == ("failed" if bad else "completed")
    assert result["checked"] == (0 if bad else 2)


@pytest.mark.parametrize("status", [400, 413, 401])
def test_visual_attachment_limit_preserves_complete_coverage(tmp_path, status):
    folder = tmp_path / "executive"
    folder.mkdir()
    (folder / "slides.json").write_text(json.dumps([{"title": str(i)} for i in range(6)]))
    for i in range(1, 7):
        (folder / f"slide-{i}.png").write_bytes(b"PNG fixture")

    class Gateway:
        settings = SimpleNamespace(mode="api", visual_review=True, model_id="vl")
        counts = []

        async def json_request(self, prompt, payload, **kwargs):
            self.counts.append(len(kwargs["images"]))
            if len(kwargs["images"]) > 2:
                request = httpx.Request("POST", "https://example.test/v1/chat/completions")
                raise httpx.HTTPStatusError(
                    "rejected", request=request, response=httpx.Response(status, request=request)
                )
            return {"checked_slides": payload["image_order"], "findings": []}

    gateway = Gateway()
    results = [
        {"key": "executive", "title": "Главное", "slides": 6, "rendering": {"native_render": True}}
    ]
    report = asyncio.run(review_visuals(results, tmp_path, gateway, 10))
    if status == 401:
        assert report["status"] == "failed" and report["http_status"] == 401
        assert gateway.counts == [4]
    else:
        assert report["status"] == "completed" and report["checked"] == 6
        assert gateway.counts == [4, 2, 2, 2]
        assert [i for batch in report["batches"] for i in batch["slides"]] == list(range(1, 7))
        assert len(report["request_adjustments"]) == 1


def test_identical_render_and_context_reuses_verified_visual_result(tmp_path):
    results = []
    for key in ("executive", "analytical", "story"):
        folder = tmp_path / key
        folder.mkdir()
        (folder / "slides.json").write_text(json.dumps([{"title": "Одинаковый слайд"}]))
        (folder / "slide-1.png").write_bytes(b"same" if key != "story" else b"different")
        results.append({"key": key, "slides": 1, "rendering": {"native_render": True}})

    class Gateway:
        settings = SimpleNamespace(mode="api", visual_review=True, model_id="vl")
        calls = 0

        async def json_request(self, *args, **kwargs):
            self.calls += 1
            return {"checked_slides": [1], "findings": []}

    gateway = Gateway()
    report = asyncio.run(review_visuals(results, tmp_path, gateway, None))
    assert report["status"] == "completed" and report["checked"] == 3 and gateway.calls == 2
    assert [b["cache_hit"] for b in report["batches"]] == [False, True, False]
