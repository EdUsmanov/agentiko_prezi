"""Synthetic OOXML checks for static OLE previews in template backgrounds."""

from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from studio.composition.powerpoint import open_presentation
from studio.templates.portable_templates import (
    _normalize_ole_previews,
    clean_editable_source,
    extract_backgrounds,
)

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"


def test_ole_preview_uses_picture_classifier_in_both_template_paths(tmp_path):
    icon = BytesIO()
    Image.new("RGB", (96, 96), "#24629b").save(icon, format="PNG")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    expected = {}
    for index, (name, wrapped) in enumerate(
        (("Decorative blue panel", True), ("Screenshot sample", False)), 1
    ):
        icon.seek(0)
        shape = slide.shapes.add_ole_object(
            BytesIO(b"inert synthetic payload"),
            "Package",
            Inches(index),
            Inches(1),
            Inches(1),
            Inches(1),
            icon_file=icon,
        )
        frame = shape._element
        frame.find(f"{P}nvGraphicFramePr/{P}cNvPr").set("name", name)
        expected[str(shape.shape_id)] = (
            name,
            tuple(frame.find(f"{P}xfrm/{A}off").attrib.items()),
            tuple(frame.find(f"{P}xfrm/{A}ext").attrib.items()),
        )
        if wrapped:
            wrapper = etree.Element(f"{MC}AlternateContent")
            etree.SubElement(wrapper, f"{MC}Choice", Requires="p").append(frame)
            etree.SubElement(wrapper, f"{MC}Fallback")
            slide.shapes._spTree.append(wrapper)
    source = tmp_path / "synthetic-ole.pptx"
    prs.save(source)

    profile = SimpleNamespace(patterns=[], background_source=None)
    model = extract_backgrounds(profile, source, tmp_path)
    rows = {row["id"]: row for row in model["parts"]["ppt/slides/slide1.xml"]}
    assert (rows["2"]["type"], rows["2"]["action"]) == ("pic", "keep")
    assert (rows["3"]["type"], rows["3"]["action"]) == ("pic", "remove")

    editable = open_presentation(source)
    clean_editable_source(editable, [])
    edited = tmp_path / "editable.pptx"
    editable.save(edited)
    for output in (profile.background_source, edited):
        with ZipFile(output) as archive:
            names = archive.namelist()
            assert not any(name.startswith("ppt/embeddings/") for name in names)
            root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
            pictures = root.findall(f".//{P}pic")
            assert len(pictures) == 1
            assert not root.findall(f".//{P}graphicFrame")
            picture = pictures[0]
            identity = picture.find(f"{P}nvPicPr/{P}cNvPr")
            assert identity.get("id") == "2"
            assert identity.get("name") == expected["2"][0]
            transform = picture.find(f"{P}spPr/{A}xfrm")
            assert tuple(transform.find(f"{A}off").attrib.items()) == expected["2"][1]
            assert tuple(transform.find(f"{A}ext").attrib.items()) == expected["2"][2]


def test_ole_preview_can_be_in_alternate_content_fallback():
    icon = BytesIO()
    Image.new("RGB", (16, 16), "#24629b").save(icon, format="PNG")
    icon.seek(0)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_ole_object(
        BytesIO(b"inert synthetic payload"),
        "Package",
        Inches(1),
        Inches(1),
        Inches(1),
        Inches(1),
        icon_file=icon,
    )
    frame = shape._element
    ole = frame.find(f".//{P}oleObj")
    preview = ole.find(f"{P}pic")
    ole.remove(preview)
    wrapper = etree.Element(f"{MC}AlternateContent")
    etree.SubElement(wrapper, f"{MC}Choice", Requires="p").append(frame)
    etree.SubElement(wrapper, f"{MC}Fallback").append(preview)
    slide.shapes._spTree.append(wrapper)

    assert _normalize_ole_previews(prs)
    assert len(slide._element.findall(f".//{P}pic")) == 1
    assert not slide._element.findall(f".//{P}graphicFrame")
    assert not any("oleObject" in rel.reltype for rel in slide.part.rels.values())
