"""Synthetic OOXML coverage; no user template assets or slide copy."""

from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches
import pytest

from studio.composition.powerpoint import open_presentation
from studio.templates.portable_templates import (
    clean_editable_source,
    extract_backgrounds,
    template_reference,
)


def ole_template(tmp_path):
    image = BytesIO()
    Image.new("RGBA", (24, 16), (20, 90, 180, 150)).save(image, format="PNG")
    image.seek(0)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ole = slide.shapes.add_ole_object(
        BytesIO(b"opaque synthetic embedded data"),
        "Synthetic.Drawing",
        Inches(0.3),
        Inches(0.2),
        Inches(1.2),
        Inches(0.8),
        icon_file=image,
    )
    ole.name = "Brand mark"
    # The outer frame determines placement; its preview can carry stale geometry.
    ole._element.find(f".//{qn('p:pic')}/{qn('p:spPr')}/{qn('a:xfrm')}/{qn('a:off')}").set("x", "1")
    slide.shapes.add_table(2, 2, Inches(2), Inches(2), Inches(3), Inches(1))
    path = tmp_path / "synthetic.pptx"
    prs.save(path)
    return path, ole.shape_id, (ole.left, ole.top, ole.width, ole.height)


@pytest.mark.parametrize("alternate_content", [False, True])
def test_ole_preview_keeps_identity_geometry_and_classification(tmp_path, alternate_content):
    source, identity, bounds = ole_template(tmp_path)
    if alternate_content:
        from lxml import etree

        prs = open_presentation(source)
        frame = prs.slides[0].shapes[0]._element
        ole = frame.find(f".//{qn('p:oleObj')}")
        parent = ole.getparent()
        alternate = etree.SubElement(
            parent, "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
        )
        fallback = etree.SubElement(
            alternate, "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
        )
        fallback.append(ole)
        prs.save(source)
    original = source.read_bytes()
    normalized = open_presentation(BytesIO(template_reference(source)))
    picture = normalized.slides[0].shapes[0]
    assert picture._element.tag == qn("p:pic")
    assert picture.shape_id == identity
    assert (picture.left, picture.top, picture.width, picture.height) == bounds
    profile = SimpleNamespace(patterns=[], background_source="")
    model = extract_backgrounds(profile, source, tmp_path)
    roles = model["parts"]["ppt/slides/slide1.xml"]
    assert next(row for row in roles if row["id"] == str(identity))["action"] == "keep"
    assert next(row for row in roles if row["type"] == "graphicFrame")["action"] == "remove"
    cleaned = open_presentation(profile.background_source)
    assert [shape.shape_id for shape in cleaned.slides[0].shapes] == [identity]
    with ZipFile(profile.background_source) as package:
        assert not any(name.startswith("ppt/embeddings/") for name in package.namelist())
    editable = open_presentation(source)
    clean_editable_source(editable, [])
    assert [shape.shape_id for shape in editable.slides[0].shapes] == [identity]
    assert editable.slides[0].shapes[0]._element.tag == qn("p:pic")
    assert source.read_bytes() == original


def test_ole_without_preview_remains_subject_to_structured_content_rule(tmp_path):
    source, _, _ = ole_template(tmp_path)
    prs = open_presentation(source)
    prs.slides[0].shapes[0].name = "Embedded document"
    preview = prs.slides[0].shapes[0]._element.find(f".//{qn('p:pic')}")
    preview.getparent().remove(preview)
    prs.save(source)
    extract_backgrounds(SimpleNamespace(patterns=[], background_source=""), source, tmp_path)
    cleaned = open_presentation(tmp_path / "template-layers/background-source.pptx")
    assert len(cleaned.slides[0].shapes) == 0


def test_ole_preview_with_sample_content_is_still_removed(tmp_path):
    source, _, _ = ole_template(tmp_path)
    prs = open_presentation(source)
    prs.slides[0].shapes[0].name = "Sample chart"
    prs.save(source)
    profile = SimpleNamespace(patterns=[], background_source="")
    extract_backgrounds(profile, source, tmp_path)
    assert len(open_presentation(profile.background_source).slides[0].shapes) == 0
