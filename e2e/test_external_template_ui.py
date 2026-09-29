"""Opt-in real template UI checks; downloaded third-party decks stay outside Git."""

import hashlib
from io import BytesIO
import os
from pathlib import Path

from playwright.sync_api import expect
from pptx import Presentation
import pytest

from test_support.app_server import application


TEMPLATE = os.environ.get("STUDIO_E2E_TEMPLATE")
pytestmark = pytest.mark.skipif(not TEMPLATE, reason="Set STUDIO_E2E_TEMPLATE to a real PPTX/POTX")
SOURCE = (
    "# Пилот\n## Контекст\nОбщий сервис принимает заявки.\n"
    "## Проверка\nПилот длится 12 недель.\n"
    "## Решение\nКоманда принимает решение после проверки результата."
)


@pytest.mark.parametrize("input_mode", ["content", "brief"])
def test_real_template_approval_edit_download_and_history(browser_page, tmp_path, input_mode):
    template = Path(TEMPLATE).resolve()
    original = hashlib.sha256(template.read_bytes()).hexdigest()
    page = browser_page
    with application(tmp_path / "app") as (url, _, _):
        page.goto(url)
        expect(page.locator("#prepare-button")).to_be_enabled()
        page.locator("#template").set_input_files(template)
        page.locator("#input-mode").select_option(input_mode)
        page.locator("#content").fill(SOURCE)
        page.locator("#slides").select_option("mini")
        with page.expect_response(
            lambda r: r.url.endswith("/api/prepare") and r.request.method == "POST"
        ) as response:
            page.locator("#prepare-button").click()
        assert response.value.status == 202, response.value.text()
        original_pid = response.value.json()["id"]
        expect(page.locator("#prep-state")).to_have_text("Подготовлено", timeout=120_000)

        if input_mode == "brief":
            expect(page.locator(".brief-approval")).to_be_visible()
            expect(page.locator("#generate-button")).to_be_disabled()
            old_draft = page.request.get(f"{url}/api/packages/{original_pid}/draft").json()
            assert not old_draft["approved"]
            copy = page.get_by_label("Тезис 1 слайда 1", exact=True)
            edited = copy.input_value() + " Ответственный — команда пилота."
            copy.fill(edited)
            with page.expect_response(
                lambda r: (
                    r.url.endswith(f"/api/packages/{original_pid}/draft")
                    and r.request.method == "POST"
                )
            ) as saved:
                page.get_by_role("button", name="Сохранить изменения в новом черновике").click()
            assert saved.value.status == 202, saved.value.text()
            pid = saved.value.json()["id"]
            assert pid != original_pid
            expect(page.get_by_label("Тезис 1 слайда 1", exact=True)).to_have_value(
                edited, timeout=120_000
            )
            expect(
                page.get_by_role("button", name="Утвердить план и текст", exact=True)
            ).to_be_enabled()
            expect(page.locator("#generate-button")).to_be_disabled()
            with page.expect_response(
                lambda r: (
                    r.url.endswith(f"/api/packages/{pid}/approve") and r.request.method == "POST"
                )
            ) as approved:
                page.get_by_role("button", name="Утвердить план и текст", exact=True).click()
            assert approved.value.status == 200, approved.value.text()
            # The original source and unapproved package survive editing.
            unchanged = page.request.get(f"{url}/api/packages/{original_pid}/draft").json()
            assert unchanged == old_draft
        else:
            pid = original_pid

        expect(page.locator("#generate-button")).to_be_enabled(timeout=30_000)
        assert page.request.post(f"{url}/api/packages/{pid}/auto-generation/cancel").ok
        with page.expect_response(
            lambda r: r.url.endswith("/api/generate") and r.request.method == "POST"
        ) as generated:
            page.locator("#generate-button").click()
        assert generated.value.status == 202, generated.value.text()
        gid = generated.value.json()["id"]
        expect(page.locator("#results")).to_be_visible(timeout=180_000)
        done = page.request.get(f"{url}/api/jobs/{gid}").json()
        assert done["state"] in ("completed", "needs_review"), done
        assert done["quality_report"]["errors"] == 0, done
        pptx_links = page.get_by_role("link", name="Скачать PPTX", exact=False)
        expect(pptx_links).to_have_count(3)
        with page.expect_download() as download:
            pptx_links.first.click()
        output = tmp_path / "result.pptx"
        download.value.save_as(output)
        deck = Presentation(BytesIO(output.read_bytes()))
        assert len(deck.slides) == 3
        text = "\n".join(
            shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
        )
        assert "12 недель" in text
        if input_mode == "brief":
            assert " ".join(edited.split()) in " ".join(text.split())
        page.locator("#nav-history").click()
        expect(page.locator("#history-list")).to_contain_text(template.name)
    assert hashlib.sha256(template.read_bytes()).hexdigest() == original
